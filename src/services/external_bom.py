from infrastructure.operation_progress import checkpoint, file_work
"""Read Excel into a detached comparison graph. Never accesses a database."""
from pathlib import Path
from collections import defaultdict
from copy import deepcopy
import math,re

COLUMNS=[('层级','level'),('序号','serial'),('父项物料编码','parent_base'),('父项版本号','parent_version'),
 ('物料基础编码','base_code'),('物料版本号','version'),('物料名称','name'),('规格','description'),
 ('用量','quantity'),('单位','unit'),('物料等级','grade'),('制造商','manufacturer'),('MPN','mpn'),
 ('M/P','mp_flag'),('物料位置','position'),('备注','remark')]
REQUIRED={'父项物料编码','父项版本号','物料基础编码','物料版本号','物料名称','用量','单位'}
ALT_SUFFIXES=('编码','版本号','名称','规格','制造商','MPN')
ALIASES={'父项物料编码':['父项编码'],'物料基础编码':['物料编码'],'用量':['物料数量','数量'],'规格':['物料规格描述','物料描述'],
 '物料位置':['位置'],'制造商':['物料制造商','供应商'],'MPN':['物料MPN'],'父项版本号':['父项版本']}

def text(value):
    if value is None:return ''
    if isinstance(value,float) and value.is_integer():return str(int(value))
    return str(value).strip()

class ExcelSource:
    def __init__(self,path):
        self.xls=Path(path).suffix.lower()=='.xls'
        if self.xls:
            import xlrd
            self.book=xlrd.open_workbook(str(path),on_demand=True)
            self.names=self.book.sheet_names()
        else:
            from openpyxl import load_workbook
            self.book=file_work(load_workbook,path,read_only=True,data_only=True);self.names=self.book.sheetnames
    def rows(self,name):
        if self.xls:
            sheet=self.book.sheet_by_name(name)
            return [sheet.row_values(i) for i in range(sheet.nrows)]
        return list(self.book[name].iter_rows(values_only=True))
    def close(self):
        if self.xls:self.book.release_resources()
        else:self.book.close()

def build_external_bom(rows,mapping):
    headers=[text(v) for v in rows[0]] if rows else []
    issues=[];indices={}
    for field in REQUIRED:
        checkpoint()
        if not mapping.get(field):issues.append('必选项未映射：'+field)
    used=[]
    for field,target in mapping.items():
        checkpoint()
        if not target:continue
        if headers.count(target)!=1:issues.append('源表头缺失或重名：'+target);continue
        indices[field]=headers.index(target);used.append(target)
    if len(used)!=len(set(used)):issues.append('多个字段映射到同一源列')
    if issues:raise ValueError('\n'.join(issues))
    groups=sorted({int(m.group(1)) for f in mapping for m in [re.fullmatch(r'替代料(\d+)'+r'(?:'+'|'.join(ALT_SUFFIXES)+')',f)] if m})
    for n in groups:
        checkpoint()
        pair=[mapping.get(f'替代料{n}{suffix}') for suffix in ('编码','版本号')]
        if any(mapping.get(f'替代料{n}{suffix}') for suffix in ALT_SUFFIXES) and not all(pair):issues.append(f'替代料{n}编码与版本号必须成对映射')
    edges=defaultdict(dict);roots=[];defined=set();parents=set()
    for line,values in enumerate(rows[1:],2):
        checkpoint()
        def get(field):
            i=indices.get(field);return text(values[i]) if i is not None and i<len(values) else ''
        if not any(get(f) for f in indices):continue
        missing=[f for f in REQUIRED if not get(f)]
        if missing:issues.append(f'第{line}行缺少：'+ '、'.join(sorted(missing)));continue
        data={key:get(label) for label,key in COLUMNS}
        try:
            quantity=float(data['quantity'])
            if not math.isfinite(quantity) or quantity<=0:raise ValueError()
        except ValueError:issues.append(f'第{line}行用量必须是大于0的有限数字');continue
        code=(data['base_code'],data['version']);parent=(data.pop('parent_base'),data.pop('parent_version'))
        data.update(material_id=code,quantity=quantity,children=[],alternate_ids=frozenset(),_external=True)
        alt_codes=[];alt_text=[]
        for n in groups:
            checkpoint()
            vals=[get(f'替代料{n}{suffix}') for suffix in ALT_SUFFIXES]
            if not any(vals):continue
            if not all(vals[:2]):issues.append(f'第{line}行替代料{n}缺少编码或版本号');continue
            if tuple(vals[:2])==code:issues.append(f'第{line}行物料不能作为自身替代料');continue
            alt_codes.append(tuple(vals[:2]));alt_text.append(' | '.join(vals))
        data.update(alternate_codes=frozenset(alt_codes),alternates='\n'.join(sorted(set(alt_text))))
        root_flags=[x.casefold()=='root' for x in parent]
        if any(root_flags):
            if not all(root_flags):issues.append(f'第{line}行父项编码和版本号必须同时为ROOT')
            roots.append((code,data));defined.add(code);continue
        defined.add(code);parents.add(parent)
        existing=edges[parent].get(code)
        # Imported sequence numbers/order do not define shared material structure.
        data.pop('serial',None);data.pop('level',None)
        if existing and existing!=data:issues.append(f'第{line}行同一父项的重复子物料信息不一致：'+''.join(code))
        else:edges[parent][code]=data
    if len(roots)>1:issues.append('存在多个ROOT根节点')
    if not roots:
        missing_parents=parents-defined
        if len(missing_parents)!=1:issues.append('无法确定唯一根节点，请提供父项编码及版本号为ROOT的根节点行')
        else:
            code=next(iter(missing_parents));roots=[(code,dict(base_code=code[0],version=code[1],name='',material_id=code,quantity=1,children=[],alternates='',alternate_codes=frozenset(),_external=True))]
    elif parents-defined:issues.append('父项物料没有对应记录：'+'、'.join(''.join(x) for x in sorted(parents-defined)))
    visited=set();active=set()
    def check(code):
        if code in active:issues.append('存在循环引用：'+''.join(code));return
        if code in visited:return
        active.add(code)
        for child in edges.get(code,{}):check(child)
        active.remove(code);visited.add(code)
    for code in list(edges):check(code)
    if roots:
        reachable=set();pending=[roots[0][0]]
        while pending:
            checkpoint()
            code=pending.pop()
            if code in reachable:continue
            reachable.add(code);pending.extend(edges.get(code,{}))
        if defined-reachable:issues.append('存在未连接到根节点的物料：'+'、'.join(''.join(x) for x in sorted(defined-reachable)))
    if issues:raise ValueError('外部BOM未载入：\n'+'\n'.join(issues))
    if not roots:raise ValueError('文件中没有BOM数据')
    def expand(data,parent=None,serial='1',depth=0):
        node=deepcopy(data);code=(node['base_code'],node['version'])
        node.update(level=depth,serial=serial,parent_code=''.join(parent) if parent else 'ROOT')
        node['children']=[expand(child,code,serial+'.'+str(i),depth+1) for i,child in enumerate(edges.get(code,{}).values(),1)]
        return node
    return expand(roots[0][1])
