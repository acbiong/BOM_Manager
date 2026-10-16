from domain.quantity import format_quantity
import math
"""业务服务层：物料、SKU、BOM、日志和Excel；UI不得直接访问数据库。编辑时间：2026-09-24。"""
from pathlib import Path
from datetime import datetime, timezone
from openpyxl import Workbook,load_workbook
class BomImportError(ValueError):
 def __init__(self,issues):
  self.issues=issues
  super().__init__(f'BOM扫描完成，共发现{len(issues)}个问题，未导入：\n'+'\n'.join(issues))

class MaterialHeaderError(ValueError):
 def __init__(self,headers,missing):
  self.headers=headers; self.missing=missing
  super().__init__('缺少表头：'+ '、'.join(missing))
from config.config_service import DEFAULT_MATERIAL_UNITS
MATERIAL_IMPORT_HEADERS=['物料基础编码','物料版本号','物料名称','物料规格描述','单位','物料等级','M/P','物料制造商','物料MPN']
MATERIAL_IMPORT_DEFAULTS={'物料等级':'','M/P':'M'}

class Services:
 UNITS=list(DEFAULT_MATERIAL_UNITS); MP=['M','P']
 def __init__(self,db,config_service=None):
  self.db=db; self.config_service=config_service; self.UNITS=config_service.material_units() if config_service else list(DEFAULT_MATERIAL_UNITS)
 MATERIAL_FIELDS=[('base_code','基础编码'),('version','版本号'),('name','物料名称'),('description','规格描述'),('unit','单位'),('grade','物料等级'),('manufacturer','制造商'),('mpn','MPN'),('mp_flag','M/P')]
 def material_snapshot(self,i):
  row=self.db.conn.execute('SELECT * FROM materials WHERE id=?',(i,)).fetchone()
  return dict(row) if row else {}
 def material_detail(self,before,after):
  def text(value):return '（空）' if value is None or value=='' else str(value)
  return '；'.join(f'{label}：{text(before.get(key))} → {text(after.get(key))}' for key,label in self.MATERIAL_FIELDS if before.get(key)!=after.get(key))
 def log(self,typ,obj,act,detail,snapshot=None):
  if snapshot is None and typ in ('material','bom'):snapshot=self.material_snapshot(obj)
  if snapshot is None and typ=='sku':
   sku=self.db.conn.execute('SELECT material_id FROM skus WHERE id=?',(obj,)).fetchone()
   if sku:snapshot=self.material_snapshot(sku['material_id'])
  snapshot=snapshot or {}
  code=f"{snapshot['base_code']}{snapshot['version']}" if snapshot else None
  self.db.conn.execute('INSERT INTO change_logs(object_type,object_id,action,detail,material_code,material_name,changed_at) VALUES(?,?,?,?,?,?,?)',(typ,obj,act,detail,code,snapshot.get('name'),datetime.now().astimezone().strftime('%Y-%m-%d %H:%M:%S.%f')[:-4]))
 def publish(self,event):
  from infrastructure.events import WindowRefreshBus
  if event.get('type')=='bom_changed' and getattr(self,'scope_id',0):event=dict(event,scope_id=self.scope_id)
  database=getattr(self.db,'original',self.db)
  if getattr(database,'_transaction_depth',0):
   database._pending_events=getattr(database,'_pending_events',[])+[event]
  else:WindowRefreshBus.publish(event)
 def for_sku(self,sid):
  row=self.db.conn.execute('SELECT * FROM skus WHERE id=?',(sid,)).fetchone()
  if row is None:raise ValueError('SKU不存在')
  if not row['scope_id']:return self
  from services.mirror_service import MirrorServices
  return MirrorServices(getattr(self,'primary',self),sid)
 def bom_skus(self):return [row for row in self.skus() if row['scope_id']==getattr(self,'scope_id',0)]
 def copy_sku_bom(self,sid):
  if getattr(self,'scope_id',0):return self.primary.copy_sku_bom(sid)
  from services.mirrors import create_namespace
  source=self.for_sku(sid)
  sku=self.db.conn.execute('SELECT * FROM skus WHERE id=?',(sid,)).fetchone()
  with self.db.transaction():
   version=max(sku['mirror_version']+1,self.db.conn.execute('SELECT MAX(mirror_version)+1 FROM skus WHERE material_id=?',(sku['material_id'],)).fetchone()[0])
   new_id=self.db.conn.execute('INSERT INTO skus(material_id,mirror_version) VALUES(?,?)',(sku['material_id'],version)).lastrowid
   self.db.conn.execute('UPDATE skus SET scope_id=? WHERE id=?',(new_id,new_id))
   create_namespace(self.db,new_id,source)
   copied=self.for_sku(new_id);copied.refresh_usage()
   copied.db.conn.execute("UPDATE change_logs SET object_id=? WHERE object_type='sku' AND object_id=?",(new_id,sid))
   copied.log('sku',new_id,'copy',f'复制BOM；源镜像版本：{sku["mirror_version"]}；新镜像版本：{version}')
   from services.current_catalog import refresh_latest_from_catalog
   refresh_latest_from_catalog(self)
  self.publish({'type':'sku_changed','id':new_id});return new_id
 def change_mirror_version(self,sid,value):
  try:value=int(value)
  except (TypeError,ValueError):raise ValueError('镜像版本号必须为正整数')
  if value<1:raise ValueError('镜像版本号必须为正整数')
  row=self.db.conn.execute('SELECT * FROM skus WHERE id=?',(sid,)).fetchone()
  if row is None:raise ValueError('SKU不存在')
  if self.db.conn.execute('SELECT 1 FROM skus WHERE material_id=? AND mirror_version=? AND id<>?',(row['material_id'],value,sid)).fetchone():raise ValueError('此SKU镜像版本号已存在')
  with self.db.transaction():
   stamp=datetime.now(timezone.utc).strftime('%Y-%m-%d %H:%M:%S.%f')
   self.db.conn.execute('UPDATE skus SET mirror_version=?,updated_at=? WHERE id=?',(value,stamp,sid))
   if not row['scope_id']:self.db.conn.execute('UPDATE materials SET updated_at=? WHERE id=?',(stamp,row['material_id']))
   self.for_sku(sid).log('sku',sid,'mirror_version',f'镜像版本号：{row["mirror_version"]} → {value}')
   from services.current_catalog import refresh_latest_from_catalog
   refresh_latest_from_catalog(self)
  self.publish({'type':'sku_changed','id':sid})
 def update_sku_change_description(self,sid,text):
  row=self.db.conn.execute('SELECT change_description,material_id,scope_id FROM skus WHERE id=?',(sid,)).fetchone()
  if row is None:raise ValueError('SKU不存在')
  if row['change_description']==text:return
  stamp=datetime.now().strftime('%Y-%m-%d %H:%M:%S.%f')[:-4]
  with self.db.transaction():
   self.db.conn.execute('UPDATE skus SET change_description=?,updated_at=? WHERE id=?',(text,stamp,sid))
   if not row['scope_id']:self.db.conn.execute('UPDATE materials SET updated_at=? WHERE id=?',(stamp,row['material_id']))
   self.for_sku(sid).log('sku',sid,'change_description','变更描述：'+(row['change_description'] or '（空）')+' → '+(text or '（空）'))
  self.publish({'type':'sku_changed','id':sid})
 def materials(self,query=''):
  # 自身有下阶仅决定has_bom；SKU、有效BOM主料或替代料引用均为非呆滞。
  q=f'%{query}%'; rows=self.db.conn.execute('SELECT m.id,m.base_code,m.version,m.name,m.description,m.unit,m.grade,m.manufacturer,m.mpn,m.mp_flag,m.has_bom AS stored_has_bom,m.created_at,m.updated_at,CASE WHEN EXISTS(SELECT 1 FROM material_bom_items b WHERE b.parent_material_id=m.id) OR EXISTS(SELECT 1 FROM mirror_material_usage u JOIN skus s ON s.id=u.sku_id WHERE u.parent_id=m.id AND u.role="main") THEN "Y" ELSE "N" END AS has_bom,CASE WHEN EXISTS(SELECT 1 FROM material_bom_items b WHERE b.child_material_id=m.id) OR EXISTS(SELECT 1 FROM alternates a JOIN material_bom_items b ON b.id=a.bom_item_id WHERE a.material_id=m.id) OR EXISTS(SELECT 1 FROM mirror_material_usage u WHERE u.material_id=m.id) OR EXISTS(SELECT 1 FROM skus s WHERE s.material_id=m.id) THEN "N" ELSE "Y" END AS idle FROM materials m WHERE base_code LIKE ? OR name LIKE ? ORDER BY base_code,version',(q,q)).fetchall()
  from services.snapshot_bom import enabled,latest_usage
  if enabled(self.db):
   used=latest_usage(self)
   from services.current_catalog import enabled as catalog_enabled
   if catalog_enabled(self):
    primary=getattr(self,'primary',self)
    parents={r[0] for r in primary.db.conn.execute('SELECT DISTINCT parent_material_id FROM material_bom_items')}
    from services.current_catalog import material_structure_sources
    for mid,sources in material_structure_sources(primary).items():
     if sources[0]['children']:parents.add(mid)
     else:parents.discard(mid)
    return [dict(row,latest_used='Y' if row['id'] in used else 'N',has_bom='Y' if row['id'] in parents else 'N') for row in rows]
   return [dict(row,latest_used='Y' if row['id'] in used else 'N') for row in rows]
  return rows
 def create_material(self,d):
  if not all(d.get(k) for k in ('base_code','version','name')): raise ValueError('基础编码、版本号、名称为必填')
  if d.get('unit','pcs') not in self.UNITS: raise ValueError('单位不合规')
  with self.db.transaction():
   c=self.db.conn.execute('INSERT INTO materials(base_code,version,name,description,unit,grade,manufacturer,mpn,mp_flag) VALUES(?,?,?,?,?,?,?,?,?)',(d['base_code'],d['version'],d['name'],d.get('description',''),d.get('unit','pcs'),d.get('grade',''),d.get('manufacturer',''),d.get('mpn',''),d.get('mp_flag','M'))); i=c.lastrowid; self.log('material',i,'create','新建物料；'+self.material_detail({},self.material_snapshot(i)))
  try:
   from infrastructure.events import WindowRefreshBus; self.publish({'type':'material_changed','id':i})
  except Exception: pass
  return i
 def update_material(self,i,d):
  if d.get('unit') not in self.UNITS:raise ValueError('单位不合规，允许：'+ '、'.join(self.UNITS))
  before=self.material_snapshot(i)
  if not before:raise ValueError('物料不存在')
  with self.db.transaction(): self.db.conn.execute('UPDATE materials SET name=?,description=?,unit=?,grade=?,manufacturer=?,mpn=?,mp_flag=?,updated_at=CURRENT_TIMESTAMP WHERE id=?',(d['name'],d.get('description',''),d['unit'],d.get('grade',before.get('grade','')),d.get('manufacturer',''),d.get('mpn',''),d.get('mp_flag','M'),i)); self.log('material',i,'update','修改物料；'+(self.material_detail(before,self.material_snapshot(i)) or '物料属性未发生变化'))
  self.db.conn.execute('UPDATE skus SET updated_at=CURRENT_TIMESTAMP WHERE material_id=?',(i,));self.db.conn.commit()
  from infrastructure.events import WindowRefreshBus
  self.publish({'type':'material_changed','id':i})
 def delete_material(self,i):
  if self.db.conn.execute('SELECT 1 FROM mirror_material_usage WHERE material_id=?',(i,)).fetchone():raise ValueError('物料仍被镜像BOM引用，请先删除引用')
  if self.db.conn.execute('SELECT 1 FROM skus WHERE material_id=?',(i,)).fetchone(): raise ValueError('物料是SKU，请先删除SKU引用')
  r=self.db.conn.execute('SELECT COUNT(*) n FROM material_bom_items WHERE parent_material_id=? OR child_material_id=?',(i,i)).fetchone();
  if r['n']: raise ValueError('物料仍被BOM引用，请先删除引用')
  if self.db.conn.execute('SELECT 1 FROM alternates a JOIN material_bom_items b ON b.id=a.bom_item_id WHERE a.material_id=?',(i,)).fetchone():raise ValueError('物料仍被BOM作为替代料引用，请先删除引用')
  before=self.material_snapshot(i)
  with self.db.transaction(): self.db.conn.execute('DELETE FROM material_bom_items WHERE parent_material_id=? OR child_material_id=?',(i,i)); self.db.conn.execute('DELETE FROM materials WHERE id=?',(i,)); self.log('material',i,'delete','删除物料；删除前：'+self.material_detail({},before),before)
 def skus(self): return self.db.conn.execute('SELECT s.id,m.id AS material_id,m.base_code,m.version,m.name,m.description,m.unit,m.grade,m.manufacturer,m.mpn,m.mp_flag,m.has_bom,s.mirror_version,s.scope_id,s.change_description,CASE WHEN s.scope_id=0 THEN m.updated_at ELSE s.updated_at END AS updated_at FROM skus s JOIN materials m ON m.id=s.material_id ORDER BY s.id').fetchall()
 def sku_material_quantities(self,skus=None):
  """每个SKU数量为1；仅累计主BOM，Decimal保留路径连乘精度。"""
  from decimal import Decimal
  graph={}
  for row in self.db.conn.execute('SELECT parent_material_id,child_material_id,quantity FROM material_bom_items ORDER BY sort_no,id'):
   graph.setdefault(row['parent_material_id'],[]).append((row['child_material_id'],Decimal(str(row['quantity']))))
  result={}
  for sku in self.skus() if skus is None else skus:
   if sku['scope_id'] and sku['scope_id']!=getattr(self,'scope_id',0):
    result.update(self.for_sku(sku['id']).sku_material_quantities());continue
   root=sku['material_id'];totals={root:Decimal('1')}
   def walk(parent,quantity,ancestors):
    for child,usage in graph.get(parent,[]):
     if child in ancestors:continue
     amount=quantity*usage;totals[child]=totals.get(child,Decimal('0'))+amount
     walk(child,amount,ancestors|{child})
   walk(root,Decimal('1'),{root});result[sku['id']]=totals
  return result
 def create_sku(self,mid,require_empty_bom=False):
  if not self.material_snapshot(mid):raise ValueError('物料不存在')
  if require_empty_bom and self.db.conn.execute('SELECT 1 FROM material_bom_items WHERE parent_material_id=?',(mid,)).fetchone():raise ValueError('物料已有BOM结构，不可用于新建SKU')
  if self.db.conn.execute('SELECT 1 FROM skus WHERE material_id=?',(mid,)).fetchone(): raise ValueError('已是SKU')
  from services.snapshot_bom import enabled,record_definitions
  with self.db.transaction():
   c=self.db.conn.execute('INSERT INTO skus(material_id) VALUES(?)',(mid,));sid=c.lastrowid
   self.log('sku',sid,'create','创建SKU')
   self.db.conn.execute('UPDATE sku_history_reset SET start_log_id=(SELECT MAX(id) FROM change_logs) WHERE root_id=?',(mid,))
   if enabled(self.db):
    from services.mirrors import create_namespace
    create_namespace(self.db,sid,self)
    self.db.conn.execute('UPDATE skus SET scope_id=? WHERE id=?',(sid,sid))
    scoped=self.for_sku(sid);scoped.refresh_usage();record_definitions(scoped,mid)
  from infrastructure.events import WindowRefreshBus
  self.publish({'type':'sku_changed','id':sid})
  return sid
 def material_is_bom_referenced(self,mid):
  return bool(self.db.conn.execute('SELECT 1 FROM mirror_material_usage WHERE material_id=?',(mid,)).fetchone() or self.db.conn.execute('SELECT 1 FROM material_bom_items WHERE child_material_id=? LIMIT 1',(mid,)).fetchone() or self.db.conn.execute('SELECT 1 FROM alternates a JOIN material_bom_items b ON b.id=a.bom_item_id WHERE a.material_id=? LIMIT 1',(mid,)).fetchone())
 def delete_sku(self,sid):
  row=self.db.conn.execute('SELECT * FROM skus WHERE id=?',(sid,)).fetchone()
  if row and row['scope_id']:
   from services.mirrors import drop_namespace
   with self.db.transaction():
    drop_namespace(self.db,sid);self.db.conn.execute('DELETE FROM skus WHERE id=?',(sid,))
    from services.current_catalog import enabled as catalog_enabled,refresh_latest_from_catalog,prune_public
    if catalog_enabled(self):
     self.db.conn.execute('DELETE FROM bom_revision_legacy_conflicts WHERE sku_id=?',(sid,))
     refresh_latest_from_catalog(self);prune_public(self)
   self.publish({'type':'sku_changed','id':sid});return
  with self.db.transaction():
   r=self.db.conn.execute('SELECT material_id FROM skus WHERE id=?',(sid,)).fetchone()
   if r is None:raise ValueError('SKU不存在')
   root=r['material_id']
   preserve_bom=self.material_is_bom_referenced(root)
   if not preserve_bom:
    self.db.conn.execute('DELETE FROM alternates WHERE bom_item_id IN (SELECT id FROM material_bom_items WHERE parent_material_id=?)',(root,))
    self.db.conn.execute('DELETE FROM material_bom_items WHERE parent_material_id=?',(root,))
    self.db.conn.execute('DELETE FROM bom_items WHERE bom_root_id=?',(root,))
    self.db.conn.execute('UPDATE materials SET has_bom=0 WHERE id=?',(root,))
    owned={row[0] for row in self.db.conn.execute('SELECT log_id FROM bom_log_roots WHERE root_id=?',(root,))}
    owned.update(row[0] for row in self.db.conn.execute("SELECT id FROM change_logs WHERE object_type='bom' AND object_id=?",(root,)))
    self.db.conn.execute('DELETE FROM bom_log_roots WHERE root_id=?',(root,))
    for log_id in owned:
     if not self.db.conn.execute('SELECT 1 FROM bom_log_roots WHERE log_id=?',(log_id,)).fetchone():self.db.conn.execute('DELETE FROM change_logs WHERE id=?',(log_id,))
   self.db.conn.execute("DELETE FROM change_logs WHERE object_type='sku' AND object_id=?",(sid,))
   self.db.conn.execute('INSERT OR REPLACE INTO sku_history_reset(root_id,start_log_id) VALUES(?,0)',(root,))
   self.db.conn.execute('DELETE FROM skus WHERE id=?',(sid,))
  from infrastructure.events import WindowRefreshBus
  self.publish({'type':'sku_changed','id':sid})
  if not preserve_bom:self.publish({'type':'bom_changed','root':r['material_id']})
 def bom_graph_edges(self):return self.db.conn.execute('SELECT parent_material_id,child_material_id FROM material_bom_items').fetchall()
 def bom(self,root):
  """递归展开物料关系表，返回兼容BOM编辑器的扁平树行。"""
  result=[]
  def walk(parent_id,level,parent_item_id,parent_node_id,ancestors):
   if parent_id in ancestors:return
   ancestors=ancestors|{parent_id}
   rows=self.db.conn.execute('SELECT b.*,m.base_code,m.version,m.name,m.description,m.unit AS material_unit,m.grade,m.manufacturer,m.mpn,m.mp_flag FROM material_bom_items b JOIN materials m ON m.id=b.child_material_id WHERE b.parent_material_id=? ORDER BY b.sort_no,b.id',(parent_id,)).fetchall()
   for r in rows:
    if r['child_material_id'] in ancestors:continue
    node_id=f"{parent_node_id or ''}/{r['id']}"
    d=dict(r); d.update(material_id=r['child_material_id'],unit=r['material_unit'],level_no=level,parent_item_id=parent_item_id,bom_root_id=root,node_id=node_id,parent_node_id=parent_node_id)
    result.append(d); walk(r['child_material_id'],level+1,r['id'],node_id,ancestors)
  walk(root,1,None,None,set()); return result
 def save_bom(self,root,items,preview=False,progress=None):
  """把树中的重复子BOM归并为物料关系，在单个事务中保存。"""
  from services.snapshot_bom import revision_conflicts,ParentRevisionRequired,enabled,record_definitions
  violations=[] if preview or getattr(self,'_catalog_sync',False) else revision_conflicts(self,root,items)
  if violations:raise ParentRevisionRequired(violations)
  report=progress or (lambda value,message:None)
  report(5,'读取现有BOM…')
  old_rows=self.bom(root)
  retained={root}|{row['material_id'] for row in items}
  parents={root}|{row['parent_material_id'] for row in old_rows if row['parent_material_id'] in retained}
  # 编辑器提供完整子树；新增父物料的全部子项被标删时也需要清空其原共享BOM。
  parents.update(row['material_id'] for row in items if row.get('replace_children'))
  by_id={row['id']:row for row in items if row.get('id') is not None}
  desired={}; alternate_updates={}
  old={ (row['parent_material_id'],row['child_material_id'],row['position'] or ''):dict(row)
        for row in self.db.conn.execute('SELECT * FROM material_bom_items') }
  def payload(row):
   return (float(row['quantity']),row['unit'],row.get('remark','') or '')
  for row in items:
   parent=row.get('parent_material_id')
   if parent is None:
    pid=row.get('parent_item_id'); parent_row=by_id.get(pid)
    if pid is not None and parent_row is None:
     if isinstance(pid,int) and 1<=pid<=len(items): parent_row=items[pid-1]
     else: raise ValueError('BOM结构存在无效的父级关系')
    parent=parent_row['material_id'] if parent_row else root
   key=(parent,row['material_id'],row.get('position','') or '')
   if 'alternate_material_ids' in row:
    alternates=list(dict.fromkeys(row['alternate_material_ids']))
    if key in alternate_updates and alternate_updates[key]!=alternates:raise ValueError('共享BOM的替代料修改不一致')
    if row['material_id'] in alternates:raise ValueError('物料不能作为自身的替代料')
    for mid in alternates:
     if not self.db.conn.execute('SELECT 1 FROM materials WHERE id=?',(mid,)).fetchone():raise ValueError('替代料不存在，请重新选择')
    alternate_updates[key]=alternates
   values=payload(row)
   if not math.isfinite(values[0]) or values[0]<=0: raise ValueError('物料数量必须为大于0的有限数字')
   parents.add(parent)
   if key in desired and desired[key]!=values:
    baseline=payload(old[key]) if key in old else None
    if values==baseline: continue
    if desired[key]!=baseline:
     material=self.db.conn.execute('SELECT base_code,version,name FROM materials WHERE id=?',(row['material_id'],)).fetchone()
     raise ValueError(f'共享子BOM的同一关系存在不同修改：{material[0]}{material[1]} {material[2]}。请统一数量、位置和备注后保存。')
   desired[key]=values
  # 替代料属于父子关系。只通知真正改变的父级BOM，不能把整棵树都当成已更新。
  changed_parents=set()
  old_groups={}; new_groups={}
  for key,row in sorted(old.items(),key=lambda pair:(pair[1]['sort_no'],pair[1]['id'])):
   old_groups.setdefault(key[0],[]).append((key,payload(row)))
  for key,values in desired.items():new_groups.setdefault(key[0],[]).append((key,values))
  for parent in parents:
   if old_groups.get(parent,[])!=new_groups.get(parent,[]):changed_parents.add(parent)
  for key,alternates in alternate_updates.items():
   previous=[row['material_id'] for row in self.alternates(old[key]['id'])] if key in old else []
   if previous!=alternates:changed_parents.add(key[0])
  # 同一个物料关系可能在树的多条路径出现，但数据库只保存一份。
  graph={}
  for key in old:
   if key[0] not in parents: graph.setdefault(key[0],set()).add(key[1])
  for parent,child,_ in desired: graph.setdefault(parent,set()).add(child)
  visiting=set(); visited=set()
  def check(node):
   if node in visiting: raise ValueError('BOM包含循环引用，请检查父级和子级物料')
   if node in visited: return
   visiting.add(node)
   for child in graph.get(node,()): check(child)
   visiting.remove(node); visited.add(node)
  for parent in parents: check(parent)
  if preview:
   # 使用与正式保存完全相同的关系集合生成编辑树，不写库、不发送刷新通知。
   relations={key:payload(row) for key,row in old.items() if key[0] not in parents}
   relations.update(desired)
   groups={}; materials={r['id']:dict(r) for r in self.materials()}
   for number,(key,values) in enumerate(relations.items(),1):
    parent,child,position=key; quantity,unit,remark=values
    previous=old.get(key); relation_id=previous['id'] if previous else -number
    data=dict(materials[child]); data.update(id=relation_id,material_id=child,parent_material_id=parent,quantity=quantity,unit=unit,position=position,remark=remark)
    data['_db_alternates']=[r['material_id'] for r in self.alternates(relation_id)] if previous else []
    if key in alternate_updates:data['alternate_material_ids']=alternate_updates[key]
    groups.setdefault(parent,[]).append(data)
   def expand(parent):
    return [{'data':dict(data),'children':expand(data['material_id'])} for data in groups.get(parent,[])]
   return expand(root)
  report(25,'检查结构并计算联动范围…')
  def reachable(root,relations):
   found={root};pending=[root]
   while pending:
    for child in relations.get(pending.pop(),()):
     if child not in found:found.add(child);pending.append(child)
   return found
  old_graph={}
  for parent,child,_ in old:old_graph.setdefault(parent,set()).add(child)
  log_roots={sku['material_id']:reachable(sku['material_id'],old_graph) for sku in self.bom_skus()}
  log_roots.setdefault(root,{root}|{row['material_id'] for row in old_rows})
  def alternate_map():
   mapping={}
   for row in self.db.conn.execute('SELECT bom_item_id,material_id FROM alternates ORDER BY sort_no,id'):
    mapping.setdefault(row[0],[]).append(row[1])
   return mapping
  old_alternates=alternate_map()
  report(40,'正在写入BOM关系…')
  saved_at=datetime.now(timezone.utc).strftime('%Y-%m-%d %H:%M:%S.%f')
  affected_skus=[]
  with self.db.transaction():
   for key,row in old.items():
    if key[0] in parents and key not in desired:
     self.db.conn.execute('DELETE FROM alternates WHERE bom_item_id=?',(row['id'],))
     self.db.conn.execute('DELETE FROM material_bom_items WHERE id=?',(row['id'],))
   for order,(key,values) in enumerate(desired.items()):
    parent,child,position=key; quantity,unit,remark=values
    if key in old and payload(old[key])==values and old[key]['sort_no']==order:continue
    self.db.conn.execute('INSERT INTO material_bom_items(parent_material_id,child_material_id,quantity,unit,position,remark,sort_no) VALUES(?,?,?,?,?,?,?) ON CONFLICT(parent_material_id,child_material_id,position) DO UPDATE SET quantity=excluded.quantity,unit=excluded.unit,remark=excluded.remark,sort_no=excluded.sort_no,updated_at=CURRENT_TIMESTAMP',(parent,child,quantity,unit,position,remark,order))
   # 只清理本次移除分支中不再被外部BOM或SKU使用的父物料结构。
   candidates={row['material_id'] for row in old_rows}-retained
   protected={r[0] for r in self.db.conn.execute('SELECT material_id FROM skus')}
   remaining=list(self.db.conn.execute('SELECT parent_material_id,child_material_id FROM material_bom_items'))
   protected.update(row[1] for row in remaining if row[0] not in candidates)
   pending=list(protected)
   graph={}
   for parent,child in remaining:graph.setdefault(parent,[]).append(child)
   while pending:
    parent=pending.pop()
    for child in graph.get(parent,[]):
     if child not in protected:protected.add(child);pending.append(child)
   for parent in candidates-protected:
    self.db.conn.execute('DELETE FROM alternates WHERE bom_item_id IN (SELECT id FROM material_bom_items WHERE parent_material_id=?)',(parent,))
    self.db.conn.execute('DELETE FROM material_bom_items WHERE parent_material_id=?',(parent,))
    self.db.conn.execute('DELETE FROM bom_items WHERE bom_root_id=?',(parent,))
    changed_parents.add(parent)
   self.db.conn.execute('DELETE FROM bom_items WHERE bom_root_id=?',(root,))
   self.db.conn.execute('DELETE FROM alternates WHERE NOT EXISTS (SELECT 1 FROM material_bom_items b WHERE b.id=alternates.bom_item_id)')
   for key,alternates in alternate_updates.items():
    relation=self.db.conn.execute('SELECT id FROM material_bom_items WHERE parent_material_id=? AND child_material_id=? AND position=?',key).fetchone()[0]
    self.db.conn.execute('DELETE FROM alternates WHERE bom_item_id=?',(relation,))
    for order,mid in enumerate(alternates):self.db.conn.execute('INSERT INTO alternates(bom_item_id,material_id,sort_no) VALUES(?,?,?)',(relation,mid,order))
   if not getattr(self,'scope_id',0):self.db.conn.execute('UPDATE materials SET has_bom=CASE WHEN EXISTS(SELECT 1 FROM material_bom_items WHERE parent_material_id=materials.id) THEN 1 ELSE 0 END')
   # 使用保存前后包含的父物料判断影响范围，覆盖新增和删除共享分支。
   # 在同一事务内更新时间，保证界面收到通知时读取的是已提交数据。
   for sku in self.bom_skus():
    sku_root=sku['material_id']
    contained=log_roots.get(sku_root,{sku_root})|reachable(sku_root,graph)
    if contained & changed_parents:
     if not getattr(self,'scope_id',0):self.db.conn.execute('UPDATE materials SET updated_at=? WHERE id=?',(saved_at,sku_root))
     self.db.conn.execute('UPDATE skus SET updated_at=? WHERE id=?',(saved_at,sku['id']))
     affected_skus.append(sku['id'])
   after={ (r['parent_material_id'],r['child_material_id'],r['position'] or ''):dict(r) for r in self.db.conn.execute('SELECT * FROM material_bom_items') }
   report(70,'正在生成变更记录…')
   after_alternates=alternate_map()
   material_cache={}
   def label(mid):
    if mid not in material_cache:material_cache[mid]=self.material_snapshot(mid)
    material=material_cache[mid]
    return material.get('base_code','')+material.get('version','')+' '+material.get('name','')
   # 按父项+子物料配对位置变化，避免一处位置修改被写成删除和新增两条。
   remaining=set(after); pairs=[]
   for key,before in old.items():
    target=key if key in remaining else None
    if target is None:
     candidates=[candidate for candidate in remaining if candidate[:2]==key[:2] and candidate not in old]
     old_candidates=[candidate for candidate in old if candidate[:2]==key[:2] and candidate not in after]
     if len(candidates)==len(old_candidates)==1:target=candidates[0]
    if target is not None:remaining.remove(target)
    pairs.append((key,before,target,after.get(target)))
   pairs.extend((key,None,key,after[key]) for key in after if key in remaining)
   for key,before,target,current in pairs:
    detail=[]
    if before is None:detail.append(f"新增物料；数量：{format_quantity(current['quantity'])}；单位：{current['unit']}；备注：{current['remark']}")
    elif current is None:detail.append('删除物料（含关系属性及替代料）')
    else:
     for field,title in [('position','位置'),('quantity','数量'),('unit','单位'),('remark','备注')]:
      if before[field]!=current[field]:detail.append(f'{title}：{before[field] or "（空）"} → {current[field] or "（空）"}')
    previous=old_alternates.get(before['id'],[]) if before else []
    following=after_alternates.get(current['id'],[]) if current else []
    for mid in following:
     if mid not in previous:detail.append('增加替代料：'+label(mid))
    for mid in previous:
     if mid not in following:detail.append('删除替代料：'+label(mid))
    if not detail:continue
    description='父项：'+label(key[0])+'\n物料：'+label(key[1])+'\n物料位置：'+(key[2] or '（未填写）')+'\n'+'\n'.join(detail)
    self.log('bom',key[0],'update',description,snapshot=self.material_snapshot(key[1]))
    log_id=self.db.conn.execute('SELECT last_insert_rowid()').fetchone()[0]
    for affected,contained in log_roots.items():
     if key[0] in contained or affected==root:self.db.conn.execute('INSERT OR IGNORE INTO bom_log_roots(log_id,root_id) VALUES(?,?)',(log_id,affected))
   from services.current_catalog import enabled as catalog_enabled,prune_snapshot,synchronize
   if catalog_enabled(self):prune_snapshot(self)
   if getattr(self,'scope_id',0):self.refresh_usage()
   if enabled(self.db):record_definitions(self,root)
   synchronize(self,changed_parents)
  from infrastructure.events import WindowRefreshBus
  report(90,'正在通知相关BOM窗口…')
  for parent in changed_parents: self.publish({'type':'bom_changed','root':parent})
  if affected_skus:self.publish({'type':'sku_changed','ids':affected_skus})
  report(100,'BOM保存完成')
 def validate_bom_rows(self,root,rows):
  """校验导入BOM与现有结构的子项、数量、位置一致性。"""
  old={(x['material_id'],str(x['position'])):float(x['quantity']) for x in self.bom(root)}; new={(x['material_id'],str(x.get('position',''))):float(x['quantity']) for x in rows}
  return old==new, {'missing':list(set(old)-set(new)),'extra':list(set(new)-set(old))}
 def change_sku_version(self,sid,new_version):
  r=self.db.conn.execute('SELECT material_id FROM skus WHERE id=?',(sid,)).fetchone()
  if not r: raise ValueError('SKU不存在')
  self.change_material_version(r['material_id'],new_version)
  self.log('sku',sid,'version',new_version)
 def delete_bom_item(self,bom_item_id,recursive=True):
  r=self.db.conn.execute('SELECT parent_material_id,child_material_id FROM material_bom_items WHERE id=?',(bom_item_id,)).fetchone()
  if not r:return
  with self.db.transaction():
   self.db.conn.execute('DELETE FROM material_bom_items WHERE id=?',(bom_item_id,))
   if not getattr(self,'scope_id',0):self.db.conn.execute('UPDATE materials SET has_bom=CASE WHEN EXISTS(SELECT 1 FROM material_bom_items WHERE parent_material_id=materials.id) THEN 1 ELSE 0 END')
   self.log('bom',r['parent_material_id'],'delete',f'删除下阶物料{r["child_material_id"]}',snapshot=self.material_snapshot(r['child_material_id']))
  try:
   from infrastructure.events import WindowRefreshBus
   self.publish({'type':'bom_changed','root':r['parent_material_id']})
  except Exception: pass
 BOM_HEADERS=['父项物料编码','父项版本号','物料基础编码','物料版本号','物料数量','位置']
 def import_bom(self,path,root=None,sheet_name=None,header_mapping=None,persist=True,conflict_resolver=None):
  import math
  wb=load_workbook(path,read_only=True,data_only=True)
  try:
   if sheet_name is None:
    if len(wb.sheetnames)!=1:raise ValueError('文件包含多个工作表，请选择BOM所在工作表')
    sheet_name=wb.sheetnames[0]
   if sheet_name not in wb.sheetnames:raise ValueError('所选工作表不存在')
   ws=wb[sheet_name]; headers=[str(x if x is not None else '').strip() for x in next(ws.iter_rows(values_only=True),())]
   mapping=header_mapping or {field:field for field in self.BOM_HEADERS}
   missing=[field for field in self.BOM_HEADERS if not mapping.get(field) or not str(mapping[field]).strip() or mapping[field] not in headers]
   if missing:raise MaterialHeaderError(headers,missing)
   if len(set(mapping[field] for field in self.BOM_HEADERS))!=len(self.BOM_HEADERS):raise ValueError('表头映射重复')
   if any(headers.count(mapping[field])!=1 for field in self.BOM_HEADERS):raise ValueError('源表头名称重复，请修改文件表头后重试')
   import re
   groups=sorted({int(match.group(1)) for field in mapping for match in [re.fullmatch(r'替代料(\d+)(?:编码|版本号)',field)] if match})
   alternate_fields=[f'替代料{number}{suffix}' for number in groups for suffix in ('编码','版本号')]
   for field in alternate_fields:
    if not mapping.get(field) or mapping[field] not in headers:raise ValueError('替代料映射不完整：'+field)
   selected=[mapping[field] for field in self.BOM_HEADERS+alternate_fields]
   if len(set(selected))!=len(selected):raise ValueError('表头映射重复')
   if any(headers.count(name)!=1 for name in selected):raise ValueError('源表头名称重复')
   alternate_index={field:headers.index(mapping[field]) for field in alternate_fields}
   index={field:headers.index(mapping[field]) for field in self.BOM_HEADERS}; rows=[]; pairs=set(); children={}; parents=set(); child_ids=set()
   materials={(r['base_code'],r['version']):r for r in self.materials()}
   issues=[]; records=[]; root_row=None
   current=self.material_snapshot(root) if root is not None else None
   if root is not None and not current:issues.append('当前BOM根物料不存在')
   for no,values in enumerate(ws.iter_rows(values_only=True),1):
    if no==1:continue
    data={field:str(values[i] if i<len(values) and values[i] is not None else '').strip() for field,i in index.items()}
    if not any(data.values()):continue
    is_root=[data[field].casefold()=='root' for field in ('父项物料编码','父项版本号')]
    if any(is_root):
     if not all(is_root):issues.append(f'第{no}行根节点标记不完整：父项物料编码和父项版本号必须同时为ROOT/root')
     if root_row is not None:issues.append(f'第{no}行存在重复根节点声明，BOM只能有一个ROOT根节点')
     root_row=no; key=(data['物料基础编码'],data['物料版本号'])
     if current:
      if key!=(current['base_code'],current['version']):
       issues.append(f"BOM选择错误：第{no}行ROOT物料为{key[0]}{key[1]}，当前BOM为{current['base_code']}{current['version']}，禁止导入")
     elif root is None:
      if key not in materials:issues.append(f'第{no}行ROOT物料不存在：{key[0]}{key[1]}')
      elif all(is_root):root=materials[key]['id'];current=materials[key]
     continue
    data.update({field:str(values[i] if i<len(values) and values[i] is not None else '').strip() for field,i in alternate_index.items()})
    records.append((no,data))
   edge_lines={}
   for no,data in records:
    start_errors=len(issues)
    absent=[field for field in self.BOM_HEADERS[:-1] if not data[field]]
    if absent:issues.append(f'第{no}行必填项缺失：'+ '、'.join(absent))
    parent=materials.get((data['父项物料编码'],data['父项版本号'])); child=materials.get((data['物料基础编码'],data['物料版本号']))
    if parent is None and data['父项物料编码'] and data['父项版本号']:issues.append(f'第{no}行父项不存在：'+data['父项物料编码']+data['父项版本号'])
    if child is None and data['物料基础编码'] and data['物料版本号']:issues.append(f'第{no}行子物料不存在：'+data['物料基础编码']+data['物料版本号'])
    quantity=None
    if data['物料数量']:
     try:
      quantity=float(data['物料数量'])
      if not math.isfinite(quantity) or quantity<=0:issues.append(f'第{no}行数量必须大于0且为有限数字')
     except ValueError:issues.append(f'第{no}行数量不是有效数字')
    alternate_ids=[]
    for number in groups:
     code=data[f'替代料{number}编码']; version=data[f'替代料{number}版本号']
     if not code and not version:continue
     if not code or not version:issues.append(f'第{no}行替代料{number}编码和版本号必须同时填写');continue
     alternate=materials.get((code,version))
     if alternate is None:issues.append(f'第{no}行替代料{number}不存在：{code}{version}');continue
     if child is not None and alternate['id']==child['id']:issues.append(f'第{no}行替代料不能是当前物料自身');continue
     if alternate['id'] in alternate_ids:issues.append(f'第{no}行替代料重复：{code}{version}');continue
     alternate_ids.append(alternate['id'])
    if parent is None or child is None:continue
    pid=parent['id']; mid=child['id']
    if (pid,mid) in pairs:issues.append(f'第{no}行同一父级重复包含相同物料（首次出现于第{edge_lines[pid,mid]}行）')
    if pid==mid:issues.append(f'第{no}行物料不能包含自身')
    pairs.add((pid,mid));edge_lines.setdefault((pid,mid),no);parents.add(pid);child_ids.add(mid)
    if mid not in children.setdefault(pid,[]):children[pid].append(mid)
    if len(issues)==start_errors:rows.append(dict(parent_material_id=pid,material_id=mid,quantity=quantity,unit=child['unit'],position=data['位置'],alternate_material_ids=alternate_ids))
   if not rows and root_row is None and not issues:issues.append('所选工作表没有可导入的BOM数据')
   if root is None:
    roots=parents-child_ids
    if len(roots)!=1:issues.append('BOM必须具有唯一根物料，或结构存在循环')
    else:root=next(iter(roots))
   visited=set();active=set()
   def walk(mid):
    if mid in visited:return
    active.add(mid)
    for child in children.get(mid,[]):
     if child in active:issues.append(f'第{edge_lines[mid,child]}行关系使BOM结构存在循环包含')
     else:walk(child)
    active.remove(mid);visited.add(mid)
   if root is not None:
    walk(root)
    unreachable=(parents|child_ids)-visited
    if unreachable:
     lines=sorted({no for (parent,child),no in edge_lines.items() if parent in unreachable or child in unreachable})
     issues.append('第'+ '、'.join(map(str,lines))+'行存在不属于当前根物料的孤立分支，请检查父项编码和版本号')
   for mid in sorted(parents|child_ids):walk(mid)
   if issues:raise BomImportError(issues)
   imported_count=len(rows)
   rows=self.resolve_import_bom_conflicts(root,rows,conflict_resolver)
   try:preview_tree=self.save_bom(root,rows,preview=True)
   except ValueError as error:raise BomImportError(['与数据库共享BOM结构校验失败：'+str(error)]) from error
   if not persist:return preview_tree
   self.save_bom(root,rows)
   return imported_count
  finally:wb.close()
 def resolve_import_bom_conflicts(self,root,rows,resolver=None):
  """Compare explicit imported sub-BOMs, then validate a single resolved shared graph."""
  from collections import defaultdict
  system=defaultdict(list); incoming=defaultdict(list)
  materials={r['id']:dict(r) for r in self.materials()}
  for record in self.db.conn.execute('SELECT * FROM material_bom_items ORDER BY parent_material_id,sort_no,id'):
   data=dict(record);data['material_id']=data['child_material_id']
   data['unit']=materials[data['material_id']]['unit']
   data['alternate_material_ids']=[a['material_id'] for a in self.alternates(data['id'])]
   system[data['parent_material_id']].append(data)
  incoming[root]=[]
  for row in rows:incoming[row['parent_material_id']].append(row)
  def label(mid):
   m=materials[mid];return m['base_code']+m['version']+' '+m['name']
  def signature(row):return (float(row['quantity']),row.get('position','') or '',row.get('remark','') or '',tuple(sorted(row.get('alternate_material_ids',[]))))
  def describe(items):
   return '\n'.join(f"数量：{format_quantity(signature(row)[0])}；位置：{signature(row)[1] or '（空）'}；备注：{signature(row)[2] or '（空）'}；替代料："+('、'.join(label(mid) for mid in signature(row)[3]) or '无') for row in items) or '不存在'
  conflicts=[];conflict_parents=set()
  # 检查明确提供子项且系统也已有子项的每个物料，包含导入根。
  for parent,new_rows in incoming.items():
   if not new_rows or not system.get(parent):continue
   old_by=defaultdict(list);new_by=defaultdict(list)
   for row in system[parent]:old_by[row['material_id']].append(row)
   for row in new_rows:new_by[row['material_id']].append(row)
   for mid in sorted(set(old_by)|set(new_by)):
    old=old_by[mid];new=new_by[mid]
    if sorted(map(signature,old))==sorted(map(signature,new)):continue
    fields=[]
    if not old:fields=['导入新增子项']
    elif not new:fields=['导入删除子项']
    elif len(old)==len(new)==1:
     fields=[name for i,name in enumerate(('数量','位置','备注','替代料')) if signature(old[0])[i]!=signature(new[0])[i]]
    else:fields=['子项关系不同']
    conflicts.append(dict(parent_id=parent,parent=label(parent),material=label(mid),difference='、'.join(fields),system=describe(old),incoming=describe(new)))
    conflict_parents.add(parent)
  choice='incoming'
  if conflicts:
   choice=resolver(conflicts) if resolver else 'stop'
   if choice not in ('system','incoming'):
    raise BomImportError(['用户停止导入，子BOM不一致，未修改编辑缓存或数据库']+[f"{c['parent']} / {c['material']}：{c['difference']}；系统：{c['system']}；导入：{c['incoming']}" for c in conflicts])
  effective=dict(system);effective.update(incoming)
  if choice=='system':
   # 保留系统整个子树，包括系统叶子节点；防止较深导入定义绕过选择。
   protected=set();pending=list(conflict_parents)
   while pending:
    parent=pending.pop()
    if parent==root and root not in conflict_parents:raise BomImportError(['保留系统子BOM会使当前导入根物料成为自身的下阶，存在循环引用，未导入'])
    if parent in protected:continue
    protected.add(parent);effective[parent]=system.get(parent,[])
    pending.extend(row['material_id'] for row in system.get(parent,[]))
  # 丢弃已被替换分支中的孤立导入定义；同一父物料只保存一套定义。
  resolved=[];visited=set();pending=[root]
  while pending:
   parent=pending.pop()
   if parent in visited:continue
   visited.add(parent)
   children=effective.get(parent,[]);resolved.extend(dict(row) for row in children)
   pending.extend(reversed([row['material_id'] for row in children]))
  return resolved
 def update_project(self,d):
  with self.db.transaction(): self.db.conn.execute('UPDATE project SET project_name=?,project_code=?,product_name=?,product_model=? WHERE id=1',(d.get('project_name',''),d.get('project_code',''),d.get('product_name',''),d.get('product_model',''))); self.log('project',1,'update',d.get('project_name',''))
 def project(self): return self.db.conn.execute('SELECT * FROM project WHERE id=1').fetchone()
 def material_version_references(self,i):
  """以实际共享父级BOM为选择单位，列出所有受影响SKU。"""
  parents=self.db.conn.execute('SELECT DISTINCT m.* FROM material_bom_items b JOIN materials m ON m.id=b.parent_material_id WHERE b.child_material_id=? ORDER BY m.base_code,m.version',(i,)).fetchall()
  from services.current_catalog import enabled as catalog_enabled
  if catalog_enabled(self) and not getattr(self,'scope_id',0):parents=[]  # 公共目录是镜像缓存，不是独立SKU引用。
  result=[]
  for parent in parents:
   usages=[]; sku_ids=[]
   for sku in self.bom_skus():
    if any(row['material_id']==i and row['parent_material_id']==parent['id'] for row in self.bom(sku['material_id'])):
     usages.append(f"{sku['base_code']}{sku['version']} {sku['name']}"); sku_ids.append(sku['id'])
   parent_label=f"{parent['base_code']}{parent['version']} {parent['name']}"
   result.append({'parent_id':parent['id'],'sku_ids':sku_ids or [None],'sku_labels':usages or ['未包含在SKU中'],'parent_label':parent_label,'labels':[f'{label}/{parent_label}' for label in (usages or ['未包含在SKU中'])]})
  if not getattr(self,'scope_id',0):
   for sku in self.skus():
    if not sku['scope_id']:continue
    service=self.for_sku(sku['id'])
    parent_ids={row['parent_material_id'] for row in service.bom(sku['material_id']) if row['material_id']==i}
    for parent_id in sorted(parent_ids):
     parent=self.material_snapshot(parent_id);label=parent['base_code']+parent['version']+' '+parent['name']
     sku_label=f"{sku['base_code']}{sku['version']} {sku['name']} 镜像{sku['mirror_version']}"
     result.append(dict(parent_id=f"mirror:{sku['id']}:{parent_id}",material_parent_id=parent_id,sku_ids=[sku['id']],sku_labels=[sku_label],parent_label=label,labels=[sku_label+'/'+label]))
  return result
 def material_usage(self,i):
  from services.current_catalog import enabled as catalog_enabled
  if catalog_enabled(self):
   from services.bom_maintenance import actual_material_usage
   return actual_material_usage(self,i)
  result=[{'sku_id':sid,'sku':sku,'parent':reference['parent_label'],'parent_id':reference.get('material_parent_id',reference['parent_id'])} for reference in self.material_version_references(i) for sid,sku in zip(reference['sku_ids'],reference['sku_labels'])]
  for sku in self.skus():
   if sku['material_id']==i: result.append({'sku_id':sku['id'],'sku':f"{sku['base_code']}{sku['version']} {sku['name']} 镜像{sku['mirror_version']}",'parent':'SKU根节点（无父级）'})
  return result
 def change_material_version(self,i,new_version,selected_parents=None,expected_parents=None):
  new_version=new_version.strip()
  if not new_version: raise ValueError('版本号不能为空')
  with self.db.transaction():
   old=self.db.conn.execute('SELECT * FROM materials WHERE id=?',(i,)).fetchone()
   if old is None: raise ValueError('物料不存在')
   if self.db.conn.execute('SELECT 1 FROM materials WHERE base_code=? AND version=?',(old['base_code'],new_version)).fetchone(): raise ValueError('目标版本已存在')
   parents={reference['parent_id'] for reference in self.material_version_references(i)}
   if expected_parents is not None and parents!=set(expected_parents): raise ValueError('物料的BOM引用已变化，请重新打开版本变更窗口')
   selected=parents if selected_parents is None else set(selected_parents)
   if not selected<=parents: raise ValueError('所选BOM引用已变化，请重新选择')
   from services.snapshot_bom import enabled
   is_sku=bool(self.db.conn.execute('SELECT 1 FROM skus WHERE material_id=?',(i,)).fetchone())
   if selected==parents and (is_sku or not (parents and enabled(self.db))):
    new_id=i
    self.db.conn.execute('UPDATE materials SET version=?,updated_at=CURRENT_TIMESTAMP WHERE id=?',(new_version,i))
   else:
    columns=['base_code','name','description','unit','grade','manufacturer','mpn','mp_flag','has_bom']
    values=[old[k] for k in columns]
    new_id=self.db.conn.execute('INSERT INTO materials('+','.join(columns)+',version) VALUES('+','.join('?' for _ in range(len(columns)+1))+')',values+[new_version]).lastrowid
    # 新版本继承原物料自己的下阶结构和替代料，新旧版本后续可独立维护。
    for child in self.db.conn.execute('SELECT * FROM material_bom_items WHERE parent_material_id=?',(i,)).fetchall():
     relation=self.db.conn.execute('INSERT INTO material_bom_items(parent_material_id,child_material_id,quantity,unit,position,remark,sort_no) VALUES(?,?,?,?,?,?,?)',(new_id,child['child_material_id'],child['quantity'],child['unit'],child['position'],child['remark'],child['sort_no'])).lastrowid
     self.db.conn.execute('INSERT INTO alternates(bom_item_id,material_id,sort_no) SELECT ?,material_id,sort_no FROM alternates WHERE bom_item_id=?',(relation,child['id']))
    cloned_mirrors=set()
    for parent in selected:
     if isinstance(parent,str) and parent.startswith('mirror:'):
      _,sid,parent_id=parent.split(':');service=self.for_sku(int(sid));connection=service.db.conn
      if sid not in cloned_mirrors:
       # New identity inherits the old identity's private descendants in this mirror.
       for child in connection.execute('SELECT * FROM material_bom_items WHERE parent_material_id=?',(i,)).fetchall():
        relation=connection.execute('INSERT INTO material_bom_items(parent_material_id,child_material_id,quantity,unit,position,remark,sort_no) VALUES(?,?,?,?,?,?,?)',(new_id,child['child_material_id'],child['quantity'],child['unit'],child['position'],child['remark'],child['sort_no'])).lastrowid
        connection.execute('INSERT INTO alternates(bom_item_id,material_id,sort_no) SELECT ?,material_id,sort_no FROM alternates WHERE bom_item_id=?',(relation,child['id']))
       cloned_mirrors.add(sid)
      connection.execute('UPDATE material_bom_items SET child_material_id=?,updated_at=CURRENT_TIMESTAMP WHERE parent_material_id=? AND child_material_id=?',(new_id,int(parent_id),i));service.refresh_usage()
     else:self.db.conn.execute('UPDATE material_bom_items SET child_material_id=?,updated_at=CURRENT_TIMESTAMP WHERE parent_material_id=? AND child_material_id=?',(new_id,parent,i))
   self.log('material',new_id,'version',f'{old["base_code"]}: {old["version"]}->{new_version}; 原物料={i}; 更新父级BOM={[str(parent) if isinstance(parent,str) else str(self.material_snapshot(parent).get('base_code',''))+str(self.material_snapshot(parent).get('version',''))+' '+str(self.material_snapshot(parent).get('name','')) for parent in sorted(selected,key=str)]}; '+('原位更新' if new_id==i else '保留旧版本并创建新物料'))
  from infrastructure.events import WindowRefreshBus
  self.publish({'type':'material_changed','id':new_id})
  for parent in selected:
   if isinstance(parent,str):
    _,sid,parent_id=parent.split(':');self.for_sku(int(sid)).publish({'type':'bom_changed','root':int(parent_id)})
   else:self.publish({'type':'bom_changed','root':parent})
  return new_id
 def add_alternate(self,bom_item_id,material_id):
  with self.db.transaction(): self.db.conn.execute('INSERT INTO alternates(bom_item_id,material_id,sort_no) VALUES(?,?,COALESCE((SELECT MAX(sort_no)+1 FROM alternates WHERE bom_item_id=?),0))',(bom_item_id,material_id,bom_item_id))
 def alternates(self,bom_item_id): return self.db.conn.execute('SELECT a.*,m.base_code,m.version,m.name,m.description,m.manufacturer,m.mpn FROM alternates a JOIN materials m ON m.id=a.material_id WHERE a.bom_item_id=? ORDER BY a.sort_no',(bom_item_id,)).fetchall()
 def remove_alternate(self,aid):
  with self.db.transaction(): self.db.conn.execute('DELETE FROM alternates WHERE id=?',(aid,))
 def import_materials(self,path,conflict_resolver=None,header_mapping=None,progress_callback=None,sheet_name=None):
  required=[field for field in MATERIAL_IMPORT_HEADERS if field not in MATERIAL_IMPORT_DEFAULTS]
  # 配置中首个大小写变体为标准写法，先归一化再做重复判断和入库。
  unit_spellings={}
  for unit in self.UNITS:unit_spellings.setdefault(unit.casefold(),unit)
  wb=load_workbook(path,read_only=True,data_only=True)
  try:
   if sheet_name is None:
    if len(wb.sheetnames)!=1:raise ValueError('文件包含多个工作表，请先选择物料库所在的工作表')
    sheet_name=wb.sheetnames[0]
   if sheet_name not in wb.sheetnames:raise ValueError('所选工作表不存在：'+sheet_name)
   ws=wb[sheet_name]; headers=[str(x if x is not None else '').strip() for x in next(ws.iter_rows(values_only=True),())]
   if header_mapping is None:
    missing=[x for x in required if x not in headers]
    if missing: raise MaterialHeaderError(headers,missing)
    header_mapping={x:x for x in MATERIAL_IMPORT_HEADERS if x in headers}
   missing=[x for x in required if not header_mapping.get(x) or header_mapping[x] not in headers]
   missing.extend(x for x in MATERIAL_IMPORT_DEFAULTS if header_mapping.get(x) and header_mapping[x] not in headers)
   if missing: raise MaterialHeaderError(headers,missing)
   idx={x:headers.index(header_mapping[x]) for x in required}; ok=[]; fail=[]; total=max((ws.max_row or 1)-1,0)
   optional_indices={field:headers.index(header_mapping[field]) if header_mapping.get(field) else None for field in MATERIAL_IMPORT_DEFAULTS}
   if progress_callback: progress_callback(0,total)
   for no,row in enumerate(ws.iter_rows(min_row=2,values_only=True),2):
    d={k:str(row[i] if row[i] is not None else '').strip() for k,i in idx.items()};
    for field,default in MATERIAL_IMPORT_DEFAULTS.items():
     index=optional_indices[field]
     d[field]=(str(row[index]).strip() if index is not None and index<len(row) and row[index] is not None else '') or default
    d['M/P']=d['M/P'].upper()
    if progress_callback: progress_callback(no-1,total)
    if not d['物料基础编码'] and not d['物料名称']:continue
    display_code=(d['物料基础编码'] or '（空基础编码）')+(d['物料版本号'] or '（空版本号）')
    missing_fields=[field for field in ('物料基础编码','物料版本号','物料名称','单位','物料制造商') if not d[field]]
    if missing_fields:
     fail.append((no,display_code,d.get('物料名称',''),'必填项缺失：'+ '、'.join(missing_fields))); continue
    canonical_unit=unit_spellings.get(d['单位'].casefold())
    if canonical_unit is None: fail.append((no,display_code,d['物料名称'],'单位不合规：'+d['单位']+'；允许：'+'、'.join(self.UNITS))); continue
    d['单位']=canonical_unit
    if d['M/P'] not in self.MP:fail.append((no,display_code,d['物料名称'],'M/P不合规：'+d['M/P']+'；仅允许M或P')); continue
    old=self.db.conn.execute('SELECT * FROM materials WHERE base_code=? AND version=?',(d['物料基础编码'],d['物料版本号'])).fetchone()
    if old:
     for field,key in (('物料等级','grade'),('M/P','mp_flag')):
      if optional_indices[field] is None:d[field]=old[key]
     same=all(str(old[k] or '')==str(d[key] or '') for k,key in [('name','物料名称'),('description','物料规格描述'),('unit','单位'),('grade','物料等级'),('mp_flag','M/P'),('manufacturer','物料制造商'),('mpn','物料MPN')])
     if same:
      fail.append((no,display_code,d['物料名称'],'重复物料信息相同，自动跳过')); continue
     choice=conflict_resolver(d,old) if conflict_resolver else 'skip'
     if choice=='cancel': return ok,fail+[(no,display_code,d['物料名称'],'用户取消')]
     if choice=='overwrite':
      self.update_material(old['id'],{'name':d['物料名称'],'description':d['物料规格描述'],'unit':d['单位'],'grade':d['物料等级'],'mp_flag':d['M/P'],'manufacturer':d['物料制造商'],'mpn':d['物料MPN']}); ok.append(d['物料基础编码']); continue
     fail.append((no,display_code,d['物料名称'],'跳过重复物料')); continue
    try: self.create_material({'base_code':d['物料基础编码'],'version':d['物料版本号'],'name':d['物料名称'],'description':d['物料规格描述'],'unit':d['单位'],'grade':d['物料等级'],'mp_flag':d['M/P'],'manufacturer':d['物料制造商'],'mpn':d['物料MPN']}); ok.append(d['物料基础编码'])
    except Exception as e: fail.append((no,display_code,d['物料名称'],str(e)))
   if progress_callback: progress_callback(total,total)
   return ok,fail
  finally:wb.close()
 def export_database(self,path):
  wb=Workbook(); ws=wb.active; ws.title='SKU清单'; ws.append(['序号','SKU编码','SKU名称','SKU描述','版本号','修改日期时间'])
  for n,r in enumerate(self.skus(),1): ws.append([n,r['base_code']+r['version'],r['name'],r['description'],r['mirror_version'],r['updated_at']])
  self.format_export_sheet(ws,{'序号'})
  self.write_material_sheet(wb.create_sheet('物料库'))
  self.write_change_sheet(wb.create_sheet('变更记录'),list(self.db.conn.execute("SELECT * FROM change_logs WHERE object_type='material' ORDER BY id")))
  for r in self.skus():
   self.for_sku(r['id']).write_bom_sheet(wb.create_sheet((r['base_code']+r['version']+(' 镜像'+str(r['mirror_version']) if sum(x['material_id']==r['material_id'] for x in self.skus())>1 else ''))[:31]),r['material_id'])
  wb.save(path)
 def export_materials(self,path):
  wb=Workbook(); ws=wb.active; ws.title='变更记录'
  self.write_change_sheet(ws,list(self.db.conn.execute("SELECT * FROM change_logs WHERE object_type='material' ORDER BY id")))
  self.write_material_sheet(wb.create_sheet('物料库'))
  wb.save(path)
 def write_material_sheet(self,sheet):
  from services.snapshot_bom import enabled
  extra=['是否被最新镜像使用'] if enabled(self.db) else []
  sheet.append(['序号','物料基础编码','物料版本号','物料名称','物料规格描述','单位','物料等级','物料制造商','物料MPN','M/P','是否存在BOM结构','是否为呆滞料']+extra)
  for number,row in enumerate(self.materials(),1):
   sheet.append([number]+[row[key] for key in ('base_code','version','name','description','unit','grade','manufacturer','mpn','mp_flag','has_bom','idle')]+([row['latest_used']] if extra else []))
  self.format_export_sheet(sheet,{'序号','单位','物料等级','M/P','是否存在BOM结构','是否为呆滞料','是否被最新镜像使用'})
 def write_change_sheet(self,sheet,records):
  import re
  sheet.append(['序号','变更日期','变更时间','变更物料编码','变更物料名称','变更类型','变更详细内容'])
  legacy_codes={}
  for historical in self.db.conn.execute("SELECT * FROM change_logs WHERE object_type='material' ORDER BY id"):
   text=historical['detail'] or ''
   base=re.search(r'基础编码：([^；\n]+)',text);version=re.search(r'版本号：([^；\n]+)',text)
   if base and version:
    base_value=base.group(1).split(' → ')[-1];version_value=version.group(1).split(' → ')[-1]
    legacy_codes[base_value+'-'+version_value]=base_value+version_value
   if historical['action']=='version':
    match=re.match(r'(.+): ([^;\n]+)->([^;\n]+);',text)
    if match:
     base_value,old,new=match.groups()
     for value in (old,new):legacy_codes[base_value+'-'+value]=base_value+value
  records=[row for row in records if not (row['object_type']=='bom' and (row['detail'] or '').startswith('父项 '))]
  visible=[]
  for row in records:
   row=dict(row)
   if row['object_type']=='bom':
    lines=(row['detail'] or '').split('\n')
    filtered=[line for line in lines if not re.fullmatch(r'排序：[0-9（空）]+ → [0-9（空）]+',line)]
    if filtered!=lines:
     if not any(line and not line.startswith(('父项：','物料：','物料位置：')) for line in filtered):continue
     row['detail']='\n'.join(filtered)
   visible.append(row)
  for number,row in enumerate(visible,1):
   code=row['material_code'];name=row['material_name'];detail=row['detail'] or ''
   if row['object_type']=='bom':
    changed_material=re.search(r'^物料：([^\n]+)',detail,re.MULTILINE)
    if changed_material:
     identity=changed_material.group(1).partition(' ')
     code=identity[0];name=identity[2]

   # 历史快照明确记录的编码/版本可规范化，绝不按连字符盲目切分基础编码。
   code=legacy_codes.get(code,code)
   if code is None:
    if row['object_type'] in ('bom','sku'):
     # 旧BOM日志没有快照，仅显示当前可定位对象；已删除对象明确标识。
     material_id=row['object_id']
     if row['object_type']=='sku':
      sku=self.db.conn.execute('SELECT material_id FROM skus WHERE id=?',(material_id,)).fetchone();material_id=sku[0] if sku else None
     current=self.material_snapshot(material_id) if material_id is not None else {}
     code=current['base_code']+current['version'] if current else '已删除对象'
     name=current.get('name','已删除对象')
    else:code='历史记录未保存';name=name or '历史记录未保存'
   detail=re.sub(r'[；;][ \t]*','\n',detail)
   object_name={'material':'物料','bom':'BOM','sku':'SKU','project':'项目'}.get(row['object_type'],row['object_type'] or '')
   action_name={'create':'新建','update':'修改','delete':'删除','version':'版本变更'}.get(row['action'],row['action'] or '')
   change_type=action_name+object_name if row['action']!='version' else object_name+action_name
   if row['action']=='copy':change_type='复制BOM'
   elif row['action']=='mirror_version':change_type='镜像版本变更'
   prefixes={'复制BOM','新建物料','修改物料','删除物料','创建SKU','新建SKU','修改BOM','删除SKU','创建物料'}
   first,separator,remainder=detail.partition('\n')
   if first.strip() in prefixes:
    change_type=first.strip();detail=remainder if separator else ''
   if row['object_type']=='bom':
    types=[]
    if any(line.startswith('新增物料') for line in detail.splitlines()):types.append('新增物料')
    if any(line.startswith('删除物料') for line in detail.splitlines()):types.append('删除物料')
    if any(line.startswith(('数量：','单位：','备注：','位置：')) and ' → ' in line for line in detail.splitlines()):types.append('修改物料')
    if '增加替代料：' in detail:types.append('新增替代料')
    if '删除替代料：' in detail:types.append('删除替代料')
    change_type='、'.join(types) or ('删除物料' if row['action']=='delete' else '修改物料')
   if row['object_type']=='material' and row['action']=='create':detail=''
   changed=row['changed_at'] or ''
   sheet.append([number,changed[:10],changed[11:],code,name,change_type,detail])
  self.format_export_sheet(sheet,{'序号','变更日期','变更时间'})
  sheet.auto_filter.ref=sheet.dimensions
 def write_bom_sheet(self,sheet,root):
  headers=['序号','层级','父项编码','父项版本','物料基础编码','物料版本号','物料名称','物料规格描述','物料位置','用量','单位','物料等级','M/P','制造商','MPN','备注']
  rows=self.bom(root)
  alternates={row['id']:list(self.alternates(row['id'])) for row in rows}
  count=max((len(values) for values in alternates.values()),default=0)
  alternate_headers=[f'替代料{number}{field}' for number in range(1,count+1) for field in ('物料编码','版本号','名称','供应商','MPN')]
  headers=headers[:-1]+alternate_headers+['备注']
  sheet.append(headers)
  materials={row['id']:row for row in self.materials()}
  # 根SKU也属于导出的BOM：ROOT标记便于导入时核对根物料身份。
  material=materials[root]
  sheet.append(['1',0,'ROOT','ROOT',material['base_code'],material['version'],material['name'],material['description'],'',1.0,material['unit'],material['grade'],material['mp_flag'],material['manufacturer'],material['mpn']]+['']*(count*5)+[''])
  sheet.cell(2,1).number_format='@'; sheet.cell(2,10).number_format='0.############'
  counters={}; numbers={}
  for row in rows:
   parent_path=row['parent_node_id']
   counters[parent_path]=counters.get(parent_path,0)+1
   number=(numbers[parent_path] if parent_path else '1')+'.'+str(counters[parent_path])
   numbers[row['node_id']]=number
   parent=materials[row['parent_material_id']]
   alternate_values=[value for alternate in alternates[row['id']] for value in (alternate['base_code'],alternate['version'],alternate['name'],alternate['manufacturer'],alternate['mpn'])]
   alternate_values+=['']*(count*5-len(alternate_values))
   sheet.append([number,row['level_no'],parent['base_code'],parent['version'],row['base_code'],row['version'],row['name'],row['description'],row['position'],row['quantity'],row['unit'],row['grade'],row['mp_flag'],row['manufacturer'],row['mpn']]+alternate_values+[row['remark']])
   sheet.cell(sheet.max_row,1).number_format='@'
   sheet.cell(sheet.max_row,10).number_format='0.############'
  centered={'层级','父项版本','物料版本号','用量','单位','物料等级','M/P'}
  centered.update(f'替代料{number}版本号' for number in range(1,count+1))
  self.format_export_sheet(sheet,centered)
 def export_purchase_bom(self,root,path):
  """汇总全部下阶物料：路径用量连乘，同版本物料合并，根物料不采购。"""
  from decimal import Decimal, ROUND_HALF_UP
  import re
  material=self.material_snapshot(root)
  if not material:raise ValueError('根物料不存在')
  totals={}; path_quantities={None:Decimal('1')}; path_labels={None:material['base_code']+material['version']}
  alternate_cache={}; snapshots={}
  for row in self.bom(root):
   quantity=path_quantities[row['parent_node_id']]*Decimal(str(row['quantity']))
   path_quantities[row['node_id']]=quantity
   label=path_labels[row['parent_node_id']]+'/'+row['base_code']+row['version']
   if row['position']:label+='（位置：'+row['position']+'）'
   path_labels[row['node_id']]=label
   entry=totals.setdefault(row['material_id'],dict(material=row,quantity=Decimal('0'),alternates={},remarks=[],occurrences=[]))
   entry['quantity']+=quantity
   if row['remark'] and row['remark'] not in entry['remarks']:entry['remarks'].append(row['remark'])
   if row['id'] not in alternate_cache:alternate_cache[row['id']]=list(self.alternates(row['id']))
   alternatives=alternate_cache[row['id']]
   for alternate in alternatives:
    mid=alternate['material_id']
    if mid not in snapshots:snapshots[mid]=self.material_snapshot(mid)
    entry['alternates'].setdefault(mid,snapshots[mid])
   entry['occurrences'].append((label,quantity,tuple(a['material_id'] for a in alternatives)))
  count=max((len(e['alternates']) for e in totals.values()),default=0)
  headers=['序号','物料编码','物料名称','规格','用量','单位','物料等级','M/P','制造商','MPN']
  headers += [f'替代料{n}{field}' for n in range(1,count+1) for field in ('物料编码','物料名称','规格','制造商','MPN')]+['备注']
  wb=Workbook(); sheet=wb.active
  sheet.title=re.sub(r'[\\/*?:\[\]]','_',material['name']).strip().strip("'")[:31] or '采购BOM'
  sheet.append(headers)
  try:
   for number,entry in enumerate(totals.values(),1):
    row=entry['material']; notes=list(entry['remarks'])
    # 合并物料不代表不同使用位置的替代关系可以互换，明确标注适用范围。
    if len({frozenset(ids) for _,_,ids in entry['occurrences']})>1:
     notes.append('替代料仅适用于对应BOM位置：')
     for label,quantity,ids in entry['occurrences']:
      codes='、'.join(snapshots[mid]['base_code']+snapshots[mid]['version'] for mid in ids) or '无替代料'
      notes.append(f'{label}（用量{format_quantity(quantity)}）：{codes}')
    values=[v for a in entry['alternates'].values() for v in (a['base_code']+a['version'],a['name'],a['description'],a['manufacturer'],a['mpn'])]
    values+=['']*(count*5-len(values))
    quantity=float(entry['quantity'])
    sheet.append([number,row['base_code']+row['version'],row['name'],row['description'],quantity,row['unit'],row['grade'],row['mp_flag'],row['manufacturer'],row['mpn']]+values+['\n'.join(notes)])
    sheet.cell(sheet.max_row,5).number_format='0.############'
   self.format_export_sheet(sheet,{'序号','用量','单位','物料等级','M/P'})
   wb.save(path)
  finally:wb.close()
 def format_export_sheet(self,sheet,centered):
  from openpyxl.styles import Alignment, Border, Side, Font
  from openpyxl.utils import get_column_letter
  import math, unicodedata
  headers=[cell.value for cell in sheet[1]]
  edge=Side(style='thin',color='FF000000'); border=Border(left=edge,right=edge,top=edge,bottom=edge)
  def width(text):return sum(2 if unicodedata.east_asian_width(char) in ('W','F') else 1 for char in text)
  def display(cell):
   if cell.value is None:return ''
   if headers[cell.column-1]=='用量' and cell.row>1:return format_quantity(cell.value)
   return str(cell.value)
  widths={}
  for column,header in enumerate(headers,1):
   longest=max(width(line) for cells in sheet.iter_rows(min_col=column,max_col=column) for line in (display(cells[0]).split('\n')))
   widths[column]=min(60,max(5,longest+2))
   sheet.column_dimensions[get_column_letter(column)].width=widths[column]
  for cells in sheet:
   lines=1
   for cell in cells:
    cell.border=border
    cell.alignment=Alignment(horizontal='center' if headers[cell.column-1] in centered else 'left',vertical='center',wrap_text=True)
    if cell.row==1:cell.font=Font(bold=True)
    lines=max(lines,sum(max(1,math.ceil(width(line)/max(1,widths[cell.column]-2))) for line in display(cell).split('\n')))
   sheet.row_dimensions[cells[0].row].height=min(409,lines*16+6)
  sheet.freeze_panes='A2'
 def export_sku(self,sid,path):
  if not getattr(self,'scope_id',0):
   scoped=self.for_sku(sid)
   if scoped is not self:return scoped.export_sku(sid,path)
  sku=self.db.conn.execute('SELECT material_id FROM skus WHERE id=?',(sid,)).fetchone()
  if sku is None:raise ValueError('SKU不存在')
  wb=Workbook(); w=wb.active; w.title='变更记录'
  material_ids={sku['material_id']}|{row['material_id'] for row in self.bom(sku['material_id'])}
  linked={row[0] for row in self.db.conn.execute('SELECT log_id FROM bom_log_roots WHERE root_id=?',(sku['material_id'],))}
  records=[row for row in self.db.conn.execute('SELECT * FROM change_logs ORDER BY id') if row['id'] in linked or (row['object_type'] in ('material','bom') and row['object_id'] in material_ids) or (row['object_type']=='sku' and row['object_id']==sid)]
  reset=self.db.conn.execute('SELECT start_log_id FROM sku_history_reset WHERE root_id=?',(sku['material_id'],)).fetchone()
  if reset:records=[row for row in records if row['id']>=reset[0]]
  self.write_change_sheet(w,records)
  import re
  name=self.material_snapshot(sku['material_id'])['name']
  # Excel工作表名称限制：非法字符替换、31字符上限，重名由openpyxl消歧。
  title=re.sub(r'[\\/*?:\[\]]','_',name).strip().strip("'")[:31] or 'BOM'
  self.write_bom_sheet(wb.create_sheet(title),sku['material_id'])
  wb.save(path)
