from ui.sku_change_description import ChangeDescriptionDialog, VersionDescriptionHover, DESCRIPTION_ROLE
from PyQt5.QtCore import QEventLoop
from domain.quantity import format_quantity
from ui.time_display import sku_timestamp
"""PyQt主窗口、SKU清单、物料库和BOM编辑窗口。输入服务对象，输出桌面交互。编辑时间：2026-09-24。"""
from PyQt5.QtWidgets import *
from ui.interactive_windows import QColorDialog,QDialog,QFileDialog,QInputDialog,QMainWindow,QMessageBox,QProgressDialog
from PyQt5.QtCore import Qt, pyqtSignal, QSettings, QTimer, QEvent, QSize, QSortFilterProxyModel, QItemSelectionModel
from PyQt5.QtGui import QColor, QBrush, QFont, QStandardItemModel, QStandardItem
from pathlib import Path
import os
from config.config_service import ConfigService, DEFAULT_MATERIAL_UNITS
from database.database import Database
from infrastructure.events import WindowRefreshBus
from services.services import MaterialHeaderError, BomImportError, MATERIAL_IMPORT_HEADERS, MATERIAL_IMPORT_DEFAULTS
from ui.theme import apply_theme, theme_icon
from ui.material_clipboard import copy_material, clipboard_material, clipboard_relation, copy_materials, clipboard_entries
from ui.material_search import MaterialSearchDialog, BomSearchDialog
from ui.table_layout import fit_columns
from ui.about_dialog import AboutDialog
from ui.frozen_table import FrozenMaterialTable
from ui.import_paths import bom_import_directory, remember_bom_import
from ui.replacement_selection import ReplacementSelectionDialog
from ui.bom_view_state import preserve_bom_view
from ui.material_enter import FirstMaterialOnEnter

class WrappingTextEdit(QPlainTextEdit):
 def __init__(self):
  super().__init__();self.setLineWrapMode(QPlainTextEdit.WidgetWidth);self.setMinimumWidth(0);self.setFixedHeight(self.fontMetrics().height()*3+12)
 def text(self):return self.toPlainText()
 def setText(self,value):self.setPlainText(value)

class QuantityEdit(QDoubleSpinBox):
 def __init__(self):
  super().__init__(); self.setDecimals(12); self.setRange(0.000000000001,1e12); self.setSingleStep(0.1); self.setValue(1.0)
 def setText(self,value):
  self.setValue(float(value or 1.0))
 def text(self):
  return format_quantity(self.value())
 def textFromValue(self,value):return format_quantity(value)

class MaterialTextEdit(WrappingTextEdit):
 def __init__(self,value='',lines=2):
  super().__init__()
  from PyQt5.QtGui import QTextOption
  self.setPlainText(str(value or ''));self.setFixedHeight(self.fontMetrics().lineSpacing()*lines+16)
  self.setTabChangesFocus(True);self.setHorizontalScrollBarPolicy(Qt.ScrollBarAlwaysOff)
  self.setVerticalScrollBarPolicy(Qt.ScrollBarAsNeeded)
  option=self.document().defaultTextOption();option.setAlignment(Qt.AlignLeft);option.setWrapMode(QTextOption.WrapAtWordBoundaryOrAnywhere);self.document().setDefaultTextOption(option)

class MaterialDialog(QDialog):
 def __init__(self,parent=None,data=None,readonly_keys=None):
  super().__init__(parent);self.setWindowTitle('物料');self.resize(640,620)
  layout=QVBoxLayout(self);scroll=QScrollArea();scroll.setWidgetResizable(True);scroll.setFrameShape(QFrame.NoFrame);layout.addWidget(scroll)
  body=QWidget();f=QFormLayout(body);f.setFieldGrowthPolicy(QFormLayout.AllNonFixedFieldsGrow);scroll.setWidget(body)
  data=data or {};self.e={};readonly_keys=readonly_keys or set()
  def field(key,lines=2):
   if key in readonly_keys:
    w=QLabel(str(data.get(key,'') or ''));w.setWordWrap(True);w.setAlignment(Qt.AlignLeft|Qt.AlignTop);w.setTextInteractionFlags(Qt.TextSelectableByMouse|Qt.TextSelectableByKeyboard)
   else:w=MaterialTextEdit(data.get(key,''),lines)
   self.e[key]=w;return w
  for k,l in [('base_code','物料基础编码'),('version','物料版本号'),('name','物料名称'),('description','物料规格描述'),('manufacturer','物料制造商'),('mpn','物料MPN')]:
   f.addRow(l,field(k,4 if k=='description' else 3 if k=='name' else 2))
  self.unit=QComboBox();self.unit.addItems(parent.svc.UNITS if parent is not None and hasattr(parent,'svc') else list(DEFAULT_MATERIAL_UNITS));self.unit.setCurrentIndex(self.unit.findText(data.get('unit',self.unit.itemText(0))));self.unit.setEnabled('unit' not in readonly_keys);f.addRow('单位',self.unit)
  f.addRow('物料等级',field('grade'))
  self.mp_flag=QComboBox();self.mp_flag.addItems(['M','P']);self.mp_flag.setCurrentText(data.get('mp_flag','M'));self.mp_flag.setEnabled('mp_flag' not in readonly_keys);f.addRow('M/P',self.mp_flag)
  b=QDialogButtonBox(QDialogButtonBox.Ok|QDialogButtonBox.Cancel);b.accepted.connect(self.accept);b.rejected.connect(self.reject);layout.addWidget(b)
 def data(self):
  d={k:w.text().strip() for k,w in self.e.items()};d['unit']=self.unit.currentText();d['mp_flag']=self.mp_flag.currentText();return d
class CreateMaterialDialog(MaterialDialog):
 def __init__(self,owner,data=None):
  super().__init__(owner,data=data);self.setWindowTitle('新建物料');self.setWindowModality(Qt.NonModal)
  self.service=owner.svc;self.database=owner.svc.db;self._saving=False;self._saved=False
 def accept(self):
  if self._saving or self._saved:return
  if self.service.db is not self.database:
   QMessageBox.warning(self,'数据库已切换','此窗口属于之前的数据库，请关闭后重新新建物料。');return
  self._saving=True
  try:
   self.service.create_material(self.data())
  except Exception as error:
   show_operation_error(self,'校验失败',error);return
  finally:self._saving=False
  self._saved=True;super().accept()

def show_create_material(owner,data=None):
 dialog=getattr(owner,'_create_material_dialog',None)
 if data is not None and dialog is not None:
  if QMessageBox.question(owner,'替换新建物料草稿','已打开新建物料窗口，是否放弃该窗口的输入，改为从选中物料类似新增？',QMessageBox.Yes|QMessageBox.No,QMessageBox.No)!=QMessageBox.Yes:
   dialog.show();dialog.raise_();dialog.activateWindow();return dialog
  dialog.reject();dialog=None
 if dialog is None:
  initial=dict(data) if data is not None else None
  if initial is not None:initial.update(base_code='',version='')
  dialog=CreateMaterialDialog(owner,data=initial);owner._create_material_dialog=dialog
  def finished(result):
   if getattr(owner,'_create_material_dialog',None) is dialog:owner._create_material_dialog=None
   dialog.deleteLater()
  dialog.finished.connect(finished)
 dialog.show();dialog.raise_();dialog.activateWindow()
 return dialog

def choose_material_sheet(path,parent=None,kind='物料库'):
 from openpyxl import load_workbook
 wb=load_workbook(path,read_only=True,data_only=True)
 try:names=wb.sheetnames
 finally:wb.close()
 if len(names)==1:return names[0]
 if not names:raise ValueError('Excel文件中没有可读取的工作表')
 sheet,accepted=QInputDialog.getItem(parent,'选择'+kind+'工作表','请选择'+kind+'所在的工作表（Sheet）：',names,0,False)
 if not accepted:raise ValueError('用户取消选择工作表，未导入任何物料')
 return sheet

class MaterialHeaderMappingDialog(QDialog):
 def __init__(self,path,headers,missing,parent=None,sheet_name=None,required=None,config_service=None,optional_fields=None,preview_rows=None,mapping_kind=None,alternate_suffixes=None):
  super().__init__(parent); self.setWindowTitle(('BOM表头映射' if required else '物料表头映射')+(' — '+sheet_name if sheet_name else '')); self.resize(1000,620); self.required=list(required) if required else list(MATERIAL_IMPORT_HEADERS); self.combos={}; self.optional=set(optional_fields or ()) if required else set(MATERIAL_IMPORT_DEFAULTS); self.alternate_suffixes=alternate_suffixes or ('编码','版本号')
  self.mapping_config=config_service or getattr(getattr(parent,'svc',None),'config_service',None)
  self.mapping_kind=mapping_kind or ('bom' if required else 'materials')
  remembered=self.mapping_config.import_mapping(self.mapping_kind) if self.mapping_config else {}
  self.source_headers=headers; self.alternate_groups=[]
  layout=QVBoxLayout(self); layout.addWidget(QLabel('请确认需求字段与文件列的映射关系（* 为必选）：'))
  if required:
   button=QPushButton('新增替代料'); button.clicked.connect(self.add_alternate_mapping); layout.addWidget(button)
  self.mapping_table=QTableWidget(2,len(self.required)); self.mapping_table.setVerticalHeaderLabels(['需求表头','目标表头']); self.mapping_table.horizontalHeader().hide(); self.mapping_table.setEditTriggers(QAbstractItemView.NoEditTriggers)
  for column,need in enumerate(self.required):
   item=QTableWidgetItem(need if need in self.optional else need+' *'); item.setToolTip(('可选；新建未映射默认：'+(MATERIAL_IMPORT_DEFAULTS.get(need) or '空值')+'；覆盖时未映射保留原值') if need in self.optional else '必需映射'); item.setTextAlignment(Qt.AlignCenter); self.mapping_table.setItem(0,column,item)
   combo=QComboBox(); combo.setSizeAdjustPolicy(QComboBox.AdjustToContents); combo.addItem('（未映射）',''); combo.addItems([header for header in headers if header.strip()])
   if need not in missing and need in headers: combo.setCurrentText(need)
   self.combos[need]=combo; self.mapping_table.setCellWidget(1,column,combo)
   self.configure_mapping_color(combo)
   combo.currentTextChanged.connect(self.schedule_mapping_fit)
   combo.installEventFilter(self)
  self.mapping_table.setHorizontalScrollBarPolicy(Qt.ScrollBarAlwaysOn); self.mapping_table.setHorizontalScrollMode(QAbstractItemView.ScrollPerPixel)
  layout.addWidget(self.mapping_table)
  layout.addWidget(QLabel('文件读取预览（前20行）：'))
  self.preview=QTableWidget(); self.preview.setColumnCount(len(headers)); self.preview.setHorizontalHeaderLabels(headers); self.preview.setWordWrap(True); self.preview.setVerticalScrollBarPolicy(Qt.ScrollBarAlwaysOn); self.preview.setHorizontalScrollBarPolicy(Qt.ScrollBarAlwaysOn)
  try:
   if preview_rows is not None:rows=preview_rows[:20]
   else:
    from openpyxl import load_workbook
    wb=load_workbook(path,read_only=True,data_only=True)
    try:
     ws=wb[sheet_name] if sheet_name is not None else wb.active; rows=list(ws.iter_rows(min_row=2,max_row=21,values_only=True))
    finally:wb.close()
   self.preview.setRowCount(len(rows))
   for r,row in enumerate(rows):
    for c in range(len(headers)): self.preview.setItem(r,c,QTableWidgetItem(str(row[c] or '') if c<len(row) else ''))
  except Exception as error: self.preview.setRowCount(1); self.preview.setItem(0,0,QTableWidgetItem('预览失败：'+str(error)))
  for column,header in enumerate(headers):
   if not header.strip():self.preview.setColumnHidden(column,True)
  layout.addWidget(self.preview,1); fit_columns(self.preview)
  self._mapping_fit_timer=QTimer(self); self._mapping_fit_timer.setSingleShot(True); self._mapping_fit_timer.timeout.connect(self.run_mapping_fit)
  self.mapping_table.viewport().installEventFilter(self); self.mapping_table.installEventFilter(self)
  self.fit_mapping_columns()
  self.preview.setEditTriggers(QAbstractItemView.NoEditTriggers)
  if required and remembered:
   import re
   numbers=sorted({int(match.group(1)) for key in remembered for match in [re.fullmatch(r'替代料([1-9][0-9]*)(?:编码|版本号)',key)] if match})
   for number in numbers:self.add_alternate_mapping(number=number)
  for field,combo in self.combos.items():
   if field in remembered:
    target=remembered[field]
    # 按列名恢复，缺失/歧义列保持未映射，不能退回同名或列序号猜测。
    combo.setCurrentIndex(combo.findText(target) if target and headers.count(target)==1 else 0)
  self.schedule_mapping_fit()
  buttons=QDialogButtonBox(QDialogButtonBox.Ok|QDialogButtonBox.Cancel); buttons.accepted.connect(self.validate_mapping); buttons.rejected.connect(self.reject); layout.addWidget(buttons)
 def configure_mapping_color(self,combo):
  def update_color(*args):
   combo.setStyleSheet('QComboBox { background-color: #d9f2df; color: #202020; }' if combo.currentIndex()>0 else '')
  combo.currentIndexChanged.connect(update_color);update_color()
 def add_alternate_mapping(self,number=None):
  if number is None or isinstance(number,bool):number=max([int(group[0].split('替代料')[1].split('编码')[0]) for group in self.alternate_groups] or [0])+1
  group=[f'替代料{number}{suffix}' for suffix in self.alternate_suffixes]; self.alternate_groups.append(group)
  for field in group:
   column=self.mapping_table.columnCount(); self.mapping_table.insertColumn(column); self.required.append(field)
   if field.endswith(('名称','规格','制造商','MPN')):self.optional.add(field)
   item=QTableWidgetItem(field if field in self.optional else field+' *'); item.setTextAlignment(Qt.AlignCenter); self.mapping_table.setItem(0,column,item)
   combo=QComboBox(); combo.addItem('（未映射）',''); combo.addItems([h for h in self.source_headers if h.strip()]); self.combos[field]=combo
   if field in self.source_headers:combo.setCurrentText(field)
   combo.setContextMenuPolicy(Qt.CustomContextMenu); combo.customContextMenuRequested.connect(lambda pos,f=field,w=combo:self.alternate_mapping_menu(f,w.mapToGlobal(pos)))
   self.configure_mapping_color(combo)
   combo.currentTextChanged.connect(self.schedule_mapping_fit); self.mapping_table.setCellWidget(1,column,combo)
  if not getattr(self,'_mapping_menu_connected',False):
   self.mapping_table.setContextMenuPolicy(Qt.CustomContextMenu); self.mapping_table.customContextMenuRequested.connect(self.mapping_context_menu); self._mapping_menu_connected=True
  self.schedule_mapping_fit()
 def mapping_context_menu(self,pos):
  column=self.mapping_table.columnAt(pos.x())
  if column>=0:self.alternate_mapping_menu(self.required[column],self.mapping_table.viewport().mapToGlobal(pos))
 def alternate_mapping_menu(self,field,pos):
  if not any(field in group for group in self.alternate_groups):return
  menu=QMenu(self); action=menu.addAction('删除替代料')
  if menu.exec_(pos)==action:self.remove_alternate_mapping(field)
 def remove_alternate_mapping(self,field):
  group=next(group for group in self.alternate_groups if field in group)
  for name in reversed(group):
   column=self.required.index(name); self.mapping_table.removeColumn(column); self.required.remove(name); self.combos.pop(name)
  self.alternate_groups.remove(group); self.schedule_mapping_fit()
 def run_mapping_fit(self):
  from PyQt5 import sip
  if sip.isdeleted(self) or not self.isVisible():return
  self.fit_mapping_columns()
 def hideEvent(self,event):
  if hasattr(self,'_mapping_fit_timer'):self._mapping_fit_timer.stop()
  super().hideEvent(event)
 def schedule_mapping_fit(self,*args):
  if hasattr(self,'_mapping_fit_timer'):self._mapping_fit_timer.start(0)
 def eventFilter(self,obj,event):
  if event.type() in (QEvent.Resize,QEvent.Show,QEvent.FontChange,QEvent.StyleChange,QEvent.ApplicationFontChange):self.schedule_mapping_fit()
  return super().eventFilter(obj,event)
 def showEvent(self,event):
  super().showEvent(event); self.schedule_mapping_fit()
 def fit_mapping_columns(self):
  from PyQt5 import sip
  if sip.isdeleted(self.mapping_table) or any(sip.isdeleted(combo) for combo in self.combos.values()):return
  try:self._fit_mapping_columns()
  except RuntimeError:
   if sip.isdeleted(self.mapping_table) or any(sip.isdeleted(combo) for combo in self.combos.values()):return
   raise
 def _fit_mapping_columns(self):
  table=self.mapping_table; table.ensurePolished(); widths=[]; heights=[]
  for need in self.required:
   combo=self.combos[need]; combo.ensurePolished(); metrics=combo.fontMetrics()
   text_width=metrics.horizontalAdvance(combo.currentText())
   combo.setToolTip(combo.currentText())
   combo.view().setMinimumWidth(max(metrics.horizontalAdvance(combo.itemText(i)) for i in range(combo.count()))+48)
   option=QStyleOptionComboBox(); combo.initStyleOption(option)
   size=combo.style().sizeFromContents(QStyle.CT_ComboBox,option,QSize(text_width+16,metrics.height()+12),combo)
   width=max(size.width(),combo.sizeHint().width(),table.fontMetrics().horizontalAdvance(need)+24)
   height=max(size.height(),combo.sizeHint().height(),metrics.height()+16)
   # 样式可能额外占用内边距；按实际文字绘制区域补足，避免主题或缩放后被截断。
   option.rect.setRect(0,0,width,height)
   field=combo.style().subControlRect(QStyle.CC_ComboBox,option,QStyle.SC_ComboBoxEditField,combo)
   width+=max(0,text_width+12-field.width()); height+=max(0,metrics.height()+8-field.height())
   combo.setMinimumSize(width,height)
   widths.append(width+4); heights.append(height+4)
  table.verticalHeader().setFixedWidth(max(table.verticalHeader().fontMetrics().horizontalAdvance(x) for x in ('需求表头','目标表头'))+24)
  extra=max(0,table.viewport().width()-sum(widths))
  table.horizontalHeader().setSectionResizeMode(QHeaderView.Fixed)
  for column,width in enumerate(widths):table.setColumnWidth(column,width+(extra*sum(widths[:column+1])//sum(widths)-extra*sum(widths[:column])//sum(widths)))
  # 不使用ResizeToContents覆盖组合框行高；需求行、目标行都按当前字体和主题计算。
  table.verticalHeader().setSectionResizeMode(QHeaderView.Fixed)
  table.setRowHeight(0,table.fontMetrics().height()+20); table.setRowHeight(1,max(heights))
  table.setFixedHeight(table.rowHeight(0)+table.rowHeight(1)+table.horizontalScrollBar().sizeHint().height()+2*table.frameWidth()+4)
 def mapping(self): return {need:self.combos[need].currentText().strip() if self.combos[need].currentIndex()>0 else '' for need in self.required}
 def validate_mapping(self):
  values=self.mapping(); missing=[x for x in self.required if x not in self.optional and not values.get(x)]
  duplicates={x for x in values.values() if x and list(values.values()).count(x)>1}
  if missing or duplicates:
   detail=[]
   if missing: detail.append('未映射：'+ '、'.join(missing))
   if duplicates: detail.append('重复映射：'+ '、'.join(sorted(duplicates)))
   QMessageBox.warning(self,'表头映射错误','；'.join(detail)+'。请继续修改映射关系。'); return
  if self.mapping_config:
   try:self.mapping_config.save_import_mapping(self.mapping_kind,values)
   except (OSError,ValueError) as error:
    show_operation_error(self,'映射保存失败',error);return
  self.accept()
class MaterialConflictDialog(QDialog):
 def __init__(self,incoming,existing,parent=None):
  super().__init__(parent); self.setWindowTitle('发现重复物料'); self.resize(1000,400); self.choice='cancel'
  layout=QVBoxLayout(self); title=QLabel('物料 '+incoming['物料基础编码']+incoming['物料版本号']+' 已存在。上行为数据库现有信息，下行为待导入信息；黄色单元格表示差异。'); title.setWordWrap(True); layout.addWidget(title)
  fields=[('base_code','物料基础编码'),('version','物料版本号'),('name','物料名称'),('description','物料规格描述'),('unit','单位'),('grade','物料等级'),('mp_flag','M/P'),('manufacturer','物料制造商'),('mpn','物料MPN')]
  self.table=QTableWidget(2,len(fields)); self.table.setHorizontalHeaderLabels([label for _,label in fields]); self.table.setVerticalHeaderLabels(['现有数据库','目标导入']); self.table.setEditTriggers(QAbstractItemView.NoEditTriggers); self.table.setHorizontalScrollBarPolicy(Qt.ScrollBarAlwaysOn)
  for column,(key,label) in enumerate(fields):
   old=str(dict(existing).get(key,'') or ''); new=incoming.get(label,'')
   for row,value in enumerate((old,new)):
    item=QTableWidgetItem(value); item.setToolTip(value)
    if old!=new:item.setBackground(QColor('#ffe082')); item.setForeground(QColor('black'))
    self.table.setItem(row,column,item)
  layout.addWidget(self.table); fit_columns(self.table)
  buttons=QDialogButtonBox()
  for label,choice in [('覆盖','overwrite'),('跳过','skip'),('取消全部','cancel')]:
   button=buttons.addButton(label,QDialogButtonBox.ActionRole); button.clicked.connect(lambda _,value=choice:self.finish(value))
  layout.addWidget(buttons)
 def finish(self,choice):self.choice=choice; self.accept()

class MaterialImportResultDialog(QDialog):
 def __init__(self,ok,fail,parent=None):
  super().__init__(parent); self.setWindowTitle('物料导入结果')
  available=self.screen().availableGeometry(); self.setMaximumHeight(min(640,int(available.height()*0.85))); self.resize(min(900,int(available.width()*0.9)),min(560,self.maximumHeight()))
  automatic=[row for row in fail if row[3]=='重复物料信息相同，自动跳过']
  other=[row for row in fail if row[3] in ('跳过重复物料','用户取消')]
  errors=[row for row in fail if row[3] not in ('重复物料信息相同，自动跳过','跳过重复物料','用户取消')]
  layout=QVBoxLayout(self)
  summary=QLabel(f'成功导入：{len(ok)}条    导入失败：{len(errors)}条    自动跳过：{len(automatic)}条    其他未导入：{len(other)}条'); summary.setWordWrap(True); layout.addWidget(summary)
  self.tabs=QTabWidget(); self.details=QTextBrowser(); self.skipped_details=QTextBrowser(); self.other_details=QTextBrowser()
  for title,rows,view in [('导入失败',errors,self.details),('自动跳过',automatic,self.skipped_details),('其他未导入',other,self.other_details)]:
   view.setOpenExternalLinks(False); view.setVerticalScrollBarPolicy(Qt.ScrollBarAsNeeded); view.setHorizontalScrollBarPolicy(Qt.ScrollBarAlwaysOff); view.setLineWrapMode(QTextEdit.WidgetWidth)
   view.setPlainText('\n\n'.join(f'第{row}行：{code or "（空编码）"} {name or "（空名称）"}\n原因：{reason}' for row,code,name,reason in rows) or '无此类记录。')
   self.tabs.addTab(view,f'{title}（{len(rows)}）')
  self.tabs.setCurrentIndex(0 if errors else 1 if automatic else 2 if other else 0)
  layout.addWidget(self.tabs,1)
  buttons=QDialogButtonBox(QDialogButtonBox.Ok); buttons.accepted.connect(self.accept); layout.addWidget(buttons)

def export_filename(text):
 import re
 return re.sub(r'[\\/:*?"<>|\x00-\x1f]', '_', text).strip().rstrip('.')+'.xlsx'

def export_start_path(filename):
 import sys
 from PyQt5.QtCore import QSettings
 settings=QSettings('Biong','BOMManager'); folder=settings.value('export_directory','')
 if not folder or not Path(folder).is_dir():
  folder=str((Path(os.environ.get('APPIMAGE')).resolve().parent if os.environ.get('APPIMAGE') else (Path(sys.executable).resolve().parent if getattr(sys,'frozen',False) else Path(__file__).resolve().parents[2])))
 return str(Path(folder)/filename)

def remember_export_path(path):
 from PyQt5.QtCore import QSettings
 QSettings('Biong','BOMManager').setValue('export_directory',str(Path(path).resolve().parent))


def export_material_structure(parent,service,material_id,root_color=None):
 material=service.material_snapshot(material_id)
 if not material:return
 sid=getattr(service,'sku_id',None)
 if sid is None:
  legacy=service.db.conn.execute('SELECT id FROM skus WHERE material_id=? AND scope_id=0',(material_id,)).fetchone()
  if legacy:sid=legacy['id']
 sku=service.db.conn.execute('SELECT material_id,mirror_version FROM skus WHERE id=?',(sid,)).fetchone() if sid else None
 suffix=' 历史版本'+str(sku['mirror_version']) if sku and getattr(service,'scope_id',0) else ''
 filename=export_filename(material['base_code']+material['version']+' '+material['name']+suffix)
 path=QFileDialog.getSaveFileName(parent,'导出BOM结构',export_start_path(filename),'Excel (*.xlsx)')[0]
 if not path:return
 try:
  remember_export_path(path)
  if sku and sku['material_id']==material_id:service.export_sku(sid,path)
  else:service.export_bom_structure(material_id,path,root_color=root_color)
  QMessageBox.information(parent,'完成','BOM结构导出成功')
 except Exception as error:show_operation_error(parent,'导出BOM结构失败',error)


class MaterialSelectionDialog(QInputDialog):
 @staticmethod
 def getItem(parent,title,label,items,current=0,editable=True):
  dialog=MaterialSelectionDialog(parent); dialog.setWindowTitle(title); dialog.setLabelText(label)
  dialog.setComboBoxItems(items); dialog.setComboBoxEditable(True); dialog.resize(650,160)
  combo=dialog.findChild(QComboBox); combo.setInsertPolicy(QComboBox.NoInsert)
  combo.setSizeAdjustPolicy(QComboBox.AdjustToMinimumContentsLengthWithIcon);combo.setMinimumContentsLength(40)
  completer=QCompleter(combo.model(),dialog); completer.setCaseSensitivity(Qt.CaseInsensitive); completer.setFilterMode(Qt.MatchContains); completer.setCompletionMode(QCompleter.PopupCompletion); combo.setCompleter(completer)
  combo.setCurrentIndex(-1); combo.lineEdit().setPlaceholderText('输入物料编码、版本或名称筛选')
  dialog._enter_selection=FirstMaterialOnEnter(combo,dialog.accept)
  dialog._valid_choices=set(items)
  return (dialog.textValue(),True) if dialog.exec_()==QDialog.Accepted else ('',False)
 def accept(self):
  if self.textValue() not in self._valid_choices:
   QMessageBox.warning(self,'请选择物料','请从筛选结果中选择物料，不能使用任意输入文字。');return
  super().accept()

class NewSkuDialog(QDialog):
 def __init__(self,svc,parent=None):
  super().__init__(parent); self.svc=svc; self.material_id=None; self.setWindowTitle('新建SKU'); self.resize(650,160)
  layout=QVBoxLayout(self); layout.addWidget(QLabel('输入基础编码、版本号或名称筛选物料；已有BOM结构的物料不可用。'))
  self.combo=QComboBox(); self.combo.setEditable(True); self.combo.setInsertPolicy(QComboBox.NoInsert); self.combo.setMinimumWidth(500)
  self.model=QStandardItemModel(self); skus={row['material_id'] for row in svc.skus()}
  for row in svc.materials():
   reason='已有BOM结构，不可用' if row['has_bom']=='Y' else '已是SKU，不可用' if row['id'] in skus else ''
   item=QStandardItem(f"{row['base_code']}{row['version']} {row['name']}"+('（'+reason+'）' if reason else '')); item.setData(row['id'],Qt.UserRole); item.setEnabled(not reason); self.model.appendRow(item)
  self.combo.setModel(self.model); self.combo.setCurrentIndex(-1); self.combo.lineEdit().setPlaceholderText('输入物料编码或名称')
  self.proxy=QSortFilterProxyModel(self); self.proxy.setSourceModel(self.model); self.proxy.setFilterCaseSensitivity(Qt.CaseInsensitive)
  self.completer=QCompleter(self.proxy,self); self.completer.setCompletionMode(QCompleter.UnfilteredPopupCompletion); self.combo.setCompleter(self.completer)
  self._enter_selection=FirstMaterialOnEnter(self.combo,self.accept)
  self.combo.lineEdit().textEdited.connect(self.filter_materials); layout.addWidget(self.combo)
  buttons=QDialogButtonBox(QDialogButtonBox.Ok|QDialogButtonBox.Cancel); buttons.accepted.connect(self.accept); buttons.rejected.connect(self.reject); layout.addWidget(buttons)
 def filter_materials(self,text):
  self.proxy.setFilterFixedString(text); self.completer.complete()
 def accept(self):
  text=self.combo.currentText(); matches=[self.model.item(i) for i in range(self.model.rowCount()) if self.model.item(i).text()==text]
  if not matches or not matches[0].isEnabled():QMessageBox.warning(self,'无法新建SKU','请选择可用的物料，不能使用任意输入文字或已有BOM结构的物料。'); return
  self.material_id=matches[0].data(Qt.UserRole); super().accept()

from ui.sku_lock import lock_icon,show_operation_error,CenteredLockDelegate
from ui.project_management import ProjectManagement

class MainWindow(ProjectManagement,QMainWindow):
 def __init__(self,svc,version='0.0.1'):
  super().__init__(); self.svc=svc; self._bom_windows={}; self._material_bom_windows={}; self._materials=None; self.setWindowTitle('BOM管理系统'); self.resize(1100,700); self.version=version; self.build(); WindowRefreshBus.subscribe(self.on_data_change)
 def on_data_change(self,event):
  if event.get('type') in ('bom_changed','material_changed','sku_changed','sku_lock_changed'):self.refresh()
 def closeEvent(self,event):
  # Resolve every owned BOM draft before hiding the main window or removing
  # its subscriptions. A rejected close must leave the application usable.
  windows=set(self._material_bom_windows.values())|set(self._bom_windows.values())
  if self._materials is not None:windows.update(self._materials._local_bom_windows.values())
  for window in windows:
   if not window.close():event.ignore();return
  if self._materials is not None and not self._materials.close():event.ignore();return
  if getattr(self,'_create_material_dialog',None):self._create_material_dialog.reject()
  if getattr(self,'_compare_dialog',None):self._compare_dialog.close()
  WindowRefreshBus.unsubscribe(self.on_data_change); super().closeEvent(event)
 def build(self):
  bar=self.menuBar(); fm=bar.addMenu('文件'); a=fm.addAction('选择SQLite数据库'); a.triggered.connect(self.select_database); a=fm.addAction('导出资料库'); a.triggered.connect(self.export_database); pm=bar.addMenu('项目'); self.setup_projects(fm,pm); a=pm.addAction('修改项目信息'); a.triggered.connect(self.edit_project); a=bar.addAction('物料库'); a.triggered.connect(self.show_materials); a=bar.addAction('新建SKU'); a.setShortcut('Ctrl+N');a.setShortcutContext(Qt.WindowShortcut); a.triggered.connect(self.new_sku);
  self.table=QTableWidget(); self.table.setColumnCount(8); self.table.setHorizontalHeaderLabels(['打开','SKU编码','SKU名称','SKU描述','版本号','修改日期时间','ID','锁定']); self.table.horizontalHeader().moveSection(7,1); self.table.hideColumn(6); self.table.setSelectionBehavior(QAbstractItemView.SelectRows); self.table.setSelectionMode(QAbstractItemView.ExtendedSelection); self.table.setWordWrap(True); self.table.cellDoubleClicked.connect(self.open_bom); self.table.setContextMenuPolicy(Qt.CustomContextMenu); self.table.customContextMenuRequested.connect(self.sku_context_menu)
  self.table._fixed_column_widths={7:max(44,self.table.horizontalHeader().minimumSectionSize(),self.table.fontMetrics().horizontalAdvance('锁定')+20)};self.table.setItemDelegateForColumn(7,CenteredLockDelegate(self.table))
  self._sku_sort=None;self._sku_filters={}
  self._description_hover=VersionDescriptionHover(self.table)
  filter_action=bar.addAction('筛选');filter_action.triggered.connect(self.filter_skus)
  compare_action=bar.addAction('BOM对比');compare_action.triggered.connect(self.compare_boms)
  a=bar.addAction('关于'); a.triggered.connect(self.show_about)
  self.table.setEditTriggers(QAbstractItemView.NoEditTriggers)
  self.table.horizontalHeader().setSectionsClickable(True);self.table.horizontalHeader().sectionClicked.connect(self.sort_skus)
  for column in (1,2,4,5):self.table.horizontalHeaderItem(column).setToolTip('点击升序/降序排序')
  corner=QWidget(); corner_layout=QHBoxLayout(corner); corner_layout.setContentsMargins(0,0,4,0)
  self.theme_button=QToolButton(); self.theme_button.setIconSize(QSize(24,24)); self.theme_button.setAutoRaise(True); self.theme_button.clicked.connect(self.toggle_theme); corner_layout.addWidget(self.theme_button); bar.setCornerWidget(corner,Qt.TopRightCorner)
  self.set_theme(bool(QApplication.instance().property('bom_dark_theme')))
  self.setCentralWidget(self.table)
  self.count_label=QLabel(''); self.version_label=QLabel('版本 '+self.version); self.version_label.setAlignment(Qt.AlignCenter)
  self.latest_bom_only=QCheckBox('只显示最新BOM'); self.latest_bom_only.setChecked(False); self.latest_bom_only.setToolTip('同一SKU仅显示历史版本号最高的BOM'); self.latest_bom_only.toggled.connect(self.apply_sku_filters)
  footer=QWidget(); footer_l=QGridLayout(footer); footer_l.setContentsMargins(8,2,8,2)
  footer_l.addWidget(self.count_label,0,0,Qt.AlignLeft|Qt.AlignVCenter); footer_l.addWidget(self.version_label,0,1); footer_l.addWidget(self.latest_bom_only,0,2,Qt.AlignRight|Qt.AlignVCenter)
  for column in range(3):footer_l.setColumnStretch(column,1)
  root_l=QVBoxLayout(); root_l.setContentsMargins(0,0,0,0); root_l.addWidget(self.table); root_l.addWidget(footer); container=QWidget(); container.setLayout(root_l); self.setCentralWidget(container); self.statusBar().setSizeGripEnabled(False); self.statusBar().hide(); self.refresh()

 def compare_boms(self):
  from ui.bom_compare import BomCompareDialog
  if not getattr(self,'_compare_dialog',None):self._compare_dialog=BomCompareDialog(self.svc,self)
  self._compare_dialog.live_timer.stop();self._compare_dialog.clear_results();self._compare_dialog.reload_choices();self._compare_dialog.show();self._compare_dialog.raise_();self._compare_dialog.activateWindow()
 def compare_sku_bom(self,sid):
  self.compare_boms()
  dialog=self._compare_dialog;index=dialog.left.findData(sid)
  if index<0:return
  dialog.left.setCurrentIndex(index)
 def show_about(self):
  AboutDialog(self.version,self).exec_()
 def set_theme(self,dark):
  apply_theme(dark); self.theme_button.setIcon(theme_icon(dark))
  label='当前深色主题（月亮），点击切换浅色主题' if dark else '当前浅色主题（太阳），点击切换深色主题'
  self.theme_button.setToolTip(label); self.theme_button.setAccessibleName(label)
 def toggle_theme(self):self.set_theme(not bool(QApplication.instance().property('bom_dark_theme')))
 def sort_skus(self,column):
  if column not in (1,2,4,5):return
  order=Qt.DescendingOrder if self._sku_sort==(column,Qt.AscendingOrder) else Qt.AscendingOrder
  self._sku_sort=(column,order);self.table.sortItems(column,order)
  self.table.horizontalHeader().setSortIndicator(column,order);self.table.horizontalHeader().setSortIndicatorShown(True)
  self.apply_sku_filters();self.publish_sku_order()
 def publish_sku_order(self):
  self.svc.sku_display_order=[int(self.table.item(row,6).text()) for row in range(self.table.rowCount())]
  WindowRefreshBus.publish({'type':'sku_order_changed'})
 def refresh(self):
  available=self.svc.db is not None
  def update_actions(menu):
   for action in menu.actions():
    if action.menu():update_actions(action.menu())
    elif action.text() in ('选择SQLite数据库','导出资料库','修改项目信息','物料库','新建SKU','筛选','BOM对比'):action.setEnabled(available)
  update_actions(self.menuBar());self.table.setEnabled(available);self.latest_bom_only.setEnabled(available)
  if not available:
   self.table.setRowCount(0);self.setWindowTitle('BOM管理系统');self.count_label.setText('未打开项目，请通过“项目”菜单打开历史项目、新建或导入项目');return
  selected_ids={self.table.item(index.row(),6).text() for index in self.table.selectionModel().selectedRows() if self.table.item(index.row(),6)}
  current=self.table.currentRow(); column=max(1,self.table.currentColumn())
  selected_id=str(self._description_target) if getattr(self,'_description_target',None) is not None else (self.table.item(current,6).text() if current>=0 and self.table.item(current,6) else None)
  scroll=(self.table.horizontalScrollBar().value(),self.table.verticalScrollBar().value())
  project=self.svc.project(); name=(project['project_name'] or '').strip(); code=(project['project_code'] or '').strip(); prefix=' '.join(x for x in (code,name) if x and x!='请输入项目名称'); title=f'{prefix} BOM管理系统' if prefix else 'BOM管理系统'
  self.setWindowTitle(title)
  rows=self.svc.skus(); self.table.setRowCount(len(rows));
  for i,r in enumerate(rows):
   btn=QPushButton('打开'); btn.clicked.connect(lambda _,sid=r['id']: self.open_bom_id(sid)); self.table.setCellWidget(i,0,btn); self.table.setItem(i,1,QTableWidgetItem(r['base_code']+r['version'])); self.table.setItem(i,2,QTableWidgetItem(r['name'])); self.table.setItem(i,3,QTableWidgetItem(r['description'])); self.table.setItem(i,4,QTableWidgetItem());self.table.item(i,4).setData(Qt.DisplayRole,r['mirror_version']); self.table.item(i,4).setTextAlignment(Qt.AlignCenter); self.table.setItem(i,5,QTableWidgetItem(sku_timestamp(r['updated_at']))); self.table.setItem(i,6,QTableWidgetItem(str(r['id'])))
   lock=QTableWidgetItem();lock.setData(Qt.UserRole,bool(r['locked']));lock.setTextAlignment(Qt.AlignCenter);lock.setToolTip('已锁定，只能浏览BOM' if r['locked'] else '未锁定')
   if r['locked']:lock.setIcon(lock_icon())
   self.table.setItem(i,7,lock)
   self.table.item(i,4).setData(DESCRIPTION_ROLE,r['change_description'])
   version_font=self.table.item(i,4).font();version_font.setBold(bool((r['change_description'] or '').strip()));self.table.item(i,4).setFont(version_font)
   self.table.item(i,1).setData(Qt.UserRole,r['material_id'])
  if self._sku_sort:self.table.sortItems(*self._sku_sort)
  fit_columns(self.table)
  self.apply_sku_filters()
  if selected_id is not None:
   selected_row=next((r for r in range(self.table.rowCount()) if self.table.item(r,6).text()==selected_id),None)
   self.table.clearSelection()
   if selected_row is not None and not self.table.isRowHidden(selected_row):
    self.table.setCurrentCell(selected_row,column);self.table.selectRow(selected_row)
   else:self.table.setCurrentItem(None)
  if selected_ids and not getattr(self,'_description_target',None):
   self.table.clearSelection()
   for r in range(self.table.rowCount()):
    if self.table.item(r,6).text() in selected_ids and not self.table.isRowHidden(r):self.table.selectionModel().select(self.table.model().index(r,1),QItemSelectionModel.Select|QItemSelectionModel.Rows)
  self.table.horizontalScrollBar().setValue(scroll[0]);self.table.verticalScrollBar().setValue(scroll[1])
  if selected_id is not None and selected_row is not None and not self.table.isRowHidden(selected_row):
   item=self.table.item(selected_row,column)
   if not self.table.viewport().rect().intersects(self.table.visualItemRect(item)):self.table.scrollToItem(item,QAbstractItemView.PositionAtCenter)
  self.publish_sku_order()
 def filter_skus(self):
  labels=['SKU编码','SKU名称','SKU描述','版本号','修改日期时间','清除全部筛选']
  label,ok=QInputDialog.getItem(self,'SKU筛选','选择筛选列',labels,0,False)
  if not ok:return
  if label==labels[-1]:self._sku_filters.clear();self.apply_sku_filters();return
  column=labels.index(label)+1
  dialog=MaterialColumnFilterDialog(label,[self.table.item(row,column).text() for row in range(self.table.rowCount())],self._sku_filters.get(column),self)
  if dialog.exec_()==QDialog.Accepted:
   condition=dialog.condition()
   if condition==('',None):self._sku_filters.pop(column,None)
   else:self._sku_filters[column]=condition
   self.apply_sku_filters()
 def apply_sku_filters(self):
  count=0; latest={}; sku_codes=set(); visible_codes=set()
  only_latest=self.latest_bom_only.isChecked()
  if only_latest:
   for row in range(self.table.rowCount()):
    material_id=self.table.item(row,1).data(Qt.UserRole); version=int(self.table.item(row,4).text())
    latest[material_id]=max(latest.get(material_id,version),version)
  for row in range(self.table.rowCount()):
   visible=all(query.casefold() in self.table.item(row,col).text().casefold() and (selected is None or self.table.item(row,col).text() in selected) for col,(query,selected) in self._sku_filters.items())
   if only_latest:visible=visible and int(self.table.item(row,4).text())==latest[self.table.item(row,1).data(Qt.UserRole)]
   code=self.table.item(row,1).text();sku_codes.add(code)
   if visible:visible_codes.add(code)
   self.table.setRowHidden(row,not visible);count+=int(visible)
  self.count_label.setText(f'共计{len(sku_codes)}个SKU，{self.table.rowCount()}个BOM')
  self.count_label.setToolTip(f'SKU按编码去重，每个历史版本计为一个BOM。当前显示{len(visible_codes)}个SKU，{count}个BOM。')
 def copy_sku_bom(self,sid):
  row=self.svc.db.conn.execute('SELECT material_id,scope_id FROM skus WHERE id=?',(sid,)).fetchone()
  window=self._bom_windows.get(sid) or (self._material_bom_windows.get(row['material_id']) if row and not row['scope_id'] else None)
  if window and window.dirty:QMessageBox.information(self,'请先保存','当前历史版本有未保存编辑，请先保存再复制BOM。');return
  try:
   new_id=self.svc.copy_sku_bom(sid);self.open_bom_id(new_id)
  except Exception as error:show_operation_error(self,'复制BOM失败',error)
 def change_mirror_version(self,sid):
  row=self.svc.db.conn.execute('SELECT mirror_version FROM skus WHERE id=?',(sid,)).fetchone()
  if row is None:return
  value,ok=QInputDialog.getInt(self,'历史版本变更','新历史版本号',row[0],1,2147483647)
  if ok:
   try:self.svc.change_mirror_version(sid,value)
   except Exception as error:show_operation_error(self,'历史版本变更失败',error)
 def select_database(self):
  current=Path(self.svc.db.path).expanduser() if self.svc.db is not None else self.svc.config_service.base
  # 使用Qt对话框而非原生保存对话框，禁止已有数据库的“是否替换”确认。
  # 这样已有文件会直接作为候选数据库打开，输入新文件名仍可创建数据库。
  chooser=QFileDialog(self,'选择SQLite数据库',str(current),'SQLite (*.db *.sqlite);;所有文件 (*)')
  chooser.setFileMode(QFileDialog.AnyFile)
  chooser.setAcceptMode(QFileDialog.AcceptSave)
  chooser.setOption(QFileDialog.DontUseNativeDialog,True)
  chooser.setOption(QFileDialog.DontConfirmOverwrite,True)
  if chooser.exec_()!=QDialog.Accepted:return
  p=chooser.selectedFiles()[0] if chooser.selectedFiles() else ''
  if not p:return
  selected=Path(p).expanduser()
  existed=selected.exists()
  if self.svc.db is not None and selected.resolve() == current.resolve():
   QMessageBox.information(self,'数据库未切换','当前已经在使用该数据库：\n'+str(selected)); return
  if self.switch_project(selected,create=not existed):
   QMessageBox.information(self,'数据库已切换',('已打开已有数据库' if existed else '已创建并打开新数据库')+'：\n'+str(selected))
 def export_database(self):
  p=QFileDialog.getSaveFileName(self,'导出资料库',export_start_path(export_filename((self.svc.project()['project_code']+' '+self.svc.project()['project_name']+' 资料库').strip())),'Excel (*.xlsx)')[0]
  if p:remember_export_path(p); self.svc.export_database(p); QMessageBox.information(self,'完成','资料库导出成功')
 def edit_project(self):
  d=ProjectDialog(self,self.svc.project())
  if d.exec_(): self.svc.update_project(d.data()); self.remember_current_project(); self.refresh()
 def activate_window(self,window):
  if window.isMinimized(): window.showNormal()
  else: window.show()
  window.raise_(); window.activateWindow()
 def show_materials(self):
  if self._materials is None:
   self._materials=MaterialWindow(self.svc,self)
   self._materials.closed.connect(lambda:setattr(self,'_materials',None))
  self.activate_window(self._materials)
 def import_materials(self):
  owner=self._materials if self._materials is not None and self._materials.isVisible() else self
  config_service=self.svc.config_service or ConfigService(Path(self.svc.db.path).parent); cfg=config_service.load()
  last=config_service.resolve_path(cfg['database']['last_import_dir']) if cfg['database'].get('last_import_dir') else Path.home()/'Desktop'
  last=str(last if last.is_dir() else Path.home()/'Desktop')
  p=QFileDialog.getOpenFileName(owner,'导入物料清单',last,'Excel (*.xlsx)')[0]
  if p:
    cfg['database']['last_import_dir']=str(Path(p).parent); config_service.save(cfg)
    try:sheet_name=choose_material_sheet(p,owner)
    except Exception as error:show_operation_error(owner,'导入失败',error); return
    def resolver(d,old):
     dialog=MaterialConflictDialog(d,old,owner); dialog.exec_(); return dialog.choice
    progress=QProgressDialog('正在导入物料，请稍候…','取消',0,100,owner); progress.setWindowTitle('导入物料'); progress.setWindowModality(Qt.NonModal); progress.setAutoClose(False); progress.show(); QApplication.processEvents()
    def update_progress(current,total):
     progress.setMaximum(max(total,1)); progress.setValue(min(current,total)); QApplication.processEvents()
    try:
     try: ok,fail=self.svc.import_materials(p,resolver,header_mapping=config_service.import_mapping('materials') or None,progress_callback=update_progress,sheet_name=sheet_name)
     except MaterialHeaderError as header_error:
      progress.close()
      mapping_dialog=MaterialHeaderMappingDialog(p,header_error.headers,header_error.missing,owner,sheet_name=sheet_name,config_service=config_service)
      if mapping_dialog.exec_()!=QDialog.Accepted:
       progress.close(); QMessageBox.warning(owner,'导入失败','用户取消表头映射：未完成必需表头对应关系。'); return
      progress=QProgressDialog('正在导入物料，请稍候…','取消',0,100,owner); progress.setWindowTitle('导入物料'); progress.setWindowModality(Qt.NonModal); progress.setAutoClose(False); progress.show(); QApplication.processEvents()
      ok,fail=self.svc.import_materials(p,resolver,mapping_dialog.mapping(),update_progress,sheet_name=sheet_name)
     progress.setValue(progress.maximum()); progress.close()
     self.refresh();
     if self._materials is not None and self._materials.isVisible(): self._materials.refresh()
     MaterialImportResultDialog(ok,fail,owner).exec_()
    except Exception as e:
     progress.close(); show_operation_error(owner,'导入失败',e)
 def new_material(self):
  owner=self._materials if self._materials is not None and self._materials.isVisible() else self
  return show_create_material(owner)
 def new_sku(self):
  dialog=NewSkuDialog(self.svc,self)
  if dialog.exec_()==QDialog.Accepted:
   try:
    sid=self.svc.create_sku(dialog.material_id,require_empty_bom=True); self.refresh(); self.open_bom_id(sid)
   except Exception as error:show_operation_error(self,'新建SKU失败',error)
 def change_sku_version(self,sid=None):
  r=self.table.currentRow() if sid is None else -1
  if sid is None and r<0:return
  if sid is None:sid=int(self.table.item(r,6).text())
  v,ok=QInputDialog.getText(self,'SKU版本变更','新版本号')
  if ok:
   try:self.svc.change_sku_version(sid,v.strip()); self.refresh()
   except Exception as e: show_operation_error(self,'版本变更失败',e)
 def delete_sku(self,sid=None):
  r=self.table.currentRow() if sid is None else -1
  if sid is None:
   if r<0:return
   sid=int(self.table.item(r,6).text())
  sku=self.svc.db.conn.execute('SELECT material_id FROM skus WHERE id=?',(sid,)).fetchone()
  if sku is None:return
  prompt='该SKU物料仍被其他BOM引用，是否仅移除SKU标记？物料及BOM结构将保留。' if self.svc.material_is_bom_referenced(sku['material_id']) else '删除当前SKU及BOM？'
  if QMessageBox.question(self,'确认',prompt)==QMessageBox.Yes:
   window=self._bom_windows.get(sid)
   if window and not window.close():return
   try:self.svc.delete_sku(sid);self.refresh()
   except Exception as error:show_operation_error(self,'删除SKU失败',error)
 def focus_sku_description(self,sid,scroll):
  row=next((r for r in range(self.table.rowCount()) if self.table.item(r,6).text()==str(sid)),None)
  if row is None or self.table.isRowHidden(row):return
  self.table.setCurrentCell(row,4);self.table.selectRow(row)
  if self.isVisible():self.activateWindow()
  self.table.setFocus(Qt.OtherFocusReason)
  self.table.horizontalScrollBar().setValue(scroll[0]);self.table.verticalScrollBar().setValue(scroll[1])
  item=self.table.item(row,4)
  if not self.table.viewport().rect().contains(self.table.visualItemRect(item)):self.table.scrollToItem(item,QAbstractItemView.PositionAtCenter)
 def edit_sku_change_description(self,sid):
  row=next((r for r in self.svc.skus() if r['id']==sid),None)
  if row is None:return
  scroll=(self.table.horizontalScrollBar().value(),self.table.verticalScrollBar().value())
  self._description_target=sid
  dialog=ChangeDescriptionDialog(row['change_description'],self)
  try:
   if dialog.exec_()==QDialog.Accepted:
    try:self.svc.update_sku_change_description(sid,dialog.description())
    except Exception as error:show_operation_error(self,'保存失败',error)
  finally:
   self._description_target=None
   self.focus_sku_description(sid,scroll)
   # Qt restores the pre-dialog focus during queued activation/menu cleanup.
   # Restore the explicitly edited mirror once more after those events settle.
   if getattr(self,'_description_focus_timer',None):self._description_focus_timer.stop();self._description_focus_timer.deleteLater()
   timer=QTimer(self);timer.setSingleShot(True);self._description_focus_timer=timer
   database=self.svc.db
   timer.timeout.connect(lambda:self.focus_sku_description(sid,scroll) if self.isVisible() and self.svc.db is database else None)
   timer.start(0)
 def toggle_sku_lock(self,sid):
  row=self.svc.db.conn.execute('SELECT material_id,locked FROM skus WHERE id=?',(sid,)).fetchone()
  if row is None:return
  window=self._bom_windows.get(sid) or self._material_bom_windows.get(row['material_id'])
  if not row['locked'] and window and window.dirty:
   QMessageBox.warning(self,'请先处理未保存内容','请先保存或放弃此BOM的未保存修改，再锁定SKU。');return
  try:self.svc.set_sku_locked(sid,not row['locked'])
  except Exception as error:show_operation_error(self,'锁定操作失败',error)
 def set_selected_sku_locks(self,sids,locked):
  if locked:
   for sid in sids:
    row=self.svc.db.conn.execute('SELECT material_id,locked FROM skus WHERE id=?',(sid,)).fetchone()
    if row is None:continue
    window=self._bom_windows.get(sid) or self._material_bom_windows.get(row['material_id'])
    if not row['locked'] and window and window.dirty:
     QMessageBox.warning(self,'请先处理未保存内容','选中的BOM包含未保存修改，请先保存或放弃修改后再锁定。');return
  try:
   with self.svc.db.transaction():
    for sid in sids:self.svc.set_sku_locked(sid,locked)
  except Exception as error:show_operation_error(self,'锁定操作失败',error)
 def sku_context_menu(self,pos):
  row=self.table.rowAt(pos.y())
  if row<0:return
  if not self.table.selectionModel().isRowSelected(row):
   index=self.table.model().index(row,1);self.table.selectionModel().setCurrentIndex(index,QItemSelectionModel.ClearAndSelect|QItemSelectionModel.Rows)
  rows=sorted(index.row() for index in self.table.selectionModel().selectedRows() if not self.table.isRowHidden(index.row()))
  sids=[int(self.table.item(r,6).text()) for r in rows]
  if not sids:return
  sid=int(self.table.item(row,6).text());multiple=len(sids)>1
  states={r['id']:bool(r['locked']) for r in self.svc.skus() if r['id'] in sids};locked=states[sid]
  menu=QMenu(self)
  description=menu.addAction('编辑变更描述');version=menu.addAction('SKU版本变更');mirror=menu.addAction('历史版本变更')
  for edit in (description,version,mirror):edit.setEnabled(not locked and not multiple)
  menu.addSeparator();open_action=menu.addAction('编辑/浏览BOM')
  menu.addSeparator()
  lock=unlock=None
  if multiple:
   lock=menu.addAction('锁定SKU');lock.setEnabled(not all(states.values()))
   unlock=menu.addAction('取消锁定');unlock.setEnabled(any(states.values()))
  else:lock=menu.addAction('取消锁定' if locked else '锁定SKU')
  menu.addSeparator();copy=menu.addAction('复制BOM');compare=menu.addAction('BOM对比')
  copy.setEnabled(not multiple);compare.setEnabled(not multiple)
  menu.addSeparator();export_action=menu.addAction('导出BOM');purchase_action=menu.addAction('采购BOM导出')
  menu.addSeparator();delete=menu.addAction('删除SKU');delete.setEnabled(not locked and not multiple)
  action=menu.exec_(self.table.viewport().mapToGlobal(pos))
  if action is None or not action.isEnabled():return
  if action==open_action:
   for selected in sids:self.open_bom_id(selected)
  elif action==lock:
   if multiple:self.set_selected_sku_locks(sids,True)
   else:self.toggle_sku_lock(sid)
  elif action==unlock:self.set_selected_sku_locks(sids,False)
  elif action==description:self.edit_sku_change_description(sid)
  elif action==copy:self.copy_sku_bom(sid)
  elif action==compare:self.compare_sku_bom(sid)
  elif action==mirror:self.change_mirror_version(sid)
  elif action==version:self.change_sku_version(sid)
  elif action in (export_action,purchase_action):
   if multiple:self.export_sku_batch(sids,purchase=action==purchase_action)
   else:
    try:self.export_sku_id(sid,purchase=action==purchase_action)
    except Exception as error:show_operation_error(self,'导出失败',error)
  elif action==delete:self.delete_sku(sid)
 def export_sku_batch(self,sids,purchase=False):
  title='批量采购BOM导出' if purchase else '批量导出BOM'
  folder=QFileDialog.getExistingDirectory(self,title+'：选择统一导出文件夹',str(Path(export_start_path('BOM.xlsx')).parent))
  if not folder:return
  remember_export_path(Path(folder)/'BOM.xlsx')
  success=[];failed=[]
  for sid in sids:
   row=self.svc.db.conn.execute('SELECT m.base_code,m.version,m.name,s.mirror_version FROM skus s JOIN materials m ON m.id=s.material_id WHERE s.id=?',(sid,)).fetchone()
   if row is None:failed.append(str(sid)+'：SKU已不存在');continue
   label=row['base_code']+row['version']+' '+row['name']+' 历史版本'+str(row['mirror_version'])
   filename=export_filename(label+(' 采购BOM' if purchase else ''))
   path=Path(folder)/filename;number=2
   while path.exists():path=Path(folder)/(Path(filename).stem+f' ({number}).xlsx');number+=1
   try:
    self.export_sku_id(sid,purchase= purchase,destination=path,notify=False);success.append(path.name)
   except Exception as error:failed.append(label+'：'+str(error))
  box=QMessageBox(self);box.setWindowTitle(title+'结果');box.setIcon(QMessageBox.Warning if failed else QMessageBox.Information)
  box.setText(f'成功导出{len(success)}个BOM，失败{len(failed)}个。\n导出目录：{folder}')
  box.setDetailedText('成功文件：\n'+'\n'.join(success)+('\n\n失败原因：\n'+'\n'.join(failed) if failed else ''))
  box.exec_()
 def export_sku_id(self,sid,purchase=False,destination=None,notify=True):
  row=self.svc.db.conn.execute('SELECT m.id AS material_id,m.base_code,m.version,m.name,s.mirror_version,s.scope_id FROM skus s JOIN materials m ON m.id=s.material_id WHERE s.id=?',(sid,)).fetchone()
  if not row:return
  title='采购BOM导出' if purchase else '导出BOM'
  filename=row['base_code']+row['version']+' '+row['name']+(' 历史版本'+str(row['mirror_version']) if row['scope_id'] else '')+(' 采购BOM' if purchase else '')
  p=str(destination) if destination is not None else QFileDialog.getSaveFileName(self,title,export_start_path(export_filename(filename)),'Excel (*.xlsx)')[0]
  if p:
   remember_export_path(p)
   if purchase:self.svc.for_sku(sid).export_purchase_bom(row['material_id'],p)
   else:self.svc.export_sku(sid,p)
   if notify:QMessageBox.information(self,'完成',title+'成功')
 def open_bom(self,row,col): self.open_bom_id(int(self.table.item(row,6).text()))
 def open_bom_id(self,sid):
  row=self.svc.db.conn.execute('SELECT material_id,scope_id FROM skus WHERE id=?',(sid,)).fetchone()
  if row is None:return
  if row['scope_id']:
   if sid not in self._bom_windows:
    window=BomWindow(self.svc.for_sku(sid),row['material_id'],self);self._bom_windows[sid]=window
    window.closed.connect(lambda:self._bom_windows.pop(sid,None))
   window=self._bom_windows[sid];self.activate_window(window);return window
  window=self.open_material_bom(row['material_id'])
  if window is not None:self._bom_windows[sid]=window
  return window
 def open_material_bom(self,material_id):
  from services.snapshot_bom import enabled,template_service
  source=self.svc
  if enabled(self.svc.db):
   skus=list(self.svc.db.conn.execute('SELECT id FROM skus WHERE material_id=? ORDER BY mirror_version DESC,id DESC',(material_id,)))
   if skus:return self.open_bom_id(skus[0]['id'])
   source=template_service(self.svc,material_id)
  if self.svc.db.conn.execute('SELECT 1 FROM materials WHERE id=?',(material_id,)).fetchone() is None:return
  if material_id not in self._material_bom_windows:
   window=BomWindow(source,material_id,self)
   self._material_bom_windows[material_id]=window
   def unregister():
    self._material_bom_windows.pop(material_id,None)
    for sid,opened in list(self._bom_windows.items()):
     if opened is window:self._bom_windows.pop(sid,None)
   window.closed.connect(unregister)
  window=self._material_bom_windows[material_id]
  self.activate_window(window)
  return window
 def export_materials(self):
  p=QFileDialog.getSaveFileName(self,'导出物料库',export_start_path(export_filename((self.svc.project()['project_code']+' '+self.svc.project()['project_name']+' 物料库').strip())),'Excel (*.xlsx)')[0]
  if p:remember_export_path(p); self.svc.export_materials(p); QMessageBox.information(self,'完成','导出成功')
 def template(self):
  p=QFileDialog.getSaveFileName(self,'导出模板',export_start_path('物料模板.xlsx'),'Excel (*.xlsx)')[0]
  if p:
   remember_export_path(p)
   from openpyxl import Workbook; w=Workbook(); w.active.append(MATERIAL_IMPORT_HEADERS); w.save(p); w.close()
class BomImportIssuesDialog(QDialog):
 def __init__(self,error,parent=None):
  super().__init__(parent); self.setWindowTitle('BOM导入失败')
  self.resize(800,520); self.setMaximumHeight(640)
  layout=QVBoxLayout(self)
  summary=QLabel(f'扫描完成，共发现{len(error.issues)}个问题。未导入，原有编辑和数据库保持不变。'); summary.setWordWrap(True); layout.addWidget(summary)
  self.details=QTextBrowser(); self.details.setPlainText('\n\n'.join(error.issues)); layout.addWidget(self.details)
  buttons=QDialogButtonBox(QDialogButtonBox.Close); buttons.rejected.connect(self.reject); layout.addWidget(buttons)

class BomImportConflictDialog(QDialog):
 def __init__(self,conflicts,parent=None):
  super().__init__(parent);self.setWindowTitle('导入子BOM与系统不一致');self.resize(1100,580);self.choice='stop'
  layout=QVBoxLayout(self)
  label=QLabel('发现以下子BOM差异。请选择停止导入、保留系统子BOM，或使用新导入子BOM。继续导入仅更新编辑缓存，点击保存后才写入数据库；共享此子BOM的其他SKU也会受到保存的影响。')
  label.setWordWrap(True);layout.addWidget(label)
  self.table=QTableWidget(len(conflicts),5);self.table.setHorizontalHeaderLabels(['父项物料','子项物料','不一致项','系统子BOM','新导入子BOM'])
  self.table.setEditTriggers(QAbstractItemView.NoEditTriggers);self.table.setWordWrap(True)
  self.table.setHorizontalScrollBarPolicy(Qt.ScrollBarAsNeeded);self.table.setVerticalScrollBarPolicy(Qt.ScrollBarAsNeeded)
  for row,conflict in enumerate(conflicts):
   for column,key in enumerate(('parent','material','difference','system','incoming')):self.table.setItem(row,column,QTableWidgetItem(conflict[key]))
  layout.addWidget(self.table);fit_columns(self.table);self.table.resizeRowsToContents()
  buttons=QDialogButtonBox()
  for text,choice in [('停止导入','stop'),('继续导入并以系统子BOM为准','system'),('继续导入并以新导入子BOM为准','incoming')]:
   button=buttons.addButton(text,QDialogButtonBox.ActionRole);button.setAutoDefault(False);button.clicked.connect(lambda _,value=choice:self.finish(value))
  layout.addWidget(buttons)
 def finish(self,choice):self.choice=choice;self.accept() if choice!='stop' else self.reject()

class BomTreeWidget(QTreeWidget):
 def __init__(self,*args):
  super().__init__(*args);self.setSelectionMode(QAbstractItemView.ExtendedSelection);self.setSelectionBehavior(QAbstractItemView.SelectRows)
 def mousePressEvent(self,event):
  item=self.itemAt(event.pos())
  if event.button()==Qt.RightButton and item is not None and item.isSelected():event.accept();return
  super().mousePressEvent(event)
  if item is not None:
   for selected in self.selectedItems():
    if selected.parent() is not item.parent():selected.setSelected(False)

 def scrollTo(self,index,hint=QAbstractItemView.EnsureVisible):
  if not getattr(self,'_hold_view',False):super().scrollTo(index,hint)
 def drawRow(self,painter,option,index):
  super().drawRow(painter,option,index)
  painter.save(); painter.setPen(self.palette().mid().color())
  rect=option.rect
  painter.drawLine(0,rect.bottom(),self.viewport().width(),rect.bottom())
  for column in range(self.columnCount()):
   x=self.header().sectionViewportPosition(column)+self.columnWidth(column)-1
   painter.drawLine(x,rect.top(),x,rect.bottom())
  painter.restore()

class BomWindow(QMainWindow):
 closed=pyqtSignal()
 def __init__(self,svc,root,parent=None):
  super().__init__(parent); self.setAttribute(Qt.WA_DeleteOnClose); self.svc=svc; self.root=root; self.current_item=None; self._deleted_relations=set(); self._alternate_edits={}; self._external_paths=set(); self._draft_baselines={}; self._loading_fields=False; self._saving=False; self.notice=QLabel(''); self.notice.setStyleSheet('color:red'); WindowRefreshBus.subscribe(self.on_external_change); self.setWindowTitle('BOM管理'); self.resize(1200,700); self.dirty=False
  self.tree=BomTreeWidget(); self.tree.setVerticalScrollMode(QAbstractItemView.ScrollPerPixel); self.tree.itemClicked.connect(self.show_item_info); self.tree.setWordWrap(True); self.tree.setContextMenuPolicy(Qt.CustomContextMenu); self.tree.customContextMenuRequested.connect(self.context_menu); self.tree.setHeaderLabels(['物料编码/名称','数量','单位','位置','备注']); self.tree.setMinimumWidth(420)
  self.delete_action=QAction('删除物料',self.tree);self.delete_action.setShortcut('Delete');self.delete_action.setShortcutContext(Qt.WidgetWithChildrenShortcut);self.delete_action.setAutoRepeat(False);self.delete_action.triggered.connect(lambda:self.confirm_delete(self.tree.currentItem()));self.tree.addAction(self.delete_action)
  import hashlib
  database_path=str(Path(self.svc.db.path).resolve())
  self._expansion_settings=QSettings('Biong','BOMManager')
  self._expansion_key='bom_expansion/'+hashlib.sha256(database_path.encode()).hexdigest()+'/'+str(self.root)+('/'+str(self.svc.scope_id) if getattr(self.svc,'scope_id',0) else '')
  import json
  try:self._expansion_state=json.loads(self._expansion_settings.value(self._expansion_key,'{}'))
  except (ValueError,TypeError):self._expansion_state={}
  self._rendering_tree=False
  self.tree.itemExpanded.connect(self.remember_expansion);self.tree.itemCollapsed.connect(self.remember_expansion)
  self.fields={}; form=QFormLayout();
  for k,l in [('base_code','物料基础编码'),('version','物料版本号'),('name','物料名称'),('description','物料描述'),('quantity','物料数量'),('unit','物料单位'),('grade','物料等级'),('position','物料位置'),('manufacturer','物料制造商'),('mpn','物料MPN'),('mp_flag','M/P'),('remark','备注')]:
   w=QuantityEdit() if k=='quantity' else (WrappingTextEdit() if k in ('position','remark') else QLabel()); self.fields[k]=w; form.addRow(l,w)
   if isinstance(w,QLabel):w.setTextInteractionFlags(Qt.TextSelectableByMouse|Qt.TextSelectableByKeyboard);w.setWordWrap(True);w.setTextFormat(Qt.PlainText);w.setMinimumWidth(0);w.setSizePolicy(QSizePolicy.Ignored,QSizePolicy.Preferred)
  self.fields['quantity'].valueChanged.connect(lambda value:self.update_tree_field('quantity',value))
  for key in ('position','remark'): self.fields[key].textChanged.connect(lambda k=key:self.update_tree_field(k,self.fields[k].text()))
  info_box=QGroupBox('物料信息'); info_box.setLayout(form); self.alt_table=QTableWidget(0,6); self.alt_table.setHorizontalHeaderLabels(['替代料编码','版本号','名称','描述','制造商','MPN']); self.alt_table.setEditTriggers(QAbstractItemView.NoEditTriggers); self.alt_table.setContextMenuPolicy(Qt.CustomContextMenu); self.alt_table.customContextMenuRequested.connect(self.alternate_menu); alt_box=QGroupBox('替代料'); alt_l=QVBoxLayout(alt_box); alt_l.addWidget(self.alt_table)
  right=QWidget(); right_l=QVBoxLayout(right); right_l.addWidget(info_box); right_l.addWidget(alt_box); self.info_scroll=QScrollArea();self.info_scroll.setWidgetResizable(True);self.info_scroll.setWidget(right);self.info_scroll.setMinimumWidth(0);spl=QSplitter();self.bom_splitter=spl; spl.addWidget(self.tree); spl.addWidget(self.info_scroll); spl.setStretchFactor(0,1);spl.setStretchFactor(1,1); box=QWidget(); lay=QVBoxLayout(box); lay.addWidget(spl); self.setCentralWidget(box); self.statusBar().addWidget(self.notice); tb=self.addToolBar('BOM'); self.save_action=QAction('保存',self); self.save_action.setShortcut('Ctrl+S'); self.save_action.setShortcutContext(Qt.WindowShortcut); self.save_action.setAutoRepeat(False); self.save_action.triggered.connect(self.save); self.save_action.setEnabled(False); tb.addAction(self.save_action); self.refresh_action=QAction('刷新',self); self.refresh_action.triggered.connect(self.refresh_external); self.refresh_action.setEnabled(False); tb.addAction(self.refresh_action); self.export_action=QAction('导出BOM',self); self.export_action.triggered.connect(self.export_bom); self.export_action.setEnabled(False); tb.addAction(self.export_action)
  self.purchase_export_action=QAction('采购BOM导出',self); self.purchase_export_action.triggered.connect(self.export_purchase_bom); self.purchase_export_action.setEnabled(self.export_action.isEnabled()); tb.addAction(self.purchase_export_action)
  self.export_action.changed.connect(lambda:self.purchase_export_action.setEnabled(self.export_action.isEnabled()))
  self.default_panel_width=self.fontMetrics().horizontalAdvance('M')*100
  screen=self.screen().availableGeometry();self.tree.setMinimumWidth(min(420,max(200,screen.width()//2-40)));self.resize(min(self.default_panel_width*2+50,screen.width()),min(850,screen.height()));spl.setSizes([self.default_panel_width,self.default_panel_width])
  self.search_action=tb.addAction('查找物料');self.search_action.setShortcut('Ctrl+F');self.search_action.setShortcutContext(Qt.WindowShortcut);self.search_action.triggered.connect(self.search_materials)
  self.load();self.update_mirror_title();self.apply_lock_state()
 def search_materials(self):
  if not hasattr(self,'_search_dialog'):self._search_dialog=BomSearchDialog(self)
  self._search_dialog.show();self._search_dialog.raise_();self._search_dialog.activateWindow()
 def is_locked(self):
  sid=getattr(self.svc,'sku_id',None)
  row=self.svc.db.conn.execute('SELECT locked FROM skus WHERE id=?',(sid,)).fetchone() if sid else self.svc.db.conn.execute('SELECT MAX(locked) FROM skus WHERE material_id=? AND scope_id=0',(self.root,)).fetchone()
  return bool(row and row[0])
 def apply_lock_state(self):
  locked=self.is_locked();self.save_action.setEnabled(self.dirty and not locked)
  title=self.windowTitle().removesuffix(' — BOM浏览')
  self.setWindowTitle(title+(' — BOM浏览' if locked else ''))
  if self.current_item:self.show_item_info(self.current_item)
 def mirror_suffix(self):
  if getattr(self.svc,'sku_id',None):
   row=self.svc.db.conn.execute('SELECT mirror_version FROM skus WHERE id=?',(self.svc.sku_id,)).fetchone()
   if row:return ' 历史版本'+str(row[0])
  return ''
 def update_mirror_title(self):
  if getattr(self.svc,'sku_id',None):
   sku=self.svc.db.conn.execute('SELECT mirror_version FROM skus WHERE id=?',(self.svc.sku_id,)).fetchone()
   material=self.svc.material_snapshot(self.root)
   if sku:self.setWindowTitle(f"{material['base_code']}{material['version']} {material['name']} 历史版本{sku[0]}")
 def mark_dirty(self):
  if self.is_locked():return False
  self.dirty=True
  self.save_action.setEnabled(True)
  self.export_action.setEnabled(False)
  self.refresh_styles()
  WindowRefreshBus.publish({'type':'bom_draft_changed'})
 def on_external_change(self,event):
  if event.get('type')=='sku_lock_changed':self.apply_lock_state();return
  if self._saving: return
  if event.get('type')=='sku_changed':
   self.update_mirror_title();self.apply_lock_state();return
  if event.get('type')=='bom_changed' and event.get('scope_id',0)!=getattr(self.svc,'scope_id',0):return
  if event.get('type') in ('bom_changed','material_changed') and self.contains_material(event.get('root') if event.get('type')=='bom_changed' else event.get('id')):
   self.notice.setText('BOM有更新，请点击刷新或重新打开窗口'); self.export_action.setEnabled(False); self.refresh_action.setEnabled(True); self.statusBar().clearMessage()
 def closeEvent(self,event):
  if not self.dirty: WindowRefreshBus.unsubscribe(self.on_external_change); self.closed.emit(); WindowRefreshBus.publish({'type':'bom_draft_changed'}); event.accept(); return
  box=QMessageBox(self); box.setWindowTitle('BOM未保存'); box.setText('BOM结构或物料信息已修改，是否保存？'); save=box.addButton('保存并退出',QMessageBox.AcceptRole); discard=box.addButton('不保存退出',QMessageBox.DestructiveRole); cancel=box.addButton('取消',QMessageBox.RejectRole); box.exec_()
  if box.clickedButton()==save:
   if self.save(): WindowRefreshBus.unsubscribe(self.on_external_change); self.closed.emit(); WindowRefreshBus.publish({'type':'bom_draft_changed'}); event.accept()
   else: event.ignore()
  elif box.clickedButton()==discard: WindowRefreshBus.unsubscribe(self.on_external_change); self.closed.emit(); WindowRefreshBus.publish({'type':'bom_draft_changed'}); event.accept()
  else: event.ignore()

 def context_menu(self,pos):
  it=self.tree.itemAt(pos)
  if not it:return
  if not it.isSelected():self.tree.clearSelection();self.tree.setCurrentItem(it)
  d=it.data(0,Qt.UserRole) or {}
  menu=QMenu(self)
  if len(self.selected_copy_items(it))>1:
   actions=self.clipboard_actions(menu,it);self.add_color_menu(menu,it);chosen=menu.exec_(self.tree.viewport().mapToGlobal(pos));self.handle_clipboard_action(chosen,actions,it);return
  if self.is_locked():
   reverse=menu.addAction('物料反查');copy=menu.addAction('复制物料');export=self.subtree_export_action(menu,it);chosen=menu.exec_(self.tree.viewport().mapToGlobal(pos))
   if chosen==reverse:self.reverse_tree_material(it)
   elif chosen==copy:self.copy_bom_material(it)
   elif export is not None and chosen==export:self.export_subtree(it)
   return
  if d.get('_deleted'):
   restore=menu.addAction('取消删除');menu.addSeparator();copy=menu.addAction('复制物料');menu.addSeparator();reverse=menu.addAction('物料反查')
   act=menu.exec_(self.tree.viewport().mapToGlobal(pos))
   if act==restore:self.restore_relation(it)
   elif act==copy:self.copy_bom_material(it)
   elif act==reverse:self.reverse_tree_material(it)
   return
  root=it is self.root_item
  edit=menu.addAction('编辑此物料');reverse=menu.addAction('物料反查')
  replace=version=same=alternate=up=down=dele=None
  if not root:
   menu.addSeparator();replace=menu.addAction('替换物料');version=menu.addAction('物料版本更新');version.setEnabled(bool(self.material_versions(it)))
  menu.addSeparator()
  if not root:same=menu.addAction('新增同级物料')
  child=menu.addAction('新增子级物料')
  if not root:alternate=menu.addAction('新增替代料')
  imp=menu.addAction('导入BOM');imp.setEnabled(it.childCount()==0 and not self.svc.bom(d.get('material_id',self.root)))
  imp.setToolTip('当前物料已有下阶物料，不能重复导入' if not imp.isEnabled() else '导入当前物料的下阶BOM')
  export=self.subtree_export_action(menu,it)
  self.add_color_menu(menu,it)
  menu.addSeparator();clipboard_actions=self.clipboard_actions(menu,it)
  if not root:
   menu.addSeparator();up=menu.addAction('物料上移');down=menu.addAction('物料下移')
   up.setEnabled(self.can_move_item(it,-1));down.setEnabled(self.can_move_item(it,1))
   menu.addSeparator();dele=menu.addAction('删除物料')
  act=menu.exec_(self.tree.viewport().mapToGlobal(pos))
  if act is None:return
  if export is not None and act==export:self.export_subtree(it);return
  if self.handle_clipboard_action(act,clipboard_actions,it):return
  if act==edit:self.edit_tree_material(it);return
  if act==reverse:self.reverse_tree_material(it);return
  if act==replace:self.replace_material(it);return
  if act==version:self.update_material_version(it);return
  if act==up:self.move_item(it,-1);return
  if act==down:self.move_item(it,1);return
  if act==imp:self.import_bom(it);return
  if act==alternate:
   self.tree.setCurrentItem(it);self.show_item_info(it);self.add_alternate();return
  if act in (same,child):self.add_relative(d,act==child);return
  if act==dele:self.confirm_delete(it)
 def add_color_menu(self,menu,item):
  if self.is_locked() or any((node.data(0,Qt.UserRole) or {}).get('_deleted') for node in self.selected_copy_items(item)):return
  from ui.bom_colors import ColorMenu
  if self.svc.config_service is None:self.svc.config_service=ConfigService(Path(self.svc.db.path).parent)
  menu.addSeparator();menu.addMenu(ColorMenu(self,self.selected_copy_items(item),menu))
 @preserve_bom_view
 def set_row_color(self,items,color):
  if self.is_locked():return
  from services.bom_colors import normalize
  color=normalize(color)
  keys={(0 if node is self.root_item else (node.parent().data(0,Qt.UserRole) or {})['material_id'],(node.data(0,Qt.UserRole) or {})['material_id']) for node in items if not (node.data(0,Qt.UserRole) or {}).get('_deleted')}
  changed=False
  for _,node in self.expansion_nodes():
   data=dict(node.data(0,Qt.UserRole) or {});key=(0 if node is self.root_item else (node.parent().data(0,Qt.UserRole) or {})['material_id'],data['material_id'])
   if key in keys and data.get('row_color','')!=color:
    data['row_color']=color;node.setData(0,Qt.UserRole,data);changed=True
    if node is self.root_item:self._root_color_dirty=True
  if changed:self.mark_dirty()
 def subtree_export_action(self,menu,item):
  if not item.childCount():return None
  menu.addSeparator();action=menu.addAction('导出BOM结构')
  action.setEnabled(self.export_action.isEnabled() and not self.dirty and not (item.data(0,Qt.UserRole) or {}).get('_deleted'))
  return action
 def export_subtree(self,item):
  if not self.export_action.isEnabled() or self.dirty or not item.childCount() or (item.data(0,Qt.UserRole) or {}).get('_deleted'):return
  export_material_structure(self,self.svc,(item.data(0,Qt.UserRole) or {})['material_id'],root_color=(item.data(0,Qt.UserRole) or {}).get('row_color',''))
 def confirm_delete(self,item):
  if self.is_locked() or item is None or item is self.root_item or (item.data(0,Qt.UserRole) or {}).get('_deleted'):return
  if QMessageBox.question(self,'确认删除','是否删除选定的物料及其子物料？\n'+item.text(0),QMessageBox.Yes|QMessageBox.No,QMessageBox.Yes)==QMessageBox.Yes:self.remove_relation(item)
 def reverse_tree_material(self,item):
  material_id=(item.data(0,Qt.UserRole) or {}).get('material_id',self.root if item is self.root_item else None)
  service=getattr(self.svc,'primary',self.svc)
  material=service.material_snapshot(material_id)
  if not material:
   QMessageBox.warning(self,'物料不存在','目标物料已不存在，请刷新BOM后重试。');return
  MaterialUsageDialog(material,service.material_usage(material_id),self).exec_()
 def edit_tree_material(self,item):
  if self.is_locked():return False
  data=item.data(0,Qt.UserRole) or {}
  if data.get('_deleted'):return
  material_id=data.get('material_id',self.root if item is self.root_item else None)
  material=self.svc.material_snapshot(material_id)
  if not material:
   QMessageBox.warning(self,'物料不存在','目标物料已不存在，请刷新BOM后重试。');return
  dialog=MaterialDialog(self,material,readonly_keys={'base_code','version'})
  dialog.setWindowTitle('编辑此物料')
  if dialog.exec_()!=QDialog.Accepted:return
  try:self.svc.update_material(material_id,dialog.data())
  except Exception as error:show_operation_error(self,'保存失败',error)
 def replacement_blocked_materials(self,item,child=False):
  # Reverse reachability computes all cyclic candidates in one graph traversal.
  target=item if child else item.parent();blocked={target.child(i).data(0,Qt.UserRole)['material_id'] for i in range(target.childCount())}
  reverse={}
  for parent,child in self.svc.bom_graph_edges():
   reverse.setdefault(child,set()).add(parent)
  for _,node in self.expansion_nodes():
   if node.parent() is not None:
    parent=(node.parent().data(0,Qt.UserRole) or {})['material_id'];child=(node.data(0,Qt.UserRole) or {})['material_id']
    reverse.setdefault(child,set()).add(parent)
  pending=[(target.data(0,Qt.UserRole) or {})['material_id']];visited=set()
  while pending:
   mid=pending.pop()
   if mid in visited:continue
   visited.add(mid);pending.extend(reverse.get(mid,()))
  return blocked|visited
 def material_versions(self,item):
  data=item.data(0,Qt.UserRole) or {}
  return list(self.svc.db.conn.execute('SELECT * FROM materials WHERE base_code=? AND id<>? ORDER BY version',(data.get('base_code'),data.get('material_id'))))
 def version_subtree_snapshot(self,item):
  nodes=[]
  for index in range(item.childCount()):
   child=item.child(index);data=dict(child.data(0,Qt.UserRole) or {})
   if data.get('_deleted'):continue
   data.update(quantity=float(child.text(1)),position=child.text(3),_db_alternates=self.alternate_ids(child))
   nodes.append(dict(data=data,children=self.version_subtree_snapshot(child)))
  return nodes
 def update_material_version(self,item):
  if self.is_locked():return False
  if item is self.root_item or (item.data(0,Qt.UserRole) or {}).get('_deleted'):return
  rows=self.material_versions(item)
  if not rows:return
  choices=[f"{row['base_code']}{row['version']} {row['name']}" for row in rows]
  choice,ok=MaterialSelectionDialog.getItem(self,'物料版本更新','选择同基础编码的其他版本物料',choices,0,False)
  if not ok:return
  material=rows[choices.index(choice)];key=self.relation_key(item)
  matches=[node for _,node in self.expansion_nodes() if node is not self.root_item and not (node.data(0,Qt.UserRole) or {}).get('_deleted') and self.relation_key(node)==key]
  for node in matches:
   reason=self.paste_block_reason(node,False,material['id'])
   if material['id'] in self.alternate_ids(node):reason='目标版本已是当前物料的替代料，不能继承为自身替代料'
   if reason:QMessageBox.warning(self,'无法更新物料版本',reason.replace('粘贴','替换'));return
  return self.replace_with_subtree_choice(item,matches,material)
 def replace_with_subtree_choice(self,item,matches,material):
  subtree=None
  original=self.version_subtree_snapshot(item)
  target=next((node for _,node in self.expansion_nodes() if not (node.data(0,Qt.UserRole) or {}).get('_deleted') and (node.data(0,Qt.UserRole) or {}).get('material_id')==material['id']),None)
  from services.current_catalog import enabled as catalog_enabled,selection_tree
  current_catalog=catalog_enabled(self.svc)
  target_tree=selection_tree(self.svc,material['id']) if current_catalog else self.version_subtree_snapshot(target) if target is not None else self.svc.bom(material['id'])
  if target_tree:
   if current_catalog or target is not None:subtree=target_tree
  elif original:
   answer=QMessageBox.question(self,'复制原子BOM','目标物料没有BOM结构，是否将原物料的子BOM复制到目标物料？\n选择“否”仅替换物料，不复制子BOM；点击保存后才写入数据库。',QMessageBox.Yes|QMessageBox.No|QMessageBox.Cancel,QMessageBox.Cancel)
   if answer==QMessageBox.Cancel:return
   if answer==QMessageBox.Yes:
    def contains(nodes):
     return any(n['data']['material_id']==material['id'] or contains(n['children']) for n in nodes)
    if contains(original):QMessageBox.warning(self,'无法更新物料版本','原子BOM包含目标版本，复制会造成循环引用。');return
    subtree=original
   else:subtree=[]
  elif target is not None:subtree=[]
  if subtree and not target_tree:
   for _,node in list(self.expansion_nodes()):
    if not (node.data(0,Qt.UserRole) or {}).get('_deleted') and (node.data(0,Qt.UserRole) or {}).get('material_id')==material['id']:
     node.takeChildren();self.clone_version_children(node,subtree)
  return self.replace_nodes(item,matches,material,version_children=subtree)
 def clone_version_children(self,parent,nodes):
  for node in nodes:
   source=node['data'];pd=parent.data(0,Qt.UserRole) or {}
   child=QTreeWidgetItem(parent,[f"{source['base_code']}{source['version']} {source['name']}",format_quantity(source['quantity']),source.get('unit',''),source.get('position',''),source.get('remark','')])
   temp=-id(child);data={k:v for k,v in source.items() if not k.startswith('_')}
   data.update(id=temp,node_id=str(temp),parent_node_id=pd.get('node_id'),parent_item_id=pd.get('id'),parent_material_id=pd['material_id'],level_no=(pd.get('level_no') or 0)+1,_added=True)
   child.setData(0,Qt.UserRole,data);self.by[str(temp)]=child
   self._alternate_edits[temp]=list(source.get('_db_alternates',[]))
   self.clone_version_children(child,node['children']);self._draft_baselines[temp]=self.capture_item(child)
  parent.setExpanded(True)
 def replace_material(self,item):
  if self.is_locked():return False
  if item is self.root_item or (item.data(0,Qt.UserRole) or {}).get('_deleted'):return
  alternates=self.alternate_ids(item)
  blocked=self.replacement_blocked_materials(item)|set(alternates)
  rows=[row for row in self.svc.db.conn.execute('SELECT * FROM materials ORDER BY base_code,version') if row['id'] not in blocked]
  if not rows:QMessageBox.information(self,'替换物料','没有可用的替换物料（需避免同级重复、循环引用以及与继承的替代料相同）。');return
  choices=[f"{row['base_code']}{row['version']} {row['name']}" for row in rows]
  choice,ok=MaterialSelectionDialog.getItem(self,'替换物料','输入筛选替换物料（继承数量、位置、备注及替代料）',choices,0,False)
  if not ok:return
  material=rows[choices.index(choice)]
  old_id=(item.data(0,Qt.UserRole) or {})['material_id']
  matches=[node for _,node in self.expansion_nodes() if node is not self.root_item and not (node.data(0,Qt.UserRole) or {}).get('_deleted') and (node.data(0,Qt.UserRole) or {}).get('material_id')==old_id]
  if len(matches)>1:
   groups={}
   for node in matches:
    key=self.relation_key(node)
    group=groups.setdefault(key,dict(nodes=[],paths=[],current=False,reason=''))
    group['nodes'].append(node);group['current']=group['current'] or node is item
    path=[];ancestor=node
    while ancestor is not None:path.append(ancestor.text(0));ancestor=ancestor.parent()
    group['paths'].append(' / '.join(reversed(path)))
    reason=self.paste_block_reason(node,False,material['id'])
    if material['id'] in self.alternate_ids(node):reason='目标物料也是此关系的替代料，不能继承为自身替代料'
    if reason:group['reason']=reason.replace('粘贴','替换')
   dialog=ReplacementSelectionDialog(list(groups.values()),self)
   if dialog.exec_()!=QDialog.Accepted:return
   matches=dialog.selected_nodes()
   if not matches:return
  return self.replace_with_subtree_choice(item,matches,material)
 @preserve_bom_view
 def replace_nodes(self,item,matches,material,swap_alternate=False,version_children=None):
  if self.is_locked():return False
  selected=None
  for old in matches:
   original=dict(old.data(0,Qt.UserRole) or {})
   added=self.insert_material(old.parent(),material,after=old)
   if version_children is not None:
    added.takeChildren();self.clone_version_children(added,version_children)
   data=dict(added.data(0,Qt.UserRole) or {})
   for field,column in [('quantity',1),('position',3),('remark',4)]:
    data[field]=original.get(field,1.0 if field=='quantity' else '')
    added.setText(column,format_quantity(data[field]) if field=='quantity' else str(data[field] or ''))
   added.setData(0,Qt.UserRole,data)
   alternates=self.alternate_ids(old)
   if swap_alternate:alternates=[original['material_id'] if mid==material['id'] else mid for mid in alternates]
   self._alternate_edits[self.alternate_key(added)]=alternates
   self._draft_baselines[data['id']]=self.capture_item(added)
   if selected is None or old is item:selected=added
  for old in matches:
   if not (old.data(0,Qt.UserRole) or {}).get('_deleted'):self.remove_relation(old)
  self.mark_dirty();fit_columns(self.tree)
  if selected is not None:self.tree.setCurrentItem(selected);self.show_item_info(selected);self.tree.scrollToItem(selected)
  return selected
 def can_move_item(self,item,direction):
  parent=item.parent()
  if parent is None or direction not in (-1,1) or (item.data(0,Qt.UserRole) or {}).get('_deleted'):return False
  return 0<=parent.indexOfChild(item)+direction<parent.childCount()
 @preserve_bom_view
 def move_item(self,item,direction):
  if self.is_locked():return False
  if not self.can_move_item(item,direction):return
  parent=item.parent(); index=parent.indexOfChild(item); expanded=item.isExpanded()
  parent.takeChild(index); parent.insertChild(index+direction,item); item.setExpanded(expanded)
  data=item.data(0,Qt.UserRole) or {}; data['_moved']=True; item.setData(0,Qt.UserRole,data)
  # 同一共享父物料的多个显示位置同步排序，避免保存时被另一副本覆盖。
  parent_id=(parent.data(0,Qt.UserRole) or {})['material_id']
  order=[(parent.child(i).data(0,Qt.UserRole) or {}).get('material_id') for i in range(parent.childCount())]
  for _,other in list(self.expansion_nodes()):
   if other is parent or (other.data(0,Qt.UserRole) or {}).get('material_id')!=parent_id:continue
   children=other.takeChildren(); rank={mid:i for i,mid in enumerate(order)}
   children.sort(key=lambda child:rank.get((child.data(0,Qt.UserRole) or {}).get('material_id'),len(rank)))
   other.addChildren(children)
   for child in children:
    d=child.data(0,Qt.UserRole) or {}
    if d.get('material_id')==data['material_id']:d['_moved']=True;child.setData(0,Qt.UserRole,d)
  self.mark_dirty(); self.tree.setCurrentItem(item); self.show_item_info(item); self.tree.scrollToItem(item); self.remember_expansion()
 def clipboard_actions(self,menu,item):
  actions={'copy':menu.addAction('复制物料')}
  if item is not self.root_item and not self.is_locked() and all(not (node.data(0,Qt.UserRole) or {}).get('_deleted') for node in self.selected_copy_items(item)):actions['cut']=menu.addAction('剪切物料')
  entries=clipboard_entries(self.svc)
  if entries and not self.is_locked():
   for key,label,child in [('same','粘贴为同级物料',False),('child','粘贴为子级物料',True)]:
    action=menu.addAction(label); reason=self.batch_paste_reason(item,child,entries)
    action.setEnabled(not reason); action.setToolTip(reason or label); actions[key]=action
  return actions
 def handle_clipboard_action(self,action,actions,item):
  if action is None:return False
  if action==actions['copy']:
   self.copy_bom_material(item); return True
  if 'cut' in actions and action==actions['cut']:
   self.cut_material(item);return True
  for key,child in [('same',False),('child',True)]:
   if key in actions and action==actions[key]:
    if actions[key].isEnabled():self.paste_material(item,child)
    return True
  return False
 def selected_copy_items(self,item):
  if item is None:return []
  selected=self.tree.selectedItems() if item.isSelected() else [item]
  selected=[node for node in selected if node.parent() is item.parent()]
  return sorted(selected,key=lambda node:node.parent().indexOfChild(node) if node.parent() else 0)
 def copy_bom_material(self,item):
  nodes=self.selected_copy_items(item);entries=[]
  for node in nodes:
   data=node.data(0,Qt.UserRole) or {}
   relation=None if node is self.root_item else dict(quantity=float(node.text(1)),position=node.text(3),remark=data.get('remark',''),alternate_material_ids=self.alternate_ids(node),row_color=data.get('row_color',''))
   entries.append((data.get('material_id',self.root),relation))
  if len(entries)==1:return copy_material(self.svc,*entries[0])
  return copy_materials(self.svc,entries)
 def cut_material(self,item):
  nodes=self.selected_copy_items(item)
  if self.is_locked() or not nodes or any(node is self.root_item or (node.data(0,Qt.UserRole) or {}).get('_deleted') for node in nodes):return False
  if self.copy_bom_material(item):
   for node in nodes:self.remove_relation(node)
   return True
  return False
 def batch_paste_reason(self,item,child,entries):
  ids=[material['id'] for material,relation in entries]
  if len(ids)!=len(set(ids)):return '剪贴板中包含重复物料'
  return next((reason for material,relation in entries if (reason:=self.paste_block_reason(item,child,material['id']))),'')
 def paste_block_reason(self,item,child,material_id):
  target=item if child else item.parent()
  if target is None:return 'SKU根节点不能增加同级物料'
  if (item.data(0,Qt.UserRole) or {}).get('_deleted'):return '待删除物料不能粘贴'
  if any(not (target.child(i).data(0,Qt.UserRole) or {}).get('_deleted') and (target.child(i).data(0,Qt.UserRole) or {}).get('material_id')==material_id for i in range(target.childCount())):return '目标父级下已存在相同物料'
  # 合并数据库和当前编辑树的关系，检查跨分支、未保存关系导致的循环。
  graph={}
  for row in self.svc.bom_graph_edges():
   graph.setdefault(row[0],set()).add(row[1])
  for _,node in self.expansion_nodes():
   if node.parent() is not None:
    parent_id=(node.parent().data(0,Qt.UserRole) or {}).get('material_id')
    graph.setdefault(parent_id,set()).add((node.data(0,Qt.UserRole) or {}).get('material_id'))
  target_id=(target.data(0,Qt.UserRole) or {}).get('material_id'); pending=[material_id]; visited=set()
  while pending:
   mid=pending.pop()
   if mid==target_id:return '粘贴会造成循环引用'
   if mid in visited:continue
   visited.add(mid); pending.extend(graph.get(mid,()))
  return ''
 @preserve_bom_view
 def paste_material(self,item,child):
  if self.is_locked():return False
  entries=clipboard_entries(self.svc)
  if not entries or self.batch_paste_reason(item,child,entries):return
  target=item if child else item.parent();anchor=item
  for material,relation in entries:
   added=self.insert_material(target,material,after=None if child else anchor)
   if relation is not None:
    data=dict(added.data(0,Qt.UserRole) or {})
    for field,column in [('quantity',1),('position',3),('remark',4)]:
     data[field]=relation[field];added.setText(column,format_quantity(relation[field]) if field=='quantity' else relation[field])
    data['row_color']=relation.get('row_color','')
    added.setData(0,Qt.UserRole,data);self._alternate_edits[self.alternate_key(added)]=relation['alternate_material_ids'];self._draft_baselines[data['id']]=self.capture_item(added);self.mark_dirty()
   anchor=added
  self.tree.setCurrentItem(added); self.show_item_info(added); self.tree.scrollToItem(added)
  return added
 def relation_key(self,item):
  data=item.data(0,Qt.UserRole) or {}; parent=item.parent()
  parent_data=parent.data(0,Qt.UserRole) if parent else {}
  return (parent_data.get('material_id'),data.get('material_id'),item.text(3))
 @preserve_bom_view
 def remove_relation(self,item):
  if self.is_locked():return False
  """将共享关系的所有显示副本标记为待删除，保存前保留可见。"""
  if item is self.root_item:return
  key=self.relation_key(item); relation_id=(item.data(0,Qt.UserRole) or {}).get('id')
  matches=[]; stack=[self.root_item]
  while stack:
   node=stack.pop(); data=node.data(0,Qt.UserRole) or {}
   if node is not self.root_item and (self.relation_key(node)==key or (relation_id and relation_id>0 and data.get('id')==relation_id)):
    matches.append(node)
   else: stack.extend(node.child(i) for i in range(node.childCount()))
  if relation_id and relation_id>0:
   original=self.svc.db.conn.execute('SELECT parent_material_id,child_material_id,position FROM material_bom_items WHERE id=?',(relation_id,)).fetchone()
   if original:self._deleted_relations.add((original[0],original[1],original[2] or ''))
  import uuid
  deletion_token=uuid.uuid4().hex
  self.current_item=None
  for node in matches:
   data=node.data(0,Qt.UserRole) or {}; current_key=self.relation_key(node)
   self._deleted_relations.add(current_key)
   self._deleted_relations.add((current_key[0],current_key[1],data.get('position','') or ''))
   pending=[node]
   while pending:
    removed=pending.pop(); removed_data=removed.data(0,Qt.UserRole) or {}
    removed_data['_delete_stack']=list(removed_data.get('_delete_stack',[]))+[deletion_token]
    removed_data['_deleted']=True; removed.setData(0,Qt.UserRole,removed_data)
    pending.extend(removed.child(i) for i in range(removed.childCount()))
  self.mark_dirty(); self.tree.setCurrentItem(item); self.show_item_info(item)
 @preserve_bom_view
 def restore_relation(self,item):
  if self.is_locked():return False
  data=item.data(0,Qt.UserRole) or {}
  if not data.get('_deleted'):return
  # An inherited deletion restores its owning branch, including shared copies.
  stack=data.get('_delete_stack',[])
  if not stack:return
  token=stack[-1]
  for _,node in self.expansion_nodes():
   current=node.data(0,Qt.UserRole) or {}
   if current.get('_delete_stack')!=[token] or node.parent() is None:continue
   parent=node.parent()
   if any(other is not node and not (other.data(0,Qt.UserRole) or {}).get('_deleted') and (other.data(0,Qt.UserRole) or {}).get('material_id')==current.get('material_id') for other in (parent.child(i) for i in range(parent.childCount()))):
    QMessageBox.warning(self,'无法取消删除','父级下已存在相同物料，请先删除重复物料。');return False
  for _,node in self.expansion_nodes():
   current=dict(node.data(0,Qt.UserRole) or {})
   deletions=current.get('_delete_stack',[])
   if token not in deletions:continue
   current['_delete_stack']=[value for value in deletions if value!=token]
   current['_deleted']=bool(current['_delete_stack']);node.setData(0,Qt.UserRole,current)
  self._deleted_relations.clear()
  for _,node in self.expansion_nodes():
   current=node.data(0,Qt.UserRole) or {}
   if node is self.root_item or not current.get('_deleted'):continue
   self._deleted_relations.add(self.relation_key(node))
   relation_id=current.get('id')
   if isinstance(relation_id,int) and relation_id>0:
    original=self.svc.db.conn.execute('SELECT parent_material_id,child_material_id,position FROM material_bom_items WHERE id=?',(relation_id,)).fetchone()
    if original:self._deleted_relations.add(tuple(original))
  self.mark_dirty();self.tree.setCurrentItem(item);self.show_item_info(item)
 def contains_material(self,material_id):
  if material_id is None:return False
  if material_id==self.root:return True
  for i in range(self.tree.topLevelItemCount()):
   stack=[self.tree.topLevelItem(i)]
   while stack:
    item=stack.pop(); d=item.data(0,Qt.UserRole) or {}
    if d.get('material_id')==material_id:return True
    stack.extend([item.child(j) for j in range(item.childCount())])
  return False
 def capture_item(self,item):
  data=dict(item.data(0,Qt.UserRole) or {})
  data.update(quantity=float(item.text(1)),position=item.text(3),remark=data.get('remark',''))
  return {'data':data,'children':[self.capture_item(item.child(n)) for n in range(item.childCount())]}
 def capture_tree(self):
  return [self.capture_item(self.root_item.child(n)) for n in range(self.root_item.childCount())]
 def database_tree(self,root=None):
  roots=[]; nodes={}; alternate_cache={}
  for row in self.svc.bom(self.root if root is None else root):
   if row['id'] not in alternate_cache:alternate_cache[row['id']]=[r['material_id'] for r in self.svc.alternates(row['id'])]
   node={'data':dict(row),'children':[]}; node['data']['_db_alternates']=alternate_cache[row['id']]; nodes[row['node_id']]=node
   parent=nodes.get(row['parent_node_id'])
   (parent['children'] if parent else roots).append(node)
  return roots
 def mark_external_differences(self,before_nodes,after_nodes,path=()):
  before=self.snapshot_map(before_nodes,path); after=self.snapshot_map(after_nodes,path)
  changed=set(before.keys()-after.keys())
  for key,node in after.items():
   old=before.get(key)
   if old is None or self.remote_signature(old)!=self.remote_signature(node):changed.add(key)
  # 只记录自身变化；子项增删由子项路径表示，不能将父级误判为变化物料。
  self._external_paths.update(changed)
 def refresh_drafts(self,nodes,conflicts,path=()):
  from copy import deepcopy
  materials={row['id']:row for row in self.svc.materials()}
  def walk(node,parent_path):
   data=node['data']; draft_id=data.get('id'); current_path=parent_path+(data['material_id'],)
   baseline=self._draft_baselines.get(draft_id)
   if baseline is not None and not data.get('_deleted'):
    fresh=deepcopy(baseline); material=materials.get(data['material_id'])
    if material is not None:
     for field in ('base_code','version','name','description','unit','grade','manufacturer','mpn','mp_flag'):fresh['data'][field]=material[field]
     fresh['children']=self.database_tree(data['material_id'])
     def added(nodes):
      for child in nodes:
       child['data']['_added']=True; added(child['children'])
     added(fresh['children'])
     self.mark_external_differences([baseline],[fresh],parent_path)
     node=self.merge_refresh([baseline],[node],[fresh],conflicts)[0]
     self._draft_baselines[draft_id]=deepcopy(fresh)
   node['children']=[walk(child,current_path) for child in node['children']]
   return node
  return [walk(node,path) for node in nodes]
 def merge_refresh(self,baseline,local,remote,conflicts):
  """三方合并：只把本地相对载入基线的修改覆盖到最新数据库结构。"""
  from copy import deepcopy
  def identity(node,all_nodes):
   mid=node['data']['material_id']
   # 正常BOM同一父级仅出现一次物料；兼容历史多位置同物料关系。
   duplicate=any(sum(n['data']['material_id']==mid for n in group)>1 for group in all_nodes)
   return (mid,node['data'].get('id')) if duplicate else (mid,)
  groups=(baseline,local,remote)
  base={identity(n,groups):n for n in baseline}; ours={identity(n,groups):n for n in local}; theirs={identity(n,groups):n for n in remote}
  def modified(node,original):
   data=node['data']; previous=original['data']
   if data.get('_moved') or data.get('_deleted') or (data.get('_added') and not previous.get('_added')):return True
   if data.get('id') in self._alternate_edits and self._alternate_edits[data['id']]!=previous.get('_db_alternates',[]):return True
   if any(data.get(k)!=previous.get(k) for k in ('quantity','position','remark','row_color')):return True
   if [n['data'].get('id') for n in node['children']]!=[n['data'].get('id') for n in original['children']]:return True
   return any(modified(child,old) for child,old in zip(node['children'],original['children']))
  merged=[]
  for key,new in theirs.items():
   previous=base.get(key); edit=ours.get(key)
   if edit is not None and edit['data'].get('_deleted'):
    merged.append(deepcopy(edit)); continue
   if previous is not None and edit is None:continue  # 本地明确删除，刷新不能恢复。
   if edit is None:merged.append(deepcopy(new)); continue
   if previous is None:
    if edit['data'].get('id')!=new['data'].get('id'):conflicts.append('双方新增相同物料')
    previous={'data':{},'children':[]}
   result=deepcopy(new)
   if edit['data'].get('_added'):result['data']['_added']=True
   if edit['data'].get('_moved'):result['data']['_moved']=True
   for field in ('quantity','position','remark','row_color'):
    value=edit['data'].get(field)
    if value!=previous['data'].get(field):
     if new['data'].get(field)!=previous['data'].get(field) and value!=new['data'].get(field):conflicts.append('同一字段同时修改')
     result['data'][field]=value
   old_id=edit['data'].get('id'); new_id=result['data'].get('id')
   if old_id in self._alternate_edits and old_id!=new_id:self._alternate_edits[new_id]=list(self._alternate_edits[old_id])
   result['children']=self.merge_refresh(previous['children'],edit['children'],new['children'],conflicts)
   merged.append(result)
  for key,edit in ours.items():
   if key in theirs:continue
   previous=base.get(key)
   if previous is None or edit['data'].get('_deleted') or modified(edit,previous):
    if previous is not None and not edit['data'].get('_deleted'):conflicts.append('外部删除了本地已编辑物料')
    merged.append(deepcopy(edit))
  if any(node['data'].get('_moved') for node in local):
   rank={node['data']['material_id']:i for i,node in enumerate(local)}
   merged.sort(key=lambda node:rank.get(node['data']['material_id'],len(rank)))
  return merged
 def refresh_external(self):
  if self.is_locked():self.load();self.notice.clear();self.apply_lock_state();return
  from copy import deepcopy
  fresh=self.database_tree(); conflicts=[]
  root=dict(self.svc.db.conn.execute('SELECT * FROM materials WHERE id=?',(self.root,)).fetchone())
  if any(root.get(k)!=self._root_baseline.get(k) for k in ('base_code','version','name','description','unit','grade','manufacturer','mpn','mp_flag')):self._external_paths.add(())
  self._root_baseline=root
  self.mark_external_differences(self._baseline_tree,fresh)
  local=self.refresh_drafts(self.capture_tree(),conflicts) if self.dirty else []
  merged=self.merge_refresh(self._baseline_tree,local,fresh,conflicts) if self.dirty else fresh
  def exclude_deleted(nodes,parent_id):
   result=[]
   for node in nodes:
    data=node['data']
    if (parent_id,data['material_id'],data.get('position','') or '') in self._deleted_relations:data['_deleted']=True
    node['children']=exclude_deleted(node['children'],data['material_id']); result.append(node)
   return result
  from services.bom_colors import colors
  root_color=(self.root_item.data(0,Qt.UserRole) or {}).get('row_color','') if getattr(self,'_root_color_dirty',False) else colors(self.svc).get((0,self.root),'')
  self.render_tree(exclude_deleted(merged,self.root));root_data=self.root_item.data(0,Qt.UserRole);root_data['row_color']=root_color;self.root_item.setData(0,Qt.UserRole,root_data); self._baseline_tree=deepcopy(fresh)
  self.notice.setText(''); self.statusBar().clearMessage(); self.refresh_action.setEnabled(False); self.save_action.setEnabled(True); self.export_action.setEnabled(False); self.refresh_styles()
  if conflicts:self.statusBar().showMessage('已加载外部更新；检测到冲突，已保留本地未保存修改，请核对后保存。')
 def load(self,preserve=None):
  self._root_color_dirty=False
  self.update_mirror_title()
  from copy import deepcopy
  fresh=self.database_tree(); self.render_tree(fresh); self._baseline_tree=deepcopy(fresh); self._root_baseline=dict(self.svc.db.conn.execute('SELECT * FROM materials WHERE id=?',(self.root,)).fetchone()); self._external_paths.clear(); self._draft_baselines.clear(); self.refresh_styles()
  self.save_action.setEnabled(False); self.refresh_action.setEnabled(False); self.export_action.setEnabled(True)
 def expansion_nodes(self):
  if not hasattr(self,'root_item'):return
  def walk(item,path):
   data=item.data(0,Qt.UserRole) or {}; key=path+'/'+str(data.get('material_id'))
   yield key,item
   for index in range(item.childCount()):yield from walk(item.child(index),key)
  yield from walk(self.root_item,'')
 def remember_expansion(self,*args):
  if self._rendering_tree:return
  import json
  self._expansion_state={key:item.isExpanded() for key,item in self.expansion_nodes()}
  self._expansion_settings.setValue(self._expansion_key,json.dumps(self._expansion_state));self._expansion_settings.sync()
 def render_tree(self,nodes):
  if hasattr(self,'root_item'):self.remember_expansion()
  self._rendering_tree=True
  self.current_item=None; self.tree.clear(); self.by={}; root=self.svc.db.conn.execute('SELECT * FROM materials WHERE id=?',(self.root,)).fetchone()
  self.setWindowTitle(f"{root['base_code']}{root['version']} {root['name']}");self.update_mirror_title()
  if self.is_locked():self.setWindowTitle(self.windowTitle()+' — BOM浏览')
  self.root_item=QTreeWidgetItem(self.tree,[f"{root['base_code']}{root['version']} {root['name']}",'','','','']); self.root_item.setData(0,Qt.UserRole,{'material_id':self.root,'root':True,'row_color':__import__('services.bom_colors',fromlist=['colors']).colors(self.svc).get((0,self.root),'')}); self.root_item.setExpanded(True)
  def append(node,parent,path,level):
   data=dict(node['data']); data.setdefault('_edit_base',{k:data.get(k) for k in ('quantity','position','remark')}); parent_data=parent.data(0,Qt.UserRole) or {}
   node_id=f"{path}/{data['id']}"
   data.update(node_id=node_id,parent_node_id=path or None,parent_material_id=parent_data['material_id'],parent_item_id=parent_data.get('id'),level_no=level)
   item=QTreeWidgetItem(parent,[f"{data['base_code']}{data['version']} {data['name']}",format_quantity(data['quantity']),data['unit'],data.get('position',''),data.get('remark','')]); item.setData(0,Qt.UserRole,data); self.by[node_id]=item
   for child in node['children']:append(child,item,node_id,level+1)
  for node in nodes:append(node,self.root_item,'',1)
  for key,item in self.expansion_nodes():
   if key in self._expansion_state:item.setExpanded(self._expansion_state[key])
  self._rendering_tree=False
  fit_columns(self.tree); self.tree.setCurrentItem(self.root_item); self.show_item_info(self.root_item)
 def snapshot_map(self,nodes,path=()):
  result={}
  for node in nodes:
   key=path+(node['data']['material_id'],); result[key]=node
   result.update(self.snapshot_map(node['children'],key))
  return result
 def remote_signature(self,node):
  fields=('material_id','quantity','position','remark','base_code','version','name','description','unit','grade','manufacturer','mpn','mp_flag','_db_alternates','row_color')
  return [node['data'].get(k,'') if k=='row_color' else node['data'].get(k) for k in fields]
 def refresh_styles(self):
  if not hasattr(self,'root_item'):return
  baseline=self.snapshot_map(getattr(self,'_baseline_tree',[]))
  def style(item,path,added=False,deleted=False,external=False):
   data=item.data(0,Qt.UserRole) or {}; root=item is self.root_item
   if not root:path=path+(data['material_id'],)
   external=external or path in self._external_paths
   added=added or data.get('_added',False); deleted=deleted or data.get('_deleted',False)
   if deleted and not data.get('_deleted'):
    data['_deleted']=True; item.setData(0,Qt.UserRole,data)
   draft=self._draft_baselines.get(data.get('id'))
   if draft is not None:
    baseline[path]=draft; baseline.update(self.snapshot_map(draft['children'],path))
   original=baseline.get(path,{}).get('data',data.get('_edit_base',{}))
   changed={column for field,column in [('quantity',1),('position',3),('remark',4)] if not root and data.get(field,'')!=original.get(field,'')}
   alternate_changed=self.alternate_key(item) in self._alternate_edits and self.alternate_ids(item)!=original.get('_db_alternates',[])
   if alternate_changed:changed.add(0)
   for column in range(self.tree.columnCount()):
    if column==2:item.setTextAlignment(column,Qt.AlignCenter)
    font=QFont(self.tree.font()); font.setStrikeOut(bool(deleted)); font.setBold(column in changed and not deleted)
    item.setFont(column,font)
    color=QColor('red') if deleted else QColor('green') if added else QColor('#9a6700') if changed else None
    item.setForeground(column,QBrush(color) if color else QBrush(QColor('white')) if external else QBrush())
    moved=data.get('_moved',False)
    if moved and not color:item.setForeground(column,QBrush(QColor('#202020')))
    from services.bom_colors import foreground
    custom=data.get('row_color','')
    # Five subtle gray overlays.  In light mode black overlays darken the
    # white base; in dark mode white overlays lighten the dark base by the
    # same amounts, so the perceived level separation remains consistent.
    alpha=(24,18,12,6,0)[min(len(path),4)]
    dark_theme=bool(QApplication.instance().property('bom_dark_theme'))
    level_tint=QColor(255,255,255,alpha) if dark_theme else QColor(0,0,0,alpha)
    background=QColor(custom) if custom and not (added or deleted or changed or moved or external) else level_tint
    if not (added or deleted or changed or moved or external):item.setForeground(column,QBrush(QColor(foreground(custom))) if custom else QBrush())
    item.setBackground(column,QBrush(QColor('#f8d7e3')) if moved else QBrush(QColor('#4a148c')) if external else QBrush(background))
   remark=data.get('remark','') if not root else ''
   display=remark
   if external:display=(remark+'\n' if remark else '')+'有其他BOM更新导入此变更'
   item.setText(4,display); item.setToolTip(4,remark)
   for n in range(item.childCount()):style(item.child(n),path,added,deleted,external)
  style(self.root_item,())
  fit_columns(self.tree)
 def locate_material(self,material_id,parent_id=None):
  # 仅展开目标的祖先路径，保留已有窗口中的编辑缓存和其他展开状态。
  for _,item in self.expansion_nodes():
   data=item.data(0,Qt.UserRole) or {}
   if data.get('material_id',self.root if data.get('root') else None)!=material_id:continue
   if parent_id is not None and data.get('parent_material_id')!=parent_id:continue
   if data.get('_deleted'):continue
   ancestor=item.parent()
   while ancestor is not None:
    ancestor.setExpanded(True); ancestor=ancestor.parent()
   self.tree.setCurrentItem(item); self.show_item_info(item); self.tree.scrollToItem(item,QAbstractItemView.PositionAtCenter)
   return True
  return False
 def show_item_info(self,item,column=0):
  self.current_item=item
  d=item.data(0,Qt.UserRole) or {}
  if d.get('root'):
   r=self.svc.db.conn.execute('SELECT * FROM materials WHERE id=?',(self.root,)).fetchone(); d=dict(r) if r else d
  self._loading_fields=True
  editable=not ((item.data(0,Qt.UserRole) or {}).get('root') or (item.data(0,Qt.UserRole) or {}).get('_deleted'))
  for field in ('quantity','position','remark'):self.fields[field].setEnabled(editable and not self.is_locked())
  for k in self.fields: self.fields[k].setText(str(d.get(k,'') or ''))
  self._loading_fields=False
  self.alt_table.setRowCount(0)
  if not (item.data(0,Qt.UserRole) or {}).get('root'):
   materials={row['id']:row for row in self.svc.materials()}
   for material_id in self.alternate_ids(item):
    r=materials.get(material_id)
    if r is None:continue
    row=self.alt_table.rowCount(); self.alt_table.insertRow(row)
    for j,k in enumerate(['base_code','version','name','description','manufacturer','mpn']):
     cell=QTableWidgetItem(str(r[k] or '')); cell.setData(Qt.UserRole,material_id); self.alt_table.setItem(row,j,cell)
  fit_columns(self.alt_table)
 def update_tree_field(self,key,value):
  if self.is_locked():return False
  if self._loading_fields or self.current_item is None or self.current_item is getattr(self,'root_item',None): return
  if (self.current_item.data(0,Qt.UserRole) or {}).get('_deleted'):return
  value=format_quantity(value) if key=='quantity' else value
  col={'quantity':1,'position':3,'remark':4}[key]; self.current_item.setText(col,value); d=self.current_item.data(0,Qt.UserRole) or {}; d[key]=float(value) if key=='quantity' and value else value; self.current_item.setData(0,Qt.UserRole,d); self.mark_dirty()
 def add_relative(self,parent,child):
  if self.is_locked():return False
  target=self.by.get(parent.get('node_id'),self.root_item) if child else (self.by.get(parent.get('parent_node_id'),self.root_item) or self.root_item)
  blocked=self.replacement_blocked_materials(target,True)
  rows=[r for r in self.svc.materials() if r['id'] not in blocked]; items=[f"{r['base_code']}{r['version']} {r['name']}" for r in rows]
  if not items: QMessageBox.information(self,'提示','当前父级下已没有可新增的物料'); return
  x,ok=MaterialSelectionDialog.getItem(self,'新建BOM物料','物料',items,0,False)
  if not ok:return
  selected=self.by.get(parent.get('node_id')) if not child else None
  self.insert_material(target,rows[items.index(x)],after=selected)
 @preserve_bom_view
 def insert_material(self,target,r,after=None):
  if self.is_locked():return False
  from services.current_catalog import enabled as catalog_enabled,selection_tree
  current_catalog=catalog_enabled(self.svc)
  r=self.svc.material_snapshot(r['id'])
  latest_children=selection_tree(self.svc,r['id']) if current_catalog else None
  if not current_catalog and hasattr(self.svc,'ensure_material_bom'):self.svc.ensure_material_bom(r['id'])
  parent=target.data(0,Qt.UserRole) or {}
  self._deleted_relations.discard((parent.get('material_id'),r['id'],''))
  item=QTreeWidgetItem(target,[f"{r['base_code']}{r['version']} {r['name']}",'1',r['unit'],'',''])
  if after is not None and after.parent() is target:
   target.takeChild(target.indexOfChild(item));target.insertChild(target.indexOfChild(after)+1,item)
  temporary_id=-id(item); data=dict(r)
  data.update(_added=True,id=temporary_id,material_id=r['id'],quantity=1.0,position='',remark='',level_no=(parent.get('level_no') or 0)+1,parent_item_id=parent.get('id'),parent_material_id=parent.get('material_id'),node_id=str(temporary_id),parent_node_id=parent.get('node_id'))
  item.setData(0,Qt.UserRole,data); self.by[str(temporary_id)]=item; target.setExpanded(True)
  if current_catalog:self.clone_version_children(item,latest_children)
  else:self.add_existing_children(item,r['id'],set())
  self._draft_baselines[temporary_id]=self.capture_item(item)
  fit_columns(self.tree); self.mark_dirty(); self.tree.setCurrentItem(item); self.show_item_info(item)
  return item
 def add_existing_children(self,parent_item,material_id,ancestors):
  if material_id in ancestors:return
  ancestors=set(ancestors); ancestors.add(material_id)
  rows=self.svc.db.conn.execute('SELECT b.*,m.base_code,m.version,m.name,m.description,m.unit AS material_unit,m.grade,m.manufacturer,m.mpn,m.mp_flag FROM material_bom_items b JOIN materials m ON m.id=b.child_material_id WHERE b.parent_material_id=? ORDER BY b.sort_no,b.id',(material_id,)).fetchall()
  parent_data=parent_item.data(0,Qt.UserRole) or {}
  for row in rows:
   if row['child_material_id'] in ancestors:continue
   child=QTreeWidgetItem(parent_item,[f"{row['base_code']}{row['version']} {row['name']}",format_quantity(row['quantity']),row['material_unit'],row['position'],row['remark']]); d=dict(row); d['_added']=parent_data.get('_added',False); d['_db_alternates']=[a['material_id'] for a in self.svc.alternates(row['id'])]; d.update({'material_id':row['child_material_id'],'unit':row['material_unit'],'parent_item_id':parent_data.get('id'),'level_no':(parent_data.get('level_no') or 0)+1}); d.update(node_id=f"{parent_data.get('node_id')}/{row['id']}",parent_node_id=parent_data.get('node_id')); child.setData(0,Qt.UserRole,d); self.by[d['node_id']]=child; self.add_existing_children(child,row['child_material_id'],ancestors)
  parent_item.setExpanded(True)
 def alternate_key(self,item):
  data=item.data(0,Qt.UserRole) or {}
  return data.get('id') or data.get('node_id')
 def alternate_ids(self,item):
  key=self.alternate_key(item)
  if key in self._alternate_edits:return list(self._alternate_edits[key])
  return [row['material_id'] for row in self.svc.alternates(key)] if isinstance(key,int) and key>0 else []
 def add_alternate(self):
  if self.is_locked():return False
  item=self.tree.currentItem()
  if item is None or item is self.root_item or (item.data(0,Qt.UserRole) or {}).get('_deleted'):return
  existing=self.alternate_ids(item); data=item.data(0,Qt.UserRole) or {}
  rows=[row for row in self.svc.db.conn.execute('SELECT id,base_code,version,name FROM materials ORDER BY base_code,version') if row['id']!=data.get('material_id') and row['id'] not in existing]
  if not rows: QMessageBox.information(self,'新增替代料','没有可选替代料'); return
  choices=[f"{row['base_code']}{row['version']} {row['name']}" for row in rows]
  choice,ok=MaterialSelectionDialog.getItem(self,'新增替代料','选择物料',choices,0,False)
  if not ok:return
  self._alternate_edits[self.alternate_key(item)]=existing+[rows[choices.index(choice)]['id']]
  self.mark_dirty(); self.show_item_info(item)
 def alternate_menu(self,pos):
  if self.is_locked():return False
  row=self.alt_table.rowAt(pos.y()); item=self.current_item
  if row<0 or item is None or item is self.root_item or (item.data(0,Qt.UserRole) or {}).get('_deleted'):return
  material_id=self.alt_table.item(row,0).data(Qt.UserRole)
  menu=QMenu(self); promote=menu.addAction('切换成主物料');menu.addSeparator();remove=menu.addAction('删除替代料')
  reason=self.paste_block_reason(item,False,material_id);promote.setEnabled(not reason);promote.setToolTip(reason.replace('粘贴','切换') if reason else '与主物料互换，保留数量、位置、备注及其他替代料')
  action=menu.exec_(self.alt_table.viewport().mapToGlobal(pos))
  if action==promote:
   if promote.isEnabled():self.promote_alternate(item,material_id)
   return
  if action!=remove:return
  if QMessageBox.question(self,'确认','删除当前替代料？')!=QMessageBox.Yes:return
  material_id=self.alt_table.item(row,0).data(Qt.UserRole)
  self._alternate_edits[self.alternate_key(item)]=[mid for mid in self.alternate_ids(item) if mid!=material_id]
  self.mark_dirty(); self.show_item_info(item)
 def promote_alternate(self,item,material_id):
  if self.is_locked():return False
  if item is None or item is self.root_item or material_id not in self.alternate_ids(item):return
  key=self.relation_key(item)
  matches=[node for _,node in self.expansion_nodes() if node is not self.root_item and not (node.data(0,Qt.UserRole) or {}).get('_deleted') and self.relation_key(node)==key]
  if not matches:return
  for node in matches:
   reason=self.paste_block_reason(node,False,material_id)
   if reason:QMessageBox.warning(self,'无法切换成主物料',reason.replace('粘贴','切换'));return
  material=next((row for row in self.svc.materials() if row['id']==material_id),None)
  if material is None:return
  return self.replace_nodes(item,matches,material,swap_alternate=True)
 def import_bom(self,selected_item=None):
  if self.is_locked():return False
  if self.dirty and QMessageBox.question(self,'BOM未保存','导入成功后将替换当前编辑缓存，是否放弃当前未保存编辑并继续？')!=QMessageBox.Yes:return
  config_service=self.svc.config_service or ConfigService(Path(self.svc.db.path).parent)
  p=QFileDialog.getOpenFileName(self,'导入BOM',bom_import_directory(config_service),'Excel (*.xlsx)')[0]
  if not p:return
  try:
   p=remember_bom_import(config_service,p)
   sheet=choose_material_sheet(p,self,'BOM')
   from openpyxl import load_workbook
   wb=load_workbook(p,read_only=True,data_only=True)
   try:headers=[str(value if value is not None else '').strip() for value in next(wb[sheet].iter_rows(values_only=True),())]
   finally:wb.close()
   dialog=MaterialHeaderMappingDialog(p,headers,[h for h in self.svc.BOM_HEADERS if h not in headers],self,sheet_name=sheet,required=self.svc.BOM_HEADERS,config_service=self.svc.config_service or ConfigService(Path(self.svc.db.path).parent))
   if dialog.exec_()!=QDialog.Accepted:raise ValueError('用户取消BOM表头映射，未导入')
   target_root=(selected_item.data(0,Qt.UserRole) or {}).get('material_id') if selected_item is not None else self.root
   def resolve_conflicts(conflicts):
    decision=BomImportConflictDialog(conflicts,self);decision.exec_();return decision.choice
   nodes=self.svc.import_bom(p,target_root,sheet_name=sheet,header_mapping=dialog.mapping(),persist=False,conflict_resolver=resolve_conflicts)
   self.stage_import(nodes) if target_root==self.root else self.stage_import_subtree(selected_item,nodes)
   QMessageBox.information(self,'导入完成','BOM已导入编辑缓存，请检查后点击“保存”写入数据库。')
  except BomImportError as error:BomImportIssuesDialog(error,self).exec_()
  except Exception as error:show_operation_error(self,'导入失败',error)
 def stage_import(self,nodes):
  if self.is_locked():return False
  from copy import deepcopy
  baseline=self.database_tree()
  self._deleted_relations.clear(); self._alternate_edits.clear(); self._draft_baselines.clear(); self._external_paths.clear()
  def compare(old,new,parent):
   old_by={n['data']['material_id']:n for n in old}
   new_ids={n['data']['material_id'] for n in new}
   for node in new:
    data=node['data']; previous=old_by.get(data['material_id'])
    if 'alternate_material_ids' in data:self._alternate_edits[data['id']]=list(data['alternate_material_ids'])
    if previous is None:
     data['_added']=True
     if data['id']<0:self._draft_baselines[data['id']]=deepcopy(node)
    compare(previous['children'] if previous else [],node['children'],data['material_id'])
   for node in old:
    if node['data']['material_id'] not in new_ids:
     deleted=deepcopy(node); deleted['data']['_deleted']=True; new.append(deleted)
     data=deleted['data']; self._deleted_relations.add((parent,data['material_id'],data.get('position','') or ''))
  compare(baseline,nodes,self.root)
  self._baseline_tree=deepcopy(baseline)
  self.render_tree(nodes); self.mark_dirty()
  self.notice.clear(); self.refresh_action.setEnabled(False); self.statusBar().clearMessage()
 def export_purchase_bom(self):
  if not self.purchase_export_action.isEnabled():return
  material=self.svc.material_snapshot(self.root)
  filename=material['base_code']+material['version']+' '+material['name']+self.mirror_suffix()+' 采购BOM'
  p=QFileDialog.getSaveFileName(self,'采购BOM导出',export_start_path(export_filename(filename)),'Excel (*.xlsx)')[0]
  if p:
   remember_export_path(p);self.svc.export_purchase_bom(self.root,p);QMessageBox.information(self,'完成','采购BOM导出成功')
 def export_bom(self):
  if not self.export_action.isEnabled():return
  export_material_structure(self,self.svc,self.root)
 def add(self):
  self.add_relative({'id':None,'parent_item_id':None,'level_no':0},False)
 def stage_import_subtree(self,item,nodes):
  if self.is_locked():return False
  from copy import deepcopy
  while item.childCount():item.takeChild(0)
  data=item.data(0,Qt.UserRole) or {}
  def append(node,parent,path,level):
   d=dict(node['data']); d['_added']=d.get('_added',True); d.update(parent_material_id=data.get('material_id'),parent_item_id=data.get('id'),node_id=f'{path}/{d["id"]}',parent_node_id=path,level_no=level)
   child=QTreeWidgetItem(parent,[f"{d['base_code']}{d['version']} {d['name']}",format_quantity(d['quantity']),d['unit'],d.get('position',''),d.get('remark','')]);child.setData(0,Qt.UserRole,d)
   for grand in node['children']:append(grand,child,d['node_id'],level+1)
  for node in nodes:append(node,item,data.get('node_id',''),(data.get('level_no') or 0)+1)
  item.setExpanded(True);self.mark_dirty();self.tree.setCurrentItem(item);self.show_item_info(item)
 def save_context(self,rows):
  sid=getattr(self.svc,'sku_id',None)
  sku=self.svc.db.conn.execute('SELECT material_id FROM skus WHERE id=?',(sid,)).fetchone() if sid else None
  if sku is None or sku['material_id']==self.root:return self.root,rows
  from services.current_catalog import groups,rows_for,reachable
  graph=groups(self.svc)
  if self.root not in reachable(graph,sku['material_id']):raise ValueError('此物料已不在所属BOM中，请关闭并从物料库重新打开。')
  retained={self.root}|{r['material_id'] for r in rows}
  parents={self.root}|{r['parent_material_id'] for r in self.svc.bom(self.root) if r['parent_material_id'] in retained}|{r['material_id'] for r in rows if r.get('replace_children')}
  updates={mid:[] for mid in parents}
  for original in rows:
   row=dict(original);parent=row['parent_material_id']
   if 'alternate_material_ids' not in row:
    row['alternate_material_ids']=[a['material_id'] for a in self.svc.alternates(row['id'])] if isinstance(row.get('id'),int) and row['id']>0 else []
   children=updates.setdefault(parent,[])
   if not any(r['material_id']==row['material_id'] and r.get('position','')==row.get('position','') for r in children):children.append(row)
  graph.update(updates)
  return sku['material_id'],rows_for(graph,sku['material_id'])
 def save(self):
  if self.is_locked():return False
  if self._saving:return False
  progress=QProgressDialog('正在整理BOM…',None,0,100,self)
  progress.setWindowTitle('保存BOM');progress.setWindowModality(Qt.NonModal);progress.setMinimumDuration(0);progress.setAutoClose(False);progress.show()
  def report(value,message):
   progress.setLabelText(message);progress.setValue(value)
   QApplication.processEvents(QEventLoop.ExcludeUserInputEvents)
  report(0,'正在整理BOM…')
  self._saving=True
  try:
   rows=[]
   def walk(it,parent_material):
    d=dict(it.data(0,Qt.UserRole) or {})
    if d.get('_deleted'):return
    d.update(quantity=float(it.text(1)),position=it.text(3),remark=d.get('remark',''),parent_material_id=parent_material,replace_children=True)
    if self.alternate_key(it) in self._alternate_edits:d['alternate_material_ids']=self.alternate_ids(it)
    rows.append(d)
    for j in range(it.childCount()): walk(it.child(j),d['material_id'])
   for j in range(self.root_item.childCount()): walk(self.root_item.child(j),self.root)
   save_root,rows=self.save_context(rows)
   from services.snapshot_bom import revision_conflicts,save_with_revisions
   from services.current_catalog import synchronization_plan,save_confirmed
   plan=synchronization_plan(self.svc,save_root,rows);selections={};versions={}
   if plan:
    progress.hide()
    from ui.bom_sync_confirmation import BomSyncConfirmationDialog
    confirmation=BomSyncConfirmationDialog(plan,self)
    if confirmation.exec_()!=QDialog.Accepted:return False
    selections=confirmation.selections()
   partial=[p for p in plan if selections.get(p['material_id'],set())!={s['id'] for s in p['other_skus']}]
   if any(self.svc.db.conn.execute('SELECT 1 FROM skus WHERE material_id=?',(p['material_id'],)).fetchone() for p in partial):
    QMessageBox.warning(self,'SKU版本需手动更新','未全部同步的父物料属于SKU，请先通过SKU版本变更手动升版，再保存此变更。编辑缓存已保留。');return False
   approved={p['material_id'] for p in plan if p not in partial}
   violations={v['material_id']:v for v in revision_conflicts(self.svc,save_root,rows) if v['material_id'] not in approved}
   violations.update({p['material_id']:p for p in partial})
   if violations:
    progress.hide()
    from ui.parent_revision import ParentRevisionDialog
    chooser=ParentRevisionDialog(self.svc,list(violations.values()),self)
    if chooser.exec_()!=QDialog.Accepted:return False
    versions=chooser.versions()
   progress.show()
   callback=lambda n,text:report(int(n*0.85),text)
   with self.svc.db.transaction():
    if plan:revisions=save_confirmed(self.svc,save_root,rows,plan,selections,versions,callback)
    elif violations:revisions=save_with_revisions(self.svc,save_root,rows,versions,progress=callback)
    else:self.svc.save_bom(save_root,rows,progress=callback);revisions={}
    from services.bom_colors import save as save_colors
    save_colors(self.svc,[(0,revisions.get(self.root,self.root),(self.root_item.data(0,Qt.UserRole) or {}).get('row_color',''))])
   self.root=revisions.get(self.root,self.root)
   report(90,'正在刷新BOM显示…')
   self.dirty=False; self._deleted_relations.clear(); self.export_action.setEnabled(True); self._alternate_edits.clear(); self.load(); self.notice.setText(''); self.statusBar().clearMessage(); self.save_action.setEnabled(False); self.refresh_action.setEnabled(False)
   report(100,'BOM保存完成')
  except Exception as error:
   from PyQt5 import sip
   if sip.isdeleted(self):return False
   self.mark_dirty()
   show_operation_error(self,'保存失败',error,'BOM未保存，编辑内容已保留。\n')
   return False
  finally:
   self._saving=False
   from PyQt5 import sip
   if not sip.isdeleted(progress):progress.close();progress.deleteLater()
  if sip.isdeleted(self):return False
  QMessageBox.information(self,'保存','BOM已保存')
  return True

class MaterialColumnFilterDialog(QDialog):
 def __init__(self,title,values,condition=None,parent=None):
  super().__init__(parent); self.setWindowTitle('筛选：'+title); self.resize(420,460)
  layout=QVBoxLayout(self); self.query=QLineEdit(); self.query.setPlaceholderText('包含文字（可留空，不区分大小写）'); layout.addWidget(self.query)
  self.values=sorted(set(values)); self.list=QListWidget(); layout.addWidget(self.list)
  query,selected=condition or ('',None); self.query.setText(query)
  for value in self.values:
   item=QListWidgetItem(value or '（空白）',self.list); item.setData(Qt.UserRole,value); item.setCheckState(Qt.Checked if selected is None or value in selected else Qt.Unchecked)
  actions=QHBoxLayout(); all_button=QPushButton('全选'); none_button=QPushButton('全不选'); actions.addWidget(all_button); actions.addWidget(none_button); layout.addLayout(actions)
  all_button.clicked.connect(lambda:self.set_checks(Qt.Checked)); none_button.clicked.connect(lambda:self.set_checks(Qt.Unchecked))
  buttons=QDialogButtonBox(QDialogButtonBox.Ok|QDialogButtonBox.Cancel); buttons.accepted.connect(self.accept); buttons.rejected.connect(self.reject); layout.addWidget(buttons)
 def set_checks(self,state):
  for index in range(self.list.count()):self.list.item(index).setCheckState(state)
 def condition(self):
  selected={self.list.item(i).data(Qt.UserRole) for i in range(self.list.count()) if self.list.item(i).checkState()==Qt.Checked}
  return self.query.text().strip(),None if selected==set(self.values) else selected

class MaterialWindow(QMainWindow):
 closed=pyqtSignal()
 def __init__(self,svc,parent=None):
  super().__init__(parent); self.setAttribute(Qt.WA_DeleteOnClose); self.svc=svc; self._filters={}; self._local_bom_windows={}; WindowRefreshBus.subscribe(self.on_event); self.setWindowTitle('物料库')
  menu=self.menuBar(); a=menu.addAction('新建物料'); a.setShortcut('Ctrl+N');a.setShortcutContext(Qt.WindowShortcut); a.triggered.connect(parent.new_material if parent else self.new_material); a=menu.addAction('导入物料'); a.triggered.connect(parent.import_materials if parent else self.import_materials); a=menu.addAction('导出物料库'); a.triggered.connect(parent.export_materials if parent else self.export_materials); a=menu.addAction('导出模板'); a.triggered.connect(parent.template if parent else self.template); a=menu.addAction('查找物料'); a.setShortcut('Ctrl+F');a.setShortcutContext(Qt.WindowShortcut); a.triggered.connect(self.search_materials); self.resize(1000,600); self.t=FrozenMaterialTable(); self.t.setHorizontalScrollBarPolicy(Qt.ScrollBarAlwaysOn); self.t.setVerticalScrollBarPolicy(Qt.ScrollBarAlwaysOn); self.t.setEditTriggers(QAbstractItemView.NoEditTriggers); self.t.setColumnCount(11); self.t.setHorizontalHeaderLabels(['基础编码','版本号','名称','规格描述','单位','物料等级','制造商','MPN','M/P','是否存在BOM结构','是否为呆滞料']); self.t.setWordWrap(True); self.setCentralWidget(self.t); self.t.setContextMenuPolicy(Qt.CustomContextMenu); self.t.customContextMenuRequested.connect(self.menu); self.t.cellDoubleClicked.connect(self.edit_current)
  self._column_labels=[self.t.horizontalHeaderItem(i).text() for i in range(self.t.columnCount())]
  from services.snapshot_bom import enabled
  if enabled(self.svc.db):self._column_labels.append('是否被最新历史版本使用')
  self._base_column_labels=list(self._column_labels);self._sku_column_ids=[];self._hidden_skus=set()
  self._column_config=self.svc.config_service or ConfigService(Path(self.svc.db.path).parent)
  self._selected_skus=self._column_config.sku_columns(self.svc.db.path)
  self.sku_columns_button=QToolButton();self.sku_columns_button.setText('SKU用量列');self.sku_columns_button.clicked.connect(self.choose_sku_columns);menu.setCornerWidget(self.sku_columns_button,Qt.TopRightCorner)
  header=self.t.horizontalHeader(); header.setSectionsClickable(True); header.sectionClicked.connect(self.filter_menu); self.refresh()
 def choose_sku_columns(self):
  dialog=QDialog(self);dialog.setWindowTitle('选择显示的SKU用量列');dialog.resize(600,480);layout=QVBoxLayout(dialog)
  choices=QListWidget();layout.addWidget(choices)
  for sid,label in zip(self._sku_column_ids,self._column_labels[len(self._base_column_labels):]):
   item=QListWidgetItem(label.replace('\n',' '),choices);item.setData(Qt.UserRole,sid);item.setCheckState(Qt.Unchecked if sid in self._hidden_skus else Qt.Checked)
  row=QHBoxLayout()
  for text,state in [('全选',Qt.Checked),('全不选',Qt.Unchecked)]:
   button=QPushButton(text);button.clicked.connect(lambda _,value=state:[choices.item(i).setCheckState(value) for i in range(choices.count())]);row.addWidget(button)
  layout.addLayout(row);buttons=QDialogButtonBox(QDialogButtonBox.Ok|QDialogButtonBox.Cancel);buttons.accepted.connect(dialog.accept);buttons.rejected.connect(dialog.reject);layout.addWidget(buttons)
  if dialog.exec_()==QDialog.Accepted:
   selected={choices.item(i).data(Qt.UserRole) for i in range(choices.count()) if choices.item(i).checkState()==Qt.Checked}
   try:self._column_config.save_sku_columns(self.svc.db.path,selected)
   except (OSError,ValueError) as error:show_operation_error(self,'保存失败',error);return
   self._selected_skus=selected;self.refresh()
 def search_materials(self):
  if not hasattr(self,'_search_dialog'):self._search_dialog=MaterialSearchDialog(self)
  self._search_dialog.show();self._search_dialog.raise_();self._search_dialog.activateWindow()
 def locate_material(self,material_id):
  self.refresh()
  for row in range(self.t.rowCount()):
   if self.t.item(row,0).data(Qt.UserRole)!=material_id:continue
   if self.t.isRowHidden(row):self._filters.clear();self.apply_filters()
   self.t.setCurrentCell(row,0);self.t.selectRow(row);self.t.scrollToItem(self.t.item(row,0),QAbstractItemView.PositionAtCenter)
   return True
  return False
 def filter_menu(self,column):
  menu=QMenu(self); choose=menu.addAction('筛选此列…');menu.addSeparator(); clear=menu.addAction('清除此列筛选'); clear.setEnabled(column in self._filters); clear_all=menu.addAction('清除全部筛选'); clear_all.setEnabled(bool(self._filters))
  header=self.t.horizontalHeader(); from PyQt5.QtCore import QPoint
  action=menu.exec_(header.mapToGlobal(QPoint(header.sectionViewportPosition(column),header.height())))
  if action==choose:
   values=[self.t.item(row,column).text() for row in range(self.t.rowCount())]
   dialog=MaterialColumnFilterDialog(self._column_labels[column],values,self._filters.get(column),self)
   if dialog.exec_()==QDialog.Accepted:
    condition=dialog.condition()
    if condition==('',None):self._filters.pop(column,None)
    else:self._filters[column]=condition
    self.apply_filters()
  elif action==clear:self._filters.pop(column,None); self.apply_filters()
  elif action==clear_all:self._filters.clear(); self.apply_filters()
 def apply_filters(self):
  visible=0
  for row in range(self.t.rowCount()):
   match=True
   for column,(query,selected) in self._filters.items():
    value=self.t.item(row,column).text()
    if query.casefold() not in value.casefold() or (selected is not None and value not in selected):match=False; break
   self.t.setRowHidden(row,not match); visible+=int(match)
  for column,label in enumerate(self._column_labels):
   item=self.t.horizontalHeaderItem(column); item.setText(label+(' [筛选] ▾' if column in self._filters else ' ▾')); item.setToolTip((label+'\n' if column>=len(self._base_column_labels) else '')+'点击筛选；多列条件同时生效')
  self.t.sync_frozen()
  self.statusBar().showMessage(f'显示 {visible} / {self.t.rowCount()} 个物料')
 def edit_current(self,row,column):
  rs=self.svc.materials();
  if row<0 or row>=len(rs): return
  r=rs[row]; dlg=MaterialDialog(self,{k:r[k] for k in ['base_code','version','name','description','unit','grade','manufacturer','mpn','mp_flag']},readonly_keys={'base_code','version'})
  if dlg.exec_():
   try:self.svc.update_material(r['id'],dlg.data()); self.refresh()
   except Exception as e: show_operation_error(self,'保存失败',e)
 def menu(self,pos):
  row=self.t.rowAt(pos.y())
  if row<0:return
  material=self.svc.materials()[row]; rid=material['id']
  m=QMenu(self);e=m.addAction('编辑物料');v=m.addAction('物料版本变更');v.setEnabled(not any(sku['material_id']==rid for sku in self.svc.skus()))
  if material['has_bom']=='Y' or material['idle']=='N':m.addSeparator()
  edit_bom=m.addAction('编辑物料BOM结构') if material['has_bom']=='Y' else None
  export=m.addAction('导出BOM结构') if material['has_bom']=='Y' else None
  reverse=m.addAction('物料反查') if material['idle']=='N' else None
  m.addSeparator();similar=m.addAction('类似新增物料');copy=m.addAction('复制物料')
  m.addSeparator();d=m.addAction('删除物料');d.setEnabled(material['idle']=='Y')
  a=m.exec_(self.t.viewport().mapToGlobal(pos))
  if a==similar:
   show_create_material(self,data=dict(self.svc.material_snapshot(rid)));return
  if a==copy:copy_material(self.svc,rid); return
  if a is None:return
  if export is not None and a==export:
   from services.snapshot_bom import enabled,template_service
   source=template_service(self.svc,rid) if enabled(self.svc.db) else self.svc
   export_material_structure(self,source,rid);return
  if edit_bom is not None and a==edit_bom:
   self.open_material_bom(rid); return
  if reverse is not None and a==reverse:
   MaterialUsageDialog(material,self.svc.material_usage(rid),self).exec_(); return
  if a==v and v.isEnabled():
   self.change_version(rid)
  elif a==e:
   r=self.svc.materials()[row]; dlg=MaterialDialog(self,{k:r[k] for k in ['base_code','version','name','description','unit','grade','manufacturer','mpn','mp_flag']},readonly_keys={'base_code','version'})
   if dlg.exec_():
    try:self.svc.update_material(rid,dlg.data()); self.refresh()
    except Exception as error:show_operation_error(self,'保存失败',error)
  elif a==d and d.isEnabled() and QMessageBox.question(self,'确认','删除物料？')==QMessageBox.Yes:
   try:self.svc.delete_material(rid); self.refresh()
   except Exception as ex: show_operation_error(self,'无法删除',ex)
 def open_material_bom(self,material_id):
  main=self.parent()
  while main is not None and not isinstance(main,MainWindow): main=main.parent()
  if main is not None:return main.open_material_bom(material_id)
  if material_id not in self._local_bom_windows:
   from services.snapshot_bom import enabled,template_service
   source=template_service(self.svc,material_id) if enabled(self.svc.db) else self.svc
   window=BomWindow(source,material_id,self); self._local_bom_windows[material_id]=window
   window.closed.connect(lambda:self._local_bom_windows.pop(material_id,None))
  window=self._local_bom_windows[material_id]
  if window.isMinimized():window.showNormal()
  else:window.show()
  window.raise_(); window.activateWindow(); return window
 def change_version(self,material_id):
  if any(sku['material_id']==material_id for sku in self.svc.skus()):
   QMessageBox.information(self,'SKU版本变更','请在SKU窗口中进行版本变更。'); return
  try:
   references=self.svc.material_version_references(material_id)
   old=next(row for row in self.svc.materials() if row['id']==material_id)
   if references:
    dialog=MaterialVersionDialog(old,references,self)
    if dialog.exec_()!=QDialog.Accepted:return
    version=dialog.version.text(); selected=dialog.selected_parents()
   else:
    version,ok=QInputDialog.getText(self,'物料版本变更','新版本号',text=old['version'])
    if not ok:return
    selected=set()
   self.svc.change_material_version(material_id,version,selected,{r['parent_id'] for r in references})
   self.refresh()
   if self.parent() is not None and hasattr(self.parent(),'refresh'): self.parent().refresh()
  except Exception as error: show_operation_error(self,'版本变更失败',error)
 def new_material(self):
  return show_create_material(self)
 def import_materials(self):
  p=QFileDialog.getOpenFileName(self,'导入物料清单','','Excel (*.xlsx)')[0]
  if p:
   try:
    sheet_name=choose_material_sheet(p,self)
    config_service=self.svc.config_service or ConfigService(Path(self.svc.db.path).parent)
    try:ok,fail=self.svc.import_materials(p,sheet_name=sheet_name,header_mapping=config_service.import_mapping('materials') or None)
    except MaterialHeaderError as error:
     dialog=MaterialHeaderMappingDialog(p,error.headers,error.missing,self,sheet_name=sheet_name,config_service=config_service)
     if dialog.exec_()!=QDialog.Accepted:raise ValueError('用户取消表头映射，未导入任何物料')
     ok,fail=self.svc.import_materials(p,header_mapping=dialog.mapping(),sheet_name=sheet_name)
    self.refresh(); MaterialImportResultDialog(ok,fail,self).exec_()
   except Exception as error:show_operation_error(self,'导入失败',error)
 def export_materials(self):
  p=QFileDialog.getSaveFileName(self,'导出物料库',export_start_path(export_filename((self.svc.project()['project_code']+' '+self.svc.project()['project_name']+' 物料库').strip())),'Excel (*.xlsx)')[0]
  if p:remember_export_path(p); self.svc.export_materials(p)
 def template(self):
  p=QFileDialog.getSaveFileName(self,'导出模板',export_start_path('物料模板.xlsx'),'Excel (*.xlsx)')[0]
  if p:
   remember_export_path(p)
   from openpyxl import Workbook; w=Workbook(); w.active.append(MATERIAL_IMPORT_HEADERS); w.save(p); w.close()
 def on_event(self,event):
  if event.get('type') in ('material_changed','bom_changed','sku_changed','sku_order_changed'): self.refresh()
 def closeEvent(self,event):
  for window in list(self._local_bom_windows.values()):
   if not window.close(): event.ignore(); return
  if getattr(self,'_create_material_dialog',None):self._create_material_dialog.reject()
  WindowRefreshBus.unsubscribe(self.on_event); self.closed.emit(); WindowRefreshBus.publish({'type':'bom_draft_changed'}); event.accept()
 def refresh(self):
  from decimal import Decimal,ROUND_HALF_UP
  project=self.svc.project();self.setWindowTitle(' '.join(str(project[key] or '').strip() for key in ('project_code','project_name') if str(project[key] or '').strip())+' 物料库')
  skus=list(self.svc.skus());order={sid:i for i,sid in enumerate(getattr(self.svc,'sku_display_order',[]))}
  skus.sort(key=lambda sku:order.get(sku['id'],len(order)+sku['id']))
  ids=[sku['id'] for sku in skus];base=len(self._base_column_labels)
  self.t.horizontalHeader().first_sku_column=base
  # 按SKU身份迁移筛选条件，排序或删除列不能把条件套到另一个SKU。
  remembered={self._sku_column_ids[col-base]:condition for col,condition in self._filters.items() if base<=col<base+len(self._sku_column_ids)}
  self._filters={col:condition for col,condition in self._filters.items() if col<base}
  self._filters.update({base+i:remembered[sid] for i,sid in enumerate(ids) if sid in remembered})
  self._sku_column_ids=ids
  self._column_labels=self._base_column_labels+[sku['base_code']+sku['version']+(' [历史版本'+str(sku['mirror_version'])+']' if sum(x['material_id']==sku['material_id'] for x in skus)>1 else '')+'\n'+sku['name'] for sku in skus]
  self.t.setColumnCount(len(self._column_labels));self.t.setHorizontalHeaderLabels(self._column_labels)
  self._hidden_skus=set(ids)-self._selected_skus
  for i,sid in enumerate(ids):
   self.t.setColumnHidden(base+i,sid in self._hidden_skus)
   if sid in self._hidden_skus:self._filters.pop(base+i,None)
  self.t.horizontalHeader().setMinimumHeight(self.t.fontMetrics().height()*2+16)
  quantities=self.svc.sku_material_quantities(skus)
  rs=self.svc.materials(); self.t.setRowCount(len(rs))
  for i,r in enumerate(rs):
   for j,v in enumerate([r['base_code'],r['version'],r['name'],r['description'],r['unit'],r['grade'],r['manufacturer'],r['mpn'],r['mp_flag'],r['has_bom'],r['idle']]+([r['latest_used']] if 'latest_used' in r.keys() else [])+[format_quantity(quantities[sku['id']][r['id']]) if r['id'] in quantities[sku['id']] else 'NA' for sku in skus]):
    item=QTableWidgetItem(str(v or '')); item.setData(Qt.UserRole,r['id'])
    if j>=base or self._column_labels[j] in ('版本号','单位','M/P','是否存在BOM结构','是否为呆滞料','是否被最新历史版本使用'):item.setTextAlignment(Qt.AlignCenter)
    if j>=base:
     item.setBackground(QColor('#eeeeee' if v=='NA' else '#d9ead3'));item.setForeground(QColor('black'))
    if r['idle']=='Y': item.setBackground(QColor('red')); item.setForeground(QColor('black'))
    self.t.setItem(i,j,item)
  self.apply_filters(); fit_columns(self.t)
  # SKU名称不参与列宽；固定编码宽度，窗口变宽时只拉伸普通物料列。
  header=self.t.horizontalHeader();metrics=header.fontMetrics()
  for column in range(base,self.t.columnCount()):
   code=self._column_labels[column].split('\n',1)[0]
   width=max(header.minimumSectionSize(),metrics.horizontalAdvance(code+' ▾')+20)
   header.setSectionResizeMode(column,QHeaderView.Fixed);header.resizeSection(column,width)
   if not self.t.isColumnHidden(column):self.t._proportional_columns.widths[column]=width
  self.t._proportional_columns.apply();self.t.sync_frozen()

class MaterialVersionDialog(QDialog):
 def __init__(self,material,references,parent=None):
  super().__init__(parent); self.setWindowTitle('物料版本变更'); self.resize(720,430)
  layout=QVBoxLayout(self)
  label=QLabel(f"{material['base_code']}{material['version']} {material['name']}"); label.setWordWrap(True); layout.addWidget(label)
  form=QFormLayout(); self.version=QLineEdit(material['version']); form.addRow('新版本号',self.version); layout.addLayout(form)
  text=QLabel('该物料被如下BOM所使用，请勾选需要同步使用新版本的BOM。'); text.setWordWrap(True); layout.addWidget(text)
  self.references=QTreeWidget(); self.references.setHeaderLabels(['','SKU','父项物料']); self.references.setRootIsDecorated(False); self.references.setWordWrap(True)
  for reference in references:
   item=QTreeWidgetItem(self.references,['','\n'.join(reference['sku_labels']),reference['parent_label']]); item.setData(0,Qt.UserRole,reference['parent_id']); item.setCheckState(0,Qt.Checked)
  fit_columns(self.references); self.references.header().setSectionResizeMode(0,QHeaderView.Fixed); self.references.setColumnWidth(0,40); layout.addWidget(self.references)
  self.note=QLabel('同一父级BOM由多个SKU共享时统一更新，下方列出全部受影响SKU。\n全部勾选直接修改原物料；部分或全部不勾选则保留旧物料并创建新版本。')
  self.note.setWordWrap(True); self.note.setAlignment(Qt.AlignLeft|Qt.AlignBottom)
  font=self.note.font(); font.setPointSizeF(max(7.0,font.pointSizeF()-2)); self.note.setFont(font); layout.addWidget(self.note)
  buttons=QDialogButtonBox(QDialogButtonBox.Ok|QDialogButtonBox.Cancel); buttons.accepted.connect(self.accept); buttons.rejected.connect(self.reject); layout.addWidget(buttons)
 def selected_parents(self):
  return {self.references.topLevelItem(n).data(0,Qt.UserRole) for n in range(self.references.topLevelItemCount()) if self.references.topLevelItem(n).checkState(0)==Qt.Checked}
 def accept(self):
  if not self.version.text().strip(): QMessageBox.warning(self,'版本号无效','版本号不能为空'); return
  super().accept()

class MaterialUsageDialog(QDialog):
 def __init__(self,material,usages,parent=None):
  super().__init__(parent); self.material_id=material['id']; self.usages=usages; self.setWindowTitle('物料反查'); self.resize(720,430)
  layout=QVBoxLayout(self); title=QLabel(f"{material['base_code']}{material['version']} {material['name']}"); title.setWordWrap(True); layout.addWidget(title)
  self.table=QTableWidget(len(usages),1); self.table.setHorizontalHeaderLabels(['包含此物料的SKU / 父级物料']); self.table.setEditTriggers(QAbstractItemView.NoEditTriggers); self.table.setWordWrap(True)
  for index,usage in enumerate(usages):
   item=QTableWidgetItem(f"SKU：{usage['sku']}\n父级物料：{usage['parent']}"); item.setData(Qt.UserRole,usage.get('sku_id')); self.table.setItem(index,0,item)
  self.table.cellDoubleClicked.connect(self.open_sku); self.table.setSelectionBehavior(QAbstractItemView.SelectRows)
  fit_columns(self.table); layout.addWidget(self.table)
  buttons=QDialogButtonBox(QDialogButtonBox.Close); buttons.rejected.connect(self.reject); layout.addWidget(buttons)
 def open_sku(self,row,column=0):
  item=self.table.item(row,0)
  if item is None:return
  sku_id=item.data(Qt.UserRole)
  if sku_id is None:
   QMessageBox.information(self,'无对应SKU','此引用属于独立BOM，未包含在任何SKU中。'); return
  main=self.parent()
  while main is not None and not isinstance(main,MainWindow): main=main.parent()
  if main is None:return
  if main.svc.db.conn.execute('SELECT 1 FROM skus WHERE id=?',(sku_id,)).fetchone() is None:
   QMessageBox.information(self,'SKU不存在','该SKU已删除，请重新进行物料反查。'); return
  self.accept()  # 先结束模态反查，避免目标BOM被反查弹窗阻挡。
  window=main.open_bom_id(sku_id)
  if window is not None:window.locate_material(self.usages[row].get('focus_material_id',self.material_id),self.usages[row].get('parent_id'))
 def resizeEvent(self,event):
  super().resizeEvent(event)
  self.table.resizeRowsToContents()

class ProjectDialog(QDialog):
 def __init__(self,parent=None,row=None):
  super().__init__(parent); self.setWindowTitle('修改项目信息'); f=QFormLayout(self); self.e={}
  for k,l in [('project_name','项目名称'),('project_code','项目代码'),('product_name','产品名称'),('product_model','产品型号')]:
   try: value=row[k] if row is not None else ''
   except (KeyError,IndexError,TypeError): value=''
   w=QLineEdit(str(value or '')); self.e[k]=w; f.addRow(l,w)
  b=QDialogButtonBox(QDialogButtonBox.Ok|QDialogButtonBox.Cancel); b.accepted.connect(self.accept); b.rejected.connect(self.reject); f.addRow(b)
 def data(self): return {k:w.text().strip() for k,w in self.e.items()}
