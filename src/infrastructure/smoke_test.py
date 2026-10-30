"""发布包自检：临时数据库、Qt窗口与Excel往返，不触碰用户数据。"""
import json
import tempfile
import traceback
from pathlib import Path


def run_smoke_test(report_path):
    report = {'passed': False}
    db = None
    windows = []
    try:
        from PyQt5.QtWidgets import QApplication
        from database.database import Database
        from services.services import Services
        from ui.main_window import MainWindow, MaterialWindow, BomWindow
        from ui.about_dialog import AboutDialog
        from version import get_version
        from infrastructure.build_info import load_build_info
        from infrastructure.project_lock import ProjectLock
        from openpyxl import load_workbook
        import xlrd
        from services.external_bom import ExcelSource,build_external_bom,ALIASES
        owns_app = QApplication.instance() is None
        app = QApplication.instance() or QApplication([])
        with tempfile.TemporaryDirectory(prefix='bom-smoke-') as tmp:
            lock_target=Path(tmp)/'lock-test.db';lock_target.touch()
            guard = ProjectLock(lock_target)
            other = ProjectLock(lock_target)
            try:
                assert guard.acquire() and not other.acquire(), '项目数据库互斥失败'
            finally:
                other.close(); guard.close()
            db = Database(str(Path(tmp)/'测试.db'))
            svc = Services(db)
            root, child = [svc.create_material(dict(base_code=code,version='1',name=code,unit='pcs')) for code in ('SKU','子物料')]
            sid = svc.create_sku(root)
            svc.save_bom(root,[dict(material_id=child,quantity=1.5,unit='pcs')])
            path = Path(tmp)/'BOM.xlsx'
            svc.export_sku(sid,path)
            wb = load_workbook(path)
            try:assert wb['SKU'].cell(3,10).value==1.5
            finally:wb.close()
            external=ExcelSource(path)
            try:external_rows=external.rows('SKU')
            finally:external.close()
            headers=external_rows[0];mapping={h:h for h in headers}
            for field,aliases in ALIASES.items():
                if field not in mapping:
                    for alias in aliases:
                        if alias in headers:mapping[field]=alias;mapping.pop(alias,None);break
            tree=build_external_bom(external_rows,mapping)
            assert tree['children'][0]['quantity']==1.5 and xlrd.__version__, '外部Excel读取依赖不可用'
            mirror_id=svc.copy_sku_bom(sid);mirror=svc.for_sku(mirror_id)
            mirror.save_bom(root,[dict(material_id=child,quantity=2.5,unit='pcs')])
            assert svc.bom(root)[0]['quantity']==1.5 and mirror.bom(root)[0]['quantity']==2.5, '历史版本BOM未隔离'
            from services.snapshot_bom import migrate
            from services.current_catalog import initialize,synchronization_plan,save_confirmed,selection_tree
            migrate(db);initialize(db);mirror=svc.for_sku(mirror_id)
            second=svc.create_material(dict(base_code='SECOND',version='1',name='同步SKU',unit='pcs'))
            second_id=svc.create_sku(second);second_service=svc.for_sku(second_id)
            second_service.save_bom(second,[dict(material_id=root,quantity=1,unit='pcs'),dict(material_id=child,parent_material_id=root,quantity=2.5,unit='pcs')])
            rows=[dict(material_id=child,quantity=3.5,unit='pcs')]
            plan=synchronization_plan(mirror,root,rows)
            assert plan and plan[0]['other_skus'][0]['id']==second_id
            save_confirmed(mirror,root,rows,plan,{root:{second_id}},{})
            assert second_service.bom(root)[0]['quantity']==3.5 and svc.for_sku(sid).bom(root)[0]['quantity']==1.5
            assert selection_tree(mirror,root)[0]['data']['quantity']==3.5
            windows = [BomWindow(mirror,root),MainWindow(svc,get_version()),MaterialWindow(svc),BomWindow(svc,root),AboutDialog(get_version())]
            for window in windows:window.show()
            app.processEvents()
            from infrastructure.ui_smoke import exercise_windows
            ui_checks=exercise_windows(svc,windows[1],windows[2])
            from services import material_history as material_history
            material_history.initialize(db)
            history_id=svc.copy_sku_bom(sid);history_scope=svc.for_sku(history_id)
            previous=svc.for_sku(sid).material_snapshot(child)['name']
            material_history.save(history_scope,root,history_scope.bom(root),{child:{**history_scope.material_snapshot(child),'name':'物料历史冒烟验证'}},'fork')
            assert svc.for_sku(sid).material_snapshot(child)['name']==previous
            assert history_scope.material_snapshot(child)['name']=='物料历史冒烟验证'
            from ui.material_history import MaterialHistoryDialog
            history_view=MaterialHistoryDialog(svc,child,material_history.versions(svc,child)[0]['history_version']);history_view.show();app.processEvents();history_view.close()
            ui_checks.extend(['material_history_migration','material_history_metadata_isolation','material_history_readonly_window'])
            report.update(passed=True,version=get_version(),build=load_build_info(),qt_platform=app.platformName(),checks=['project_ownership','sqlite','excel','sku_window','material_window','bom_window','about_window','independent_sku_mirror','external_bom_excel','xls_reader','latest_catalog_sync_confirmation','historical_snapshot_preserved']+ui_checks)
            for window in windows:window.close()
            windows=[];db.close();db=None
    except Exception:
        report['error']=traceback.format_exc()
    finally:
        for window in windows:window.close()
        if db is not None:db.close()
        # Destroy native windows while QApplication is still alive, rather than
        # relying on Python/Qt shutdown order with pending widget timers.
        if locals().get('owns_app',False):
            from PyQt5.QtCore import QEvent
            for window in app.topLevelWidgets():window.deleteLater()
            app.sendPostedEvents(None,QEvent.DeferredDelete)
        from ui.operation_progress import RECENT_OPERATIONS
        report['operation_timings']=list(RECENT_OPERATIONS)
        Path(report_path).write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
    return 0 if report['passed'] else 1
