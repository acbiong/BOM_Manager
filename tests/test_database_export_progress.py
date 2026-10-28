import os,sys,tempfile,unittest
from pathlib import Path
from unittest.mock import patch
os.environ['QT_QPA_PLATFORM']='offscreen'
sys.path.insert(0,str(Path(__file__).parents[1]/'src'))
from PyQt5.QtWidgets import QApplication
from openpyxl import Workbook,load_workbook
from database.database import Database
from services.services import Services
from infrastructure.operation_progress import _current,checkpoint,progress_stage,suspend_progress
from ui.operation_progress import OperationProgress,operation

class DatabaseExportProgressTests(unittest.TestCase):
 @classmethod
 def setUpClass(cls):cls.app=QApplication.instance() or QApplication([])
 def test_real_export_preserves_scale_through_formats_and_file_write(self):
  with tempfile.TemporaryDirectory() as tmp:
   db=Database(str(Path(tmp)/'db'));svc=Services(db);events=[];original=OperationProgress.update
   for i in range(2):svc.create_sku(svc.create_material(dict(base_code=str(i),version='1',name='SKU'+str(i),unit='pcs')))
   def observe(controller,message=None,current=None,total=None):
    original(controller,message,current,total)
    if controller.dialog is not None:events.append((message,controller.dialog.value(),controller.dialog.maximum()))
   save=Workbook.save
   def write(book,path):
    self.assertNotEqual(events[-1][1],events[-1][2]);self.assertIn('正在写入',events[-1][0]);save(book,path)
   try:
    with patch.object(OperationProgress,'threshold',0),patch.object(OperationProgress,'paint_interval',0),patch.object(OperationProgress,'update',observe),patch.object(Workbook,'save',write):
     operation('导出资料库')(lambda:svc.export_database(Path(tmp)/'out.xlsx'))()
    self.assertTrue(events);self.assertTrue(all(total==6 for _,_,total in events))
    values=[value for _,value,_ in events];self.assertEqual(values,sorted(values));self.assertEqual(values[-1],6)
    wb=load_workbook(Path(tmp)/'out.xlsx');self.assertEqual(len(wb.sheetnames),5);self.assertEqual(wb['SKU清单'].max_row,3);wb.close()
   finally:db.close()
 def test_write_error_never_reports_complete_and_releases_stage(self):
  class Recorder:
   def update(self,message,current,total):events.append((message,current,total))
  events=[];token=_current.set(Recorder())
  try:
   with tempfile.TemporaryDirectory() as tmp:
    db=Database(str(Path(tmp)/'db'));svc=Services(db)
    try:
     with patch.object(Workbook,'save',side_effect=OSError('disk full')):
      with self.assertRaisesRegex(OSError,'disk full'):svc.export_database(Path(tmp)/'out.xlsx')
    finally:db.close()
   self.assertFalse(any(total and current==total for _,current,total in events))
   checkpoint('后续操作',1,2);self.assertEqual(events[-1],('后续操作',1,2))
  finally:_current.reset(token)
 def test_other_window_is_not_given_suspended_stage(self):
  from infrastructure.operation_progress import _stage
  with progress_stage('导出',2,5):
   with suspend_progress():self.assertIsNone(_stage.get())
   self.assertEqual(_stage.get(),('导出',2,5))
