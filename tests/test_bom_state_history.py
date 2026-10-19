import os
os.environ['QT_QPA_PLATFORM']='offscreen'
import sys,tempfile,unittest
from pathlib import Path
from unittest.mock import patch
sys.path.insert(0,str(Path(__file__).parents[1]/'src'))
from PyQt5.QtWidgets import QApplication,QMessageBox
from openpyxl import load_workbook
from database.database import Database
from services.services import Services
from ui.main_window import BomWindow

class BomStateHistoryTests(unittest.TestCase):
 @classmethod
 def setUpClass(cls):cls.app=QApplication.instance() or QApplication([])
 def test_expansion_save_refresh_reopen_and_export_history(self):
  with tempfile.TemporaryDirectory() as tmp:
   db=Database(str(Path(tmp)/'db'));s=Services(db)
   ids=[s.create_material(dict(base_code=code,version='1',name=code,unit='pcs')) for code in ['R','P','C','OTHER']]
   root,parent,child,other=ids
   def rows(qty):return [dict(material_id=parent,quantity=1,unit='pcs'),dict(parent_material_id=parent,material_id=child,quantity=qty,unit='pcs',remark='备注')]
   s.save_bom(root,rows(1));sid=s.create_sku(root);s.save_bom(other,rows(1));other_sid=s.create_sku(other)
   window=BomWindow(s,root)
   try:
    window.root_item.child(0).setExpanded(True);window.root_item.setExpanded(False)
    window.show_item_info(window.root_item.child(0).child(0));window.fields['quantity'].setValue(3.5)
    with patch.object(QMessageBox,'information'):self.assertTrue(window.save())
    self.assertFalse(window.root_item.isExpanded());self.assertTrue(window.root_item.child(0).isExpanded())
    s.save_bom(root,rows(4.5));window.refresh_external()
    self.assertFalse(window.root_item.isExpanded());self.assertTrue(window.root_item.child(0).isExpanded())
    window.close();window=BomWindow(s,root)
    self.assertFalse(window.root_item.isExpanded());self.assertTrue(window.root_item.child(0).isExpanded())
    for current_sid in [sid,other_sid]:
     path=Path(tmp)/(str(current_sid)+'.xlsx');s.export_sku(current_sid,path);wb=load_workbook(path)
     details='\n'.join(str(row[-1] or '') for row in list(wb['变更记录'].values)[1:]);wb.close()
     self.assertIn('数量：1.0 → 3.5',details);self.assertIn('数量：3.5 → 4.5',details)
    count=db.conn.execute("SELECT COUNT(*) FROM change_logs WHERE object_type='bom'").fetchone()[0]
    s.save_bom(root,rows(4.5));self.assertEqual(db.conn.execute("SELECT COUNT(*) FROM change_logs WHERE object_type='bom'").fetchone()[0],count)
   finally:window.dirty=False;window.close();db.close()
 def test_single_point_position_change_and_removed_history(self):
  with tempfile.TemporaryDirectory() as tmp:
   db=Database(str(Path(tmp)/'db'));s=Services(db)
   root,parent,child=[s.create_material(dict(base_code=code,version='1',name=code,unit='pcs')) for code in ['R','P','C']]
   sid=s.create_sku(root)
   def rows(position,qty):return [dict(material_id=parent,quantity=1,unit='pcs'),dict(parent_material_id=parent,material_id=child,quantity=qty,unit='pcs',position=position)]
   try:
    s.save_bom(root,rows('左侧',1));before=db.conn.execute('SELECT MAX(id) FROM change_logs').fetchone()[0]
    s.save_bom(root,rows('右侧',2))
    logs=list(db.conn.execute('SELECT * FROM change_logs WHERE id>?',(before,)))
    self.assertEqual(len(logs),1);detail=logs[0]['detail']
    for text in ['父项：P1 P','物料：C1 C','位置：左侧 → 右侧','数量：1.0 → 2.0']:self.assertIn(text,detail)
    s.save_bom(root,[])
    # 子树已移除，历史日志通过归属索引仍可导出，且同一日志只出现一次。
    path=Path(tmp)/'export.xlsx';s.export_sku(sid,path);wb=load_workbook(path)
    values=list(wb['变更记录'].values);wb.close()
    self.assertEqual(sum('位置：左侧 → 右侧' in str(row[-1]) for row in values),1)
    self.assertTrue(any('删除物料' in str(row[-1]) and '物料：C1 C' in str(row[-1]) for row in values))
   finally:db.close()

 def test_sort_only_no_log_and_alternate_add_remove_explicit(self):
  db=Database(':memory:');s=Services(db)
  try:
   root,child,other,a,b=[s.create_material(dict(base_code=code,version='1',name=code,unit='pcs')) for code in ['R','C','O','A','B']]
   def row(mid,alts):return dict(material_id=mid,quantity=1,unit='pcs',position='左侧',alternate_material_ids=alts)
   s.save_bom(root,[row(child,[]),row(other,[])])
   def count():return db.conn.execute("SELECT COUNT(*) FROM change_logs WHERE object_type='bom'").fetchone()[0]
   initial=count();s.save_bom(root,[row(other,[]),row(child,[])]);self.assertEqual(count(),initial)
   s.save_bom(root,[row(other,[]),row(child,[a,b])]);self.assertEqual(count(),initial+1)
   detail=db.conn.execute('SELECT detail FROM change_logs ORDER BY id DESC LIMIT 1').fetchone()[0]
   for text in ['父项：R1 R','物料：C1 C','物料位置：左侧','新增替代料：A1 A','新增替代料：B1 B']:self.assertIn(text,detail)
   self.assertNotIn('排序',detail)
   s.save_bom(root,[row(other,[]),row(child,[b,a])]);self.assertEqual(count(),initial+1)
   s.save_bom(root,[row(other,[]),row(child,[b])]);self.assertEqual(count(),initial+2)
   detail=db.conn.execute('SELECT detail FROM change_logs ORDER BY id DESC LIMIT 1').fetchone()[0]
   self.assertIn('删除替代料：A1 A',detail);self.assertNotIn('新增替代料',detail)
   s.save_bom(root,[row(other,[])]);detail=db.conn.execute('SELECT detail FROM change_logs ORDER BY id DESC LIMIT 1').fetchone()[0]
   self.assertIn('删除替代料：B1 B',detail)
  finally:db.close()

 def test_export_hides_legacy_sort_only_record(self):
  with tempfile.TemporaryDirectory() as tmp:
   db=Database(str(Path(tmp)/'db'));s=Services(db)
   try:
    root=s.create_material(dict(base_code='R',version='1',name='R',unit='pcs'));sid=s.create_sku(root)
    s.log('bom',root,'update','父项：R1 R\n物料：C1 C\n物料位置：左侧\n排序：0 → 1')
    s.log('bom',root,'update','父项：R1 R\n物料：C1 C\n物料位置：左侧\n排序：1 → 2\n数量：1.0 → 2.0');db.conn.commit()
    out=Path(tmp)/'out.xlsx';s.export_sku(sid,out);wb=load_workbook(out)
    rows=[row for row in list(wb['变更记录'].values)[1:] if row[5]=='修改物料'];wb.close()
    self.assertEqual(len(rows),1);self.assertNotIn('排序',rows[0][-1]);self.assertIn('数量',rows[0][-1])
    self.assertIn('排序',db.conn.execute('SELECT detail FROM change_logs ORDER BY id DESC LIMIT 1').fetchone()[0])
   finally:db.close()
 def test_delete_recreate_resets_history_but_preserves_shared_owner(self):
  with tempfile.TemporaryDirectory() as tmp:
   db=Database(str(Path(tmp)/'db'));s=Services(db)
   try:
    r,p,c,o=[s.create_material(dict(base_code=code,version='1',name=code,unit='pcs')) for code in ['R','P','C','O']]
    sid=s.create_sku(r);other=s.create_sku(o)
    def rows(q):return [dict(material_id=p,quantity=1,unit='pcs'),dict(parent_material_id=p,material_id=c,quantity=q,unit='pcs')]
    s.save_bom(r,rows(1));s.save_bom(o,rows(1));s.save_bom(r,rows(7.3))
    s.delete_sku(sid)
    self.assertEqual(db.conn.execute("SELECT COUNT(*) FROM change_logs WHERE object_type='sku' AND object_id=?",(sid,)).fetchone()[0],0)
    self.assertEqual(db.conn.execute('SELECT COUNT(*) FROM bom_log_roots WHERE root_id=?',(r,)).fetchone()[0],0)
    recreated=s.create_sku(r);s.save_bom(r,rows(7.3))
    for identity,expect_old in [(other,True),(recreated,False)]:
     path=Path(tmp)/(str(identity)+'.xlsx');s.export_sku(identity,path);wb=load_workbook(path)
     values=list(wb['变更记录'].values)[1:];wb.close()
     self.assertEqual(any('数量：1.0 → 7.3' in str(row[-1]) for row in values),expect_old)
     self.assertFalse(any(row[5]=='修改BOM' for row in values))
    # 替代料纯增加/删除分别明确分类，混合变化仍只占一行。
    s.save_bom(r,[dict(material_id=p,quantity=1,unit='pcs',alternate_material_ids=[c]),dict(parent_material_id=p,material_id=c,quantity=7.3,unit='pcs')])
    path=Path(tmp)/'types.xlsx';s.export_sku(recreated,path);wb=load_workbook(path)
    self.assertTrue(any(row[5]=='新增替代料' for row in list(wb['变更记录'].values)[1:]));wb.close()
   finally:db.close()

 def test_changed_material_identity_is_child_not_parent(self):
  with tempfile.TemporaryDirectory() as tmp:
   db=Database(str(Path(tmp)/'db'));s=Services(db)
   try:
    root=s.create_material(dict(base_code='ROOT',version='1',name='父级名称',unit='pcs'))
    child=s.create_material(dict(base_code='CHILD',version='100',name='实际变更物料 名称',unit='pcs'))
    sid=s.create_sku(root)
    def save(q):s.save_bom(root,[dict(material_id=child,quantity=q,unit='pcs',position='左侧')])
    save(1);save(2)
    logs=list(db.conn.execute("SELECT * FROM change_logs WHERE object_type='bom'"))
    self.assertEqual(len(logs),2)
    for row in logs:
     self.assertEqual(row['object_id'],root)
     self.assertEqual((row['material_code'],row['material_name']),('CHILD100','实际变更物料 名称'))
    # 兼容旧版已将父项身份写入快照的日志，按明细的实际物料修正导出。
    db.conn.execute("UPDATE change_logs SET material_code='ROOT1',material_name='父级名称' WHERE object_type='bom'");db.conn.commit()
    s.save_bom(root,[])
    out=Path(tmp)/'bom.xlsx';s.export_sku(sid,out);wb=load_workbook(out)
    rows=[row for row in list(wb['变更记录'].values)[1:] if row[5] in ('新增物料','修改物料','删除物料')];wb.close()
    self.assertEqual(len(rows),3)
    for row in rows:
     self.assertEqual(row[3:5],('CHILD100','实际变更物料 名称'));self.assertIn('父项：ROOT1 父级名称',row[-1])
   finally:db.close()
