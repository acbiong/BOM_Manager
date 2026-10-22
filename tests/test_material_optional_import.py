import os
os.environ['QT_QPA_PLATFORM']='offscreen'
import sys,tempfile,unittest
from pathlib import Path
from unittest.mock import Mock,patch
sys.path.insert(0,str(Path(__file__).parents[1]/'src'))
from PyQt5.QtWidgets import QApplication,QDialog,QFileDialog
from ui.interactive_windows import QFileDialog
from openpyxl import Workbook,load_workbook
from database.database import Database
from services.services import Services,MATERIAL_IMPORT_HEADERS,MATERIAL_IMPORT_DEFAULTS,MaterialHeaderError
from ui.main_window import MainWindow,MaterialWindow,MaterialHeaderMappingDialog,MaterialConflictDialog

class OptionalMaterialImportTests(unittest.TestCase):
 @classmethod
 def setUpClass(cls):cls.app=QApplication.instance() or QApplication([])
 def setUp(self):
  self.tmp=tempfile.TemporaryDirectory();self.base=Path(self.tmp.name);self.db=Database(str(self.base/'db'));self.s=Services(self.db)
  self.required=[h for h in MATERIAL_IMPORT_HEADERS if h not in MATERIAL_IMPORT_DEFAULTS]
 def tearDown(self):self.db.close();self.tmp.cleanup()
 def workbook(self,code='A',grade='医疗级',mp='P',headers=None):
  headers=headers or MATERIAL_IMPORT_HEADERS
  data={'物料基础编码':code,'物料版本号':'1','物料名称':'物料','物料规格描述':'','单位':'pcs','物料等级':grade,'M/P':mp,'物料制造商':'厂','物料MPN':''}
  wb=Workbook();wb.active.append(headers);wb.active.append([data[h] for h in headers]);path=self.base/'in.xlsx';wb.save(path);wb.close();return path
 def current(self):return dict(self.s.materials()[0])
 def test_auto_mapping_and_mp_only_change_are_imported(self):
  path=self.workbook(mp='p');self.assertEqual(self.s.import_materials(path),(['A'],[]))
  self.assertEqual((self.current()['grade'],self.current()['mp_flag']),('医疗级','P'))
  self.workbook(mp='M')
  def overwrite(incoming,existing):
   dialog=MaterialConflictDialog(incoming,existing)
   try:
    col=next(i for i in range(dialog.table.columnCount()) if dialog.table.horizontalHeaderItem(i).text()=='M/P')
    self.assertEqual(dialog.table.item(0,col).text(),'P');self.assertEqual(dialog.table.item(1,col).text(),'M')
   finally:dialog.close()
   return 'overwrite'
  self.assertEqual(self.s.import_materials(path,overwrite),(['A'],[]));self.assertEqual(self.current()['mp_flag'],'M')
  resolver=Mock();ok,fail=self.s.import_materials(path,resolver);resolver.assert_not_called();self.assertEqual(ok,[]);self.assertIn('信息相同',fail[0][3])
 def test_optional_missing_explicit_unmapped_and_blank_cells_default(self):
  for index,kind in enumerate(('missing','unmapped','blank')):
   path=self.workbook(code=str(index),headers=self.required if kind=='missing' else None,grade='' if kind=='blank' else '医疗级',mp='' if kind=='blank' else 'P')
   mapping={h:h for h in self.required}|{'物料等级':'','M/P':''} if kind=='unmapped' else None
   ok,fail=self.s.import_materials(path,header_mapping=mapping);self.assertEqual(ok,[str(index)]);self.assertEqual(fail,[])
  self.assertTrue(all((r['grade'],r['mp_flag'])==('','M') for r in self.s.materials()))
 def test_unmapped_fields_preserve_existing_on_overwrite(self):
  path=self.workbook();self.s.import_materials(path)
  wb=load_workbook(path);wb.active.cell(2,3,'修改名称');wb.save(path);wb.close()
  mapping={h:h for h in self.required}|{'物料等级':'','M/P':''}
  self.s.import_materials(path,lambda *_:'overwrite',header_mapping=mapping)
  self.assertEqual((self.current()['name'],self.current()['grade'],self.current()['mp_flag']),('修改名称','医疗级','P'))
 def test_optional_mapping_validation_and_invalid_mp(self):
  path=self.workbook(mp='INVALID');ok,fail=self.s.import_materials(path)
  self.assertEqual(ok,[]);self.assertIn('M/P不合规',fail[0][3]);self.assertEqual(len(self.s.materials()),0)
  dialog=MaterialHeaderMappingDialog(path,MATERIAL_IMPORT_HEADERS,[])
  try:
   for name in MATERIAL_IMPORT_DEFAULTS:dialog.combos[name].setCurrentIndex(0)
   dialog.validate_mapping();self.assertEqual(dialog.result(),QDialog.Accepted)
   mapping=dialog.mapping();self.assertEqual(mapping['物料等级'],'');self.assertEqual(mapping['M/P'],'')
   self.assertEqual(self.s.import_materials(path,header_mapping=mapping),(['A'],[]))
   mapping['M/P']='不存在的列'
   with self.assertRaises(MaterialHeaderError):self.s.import_materials(path,header_mapping=mapping)
  finally:dialog.close()
 def test_custom_header_mapping_and_both_template_entries(self):
  path=self.workbook();wb=load_workbook(path)
  wb.active.cell(1,6,'等级分类');wb.active.cell(1,7,'采购制造');wb.save(path);wb.close()
  mapping={h:h for h in self.required}|{'物料等级':'等级分类','M/P':'采购制造'}
  self.assertEqual(self.s.import_materials(path,header_mapping=mapping),(['A'],[]));self.assertEqual(self.current()['mp_flag'],'P')
  main=MainWindow(self.s);material=MaterialWindow(self.s)
  try:
   for i,window in enumerate((main,material)):
    output=self.base/f'template{i}.xlsx'
    with patch.object(QFileDialog,'getSaveFileName',return_value=(str(output),'')):window.template()
    wb=load_workbook(output)
    try:self.assertEqual([c.value for c in wb.active[1]],MATERIAL_IMPORT_HEADERS)
    finally:wb.close()
  finally:material.close();main.close()
