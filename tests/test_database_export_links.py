import tempfile,sys,unittest
from pathlib import Path
sys.path.insert(0,str(Path(__file__).parents[1]/'src'))
from openpyxl import load_workbook
from openpyxl.utils import quote_sheetname
from database.database import Database
from services.services import Services

class DatabaseExportLinkTests(unittest.TestCase):
 def test_links_target_exact_history_and_safe_unique_sheet_names(self):
  with tempfile.TemporaryDirectory() as tmp:
   db=Database(str(Path(tmp)/'db'));svc=Services(db)
   try:
    # Long and special codes exercise truncated/colliding titles and apostrophe escaping.
    ids=[]
    for code in ("ABC'中文",'L'*35+'A','L'*35+'B','BAD/NAME'):
     mid=svc.create_material(dict(base_code=code,version='100',name=code,unit='pcs'));ids.append(svc.create_sku(mid))
    svc.copy_sku_bom(ids[1])
    skus=svc.skus();path=Path(tmp)/'all.xlsx';svc.export_database(path)
    book=load_workbook(path)
    try:
     seen=set()
     for i,sku in enumerate(skus,2):
      cell=book['SKU清单'].cell(i,2);link=cell.hyperlink;sheet=book.worksheets[i+1]
      self.assertEqual(link.location,quote_sheetname(sheet.title)+'!A1');self.assertIsNone(link.target)
      self.assertEqual(cell.value,sku['base_code']+sku['version']);self.assertNotIn(link.location,seen);seen.add(link.location)
      self.assertEqual(sheet.cell(2,5).value,sku['base_code']);self.assertEqual(sheet.cell(2,6).value,sku['version'])
      self.assertLessEqual(len(sheet.title),31);self.assertTrue(cell.alignment.wrap_text);self.assertEqual(cell.border.left.style,'thin')
      self.assertEqual(cell.font.underline,'single')
     self.assertIn("''",book['SKU清单']['B2'].hyperlink.location)
    finally:book.close()
    standalone=Path(tmp)/'sku.xlsx';svc.export_sku_list(standalone);book=load_workbook(standalone)
    try:self.assertTrue(all(row[1].hyperlink is None for row in book.active.iter_rows(min_row=2)))
    finally:book.close()
   finally:db.close()
