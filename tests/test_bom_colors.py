import os,sys,tempfile,unittest
os.environ['QT_QPA_PLATFORM']='offscreen'
from pathlib import Path
from unittest.mock import patch
sys.path.insert(0,str(Path(__file__).parents[1]/'src'))
from PyQt5.QtCore import Qt,QEvent,QPoint
from PyQt5.QtGui import QColor
from PyQt5.QtWidgets import QApplication,QMenu,QMessageBox,QColorDialog
from openpyxl import load_workbook
from database.database import Database
from config.config_service import ConfigService
from services.services import Services
from services.snapshot_bom import migrate
from services.current_catalog import initialize
from services.bom_colors import colors,save,foreground
from services.sku_lock import LockedBomError
from ui.main_window import MainWindow
from ui.bom_colors import ColorMenu,palette,save_palette,DEFAULTS

class BomColorTests(unittest.TestCase):
 @classmethod
 def setUpClass(cls):cls.app=QApplication.instance() or QApplication([])
 def setUp(self):
  self.tmp=tempfile.TemporaryDirectory();self.cfg=ConfigService(self.tmp.name);self.db=Database(str(Path(self.tmp.name)/'db'));self.s=Services(self.db,self.cfg)
  self.r,self.p,self.c=[self.s.create_material(dict(base_code=x,version='100',name=x,unit='pcs')) for x in ('R','P','C')]
  self.s.save_bom(self.p,[dict(material_id=self.c,quantity=2,unit='pcs')]);self.s.save_bom(self.r,[dict(material_id=self.p,quantity=1,unit='pcs')]);self.sid=self.s.create_sku(self.r)
  migrate(self.db);initialize(self.db);self.main=MainWindow(self.s);self.w=self.main.open_bom_id(self.sid)
 def tearDown(self):
  for w in list(self.main._bom_windows.values())+list(self.main._material_bom_windows.values()):w.dirty=False;w.close()
  self.main.close();self.main.deleteLater();self.app.sendPostedEvents(None,QEvent.DeferredDelete);self.db.close();self.tmp.cleanup()
 def save_window(self):
  with patch.object(QMessageBox,'information'):self.assertTrue(self.w.save())
 def test_draft_save_reload_excel_root_and_child_colors(self):
  parent=self.w.root_item.child(0);child=parent.child(0)
  self.assertNotEqual(parent.background(0).color(),child.background(0).color())
  self.w.set_row_color([parent],'#102030');self.w.set_row_color([self.w.root_item],'#fff3cd')
  self.assertEqual(parent.background(0).color().name(),'#102030');self.assertEqual(parent.foreground(0).color().name(),'#ffffff');self.assertEqual(colors(self.w.svc),{})
  self.save_window();self.w.load();parent=self.w.root_item.child(0)
  self.assertEqual(parent.background(0).color().name(),'#102030');self.assertEqual(colors(self.w.svc)[(0,self.r)],'#fff3cd')
  path=Path(self.tmp.name)/'out.xlsx';self.s.export_sku(self.sid,path);book=load_workbook(path)
  try:
   sheet=book['R'];self.assertEqual(sheet['A2'].fill.fgColor.rgb[-6:],'fff3cd');self.assertEqual(sheet['A3'].fill.fgColor.rgb[-6:],'102030');self.assertEqual(sheet['A3'].font.color.rgb[-6:],'ffffff');self.assertEqual(sheet['A4'].fill.patternType,None)
  finally:book.close()
  self.w.set_row_color([parent],'');self.save_window();self.assertNotIn((self.r,self.p),colors(self.w.svc))
 def test_change_markers_override_custom_and_refresh_preserves_draft(self):
  parent=self.w.root_item.child(0);self.w.set_row_color([parent],'#102030');self.save_window();parent=self.w.root_item.child(0)
  self.w.show_item_info(parent);self.w.fields['quantity'].setText('3')
  self.assertNotEqual(parent.background(0).color().name(),'#102030');self.assertEqual(parent.foreground(0).color().name(),'#9a6700')
  self.w.set_row_color([parent],'#123456');self.w.refresh_external();parent=self.w.root_item.child(0)
  self.assertEqual(parent.data(0,Qt.UserRole)['row_color'],'#123456')
  self.w.remove_relation(parent);self.assertEqual(parent.foreground(0).color().name(),'#ff0000')
 def test_history_copy_isolation_lock_and_deleted_cleanup(self):
  parent=self.w.root_item.child(0);self.w.set_row_color([parent],'#d4edda');self.save_window()
  newer=self.s.copy_sku_bom(self.sid);self.assertEqual(colors(self.s.for_sku(newer))[(self.r,self.p)],'#d4edda')
  with self.db.transaction():save(self.s.for_sku(newer),[(self.r,self.p,'#112233')])
  self.assertEqual(colors(self.w.svc)[(self.r,self.p)],'#d4edda')
  self.s.set_sku_locked(self.sid,True)
  self.w.set_row_color([self.w.root_item],'#000000');self.assertNotIn((0,self.r),colors(self.w.svc))
  with self.assertRaises(LockedBomError):
   with self.db.transaction():save(self.w.svc,[(self.r,self.p,'#000000')])
  self.assertEqual(colors(self.w.svc)[(self.r,self.p)],'#d4edda')
  target=self.s.for_sku(newer);target.save_bom(self.r,[]);self.assertNotIn((self.r,self.p),colors(target))
 def test_palette_defaults_custom_save_delete_and_cancel(self):
  self.assertEqual(palette(self.cfg),DEFAULTS);node=self.w.root_item.child(0);menu=ColorMenu(self.w,[node],self.w)
  with patch.object(QColorDialog,'getColor',return_value=QColor('#123456')),patch.object(QMessageBox,'question',return_value=QMessageBox.Yes):menu.more()
  self.assertIn(['#123456','#123456'],palette(ConfigService(self.tmp.name)));self.assertEqual(node.background(0).color().name(),'#123456')
  menu.populate();action=next(a for a in menu.actions() if a.data()=='#123456')
  with patch.object(menu,'actionAt',return_value=action),patch.object(QMenu,'exec_',lambda m,*a:m.actions()[0]):menu.remove_color(QPoint())
  self.assertNotIn(['#123456','#123456'],palette(self.cfg));self.assertEqual(node.data(0,Qt.UserRole)['row_color'],'#123456')
  with patch.object(QColorDialog,'getColor',return_value=QColor()):menu.more()
  save_palette(self.cfg,[]);self.assertEqual(palette(self.cfg),[])
  self.assertEqual(foreground('#000000'),'#ffffff');self.assertEqual(foreground('#ffffff'),'#000000');menu.deleteLater()
 def test_context_submenu_apply_and_cancel_discards_color(self):
  node=self.w.root_item.child(0);child=node.child(0)
  def choose(menu,*args):
   submenu=next(a.menu() for a in menu.actions() if a.text()=='设置颜色')
   self.assertEqual([a.text() for a in submenu.actions() if a.data()],['浅红','浅黄','浅蓝','浅绿'])
   submenu.actions()[0].trigger();return None
  with patch.object(self.w.tree,'itemAt',return_value=node),patch.object(QMenu,'exec_',choose):self.w.context_menu(QPoint())
  self.assertEqual(node.data(0,Qt.UserRole)['row_color'],'#f8d7da');self.assertEqual(child.data(0,Qt.UserRole).get('row_color',''),'')
  self.assertEqual(colors(self.w.svc),{});self.w.dirty=False;self.w.load()
  self.assertEqual(self.w.root_item.child(0).data(0,Qt.UserRole).get('row_color',''),'')
  self.s.set_sku_locked(self.sid,True);menu=QMenu(self.w);self.w.add_color_menu(menu,self.w.root_item);self.assertEqual(menu.actions(),[])
 def test_subtree_export_uses_selected_root_color_and_reopen_database(self):
  node=self.w.root_item.child(0);self.w.set_row_color([node],'#112233');self.save_window()
  path=Path(self.tmp.name)/'subtree.xlsx'
  self.w.svc.export_bom_structure(self.p,path,root_color='#112233');book=load_workbook(path)
  try:self.assertEqual(book['P']['A2'].fill.fgColor.rgb[-6:],'112233');self.assertEqual(book['P']['A2'].font.color.rgb[-6:],'ffffff')
  finally:book.close()
  other=Database(self.db.path)
  try:self.assertEqual(colors(Services(other).for_sku(self.sid))[(self.r,self.p)],'#112233')
  finally:other.close()
