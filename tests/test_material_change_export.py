import sys,tempfile,unittest,sqlite3
from pathlib import Path
sys.path.insert(0,str(Path(__file__).parents[1]/'src'))
from database.database import Database
from services.services import Services
from openpyxl import load_workbook

class MaterialChangeExportTests(unittest.TestCase):
 def setUp(self):
  self.tmp=tempfile.TemporaryDirectory();self.path=Path(self.tmp.name);self.db=Database(str(self.path/'db.sqlite'));self.s=Services(self.db)
 def tearDown(self):self.db.close();self.tmp.cleanup()
 def exported(self):
  path=self.path/'materials.xlsx';self.s.export_materials(path);wb=load_workbook(path);rows=list(wb['变更记录'].values);wb.close();return rows
 def test_history_keeps_identity_and_field_diffs_after_delete_and_id_reuse(self):
  mid=self.s.create_material(dict(base_code='A',version='1',name='旧名称',unit='pcs',manufacturer='旧厂',mpn='旧MPN',description='旧描述',mp_flag='M'))
  self.s.update_material(mid,dict(name='新名称',unit='米',manufacturer='新厂',mpn='新MPN',description='',mp_flag='P'))
  self.s.change_material_version(mid,'2');self.s.delete_material(mid)
  self.s.create_material(dict(base_code='B',version='1',name='另一物料',unit='pcs'))
  rows=self.exported();self.assertEqual(rows[0],('序号','变更日期','变更时间','变更物料编码','变更物料名称','变更类型','变更详细内容'))
  self.assertEqual(rows[1][3:5],('A1','旧名称'));self.assertEqual(rows[2][3:5],('A1','新名称'))
  for change in ['物料名称：旧名称 → 新名称','单位：pcs → 米','制造商：旧厂 → 新厂','MPN：旧MPN → 新MPN','规格描述：旧描述 → （空）','M/P：M → P']:self.assertIn(change,rows[2][6])
  self.assertEqual(rows[3][3],'A2');self.assertIn('1->2',rows[3][6]);self.assertEqual(rows[4][3:5],('A2','新名称'));self.assertEqual(rows[4][5],'删除物料')
  self.assertEqual(len(rows[2][1]),10);self.assertEqual(len(rows[2][2]),8)
 def test_old_schema_migrates_and_legacy_values_are_not_fabricated(self):
  other=self.path/'old.sqlite';conn=sqlite3.connect(other);conn.execute('CREATE TABLE change_logs(id INTEGER PRIMARY KEY,object_type TEXT,object_id INTEGER,action TEXT,detail TEXT,changed_at TEXT)');conn.execute("INSERT INTO change_logs VALUES(1,'material',1,'update','旧名字','2020-01-02 03:04:05')");conn.commit();conn.close()
  db=Database(str(other))
  try:
   service=Services(db);service.create_material(dict(base_code='NEW',version='1',name='无关',unit='pcs'));out=self.path/'old.xlsx';service.export_materials(out);wb=load_workbook(out);row=list(wb['变更记录'].values)[1];wb.close()
   self.assertEqual(row[1:3],('2020-01-02','03:04:05'));self.assertEqual(row[3],'历史记录未保存');self.assertIn('旧名字',row[6])
  finally:db.close()
 def test_import_overwrite_records_differences_and_failed_version_no_log(self):
  from openpyxl import Workbook
  mid=self.s.create_material(dict(base_code='A',version='1',name='旧',unit='pcs',manufacturer='厂'))
  wb=Workbook();wb.active.append(['物料基础编码','物料版本号','物料名称','物料规格描述','单位','物料制造商','物料MPN']);wb.active.append(['A','1','新','','米','厂','']);source=self.path/'input.xlsx';wb.save(source);wb.close()
  self.s.import_materials(source,lambda *_:'overwrite');self.assertIn('单位：pcs → 米',self.exported()[-1][6])
  count=self.db.conn.execute('SELECT COUNT(*) FROM change_logs').fetchone()[0]
  with self.assertRaises(ValueError):self.s.change_material_version(mid,'1')
  self.assertEqual(self.db.conn.execute('SELECT COUNT(*) FROM change_logs').fetchone()[0],count)

 def test_material_sheet_format_newlines_and_legacy_codes(self):
  mid=self.s.create_material(dict(base_code='BC3-0000197',version='100',name='长名称'*30,unit='pcs',grade='工业级'))
  self.s.update_material(mid,dict(name='新名称',unit='pcs',description='描述',grade='A级'))
  self.db.conn.execute("UPDATE change_logs SET material_code='BC3-0000197-100' WHERE object_type='material'");self.db.conn.commit()
  sid=self.s.create_sku(mid)
  for kind in ('material','database','sku'):
   path=self.path/(kind+'.xlsx')
   if kind=='material':self.s.export_materials(path)
   elif kind=='database':self.s.export_database(path)
   else:self.s.export_sku(sid,path)
   wb=load_workbook(path)
   try:
    logs=wb['变更记录'];self.assertEqual(logs.cell(2,4).value,'BC3-0000197100')
    self.assertIn('\n',logs.cell(3,7).value);self.assertNotIn('；',logs.cell(3,7).value);self.assertTrue(logs.cell(3,7).alignment.wrap_text)
    if kind!='sku':
     sheet=wb['物料库'];self.assertEqual(sheet.cell(1,1).value,'序号');self.assertEqual(sheet.cell(2,1).value,1)
     self.assertTrue(all(d.width<=60 for d in sheet.column_dimensions.values()))
     for cells in sheet:
      for cell in cells:
       self.assertTrue(cell.alignment.wrap_text)
       for edge in (cell.border.left,cell.border.right,cell.border.top,cell.border.bottom):self.assertEqual(edge.color.rgb,'FF000000');self.assertEqual(edge.style,'thin')
       if sheet.cell(1,cell.column).value in ('序号','单位','物料等级','M/P','是否存在BOM结构','是否为呆滞料'):self.assertEqual(cell.alignment.horizontal,'center')
   finally:wb.close()
  self.assertEqual(self.db.conn.execute("SELECT material_code FROM change_logs WHERE object_type='material' LIMIT 1").fetchone()[0],'BC3-0000197-100')
