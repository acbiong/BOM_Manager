import os,sys,tempfile,unittest
os.environ['QT_QPA_PLATFORM']='offscreen'
from pathlib import Path
sys.path.insert(0,str(Path(__file__).parents[1]/'src'))
from unittest.mock import patch
from PyQt5.QtWidgets import QApplication
from PyQt5.QtCore import QSettings
from openpyxl import load_workbook
from database.database import Database
from services.services import Services
from ui.main_window import MainWindow

class SkuListExportTests(unittest.TestCase):
 @classmethod
 def setUpClass(cls):cls.app=QApplication.instance() or QApplication([])
 def test_export_reuses_database_sheet_and_remembers_directory(self):
  with tempfile.TemporaryDirectory() as tmp:
   db=Database(str(Path(tmp)/'db'));s=Services(db)
   mid=s.create_material(dict(base_code='ABC',version='100',name='SKU名称',description='较长描述'*40,unit='pcs'));s.create_sku(mid)
   db.conn.execute("UPDATE project SET project_code='P01',project_name='项目名称'");db.conn.commit();w=MainWindow(s)
   settings=QSettings('Biong','BOMManager');old=settings.value('export_directory');settings.remove('export_directory')
   try:
    self.assertEqual([a.text() for a in w.sku_menu.actions() if not a.isSeparator()],['新建SKU','SKU筛选','导出SKU清单'])
    self.assertEqual(w.sku_menu.actions()[0].shortcut().toString(),'Ctrl+N')
    target=Path(tmp)/'chosen.xlsx'
    with patch('ui.main_window.QFileDialog.getSaveFileName',return_value=(str(target),'')) as choose,patch('ui.main_window.QMessageBox.information'):
     w.export_sku_list()
     self.assertEqual(Path(choose.call_args.args[2]).name,'P01 项目名称 SKU清单.xlsx')
    with patch('ui.main_window.QFileDialog.getSaveFileName',return_value=('','')) as choose:
     w.export_sku_list();self.assertEqual(Path(choose.call_args.args[2]).parent,Path(tmp))
    whole=Path(tmp)/'all.xlsx';s.export_database(whole);a=load_workbook(target);b=load_workbook(whole)
    try:
     self.assertEqual(a.sheetnames,['SKU清单']);self.assertEqual(list(a.active.values),list(b['SKU清单'].values))
     self.assertEqual(a.active['A2'].alignment.horizontal,'center')
     for row in a.active:
      for cell in row:
       self.assertEqual(cell.border.left.style,'thin');self.assertTrue(cell.alignment.wrap_text)
     for col in a.active.column_dimensions.values():self.assertLessEqual(col.width,60)
    finally:a.close();b.close()
    with patch.object(s,'export_sku_list',side_effect=OSError('不可写')),patch('ui.main_window.QFileDialog.getSaveFileName',return_value=(str(Path(tmp)/'bad.xlsx'),'')),patch('ui.main_window.show_operation_error') as error:
     w.export_sku_list();error.assert_called_once()
   finally:
    if old is None:settings.remove('export_directory')
    else:settings.setValue('export_directory',old)
    w.close();db.close()
