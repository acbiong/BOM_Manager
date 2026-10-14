"""联动导出锁定、反查定位及Excel导出回归。"""
import os
os.environ['QT_QPA_PLATFORM'] = 'offscreen'
import sys
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch
sys.path.insert(0, str(Path(__file__).parents[1] / 'src'))
from PyQt5.QtCore import Qt
from PyQt5.QtWidgets import QApplication, QFileDialog, QMessageBox
from openpyxl import load_workbook
from database.database import Database
from services.services import Services
from ui.main_window import MainWindow, MaterialUsageDialog


class ExportRefreshUsageTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.app = QApplication.instance() or QApplication([])

    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.db = Database(str(Path(self.tmp.name) / 'db'))
        self.s = Services(self.db)
        self.r, self.o, self.p, self.c, self.q = [self.s.create_material(dict(
            base_code=x, version='1', name=x, unit='pcs', description='描述'*80))
            for x in ('R', 'O', 'P', 'C', 'Q')]
        self.sid = self.s.create_sku(self.r)
        self.other_sid = self.s.create_sku(self.o)
        self.s.save_bom(self.p, [dict(material_id=self.c, quantity=1, unit='pcs')])
        for root in (self.r, self.o):
            self.s.save_bom(root, [dict(material_id=self.p, quantity=1, unit='pcs')])
        self.main = MainWindow(self.s)

    def tearDown(self):
        for w in list(self.main._material_bom_windows.values()):
            w.dirty = False
            w.close()
        self.main.close()
        self.db.close()
        self.tmp.cleanup()

    def test_linked_notice_refresh_and_save_export_states(self):
        source = self.main.open_bom_id(self.sid)
        target = self.main.open_bom_id(self.other_sid)
        target.show_item_info(target.root_item.child(0))
        target.fields['quantity'].setValue(3)
        source.show_item_info(source.root_item.child(0).child(0))
        source.fields['quantity'].setValue(2)
        with patch.object(QMessageBox, 'information'):
            self.assertTrue(source.save())
        self.assertTrue(target.refresh_action.isEnabled())
        self.assertFalse(target.export_action.isEnabled())
        target.refresh_external()
        self.assertFalse(target.export_action.isEnabled())
        self.assertTrue(target.save_action.isEnabled())
        self.assertEqual(target.root_item.child(0).text(1), '3')
        self.assertEqual(target.root_item.child(0).child(0).text(1), '2')
        with patch.object(QFileDialog, 'getSaveFileName') as chooser:
            target.export_bom()
            chooser.assert_not_called()
        with patch.object(QMessageBox, 'information'):
            self.assertTrue(target.save())
        self.assertTrue(target.export_action.isEnabled())

    def test_clean_external_refresh_requires_save_and_failed_save_stays_locked(self):
        target = self.main.open_bom_id(self.sid)
        self.assertTrue(target.export_action.isEnabled())
        self.s.save_bom(self.p, [dict(material_id=self.c, quantity=2, unit='pcs')])
        self.assertFalse(target.export_action.isEnabled())
        target.refresh_external()
        self.assertFalse(target.export_action.isEnabled())
        self.assertTrue(target.save_action.isEnabled())
        with patch.object(self.s, 'save_bom', side_effect=ValueError('失败')), patch.object(QMessageBox, 'warning'):
            self.assertFalse(target.save())
        self.assertFalse(target.export_action.isEnabled())
        with patch.object(QMessageBox, 'information'):
            self.assertTrue(target.save())
        self.assertTrue(target.export_action.isEnabled())

    def test_reverse_selects_matching_parent_expands_and_reuses_window(self):
        self.s.save_bom(self.q, [dict(material_id=self.c, quantity=1, unit='pcs')])
        self.s.save_bom(self.r, [dict(material_id=x, quantity=1, unit='pcs') for x in (self.p,self.q)])
        window = self.main.open_bom_id(self.sid)
        window.tree.collapseAll()
        usages = self.s.material_usage(self.c)
        index = next(i for i,u in enumerate(usages) if u['sku_id']==self.sid and u['parent_id']==self.q)
        dialog = MaterialUsageDialog(self.s.material_snapshot(self.c), usages, self.main)
        dialog.table.cellDoubleClicked.emit(index, 0)
        self.assertIs(self.main._bom_windows[self.sid], window)
        self.assertEqual(len(self.main._material_bom_windows), 1)
        selected = window.tree.currentItem()
        self.assertEqual(selected.data(0,Qt.UserRole)['material_id'],self.c)
        self.assertEqual(selected.parent().data(0,Qt.UserRole)['material_id'],self.q)
        self.assertIs(window.current_item, selected)
        self.assertTrue(selected.parent().isExpanded())
        self.assertTrue(window.root_item.isExpanded())
        dialog.close()

    def test_database_sku_format_and_material_only_logs(self):
        path = Path(self.tmp.name)/'all.xlsx'
        self.s.export_database(path)
        wb = load_workbook(path)
        try:
            sheet = wb['SKU清单']
            self.assertLessEqual(max(d.width for d in sheet.column_dimensions.values()),60)
            self.assertGreater(sheet.row_dimensions[2].height,22)
            for row in sheet:
                self.assertEqual(row[0].alignment.horizontal,'center')
                for cell in row:
                    self.assertTrue(cell.alignment.wrap_text)
                    for side in ('left','right','top','bottom'):
                        self.assertEqual(getattr(cell.border,side).color.rgb,'FF000000')
                        self.assertEqual(getattr(cell.border,side).style,'thin')
            self.assertEqual(wb['变更记录'].max_row,6)
            self.assertTrue(all(row[5]=='新建物料' for row in list(wb['变更记录'].values)[1:]))
        finally:
            wb.close()

    def test_bom_sheet_name_is_material_name_with_excel_limits(self):
        for name, expected in [('SKU成品','SKU成品'),('A/B:*?[]\\','A_B______'),('长'*50,'长'*31),('变更记录','变更记录1')]:
            with self.subTest(name=name):
                self.s.update_material(self.r,dict(name=name,unit='pcs'))
                path=Path(self.tmp.name)/'bom.xlsx'
                self.s.export_sku(self.sid,path)
                wb=load_workbook(path)
                try:self.assertEqual(wb.sheetnames,['变更记录',expected])
                finally:wb.close()
