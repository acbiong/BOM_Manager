import os
os.environ['QT_QPA_PLATFORM']='offscreen'
import sys,tempfile,unittest
from pathlib import Path
sys.path.insert(0,str(Path(__file__).parents[1]/'src'))
from PyQt5.QtWidgets import QApplication
from openpyxl import Workbook
from config.config_service import ConfigService,DEFAULT_MATERIAL_UNITS
from database.database import Database
from services.services import Services
from ui.main_window import MainWindow,MaterialWindow,MaterialDialog

class MaterialUnitsTests(unittest.TestCase):
 @classmethod
 def setUpClass(cls):cls.app=QApplication.instance() or QApplication([])
 def setUp(self):
  self.tmp=tempfile.TemporaryDirectory(); self.base=Path(self.tmp.name); self.config=ConfigService(self.base); self.db=Database(str(self.base/'test.db'))
 def tearDown(self):self.db.close(); self.tmp.cleanup()
 def test_defaults_migration_and_restoration_preserve_database(self):
  self.config.path.write_text('[database]\npath = /custom/data.db\n[other]\nkey = keep\n',encoding='utf-8')
  self.assertEqual(self.config.material_units(),list(DEFAULT_MATERIAL_UNITS))
  cfg=self.config.load(); cfg['materials']['units']='箱,包'; self.config.save(cfg)
  self.assertEqual(self.config.material_units(),['箱','包'])
  self.config.restore_material_units(); cfg=self.config.load()
  self.assertEqual(cfg['database']['path'],'/custom/data.db'); self.assertEqual(cfg['other']['key'],'keep')
  cfg.remove_option('materials','units'); self.config.save(cfg); self.assertEqual(self.config.material_units(),list(DEFAULT_MATERIAL_UNITS))
  cfg=self.config.load();cfg['materials']['units']=' , ， ';self.config.save(cfg);self.assertEqual(self.config.material_units(),list(DEFAULT_MATERIAL_UNITS))
 def test_default_new_units_create_and_import(self):
  service=Services(self.db,self.config)
  for index,unit in enumerate(['米','M','克','g']):service.create_material(dict(base_code=f'A{index}',version='1',name='物料',unit=unit))
  workbook=Workbook(); workbook.active.append(['物料基础编码','物料版本号','物料名称','物料规格描述','单位','物料制造商','物料MPN'])
  for index,unit in enumerate(['米','M','克','g']):workbook.active.append([f'B{index}','1','物料','',unit,'厂家',''])
  path=self.base/'units.xlsx';workbook.save(path);workbook.close()
  ok,fail=service.import_materials(path);self.assertEqual(len(ok),4);self.assertEqual(fail,[])
 def test_custom_units_shared_by_ui_create_update_and_import(self):
  cfg=self.config.load();cfg['materials']['units']='箱, 箱，包, M, m';self.config.save(cfg)
  service=Services(self.db,self.config);self.assertEqual(service.UNITS,['箱','包','M','m'])
  main=MainWindow(service); material=MaterialWindow(service)
  try:
   for parent in (main,material):
    dialog=MaterialDialog(parent);self.assertEqual([dialog.unit.itemText(i) for i in range(dialog.unit.count())],service.UNITS);dialog.close()
   mid=service.create_material(dict(base_code='X',version='1',name='物料',unit='箱'))
   service.update_material(mid,dict(name='物料',unit='包'))
   with self.assertRaises(ValueError):service.update_material(mid,dict(name='物料',unit='pcs'))
   with self.assertRaises(ValueError):service.create_material(dict(base_code='Z',version='1',name='物料',unit='pcs'))
   wb=Workbook();wb.active.append(['物料基础编码','物料版本号','物料名称','物料规格描述','单位','物料制造商','物料MPN']);wb.active.append(['OK','1','物料','','箱','厂家','']);wb.active.append(['BAD','1','物料','','pcs','厂家','']);path=self.base/'custom.xlsx';wb.save(path);wb.close()
   ok,fail=service.import_materials(path);self.assertEqual(ok,['OK']);self.assertEqual(fail[0][1],'BAD1')
  finally:material.close();main.close()

 def write_units_workbook(self,rows):
  wb=Workbook();wb.active.append(['物料基础编码','物料版本号','物料名称','物料规格描述','单位','物料制造商','物料MPN'])
  for code,name,unit in rows:wb.active.append([code,'1',name,'',unit,'厂家',''])
  path=self.base/'case_units.xlsx';wb.save(path);wb.close();return path
 def test_import_normalizes_default_unit_case_before_storage(self):
  service=Services(self.db,self.config)
  units=['PCS','pCs','ML','Ml','l','MM','CM','M','m','G','米','克','  PCS  ']
  expected=['pcs','pcs','ml','ml','L','mm','cm','m','m','g','米','克','pcs']
  path=self.write_units_workbook([(f'U{i}','物料',unit) for i,unit in enumerate(units)])
  ok,fail=service.import_materials(path)
  self.assertEqual(len(ok),len(units));self.assertEqual(fail,[])
  for i,unit in enumerate(expected):
   self.assertEqual(self.db.conn.execute('SELECT unit FROM materials WHERE base_code=?',(f'U{i}',)).fetchone()['unit'],unit)
 def test_import_case_only_duplicate_skips_and_overwrite_uses_canonical_unit(self):
  service=Services(self.db,self.config)
  mid=service.create_material(dict(base_code='X',version='1',name='物料',unit='pcs',manufacturer='厂家'))
  path=self.write_units_workbook([('X','物料','PCS')])
  from unittest.mock import Mock
  resolver=Mock(return_value='overwrite')
  ok,fail=service.import_materials(path,conflict_resolver=resolver)
  self.assertEqual(ok,[]);self.assertIn('信息相同',fail[0][3]);resolver.assert_not_called()
  path=self.write_units_workbook([('X','新名称','pCs')])
  ok,fail=service.import_materials(path,conflict_resolver=resolver)
  self.assertEqual(ok,['X']);self.assertEqual(fail,[])
  self.assertEqual(resolver.call_args.args[0]['单位'],'pcs')
  row=service.material_snapshot(mid);self.assertEqual(row['unit'],'pcs');self.assertEqual(row['name'],'新名称')
 def test_import_configured_spelling_wins_and_unknown_units_still_fail(self):
  cfg=self.config.load();cfg['materials']['units']='KG, Box, M, m';self.config.save(cfg)
  service=Services(self.db,self.config)
  path=self.write_units_workbook([('A','物料','kg'),('B','物料','BOX'),('C','物料','m'),('D','物料','pcs'),('E','物料','')])
  ok,fail=service.import_materials(path)
  self.assertEqual(ok,['A','B','C']);self.assertEqual(len(fail),2)
  self.assertIn('单位不合规：pcs',fail[0][3]);self.assertIn('必填项缺失：单位',fail[1][3])
  self.assertEqual([r['unit'] for r in self.db.conn.execute('SELECT unit FROM materials ORDER BY base_code')],['KG','Box','M'])
