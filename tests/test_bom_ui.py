import os
os.environ['QT_QPA_PLATFORM']='offscreen'
import sys, tempfile, unittest
from pathlib import Path
from unittest.mock import patch
sys.path.insert(0,str(Path(__file__).parents[1]/'src'))
from PyQt5.QtWidgets import QApplication, QInputDialog, QMessageBox
from ui.interactive_windows import QMessageBox
from PyQt5.QtCore import Qt
from database.database import Database
from services.services import Services
from ui.main_window import BomWindow
from infrastructure.events import WindowRefreshBus

from ui.main_window import MaterialSelectionDialog
class BomUiTests(unittest.TestCase):
 @classmethod
 def setUpClass(cls): cls.app=QApplication.instance() or QApplication([])
 def setUp(self):
  self.tmp=tempfile.TemporaryDirectory(); self.db=Database(str(Path(self.tmp.name)/'test.db')); self.s=Services(self.db)
  self.ids=[self.s.create_material(dict(base_code=c,version='1',name=c,unit='pcs')) for c in ['ROOT','PARENT','CHILD','OTHER','CYCLE']]
  r,p,c,o,x=self.ids
  self.s.save_bom(r,[dict(id=10,material_id=p,quantity=1,unit='pcs'),dict(id=11,parent_item_id=10,material_id=c,quantity=2.3,unit='pcs')])
  self.s.save_bom(x,[dict(material_id=r,quantity=1,unit='pcs')])
  self.w=BomWindow(self.s,r)
 def tearDown(self):
  self.w.dirty=False; self.w.close(); self.db.conn.close(); self.tmp.cleanup()
 def test_material_update_refresh_and_quantity_cache(self):
  r,p,c,o,x=self.ids; parent=self.w.root_item.child(0)
  self.assertEqual(parent.text(1),'1')
  self.w.show_item_info(parent); self.w.fields['quantity'].setValue(2.34)
  self.assertEqual(parent.text(1),'2.34')
  self.s.update_material(p,dict(name='修改后的名称',unit='pcs'))
  self.assertTrue(self.w.refresh_action.isEnabled()); self.assertEqual(self.w.statusBar().currentMessage(),''); self.assertEqual(self.w.notice.text(),'BOM有更新，请点击刷新或重新打开窗口')
  self.w.refresh_external()
  self.assertEqual(self.w.statusBar().currentMessage(),''); self.assertEqual(self.w.notice.text(),'')
  self.assertFalse(self.w.refresh_action.isEnabled()); self.assertTrue(self.w.save_action.isEnabled())
  parent=self.w.root_item.child(0); self.assertIn('修改后的名称',parent.text(0)); self.assertEqual(parent.text(1),'2.34')
  self.w.show_item_info(parent); self.assertEqual(self.w.fields['quantity'].text(),'2.34')
  self.s.update_material(o,dict(name='无关',unit='pcs')); self.assertFalse(self.w.refresh_action.isEnabled())
 def test_sibling_and_child_candidates_block_cycles(self):
  parent=self.w.root_item.child(0); child=parent.child(0)
  for selected,is_child in [(child,False),(parent,True)]:
   with patch.object(MaterialSelectionDialog,'getItem',return_value=('',False)) as dialog:
    self.w.add_relative(selected.data(0,Qt.UserRole),is_child)
    options=dialog.call_args.args[3]
    self.assertEqual([s.split()[0] for s in options],['OTHER1'])
 def test_new_child_quantity_and_export(self):
  parent=self.w.root_item.child(0)
  with patch.object(MaterialSelectionDialog,'getItem',return_value=('OTHER1 OTHER',True)):
   self.w.add_relative(parent.data(0,Qt.UserRole),True)
  self.assertEqual(self.w.tree.currentItem().text(1),'1'); self.assertEqual(self.w.fields['quantity'].text(),'1')
  self.w.fields['quantity'].setValue(3.27)
  with patch.object(QMessageBox,'information'):
   self.w.save()
  self.assertEqual([row['quantity'] for row in self.s.bom(self.ids[0]) if row['material_id']==self.ids[3]],[3.27])
  self.w.load(); self.assertEqual(self.w.root_item.child(0).child(1).text(1),'3.27')
  from openpyxl import load_workbook
  sid=self.s.create_sku(self.ids[0]); path=Path(self.tmp.name)/'bom.xlsx'; self.s.export_sku(sid,path)
  wb=load_workbook(path); self.assertEqual(wb[self.s.material_snapshot(self.ids[0])['name']].cell(2,10).number_format,'0.############'); wb.close()
 def test_shared_sub_bom_save_twice_and_edit(self):
  # CHILD已有下阶物料，添加为ROOT的另一个子项时再次展开同一关系。
  r,p,c,o,x=self.ids
  self.s.save_bom(c,[dict(material_id=o,quantity=1,unit='pcs')]); self.w.load()
  with patch.object(MaterialSelectionDialog,'getItem',return_value=('CHILD1 CHILD',True)):
   self.w.add_relative(self.w.root_item.data(0,Qt.UserRole),True)
  self.assertEqual(self.w.tree.currentItem().childCount(),1)
  with patch.object(QMessageBox,'information'), patch.object(QMessageBox,'warning') as warning:
   self.assertTrue(self.w.save()); self.assertTrue(self.w.save()); warning.assert_not_called()
  count=self.db.conn.execute('SELECT count(*) FROM material_bom_items WHERE parent_material_id=? AND child_material_id=?',(c,o)).fetchone()[0]
  self.assertEqual(count,1)
  self.w.show_item_info(self.w.root_item.child(0)); self.w.fields['quantity'].setValue(4.2)
  with patch.object(QMessageBox,'information'): self.assertTrue(self.w.save())
  self.assertEqual(self.s.bom(r)[0]['quantity'],4.2)
 def test_save_failure_keeps_edit_and_close_cancelled(self):
  import sqlite3
  from PyQt5.QtGui import QCloseEvent
  self.w.show_item_info(self.w.root_item.child(0)); self.w.fields['quantity'].setValue(5.6)
  with patch.object(self.s,'save_bom',side_effect=sqlite3.IntegrityError('constraint test')), patch.object(QMessageBox,'warning') as warning:
   self.assertFalse(self.w.save()); warning.assert_called_once()
   self.assertTrue(self.w.dirty); self.assertFalse(self.w._saving); self.assertTrue(self.w.save_action.isEnabled())
   self.assertEqual(self.w.root_item.child(0).text(1),'5.6')
   with patch('ui.main_window.QMessageBox') as box:
    save,discard,cancel=object(),object(),object()
    box.return_value.addButton.side_effect=[save,discard,cancel]; box.return_value.clickedButton.return_value=save
    event=QCloseEvent(); self.w.closeEvent(event); self.assertFalse(event.isAccepted())
 def test_conflicting_shared_edits_do_not_write(self):
  r,p,c,o,x=self.ids
  before=[dict(row) for row in self.db.conn.execute('SELECT * FROM material_bom_items')]
  rows=[dict(material_id=p,parent_material_id=r,quantity=q,unit='pcs') for q in (2,3)]
  with self.assertRaisesRegex(ValueError,'共享子BOM'):
   self.s.save_bom(r,rows)
  self.assertEqual(before,[dict(row) for row in self.db.conn.execute('SELECT * FROM material_bom_items')])
 def test_unchanged_shared_occurrence_does_not_override_edit(self):
  r,p,c,o,x=self.ids
  for quantities in [(1,2),(3,2)]:
   self.s.save_bom(r,[dict(material_id=p,parent_material_id=r,quantity=q,unit='pcs') for q in quantities])
  self.assertEqual(self.s.bom(r)[0]['quantity'],3)
 def test_sqlite_failure_rolls_back_database(self):
  r,p,c,o,x=self.ids
  before=[dict(row) for row in self.db.conn.execute('SELECT * FROM material_bom_items')]
  self.db.conn.execute("CREATE TEMP TRIGGER fail_bom_update BEFORE UPDATE ON material_bom_items BEGIN SELECT RAISE(ABORT,'模拟数据库写入失败'); END")
  self.w.show_item_info(self.w.root_item.child(0)); self.w.fields['quantity'].setValue(6.7)
  with patch.object(QMessageBox,'warning') as warning:
   self.assertFalse(self.w.save()); self.assertIn('模拟数据库写入失败',warning.call_args.args[2])
  self.assertEqual(before,[dict(row) for row in self.db.conn.execute('SELECT * FROM material_bom_items')])
  self.assertEqual(self.w.root_item.child(0).text(1),'6.7')
  self.db.conn.execute('DROP TRIGGER fail_bom_update')
  with patch.object(QMessageBox,'information'): self.assertTrue(self.w.save())
 def test_save_preserves_relation_id_and_alternates(self):
  r,p,c,o,x=self.ids; relation=self.s.bom(r)[0]['id']
  self.s.add_alternate(relation,o)
  with patch.object(QMessageBox,'information'): self.assertTrue(self.w.save())
  self.assertEqual(self.s.bom(r)[0]['id'],relation)
  self.assertEqual(self.s.alternates(relation)[0]['material_id'],o)
 def test_shared_branches_have_distinct_nodes_and_full_children(self):
  r,p,c,o,x=self.ids
  self.s.save_bom(c,[dict(material_id=o,quantity=1,unit='pcs')]); self.w.load()
  with patch.object(MaterialSelectionDialog,'getItem',return_value=('CHILD1 CHILD',True)):
   self.w.add_relative(self.w.root_item.data(0,Qt.UserRole),True)
  with patch.object(QMessageBox,'information'): self.assertTrue(self.w.save())
  first=self.w.root_item.child(0).child(0); second=self.w.root_item.child(1)
  self.assertEqual(first.childCount(),1); self.assertEqual(second.childCount(),1)
  a,b=first.child(0),second.child(0)
  self.assertEqual(a.data(0,Qt.UserRole)['id'],b.data(0,Qt.UserRole)['id'])
  self.assertNotEqual(a.data(0,Qt.UserRole)['node_id'],b.data(0,Qt.UserRole)['node_id'])
  self.assertIs(self.w.by[a.data(0,Qt.UserRole)['node_id']],a)
  self.assertIs(self.w.by[b.data(0,Qt.UserRole)['node_id']],b)
  self.s.add_alternate(a.data(0,Qt.UserRole)['id'],p)
  self.w.show_item_info(b); self.assertEqual(self.w.alt_table.rowCount(),1)
  with patch.object(QMessageBox,'information'): self.assertTrue(self.w.save())
  self.assertEqual(self.w.root_item.child(1).childCount(),1)
 def test_cycle_guard_is_path_local(self):
  r,p,c,o,x=self.ids
  self.db.conn.execute('INSERT INTO material_bom_items(parent_material_id,child_material_id,quantity,unit) VALUES(?,?,1,?)',(c,r,'pcs')); self.db.conn.commit()
  self.assertEqual([row['material_id'] for row in self.s.bom(r)],[p,c])
 def test_single_windows_reactivate_and_reopen(self):
  from ui.main_window import MainWindow
  main=MainWindow(self.s); one=self.s.create_sku(self.ids[0]); two=self.s.create_sku(self.ids[4])
  try:
   main.open_bom_id(one); original=main._bom_windows[one]
   original.show_item_info(original.root_item.child(0)); original.fields['quantity'].setValue(7.8)
   original.showMinimized(); main.open_bom_id(one)
   self.assertIs(main._bom_windows[one],original); self.assertFalse(original.isMinimized())
   self.assertEqual(original.root_item.child(0).text(1),'7.8')
   main.open_bom_id(two); self.assertEqual(len(main._bom_windows),2)
   with patch('ui.main_window.QMessageBox') as box:
    save,discard,cancel=object(),object(),object()
    box.return_value.addButton.side_effect=[save,discard,cancel]; box.return_value.clickedButton.return_value=cancel
    self.assertFalse(original.close())
   main.open_bom_id(one); self.assertIs(main._bom_windows[one],original)
   original.dirty=False; original.close(); self.assertNotIn(one,main._bom_windows)
   main.open_bom_id(one); self.assertIsNot(main._bom_windows[one],original)
   main.show_materials(); materials=main._materials; materials.showMinimized(); main.show_materials()
   self.assertIs(main._materials,materials); self.assertFalse(materials.isMinimized())
   materials.close(); self.assertIsNone(main._materials)
   main.show_materials(); self.assertIsNot(main._materials,materials)
  finally:
   for window in list(main._bom_windows.values()): window.dirty=False; window.close()
   if main._materials is not None: main._materials.close()
   main.close()
 def test_delete_shared_relation_save_and_other_sku_refresh(self):
  from ui.main_window import MaterialWindow
  r,p,c,o,x=self.ids
  self.s.save_bom(c,[dict(material_id=o,quantity=1,unit='pcs')]); self.w.load()
  with patch.object(MaterialSelectionDialog,'getItem',return_value=('CHILD1 CHILD',True)):
   self.w.add_relative(self.w.root_item.data(0,Qt.UserRole),True)
  with patch.object(QMessageBox,'information'): self.assertTrue(self.w.save())
  other=BomWindow(self.s,x); materials=MaterialWindow(self.s)
  try:
   child=self.w.root_item.child(1).child(0)
   self.w.remove_relation(child)
   self.assertTrue(self.w.root_item.child(0).child(0).child(0).data(0,Qt.UserRole)['_deleted'])
   self.assertTrue(self.w.root_item.child(1).child(0).data(0,Qt.UserRole)['_deleted'])
   self.assertEqual(len(self.s.bom(c)),1)  # 未保存不改数据库
   self.assertIs(self.w.tree.currentItem(),child)
   self.w.refresh_external()  # 未保存删除在外部刷新后仍然保留
   self.assertTrue(self.w.root_item.child(1).child(0).data(0,Qt.UserRole)['_deleted'])
   with patch.object(QMessageBox,'information'): self.assertTrue(self.w.save())
   self.assertEqual(self.s.bom(c),[])
   self.assertTrue(other.refresh_action.isEnabled()); other.refresh_external()
   self.assertFalse(any(row['material_id']==o for row in self.s.bom(x)))
   flags={materials.t.item(i,0).text():materials.t.item(i,9).text() for i in range(materials.t.rowCount())}
   self.assertEqual(flags['CHILD'],'N')
   self.w.load(); self.assertEqual(self.w.root_item.child(1).childCount(),0)
  finally:
   other.dirty=False; other.close(); materials.close()
 def test_context_menu_includes_material_edit(self):
  from PyQt5.QtWidgets import QMenu
  from PyQt5.QtCore import QPoint
  captured=[]
  def inspect(menu,*args):captured.extend(a.text() for a in menu.actions() if not a.isSeparator());return None
  with patch.object(self.w.tree,'itemAt',return_value=self.w.root_item.child(0)),patch.object(QMenu,'exec_',inspect):self.w.context_menu(QPoint())
  self.assertEqual(captured,['编辑此物料','物料反查','替换物料','物料版本更新','新增同级物料','新增子级物料','新增替代料','导入BOM','导出BOM结构','设置颜色','复制物料','剪切物料','物料上移','物料下移','删除物料'])
 def assert_level_background(self,brush):
  self.assertEqual(brush.style(),Qt.SolidPattern)
  color=brush.color();self.assertEqual(color.red(),color.green());self.assertEqual(color.green(),color.blue())
 def test_idle_flags_recursive_shared_skus_and_red_rows(self):
  from ui.main_window import MaterialWindow
  from PyQt5.QtGui import QColor
  r,p,c,o,x=self.ids
  materials=MaterialWindow(self.s)
  def flags(): return {row['id']:row['idle'] for row in self.s.materials()}
  def assert_color(code,idle):
   index=next(i for i in range(materials.t.rowCount()) if materials.t.item(i,0).text()==code)
   self.assertEqual(materials.t.item(index,10).text(),idle)
   for col in range(materials.t.columnCount()):
    self.assertEqual(materials.t.item(index,col).background().color()==QColor('red'),idle=='Y')
  try:
   self.assertEqual(flags(),{r:'N',p:'N',c:'N',o:'Y',x:'Y'})
   assert_color('CHILD','N')
   first=self.s.create_sku(r)
   self.assertEqual(flags(),{r:'N',p:'N',c:'N',o:'Y',x:'Y'})
   assert_color('CHILD','N'); assert_color('OTHER','Y'); assert_color('ROOT','N')
   second=self.s.create_sku(x)
   self.s.save_bom(r,[])
   self.assertEqual(flags(),{r:'N',p:'Y',c:'Y',o:'Y',x:'N'})
   assert_color('PARENT','Y'); assert_color('CHILD','Y')
   self.s.delete_sku(first)
   self.assertEqual(flags()[r],'N')  # 仍在另一个SKU下
   self.s.delete_sku(second)
   self.assertEqual(flags(),{r:'Y',p:'Y',c:'Y',o:'Y',x:'Y'})
   assert_color('ROOT','Y')
   self.s.save_bom(p,[])
   self.assertEqual(flags()[c],'Y'); assert_color('CHILD','Y')
  finally: materials.close()
 def test_idle_query_handles_legacy_cycles_and_filter(self):
  r,p,c,o,x=self.ids; self.s.create_sku(r)
  self.db.conn.execute('INSERT INTO material_bom_items(parent_material_id,child_material_id,quantity,unit) VALUES(?,?,1,?)',(c,r,'pcs')); self.db.conn.commit()
  self.assertEqual(self.s.materials('CHILD')[0]['idle'],'N')
  self.assertEqual(self.s.materials('OTHER')[0]['idle'],'Y')
 def test_material_export_includes_idle_flag(self):
  from openpyxl import load_workbook
  self.s.create_sku(self.ids[0])
  for method in (self.s.export_materials,self.s.export_database):
   path=Path(self.tmp.name)/'materials.xlsx'; method(path)
   wb=load_workbook(path); rows=list(wb['物料库'].values)
   self.assertEqual(rows[0][-1],'是否为呆滞料')
   flags={row[1]:row[-1] for row in rows[1:]}
   self.assertEqual(flags['CHILD'],'N'); self.assertEqual(flags['OTHER'],'Y'); wb.close()
 def test_unreferenced_non_sku_bom_is_idle_and_editable(self):
  from ui.main_window import MainWindow,MaterialWindow
  from PyQt5.QtWidgets import QMenu
  from PyQt5.QtCore import QPoint
  from PyQt5.QtGui import QColor
  r,p,c,o,x=self.ids
  main=MainWindow(self.s); main.show_materials(); materials=main._materials
  try:
   flags={row['id']:row for row in self.s.materials()}
   self.assertEqual(flags[x]['has_bom'],'Y'); self.assertEqual(flags[x]['idle'],'Y')
   row=next(i for i,m in enumerate(self.s.materials()) if m['id']==x)
   self.assertEqual(materials.t.item(row,10).text(),'Y')
   for col in range(materials.t.columnCount()):self.assertEqual(materials.t.item(row,col).background().color(),QColor('red'))
   def choose(menu,pos):return next(action for action in menu.actions() if action.text()=='编辑物料BOM结构')
   with patch.object(materials.t,'rowAt',return_value=row),patch.object(QMenu,'exec_',choose):materials.menu(QPoint())
   window=main._material_bom_windows[x]
   self.assertEqual(window.root_item.data(0,Qt.UserRole)['material_id'],x)
   self.assertEqual(window.root_item.child(0).child(0).child(0).data(0,Qt.UserRole)['material_id'],c)
   observer=BomWindow(self.s,r)
   try:
    window.show_item_info(window.root_item.child(0).child(0).child(0)); window.fields['quantity'].setValue(8.5)
    with patch.object(QMessageBox,'information'):self.assertTrue(window.save())
    self.assertEqual(self.s.bom(p)[0]['quantity'],8.5); self.assertTrue(observer.refresh_action.isEnabled())
    observer.refresh_external(); self.assertEqual(observer.root_item.child(0).child(0).text(1),'8.5')
    self.assertEqual({m['id']:m['idle'] for m in self.s.materials()}[x],'Y')
   finally:observer.close()
  finally:
   for window in list(main._material_bom_windows.values()):window.dirty=False; window.close()
   materials.close(); main.close()
 def test_material_and_sku_entries_share_one_bom_window(self):
  from ui.main_window import MainWindow
  r,p,c,o,x=self.ids; sid=self.s.create_sku(r); main=MainWindow(self.s)
  try:
   main.show_materials(); first=main._materials.open_material_bom(r)
   main.open_bom_id(sid); self.assertIs(main._bom_windows[sid],first)
   self.assertIs(main._materials.open_material_bom(r),first)
   self.assertEqual(len(main._material_bom_windows),1)
   first.close(); self.assertNotIn(sid,main._bom_windows); self.assertNotIn(r,main._material_bom_windows)
   main.open_bom_id(sid); second=main._bom_windows[sid]
   self.assertIs(main._materials.open_material_bom(r),second)
  finally:
   for window in list(main._material_bom_windows.values()):window.dirty=False; window.close()
   main._materials.close(); main.close()
 def test_inline_alternates_add_remove_and_save(self):
  from PyQt5.QtWidgets import QMenu
  from PyQt5.QtCore import QPoint
  r,p,c,o,x=self.ids; node=self.w.root_item.child(0); relation=node.data(0,Qt.UserRole)['id']
  self.w.tree.setCurrentItem(node); self.w.show_item_info(node)
  with patch.object(MaterialSelectionDialog,'getItem',return_value=('OTHER1 OTHER',True)):self.w.add_alternate()
  self.assertEqual(self.w.alt_table.rowCount(),1); self.assertTrue(self.w.dirty)
  self.assertEqual(self.s.alternates(relation),[])
  self.w.show_item_info(self.w.root_item); self.w.show_item_info(node)
  self.assertEqual(self.w.alt_table.rowCount(),1)
  self.w.refresh_external(); node=self.w.root_item.child(0); self.w.show_item_info(node)
  self.assertEqual(self.w.alt_table.rowCount(),1)
  with patch.object(QMessageBox,'information'):self.assertTrue(self.w.save())
  self.assertEqual(self.s.alternates(relation)[0]['material_id'],o)
  node=self.w.root_item.child(0); self.w.tree.setCurrentItem(node); self.w.show_item_info(node)
  def choose(menu,pos):
   self.assertEqual([a.text() for a in menu.actions()],['切换成主物料','','删除替代料']); self.assertTrue(menu.actions()[1].isSeparator()); return menu.actions()[2]
  with patch.object(self.w.alt_table,'rowAt',return_value=0),patch.object(QMenu,'exec_',choose),patch.object(QMessageBox,'question',return_value=QMessageBox.Yes):self.w.alternate_menu(QPoint())
  self.assertEqual(self.w.alt_table.rowCount(),0); self.assertEqual(len(self.s.alternates(relation)),1)
  with patch.object(QMessageBox,'information'):self.assertTrue(self.w.save())
  self.assertEqual(self.s.alternates(relation),[])
 def test_new_unsaved_node_alternates_persist_transactionally(self):
  r,p,c,o,x=self.ids
  with patch.object(MaterialSelectionDialog,'getItem',return_value=('OTHER1 OTHER',True)):self.w.add_relative(self.w.root_item.data(0,Qt.UserRole),True)
  with patch.object(MaterialSelectionDialog,'getItem',return_value=('CHILD1 CHILD',True)):self.w.add_alternate()
  self.assertEqual(self.w.alt_table.rowCount(),1)
  self.db.conn.execute("CREATE TEMP TRIGGER fail_alt BEFORE INSERT ON alternates BEGIN SELECT RAISE(ABORT,'替代料保存失败'); END")
  with patch.object(QMessageBox,'warning'):self.assertFalse(self.w.save())
  self.assertFalse(any(row['material_id']==o for row in self.s.bom(r)))
  self.assertTrue(self.w.dirty); self.assertEqual(self.w.alt_table.rowCount(),1)
  self.db.conn.execute('DROP TRIGGER fail_alt')
  with patch.object(QMessageBox,'information'):self.assertTrue(self.w.save())
  relation=next(row['id'] for row in self.s.bom(r) if row['material_id']==o)
  self.assertEqual(self.s.alternates(relation)[0]['material_id'],c)
 def test_notice_in_status_bar_and_material_menu_buttons(self):
  from ui.main_window import MaterialWindow
  from PyQt5.QtWidgets import QToolBar
  self.assertIsNone(self.w.menuBar().cornerWidget(Qt.TopRightCorner))
  self.assertTrue(self.w.statusBar().isAncestorOf(self.w.notice)); self.assertIn('red',self.w.notice.styleSheet())
  actions=[a.text() for bar in self.w.findChildren(QToolBar) for a in bar.actions()]
  self.assertNotIn('替代料',actions)
  self.w.on_external_change({'type':'bom_changed','root':self.ids[0]})
  self.assertEqual(self.w.notice.text(),'BOM有更新，请点击刷新或重新打开窗口'); self.assertEqual(self.w.statusBar().currentMessage(),'')
  self.w.refresh_external(); self.assertEqual(self.w.notice.text(),'')
  materials=MaterialWindow(self.s)
  try:
   actions=materials.menuBar().actions()
   self.assertEqual([a.text() for a in actions],['新建物料','导入物料','导出物料库','导出模板','查找物料'])
   self.assertTrue(all(a.menu() is None for a in actions))
  finally:materials.close()
 def test_alternate_notifications_follow_owning_bom(self):
  r,p,c,o,x=self.ids
  # 两个SKU均直接包含PARENT，但ROOT->PARENT与CYCLE->PARENT属于各自的BOM。
  subtree=[dict(row) for row in self.s.bom(p)]
  self.s.save_bom(x,[dict(material_id=p,parent_material_id=x,quantity=1,unit='pcs')]+subtree)
  self.s.create_sku(r); self.s.create_sku(x)
  other=BomWindow(self.s,x); events=[]; WindowRefreshBus.subscribe(events.append)
  try:
   node=self.w.root_item.child(0); self.w.tree.setCurrentItem(node); self.w.show_item_info(node)
   with patch.object(MaterialSelectionDialog,'getItem',return_value=('OTHER1 OTHER',True)):self.w.add_alternate()
   with patch.object(QMessageBox,'information'):self.assertTrue(self.w.save())
   self.assertEqual({e['root'] for e in events if e['type']=='bom_changed'},{r})
   self.assertFalse(other.refresh_action.isEnabled()); self.assertEqual(other.notice.text(),'')
   self.assertEqual(self.s.alternates(self.s.bom(x)[0]['id']),[])
   self.assertEqual(self.s.alternates(self.s.bom(r)[0]['id'])[0]['material_id'],o)
   # 修改共享PARENT下CHILD的替代料，真正变更的是PARENT的BOM。
   events.clear(); leaf=self.w.root_item.child(0).child(0); self.w.tree.setCurrentItem(leaf); self.w.show_item_info(leaf)
   with patch.object(MaterialSelectionDialog,'getItem',return_value=('OTHER1 OTHER',True)):self.w.add_alternate()
   with patch.object(QMessageBox,'information'):self.assertTrue(self.w.save())
   self.assertEqual({e['root'] for e in events if e['type']=='bom_changed'},{p})
   self.assertTrue(other.refresh_action.isEnabled()); other.refresh_external()
   other.show_item_info(other.root_item.child(0).child(0)); self.assertEqual(other.alt_table.item(0,0).text(),'OTHER')
   self.assertFalse(self.w.refresh_action.isEnabled()); self.assertEqual(self.w.notice.text(),'')
   # 未修改再次保存、或重复提交相同替代料，不应产生无关更新事件。
   events.clear()
   with patch.object(QMessageBox,'information'):self.assertTrue(self.w.save())
   self.assertEqual(events,[]); self.assertFalse(other.refresh_action.isEnabled())
   leaf=self.w.root_item.child(0).child(0); self.w._alternate_edits[self.w.alternate_key(leaf)]=[o]
   with patch.object(QMessageBox,'information'):self.assertTrue(self.w.save())
   self.assertEqual(events,[])
   # 删除共享子BOM替代料同样通知其他SKU，并同步空列表。
   leaf=self.w.root_item.child(0).child(0)
   self.w._alternate_edits[self.w.alternate_key(leaf)]=[]
   with patch.object(QMessageBox,'information'):self.assertTrue(self.w.save())
   self.assertEqual({e['root'] for e in events if e['type']=='bom_changed'},{p})
   other.refresh_external(); other.show_item_info(other.root_item.child(0).child(0)); self.assertEqual(other.alt_table.rowCount(),0)
  finally:
   WindowRefreshBus.unsubscribe(events.append); other.close()
 def test_refresh_preserves_changed_position_and_merges_external_fields(self):
  r,p,c,o,x=self.ids
  parent=self.w.root_item.child(0); self.w.show_item_info(parent)
  self.w.fields['quantity'].setValue(4.5); self.w.fields['position'].setText('本地位置'); self.w.fields['remark'].setText('本地备注')
  rows=[dict(row) for row in self.s.bom(r)]
  rows[0]['quantity']=9.0  # 同字段冲突保留本地值
  rows[1]['quantity']=6.7; rows[1]['remark']='外部下阶备注'
  self.s.save_bom(r,rows)
  self.assertTrue(self.w.refresh_action.isEnabled()); self.w.refresh_external()
  parent=self.w.root_item.child(0); self.w.show_item_info(parent)
  self.assertEqual((parent.text(1),parent.text(3),parent.data(0,Qt.UserRole)['remark']),('4.5','本地位置','本地备注'))
  self.assertEqual(self.w.fields['position'].text(),'本地位置')
  self.assertEqual((parent.child(0).text(1),parent.child(0).data(0,Qt.UserRole)['remark']),('6.7','外部下阶备注'))
  self.assertTrue(self.w.save_action.isEnabled()); self.assertIn('保留本地',self.w.statusBar().currentMessage())
  self.w.refresh_external(); parent=self.w.root_item.child(0)
  self.assertEqual(parent.text(3),'本地位置'); self.assertEqual(parent.child(0).text(1),'6.7')
  with patch.object(QMessageBox,'information'):self.assertTrue(self.w.save())
  saved=self.s.bom(r); self.assertEqual(saved[0]['position'],'本地位置'); self.assertEqual(saved[1]['quantity'],6.7)
 def test_refresh_keeps_unsaved_nodes_descendants_and_alternates(self):
  r,p,c,o,x=self.ids
  with patch.object(MaterialSelectionDialog,'getItem',return_value=('OTHER1 OTHER',True)):self.w.add_relative(self.w.root_item.data(0,Qt.UserRole),True)
  self.w.fields['position'].setText('新增位置')
  with patch.object(MaterialSelectionDialog,'getItem',return_value=('CHILD1 CHILD',True)):self.w.add_alternate()
  added=self.w.tree.currentItem()
  with patch.object(MaterialSelectionDialog,'getItem',return_value=('CHILD1 CHILD',True)):self.w.add_relative(added.data(0,Qt.UserRole),True)
  self.w.fields['quantity'].setValue(3.2)
  rows=[dict(row) for row in self.s.bom(p)]; rows[0]['quantity']=8.1; self.s.save_bom(p,rows)
  self.w.refresh_external(); self.w.refresh_external()
  self.assertEqual(self.w.root_item.child(0).child(0).text(1),'8.1')
  added=self.w.root_item.child(1); self.assertEqual(added.text(3),'新增位置'); self.assertEqual(added.child(0).text(1),'3.2')
  self.w.show_item_info(added); self.assertEqual(self.w.alt_table.rowCount(),1)
  self.assertEqual(len(self.s.bom(r)),2)  # 刷新没有提前提交新增节点
  with patch.object(QMessageBox,'information'):self.assertTrue(self.w.save())
  relation=next(row['id'] for row in self.s.bom(r) if row['material_id']==o)
  self.assertEqual(self.s.alternates(relation)[0]['material_id'],c)
  self.assertEqual(self.s.bom(o)[0]['quantity'],3.2)
 def test_refresh_keeps_local_deletion_but_accepts_external_addition(self):
  r,p,c,o,x=self.ids
  self.w.remove_relation(self.w.root_item.child(0).child(0))
  rows=[dict(row) for row in self.s.bom(p)]+[dict(material_id=o,parent_material_id=p,quantity=2,unit='pcs')]
  self.s.save_bom(p,rows); self.w.refresh_external()
  parent=self.w.root_item.child(0)
  self.assertEqual(parent.childCount(),2); self.assertTrue(parent.child(0).data(0,Qt.UserRole)['_deleted']); self.assertEqual(parent.child(1).data(0,Qt.UserRole)['material_id'],o)
  with patch.object(QMessageBox,'information'):self.assertTrue(self.w.save())
  self.assertEqual([row['material_id'] for row in self.s.bom(p)],[o])
 def test_external_delete_does_not_discard_edited_local_subtree(self):
  r,p,c,o,x=self.ids
  self.w.show_item_info(self.w.root_item.child(0).child(0)); self.w.fields['remark'].setText('尚未保存')
  self.s.save_bom(r,[]); self.w.refresh_external()
  self.assertEqual(self.w.root_item.child(0).child(0).data(0,Qt.UserRole)['remark'],'尚未保存')
  self.assertEqual(self.w.root_item.child(0).child(0).text(4),'尚未保存\n有其他BOM更新导入此变更')
  self.assertIn('冲突',self.w.statusBar().currentMessage()); self.assertTrue(self.w.dirty)
 def test_sku_menu_material_button_opens_existing_library(self):
  from ui.main_window import MainWindow
  main=MainWindow(self.s)
  try:
   actions={a.text():a for a in main.menuBar().actions()}
   self.assertIn('物料库',actions); self.assertIsNone(actions['物料库'].menu()); self.assertNotIn('刷新',actions)
   actions['物料库'].trigger(); first=main._materials; actions['物料库'].trigger(); self.assertIs(main._materials,first)
  finally:
   if main._materials is not None:main._materials.close()
   main.close()
 def test_refresh_retains_descendant_alternate_when_parent_removed_externally(self):
  r,p,c,o,x=self.ids; leaf=self.w.root_item.child(0).child(0)
  self.w.tree.setCurrentItem(leaf); self.w.show_item_info(leaf)
  with patch.object(MaterialSelectionDialog,'getItem',return_value=('OTHER1 OTHER',True)):self.w.add_alternate()
  self.s.save_bom(r,[]); self.w.refresh_external()
  leaf=self.w.root_item.child(0).child(0); self.w.show_item_info(leaf)
  self.assertEqual(self.w.alt_table.rowCount(),1); self.assertIn('冲突',self.w.statusBar().currentMessage())
 def test_pending_add_delete_colors_and_commit(self):
  from PyQt5.QtGui import QColor
  r,p,c,o,x=self.ids
  # 新增有子BOM的物料：自身和递归子孙全部绿色。
  self.s.save_bom(o,[dict(material_id=c,quantity=1,unit='pcs')])
  with patch.object(MaterialSelectionDialog,'getItem',return_value=('OTHER1 OTHER',True)):self.w.add_relative(self.w.root_item.data(0,Qt.UserRole),True)
  added=self.w.tree.currentItem()
  for node in [added,added.child(0)]:
   for column in range(self.w.tree.columnCount()):self.assertEqual(node.foreground(column).color(),QColor('green'))
  # 待删除物料及子孙保留显示，但不提前写入数据库。
  removed=self.w.root_item.child(0); self.w.remove_relation(removed)
  self.assertEqual(self.w.root_item.childCount(),2); self.assertEqual(removed.childCount(),1)
  for node in [removed,removed.child(0)]:
   for column in range(self.w.tree.columnCount()):
    self.assertEqual(node.foreground(column).color(),QColor('red')); self.assertTrue(node.font(column).strikeOut())
  self.assertEqual(self.s.bom(r)[0]['material_id'],p)
  self.w.show_item_info(removed); self.assertFalse(self.w.fields['quantity'].isEnabled()); self.assertFalse(self.w.fields['remark'].isEnabled())
  self.w.refresh_external()
  removed=self.w.root_item.child(0); self.assertTrue(removed.font(0).strikeOut()); self.assertEqual(self.w.root_item.child(1).foreground(0).color(),QColor('green'))
  with patch.object(QMessageBox,'information'):self.assertTrue(self.w.save())
  self.assertEqual(self.w.root_item.childCount(),1); added=self.w.root_item.child(0)
  self.assertEqual(added.data(0,Qt.UserRole)['material_id'],o)
  for node in [added,added.child(0)]:
   for column in range(self.w.tree.columnCount()):
    self.assertEqual(node.foreground(column).style(),Qt.NoBrush); self.assert_level_background(node.background(column))
    self.assertFalse(node.font(column).strikeOut()); self.assertFalse(node.font(column).bold())
 def test_modified_row_yellow_only_changed_cells_bold(self):
  from PyQt5.QtGui import QColor
  node=self.w.root_item.child(0); self.w.show_item_info(node)
  self.w.fields['quantity'].setValue(5.4); self.w.fields['remark'].setText('真实修改备注')
  for column in range(self.w.tree.columnCount()):
   self.assertEqual(node.foreground(column).color(),QColor('#9a6700'))
   self.assertEqual(node.font(column).bold(),column in (1,4))
  self.w.fields['quantity'].setValue(1.0); self.assertFalse(node.font(1).bold())
  self.w.fields['remark'].setText('')
  self.assertEqual(node.foreground(0).style(),Qt.NoBrush)
 def test_external_update_purple_note_not_persisted(self):
  from PyQt5.QtGui import QColor
  r,p,c,o,x=self.ids; rows=[dict(row) for row in self.s.bom(p)]
  rows[0]['quantity']=9.2; rows[0]['remark']='数据库真实备注'; self.s.save_bom(p,rows)
  self.w.refresh_external(); node=self.w.root_item.child(0).child(0)
  for column in range(self.w.tree.columnCount()):self.assertEqual(node.background(column).color(),QColor('#4a148c'))
  self.assertIn('有其他BOM更新导入此变更',node.text(4)); self.assertTrue(self.w.save_action.isEnabled())
  self.w.show_item_info(node); self.assertEqual(self.w.fields['remark'].text(),'数据库真实备注')
  self.w.fields['position'].setText('本地位置'); self.w.refresh_external(); self.w.refresh_external()
  node=self.w.root_item.child(0).child(0)
  self.assertIn('有其他BOM更新导入此变更',node.text(4)); self.assertEqual(node.text(3),'本地位置')
  self.assertEqual(node.background(0).color(),QColor('#4a148c')); self.assertEqual(node.foreground(0).color(),QColor('#9a6700'))
  self.assertTrue(node.font(3).bold()); self.assertFalse(node.font(4).bold())
  with patch.object(QMessageBox,'information'):self.assertTrue(self.w.save())
  saved=self.s.bom(p)[0]; self.assertEqual(saved['remark'],'数据库真实备注'); self.assertEqual(saved['position'],'本地位置')
  node=self.w.root_item.child(0).child(0)
  self.assertEqual(node.text(4),'数据库真实备注'); self.assert_level_background(node.background(0)); self.assertEqual(node.foreground(0).style(),Qt.NoBrush)
 def test_external_alternate_update_marks_owner_row_purple(self):
  from PyQt5.QtGui import QColor
  r,p,c,o,x=self.ids; rows=[dict(row) for row in self.s.bom(p)]; rows[0]['alternate_material_ids']=[o]
  self.s.save_bom(p,rows); self.w.refresh_external()
  node=self.w.root_item.child(0).child(0)
  self.assertEqual(node.background(0).color(),QColor('#4a148c')); self.assertIn('有其他BOM更新导入此变更',node.text(4))
  with patch.object(QMessageBox,'information'):self.assertTrue(self.w.save())
  self.assertEqual(self.s.bom(p)[0]['remark'],'')
 def test_failed_save_keeps_delete_visuals(self):
  from PyQt5.QtGui import QColor
  node=self.w.root_item.child(0); self.w.remove_relation(node)
  with patch.object(self.s,'save_bom',side_effect=ValueError('模拟失败')),patch.object(QMessageBox,'warning'):self.assertFalse(self.w.save())
  self.assertEqual(node.foreground(0).color(),QColor('red')); self.assertTrue(node.font(0).strikeOut()); self.assertTrue(self.w.dirty)
 def test_root_material_update_is_highlighted_on_refresh(self):
  from PyQt5.QtGui import QColor
  self.s.update_material(self.ids[0],dict(name='新的根名称',unit='pcs'))
  self.w.refresh_external()
  self.assertEqual(self.w.root_item.background(0).color(),QColor('#4a148c'))
  self.assertIn('新的根名称',self.w.root_item.text(0)); self.assertIn('有其他BOM更新导入此变更',self.w.root_item.text(4))
 def test_unsaved_added_material_imports_external_updates_with_colors(self):
  from PyQt5.QtGui import QColor
  r,p,c,o,x=self.ids
  self.s.save_bom(o,[dict(material_id=c,quantity=1,unit='pcs')])
  with patch.object(MaterialSelectionDialog,'getItem',return_value=('OTHER1 OTHER',True)):self.w.add_relative(self.w.root_item.data(0,Qt.UserRole),True)
  added=self.w.tree.currentItem(); self.w.fields['quantity'].setValue(4.2)
  child=added.child(0); self.w.show_item_info(child); self.w.fields['position'].setText('本地位置')
  self.w.show_item_info(self.w.root_item.child(0)); self.w.fields['quantity'].setValue(8.7)
  # 新增关系尚未写入ROOT，但其物料自身BOM已由其他窗口更改。
  self.s.update_material(o,dict(name='OTHER外部新名称',unit='pcs'))
  rows=[dict(row) for row in self.s.bom(o)]; rows[0].update(quantity=7.6,remark='外部真实备注',alternate_material_ids=[p])
  self.s.save_bom(o,rows); self.assertTrue(self.w.refresh_action.isEnabled())
  self.w.refresh_external(); self.w.refresh_external()
  added=self.w.root_item.child(1); child=added.child(0)
  self.assertIn('OTHER外部新名称',added.text(0)); self.assertEqual(added.text(1),'4.2')
  self.assertEqual((child.text(1),child.text(3)),('7.6','本地位置'))
  for node in (added,child):
   self.assertEqual(node.foreground(0).color(),QColor('green')); self.assertEqual(node.background(0).color(),QColor('#4a148c'))
  self.assertEqual(self.w.root_item.child(0).foreground(0).color(),QColor('#9a6700'))
  self.w.show_item_info(child); self.assertEqual(self.w.alt_table.rowCount(),1); self.assertEqual(self.w.fields['remark'].text(),'外部真实备注')
  self.assertNotIn(o,[row['material_id'] for row in self.s.bom(r)])
  with patch.object(QMessageBox,'information'):self.assertTrue(self.w.save())
  self.assertEqual(self.s.bom(o)[0]['position'],'本地位置'); self.assertEqual(self.s.bom(o)[0]['quantity'],7.6)
  self.assertEqual(self.s.bom(o)[0]['remark'],'外部真实备注')
  added=self.w.root_item.child(1); self.assertEqual(added.foreground(0).style(),Qt.NoBrush); self.assert_level_background(added.background(0))
 def test_draft_subtree_external_add_remove_and_no_false_highlights(self):
  from PyQt5.QtGui import QColor
  r,p,c,o,x=self.ids
  self.s.save_bom(o,[dict(material_id=c,quantity=1,unit='pcs')])
  with patch.object(MaterialSelectionDialog,'getItem',return_value=('OTHER1 OTHER',True)):self.w.add_relative(self.w.root_item.data(0,Qt.UserRole),True)
  self.w.refresh_external()
  self.assert_level_background(self.w.root_item.child(1).background(0))
  self.assert_level_background(self.w.root_item.child(1).child(0).background(0))
  # 外部新增共享子项（不是当前SKU新增的本地关系）。
  rows=[dict(row) for row in self.s.bom(o)]+[dict(material_id=p,parent_material_id=o,quantity=3,unit='pcs')]
  rows.extend(dict(row) for row in self.s.bom(p)); self.s.save_bom(o,rows); self.w.refresh_external()
  added=self.w.root_item.child(1); self.assertEqual(added.childCount(),2)
  new_child=added.child(1); self.assertEqual(new_child.data(0,Qt.UserRole)['material_id'],p)
  self.assertEqual(new_child.background(0).color(),QColor('#4a148c')); self.assertEqual(new_child.foreground(0).color(),QColor('green'))
  # 多次刷新后，未被本地编辑的外部删除应生效，不能因绿字误认为本地编辑。
  self.w.refresh_external(); rows=[dict(row) for row in self.s.bom(o) if row['parent_material_id']!=o or row['material_id']!=c]
  self.s.save_bom(o,rows); self.w.refresh_external()
  added=self.w.root_item.child(1); self.assertEqual(added.childCount(),1); self.assertEqual(added.child(0).data(0,Qt.UserRole)['material_id'],p)
 def test_pending_deleted_draft_child_keeps_red_when_external_changes(self):
  from PyQt5.QtGui import QColor
  r,p,c,o,x=self.ids
  self.s.save_bom(o,[dict(material_id=c,quantity=1,unit='pcs')])
  with patch.object(MaterialSelectionDialog,'getItem',return_value=('OTHER1 OTHER',True)):self.w.add_relative(self.w.root_item.data(0,Qt.UserRole),True)
  self.w.remove_relation(self.w.tree.currentItem().child(0))
  rows=[dict(row) for row in self.s.bom(o)]; rows[0]['quantity']=6.3; self.s.save_bom(o,rows)
  self.w.refresh_external(); deleted=self.w.root_item.child(1).child(0)
  self.assertEqual(deleted.foreground(0).color(),QColor('red')); self.assertTrue(deleted.font(0).strikeOut()); self.assertEqual(deleted.background(0).color(),QColor('#4a148c'))
  with patch.object(QMessageBox,'information'):self.assertTrue(self.w.save())
  self.assertEqual(self.s.bom(o),[])
 def test_external_deep_change_marks_shared_occurrences_not_ancestors(self):
  from PyQt5.QtGui import QColor
  r,p,c,o,x=self.ids
  self.s.save_bom(o,[dict(material_id=p,quantity=1,unit='pcs')]+[dict(row) for row in self.s.bom(p)])
  with patch.object(MaterialSelectionDialog,'getItem',return_value=('OTHER1 OTHER',True)):self.w.add_relative(self.w.root_item.data(0,Qt.UserRole),True)
  rows=[dict(row) for row in self.s.bom(p)]; rows[0]['quantity']=7.4; self.s.save_bom(p,rows)
  self.w.refresh_external(); self.w.refresh_external()
  saved_parent=self.w.root_item.child(0); draft=self.w.root_item.child(1)
  for node in (saved_parent.child(0),draft.child(0).child(0)):
   for column in range(self.w.tree.columnCount()):self.assertEqual(node.background(column).color(),QColor('#4a148c'))
   self.assertIn('有其他BOM更新导入此变更',node.text(4))
  for node in (self.w.root_item,saved_parent,draft,draft.child(0)):
   for column in range(self.w.tree.columnCount()):self.assert_level_background(node.background(column))
   self.assertNotIn('有其他BOM更新导入此变更',node.text(4))
  self.assertEqual(draft.child(0).child(0).text(1),'7.4')
  self.assertEqual(draft.child(0).child(0).foreground(0).color(),QColor('green'))
  with patch.object(QMessageBox,'information'):self.assertTrue(self.w.save())
  for node in self.w.by.values():
   self.assert_level_background(node.background(0))
   self.assertNotIn('有其他BOM更新导入此变更',node.text(4))
  self.assertTrue(all(row['remark']=='' for row in self.s.bom(r)))
 def test_external_deleted_locally_edited_subtree_keeps_purple_note(self):
  from PyQt5.QtGui import QColor
  parent=self.w.root_item.child(0); child=parent.child(0)
  self.w.show_item_info(child); self.w.fields['remark'].setText('保留本地备注')
  self.s.save_bom(self.ids[0],[]); self.w.refresh_external()
  parent=self.w.root_item.child(0); child=parent.child(0)
  for node in (parent,child):
   self.assertEqual(node.background(0).color(),QColor('#4a148c'))
   self.assertIn('有其他BOM更新导入此变更',node.text(4))
  self.assertEqual(child.foreground(0).color(),QColor('#9a6700'))
  self.w.show_item_info(child); self.assertEqual(self.w.fields['remark'].text(),'保留本地备注')
 def test_external_highlight_does_not_spread_to_unchanged_siblings(self):
  r,p,c,o,x=self.ids
  rows=[dict(row) for row in self.s.bom(r)]+[dict(material_id=o,quantity=1,unit='pcs')]
  self.s.save_bom(r,rows); self.w.load()
  rows=[dict(row) for row in self.s.bom(p)]; rows[0]['quantity']=4.6; self.s.save_bom(p,rows)
  self.w.refresh_external(); self.w.refresh_external()
  sibling=self.w.root_item.child(1)
  for column in range(self.w.tree.columnCount()):self.assert_level_background(sibling.background(column))
  self.assertEqual(sibling.text(4),'')
 def test_changed_material_highlights_descendants_only(self):
  from PyQt5.QtGui import QColor
  r,p,c,o,x=self.ids
  self.s.save_bom(r,[dict(row) for row in self.s.bom(r)]+[dict(material_id=o,quantity=1,unit='pcs')]); self.w.load()
  self.w.show_item_info(self.w.root_item.child(0).child(0)); self.w.fields['remark'].setText('本地备注')
  self.s.update_material(p,dict(name='父物料外部修改',unit='pcs')); self.w.refresh_external(); self.w.refresh_external()
  parent=self.w.root_item.child(0); child=parent.child(0)
  for node in (parent,child):
   for column in range(self.w.tree.columnCount()):self.assertEqual(node.background(column).color(),QColor('#4a148c'))
   self.assertIn('有其他BOM更新导入此变更',node.text(4))
  self.assertEqual(child.foreground(0).color(),QColor('#9a6700'))
  for node in (self.w.root_item,self.w.root_item.child(1)):
   self.assert_level_background(node.background(0)); self.assertEqual(node.text(4),'')
  with patch.object(QMessageBox,'information'):self.assertTrue(self.w.save())
  for node in self.w.by.values():self.assert_level_background(node.background(0))
  self.assertEqual(self.s.bom(p)[0]['remark'],'本地备注')
 def test_external_child_add_remove_does_not_mark_parent_or_sibling(self):
  from PyQt5.QtGui import QColor
  r,p,c,o,x=self.ids
  self.s.save_bom(p,[dict(row) for row in self.s.bom(p)]+[dict(material_id=o,quantity=1,unit='pcs')])
  self.w.refresh_external(); parent=self.w.root_item.child(0)
  self.assertEqual(parent.child(1).background(0).color(),QColor('#4a148c'))
  for node in (self.w.root_item,parent,parent.child(0)):
   self.assert_level_background(node.background(0)); self.assertEqual(node.text(4),'')
  self.s.save_bom(p,[dict(row) for row in self.s.bom(p) if row['material_id']==c]); self.w.refresh_external()
  parent=self.w.root_item.child(0); self.assertEqual(parent.childCount(),1)
  for node in (self.w.root_item,parent,parent.child(0)):
   self.assert_level_background(node.background(0)); self.assertEqual(node.text(4),'')

 def test_new_sibling_inserted_after_selection_and_persisted(self):
  r,p,c,o,x=self.ids
  extra=self.s.create_material(dict(base_code='EXTRA',version='1',name='EXTRA',unit='pcs'))
  last=self.s.create_material(dict(base_code='LAST',version='1',name='LAST',unit='pcs'))
  for target in (r,p):
   with self.subTest(parent=target):
    self.s.save_bom(target,[dict(material_id=c,quantity=1,unit='pcs'),dict(material_id=o,quantity=1,unit='pcs'),dict(material_id=last,quantity=1,unit='pcs')])
    self.w.load()
    container=self.w.root_item if target==r else next(item for _,item in self.w.expansion_nodes() if (item.data(0,Qt.UserRole) or {}).get('material_id')==p)
    selected=container.child(1)
    with patch.object(MaterialSelectionDialog,'getItem',return_value=('EXTRA1 EXTRA',True)):
     self.w.add_relative(selected.data(0,Qt.UserRole),False)
    expected=[c,o,extra,last]
    self.assertEqual([container.child(i).data(0,Qt.UserRole)['material_id'] for i in range(container.childCount())],expected)
    added=container.child(2);self.assertIs(self.w.tree.currentItem(),added);self.assertTrue(added.data(0,Qt.UserRole)['_added'])
    self.assertEqual(self.w.fields['base_code'].text(),'EXTRA')
    with patch.object(QMessageBox,'information'),patch.object(QMessageBox,'warning') as warning:
     self.assertTrue(self.w.save());warning.assert_not_called()
    self.w.load()
    self.assertEqual([row['child_material_id'] for row in self.s.db.conn.execute('SELECT child_material_id FROM material_bom_items WHERE parent_material_id=? ORDER BY sort_no,id',(target,))],expected)
    # Restore the root so the next subcase tests a nested parent.
    if target==r:self.s.save_bom(r,[dict(material_id=p,quantity=1,unit='pcs')])
