import os,sys,tempfile,unittest
os.environ['QT_QPA_PLATFORM']='offscreen'
from pathlib import Path
from unittest.mock import patch
sys.path.insert(0,str(Path(__file__).parents[1]/'src'))
from PyQt5.QtCore import Qt,QItemSelectionModel,QPoint,QEvent
from PyQt5.QtWidgets import QApplication,QMenu,QMessageBox,QAbstractItemView
from PyQt5.QtTest import QTest
from database.database import Database
from services.services import Services
from ui.main_window import MainWindow,MaterialWindow

class MultiSelectTests(unittest.TestCase):
 @classmethod
 def setUpClass(cls):cls.app=QApplication.instance() or QApplication([])
 def setUp(self):
  self.tmp=tempfile.TemporaryDirectory();self.db=Database(str(Path(self.tmp.name)/'db'));self.s=Services(self.db)
  self.ids=[self.s.create_sku(self.s.create_material(dict(base_code=str(i),version='100',name='SKU'+str(i),unit='pcs'))) for i in range(3)]
  self.w=MainWindow(self.s);self.w.show();self.app.processEvents()
 def tearDown(self):
  for w in list(self.w._bom_windows.values()):w.dirty=False;w.close()
  self.w.close();self.w.deleteLater();self.app.sendPostedEvents(None,QEvent.DeferredDelete);self.db.close();self.tmp.cleanup()
 def select(self,*rows):
  self.w.table.clearSelection()
  for r in rows:self.w.table.selectionModel().select(self.w.table.model().index(r,1),QItemSelectionModel.Select|QItemSelectionModel.Rows)
 def menu(self,row,label=None):
  found={}
  def choose(menu,*args):
   found.update({a.text():a for a in menu.actions() if not a.isSeparator()})
   return found.get(label)
  with patch.object(self.w.table,'rowAt',return_value=row),patch.object(QMenu,'exec_',choose):self.w.sku_context_menu(QPoint())
  return {k:v.isEnabled() for k,v in found.items()}
 def test_ctrl_shift_and_context_selection(self):
  t=self.w.table
  def click(row,mod=Qt.NoModifier):QTest.mouseClick(t.viewport(),Qt.LeftButton,mod,t.visualItemRect(t.item(row,2)).center())
  click(0);click(2,Qt.ControlModifier);self.assertEqual({i.row() for i in t.selectionModel().selectedRows()},{0,2})
  actions=self.menu(0)
  for label in ('编辑变更描述','SKU版本变更','历史版本变更','复制BOM','BOM对比','删除SKU'):self.assertFalse(actions[label])
  self.assertEqual({i.row() for i in t.selectionModel().selectedRows()},{0,2})
  self.menu(1);self.assertEqual([i.row() for i in t.selectionModel().selectedRows()],[1])
  click(0);click(2,Qt.ShiftModifier);self.assertEqual(len(t.selectionModel().selectedRows()),3)
 def test_batch_open_export_and_lock_preserve_selection(self):
  self.select(0,2)
  with patch.object(self.w,'open_bom_id') as opened:self.menu(0,'编辑/浏览BOM');self.assertEqual([c.args[0] for c in opened.call_args_list],[self.ids[0],self.ids[2]])
  for label,purchase in [('导出BOM',False),('采购BOM导出',True)]:
   with patch.object(self.w,'export_sku_batch') as exported:
    self.menu(0,label);exported.assert_called_once_with([self.ids[0],self.ids[2]],purchase=purchase)
  self.menu(0,'锁定SKU');self.assertEqual([r['locked'] for r in self.s.skus()],[1,0,1]);self.assertEqual(len(self.w.table.selectionModel().selectedRows()),2)
  self.menu(0,'取消锁定');self.assertEqual([r['locked'] for r in self.s.skus()],[0,0,0])
 def test_mixed_locks_and_dirty_preflight(self):
  self.s.set_sku_locked(self.ids[0],True);self.select(0,1);actions=self.menu(0)
  self.assertTrue(actions['锁定SKU']);self.assertTrue(actions['取消锁定'])
  editor=self.w.open_bom_id(self.ids[1]);editor.dirty=True
  with patch.object(QMessageBox,'warning') as warning:self.menu(0,'锁定SKU');warning.assert_called_once()
  self.assertEqual([r['locked'] for r in self.s.skus()],[1,0,0]);editor.dirty=False
 def test_library_row_selection_in_frozen_and_regular_cells(self):
  library=MaterialWindow(self.s);library.show();self.app.processEvents()
  try:
   table=library.t
   for view,col in [(table.frozen,2),(table,4)]:
    self.assertEqual(view.selectionBehavior(),QAbstractItemView.SelectRows)
    index=view.model().index(1,col);QTest.mouseClick(view.viewport(),Qt.LeftButton,Qt.NoModifier,view.visualRect(index).center())
    self.assertTrue(table.selectionModel().isRowSelected(1))
  finally:library.close();library.deleteLater()

 def test_batch_export_one_folder_real_files_and_collision(self):
  from PyQt5.QtWidgets import QFileDialog
  from openpyxl import load_workbook
  folder=Path(self.tmp.name)/'exports';folder.mkdir()
  for purchase in (False,True):
   with patch.object(QFileDialog,'getExistingDirectory',return_value=str(folder)) as choose,patch.object(QFileDialog,'getSaveFileName') as save,patch.object(QMessageBox,'exec_',return_value=QMessageBox.Ok),patch('ui.main_window.remember_export_path'):
    self.w.export_sku_batch(self.ids[:2],purchase);choose.assert_called_once();save.assert_not_called()
   self.assertEqual(len(list(folder.glob('*.xlsx'))),4 if purchase else 2)
  originals={p.name:p.read_bytes() for p in folder.glob('*.xlsx')}
  with patch.object(QFileDialog,'getExistingDirectory',return_value=str(folder)),patch.object(QMessageBox,'exec_',return_value=QMessageBox.Ok),patch('ui.main_window.remember_export_path'):
   self.w.export_sku_batch(self.ids[:2])
  self.assertEqual(len(list(folder.glob('*.xlsx'))),6)
  for name,data in originals.items():self.assertEqual((folder/name).read_bytes(),data)
  for path in folder.glob('*.xlsx'):
   book=load_workbook(path);self.assertTrue(book.sheetnames);book.close()
 def test_batch_export_cancel_and_failure_continues(self):
  from PyQt5.QtWidgets import QFileDialog
  with patch.object(QFileDialog,'getExistingDirectory',return_value=''),patch.object(self.w,'export_sku_id') as export:
   self.w.export_sku_batch(self.ids);export.assert_not_called()
  with patch.object(QFileDialog,'getExistingDirectory',return_value=self.tmp.name),patch.object(self.w,'export_sku_id',side_effect=[OSError('无法写入'),None,None]) as export,patch.object(QMessageBox,'exec_',return_value=QMessageBox.Ok),patch('ui.main_window.remember_export_path'):
   self.w.export_sku_batch(self.ids);self.assertEqual(export.call_count,3)
