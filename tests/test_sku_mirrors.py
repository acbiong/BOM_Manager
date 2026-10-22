import os
os.environ['QT_QPA_PLATFORM']='offscreen'
import sys,tempfile,unittest,re
from pathlib import Path
from unittest.mock import patch
sys.path.insert(0,str(Path(__file__).parents[1]/'src'))
from PyQt5.QtWidgets import QApplication,QMessageBox
from ui.interactive_windows import QMessageBox
from services.services import Services
from database.database import Database
from ui.main_window import MainWindow
from openpyxl import load_workbook

class MirrorTests(unittest.TestCase):
 @classmethod
 def setUpClass(cls):cls.app=QApplication.instance() or QApplication([])
 def test_independent_graph_alternates_history_windows_export_delete_and_reopen(self):
  with tempfile.TemporaryDirectory() as tmp:
   path=str(Path(tmp)/'db');db=Database(path);s=Services(db)
   r,p,c,a=[s.create_material(dict(base_code=x,version='1',name=x,unit='pcs')) for x in ('R','P','C','A')]
   sid=s.create_sku(r);s.save_bom(r,[dict(material_id=p,quantity=2,unit='pcs')]);s.save_bom(p,[dict(material_id=c,quantity=3,unit='pcs',remark='原备注')]);s.add_alternate(s.bom(p)[0]['id'],a)
   copy=s.copy_sku_bom(sid);mirror=s.for_sku(copy)
   self.assertEqual([row['mirror_version'] for row in s.skus()],[100,101])
   self.assertEqual([(row['material_id'],row['quantity'],row['remark']) for row in s.bom(r)],[(row['material_id'],row['quantity'],row['remark']) for row in mirror.bom(r)])
   self.assertEqual(mirror.alternates(mirror.bom(p)[0]['id'])[0]['material_id'],a)
   main=MainWindow(s);w=main.open_bom_id(sid);other=main.open_bom_id(copy)
   try:
    self.assertIsNot(w,other);self.assertIs(main.open_bom_id(copy),other)
    self.assertEqual(main.table.item(0,4).text(),'100');self.assertEqual(main.table.item(1,4).text(),'101')
    mirror.save_bom(p,[dict(material_id=c,quantity=5,unit='pcs',remark='镜像备注')])
    self.assertEqual(s.bom(p)[0]['quantity'],3);self.assertEqual(mirror.bom(p)[0]['quantity'],5)
    self.assertFalse(w.refresh_action.isEnabled())
    s.save_bom(p,[dict(material_id=c,quantity=7,unit='pcs')]);self.assertEqual(mirror.bom(p)[0]['quantity'],5)
    main._sku_filters={4:('',{'101'})};main.apply_sku_filters();self.assertTrue(main.table.isRowHidden(0));self.assertFalse(main.table.isRowHidden(1))
    totals=s.sku_material_quantities();self.assertEqual(totals[sid][c],14);self.assertEqual(totals[copy][c],10)
    s.change_mirror_version(copy,105);self.assertIn('105',other.windowTitle())
    with self.assertRaises(ValueError):s.change_mirror_version(copy,100)
    next_id=s.copy_sku_bom(copy);self.assertEqual(s.skus()[-1]['mirror_version'],106)
    export=Path(tmp)/'mirror.xlsx';s.export_sku(copy,export);wb=load_workbook(export)
    self.assertEqual(wb['R'].cell(4,10).value,5);wb.close()
    logs=list(mirror.db.conn.execute("SELECT changed_at,detail FROM change_logs WHERE object_type='bom' ORDER BY id"));self.assertTrue(any('镜像备注' in row['detail'] for row in logs));self.assertRegex(logs[-1]['changed_at'],r'\d{4}-\d\d-\d\d \d\d:\d\d:\d\d\.\d{2}$')
    other.close();s.delete_sku(copy);self.assertEqual(s.bom(p)[0]['quantity'],7);self.assertEqual(s.for_sku(next_id).bom(p)[0]['quantity'],5)
   finally:
    for window in list(main._bom_windows.values()):window.dirty=False;window.close()
    main.close();db.close()
   db=Database(path);s=Services(db);self.assertEqual(s.for_sku(next_id).bom(p)[0]['quantity'],5);db.close()
 def test_migration_backup_default_versions_and_local_centisecond_logs(self):
  import sqlite3
  from datetime import datetime,timezone
  with tempfile.TemporaryDirectory() as tmp:
   path=str(Path(tmp)/'legacy.db');conn=sqlite3.connect(path)
   conn.executescript("CREATE TABLE skus(id INTEGER PRIMARY KEY,material_id INTEGER UNIQUE,updated_at TEXT DEFAULT CURRENT_TIMESTAMP);INSERT INTO skus(id,material_id) VALUES(7,9);CREATE TABLE change_logs(id INTEGER PRIMARY KEY,object_type TEXT,object_id INTEGER,action TEXT,detail TEXT,changed_at TEXT);INSERT INTO change_logs VALUES(1,'material',9,'update','旧记录','2020-01-02 03:04:05');")
   conn.commit();conn.close();db=Database(path);s=Services(db)
   self.assertEqual(tuple(db.conn.execute('SELECT id,material_id,mirror_version,scope_id FROM skus').fetchone()),(7,9,100,0))
   backups=list(Path(tmp).glob('*.pre-mirrors-*.db'));self.assertEqual(len(backups),1)
   expected=datetime(2020,1,2,3,4,5,tzinfo=timezone.utc).astimezone().strftime('%Y-%m-%d %H:%M:%S.%f')[:-4]
   self.assertEqual(db.conn.execute('SELECT changed_at FROM change_logs WHERE id=1').fetchone()[0],expected)
   mid=s.create_material(dict(base_code='M',version='1',name='本机时间',unit='pcs'))
   stamp=db.conn.execute('SELECT changed_at FROM change_logs ORDER BY id DESC').fetchone()[0]
   self.assertRegex(stamp,r'\d{4}-\d\d-\d\d \d\d:\d\d:\d\d\.\d\d$')
   self.assertLess(abs((datetime.now()-datetime.strptime(stamp,'%Y-%m-%d %H:%M:%S.%f')).total_seconds()),2)
   db.close();db=Database(path)
   self.assertEqual(db.conn.execute('SELECT changed_at FROM change_logs WHERE id=1').fetchone()[0],expected);self.assertEqual(len(list(Path(tmp).glob('*.pre-mirrors-*.db'))),1);db.close()
 def test_mirror_only_references_protect_material_and_version_updates_are_selective(self):
  with tempfile.TemporaryDirectory() as tmp:
   db=Database(str(Path(tmp)/'db'));s=Services(db)
   r,c,a=[s.create_material(dict(base_code=x,version='1',name=x,unit='pcs')) for x in ('R','C','A')]
   sid=s.create_sku(r);s.save_bom(r,[dict(material_id=c,quantity=2,unit='pcs')]);s.add_alternate(s.bom(r)[0]['id'],a)
   copied=s.copy_sku_bom(sid);s.save_bom(r,[])
   self.assertEqual({row['id']:row['idle'] for row in s.materials()}[c],'N')
   self.assertEqual({row['id']:row['idle'] for row in s.materials()}[a],'N')
   with self.assertRaises(ValueError):s.delete_material(c)
   refs=s.material_version_references(c);self.assertEqual(refs[0]['sku_ids'],[copied])
   new=s.change_material_version(c,'2',set(),{row['parent_id'] for row in refs})
   self.assertNotEqual(c,new);self.assertEqual(s.for_sku(copied).bom(r)[0]['material_id'],c)
   s.delete_sku(copied);self.assertEqual({row['id']:row['idle'] for row in s.materials()}[c],'Y');db.close()
 def test_new_material_subtree_capture_and_atomic_copy_failure(self):
  from ui.main_window import BomWindow
  with tempfile.TemporaryDirectory() as tmp:
   db=Database(str(Path(tmp)/'db'));s=Services(db)
   r,p,c=[s.create_material(dict(base_code=x,version='1',name=x,unit='pcs')) for x in ('R','P','C')]
   sid=s.create_sku(r);copy=s.copy_sku_bom(sid);mirror=s.for_sku(copy)
   s.save_bom(p,[dict(material_id=c,quantity=3,unit='pcs')])
   w=BomWindow(mirror,r)
   try:
    node=w.insert_material(w.root_item,s.material_snapshot(p));self.assertEqual(node.childCount(),1);self.assertEqual(node.child(0).text(1),'3')
    self.assertEqual(mirror.bom(r),[])
    with patch.object(QMessageBox,'information'):self.assertTrue(w.save())
    s.save_bom(p,[dict(material_id=c,quantity=9,unit='pcs')]);self.assertEqual(mirror.bom(p)[0]['quantity'],3)
    before=len(s.skus())
    with patch('services.mirrors.create_namespace',side_effect=RuntimeError('test failure')):
     with self.assertRaises(RuntimeError):s.copy_sku_bom(copy)
    self.assertEqual(len(s.skus()),before)
   finally:w.dirty=False;w.close();db.close()
 def test_selective_material_version_only_updates_selected_mirror(self):
  with tempfile.TemporaryDirectory() as tmp:
   db=Database(str(Path(tmp)/'db'));s=Services(db)
   r,c=[s.create_material(dict(base_code=x,version='1',name=x,unit='pcs')) for x in ('R','C')]
   sid=s.create_sku(r);s.save_bom(r,[dict(material_id=c,quantity=2,unit='pcs')]);copy=s.copy_sku_bom(sid)
   refs=s.material_version_references(c);selected={ref['parent_id'] for ref in refs if isinstance(ref['parent_id'],str)}
   new=s.change_material_version(c,'2',selected,{ref['parent_id'] for ref in refs})
   self.assertEqual(s.bom(r)[0]['material_id'],c);self.assertEqual(s.for_sku(copy).bom(r)[0]['material_id'],new);db.close()
