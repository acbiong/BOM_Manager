import os,sys,tempfile,unittest
os.environ['QT_QPA_PLATFORM']='offscreen'
from pathlib import Path
from unittest.mock import patch
sys.path.insert(0,str(Path(__file__).parents[1]/'src'))
from PyQt5.QtCore import Qt,QPoint,QEvent
from PyQt5.QtWidgets import QApplication,QMenu,QFileDialog,QMessageBox
from openpyxl import load_workbook
from database.database import Database
from services.services import Services
from services.snapshot_bom import migrate
from services.current_catalog import initialize
from ui.main_window import MainWindow

class SubtreeExportTests(unittest.TestCase):
 @classmethod
 def setUpClass(cls):cls.app=QApplication.instance() or QApplication([])
 def setUp(self):
  self.tmp=tempfile.TemporaryDirectory();self.db=Database(str(Path(self.tmp.name)/'db'));self.s=Services(self.db)
  self.r,self.p,self.c,self.a,self.other=[self.s.create_material(dict(base_code=k,version='100',name=k,unit='pcs')) for k in ('ROOT','PARENT','CHILD','ALT','OTHER')]
  self.s.save_bom(self.p,[dict(material_id=self.c,quantity=2.5,unit='pcs',position='U1',remark='备注',alternate_material_ids=[self.a])])
  self.s.save_bom(self.r,[dict(material_id=self.p,quantity=1,unit='pcs'),dict(material_id=self.other,quantity=1,unit='pcs')])
  self.sid=self.s.create_sku(self.r);migrate(self.db);initialize(self.db);self.main=MainWindow(self.s);self.window=self.main.open_bom_id(self.sid)
 def tearDown(self):
  for w in list(self.main._bom_windows.values())+list(self.main._material_bom_windows.values()):w.dirty=False;w.close()
  self.main.close();self.main.deleteLater();self.app.sendPostedEvents(None,QEvent.DeferredDelete);self.db.close();self.tmp.cleanup()
 def check_file(self,path):
  book=load_workbook(path)
  try:
   self.assertIn('变更记录',book.sheetnames);sheet=book['PARENT'];rows=list(sheet.values);headers=rows[0]
   self.assertEqual([r[4] for r in rows[1:]],['PARENT','CHILD']);self.assertEqual(rows[1][:4],('1',0,'ROOT','ROOT'));self.assertEqual(rows[2][:2],('1.1',1))
   self.assertEqual(rows[2][headers.index('用量')],2.5);self.assertEqual(rows[2][headers.index('替代料1物料编码')],'ALT');self.assertEqual(rows[2][-1],'备注')
   self.assertTrue(sheet['A1'].font.bold);self.assertTrue(sheet['G3'].alignment.wrap_text);self.assertEqual(sheet['A2'].border.left.style,'thin')
   self.assertTrue(all(d.width<=60 for d in sheet.column_dimensions.values()))
  finally:book.close()
 def test_tree_edit_and_browse_menu_exports_only_selected_root(self):
  item=self.window.root_item.child(0)
  for locked in (False,True):
   self.s.set_sku_locked(self.sid,locked);path=Path(self.tmp.name)/f'{locked}.xlsx'
   def choose(menu,*args):return next(a for a in menu.actions() if a.text()=='导出BOM结构')
   with patch.object(self.window.tree,'itemAt',return_value=item),patch.object(QMenu,'exec_',choose),patch.object(QFileDialog,'getSaveFileName',return_value=(str(path),'')),patch.object(QMessageBox,'information'),patch('ui.main_window.remember_export_path'):
    self.window.context_menu(QPoint())
   self.check_file(path)
 def test_library_menu_and_material_root_toolbar(self):
  self.main.show_materials();library=self.main._materials
  row=next(i for i,r in enumerate(self.s.materials()) if r['id']==self.p);path=Path(self.tmp.name)/'library.xlsx'
  def choose(menu,*args):return next(a for a in menu.actions() if a.text()=='导出BOM结构')
  with patch.object(library.t,'rowAt',return_value=row),patch.object(QMenu,'exec_',choose),patch.object(QFileDialog,'getSaveFileName',return_value=(str(path),'')),patch.object(QMessageBox,'information'),patch('ui.main_window.remember_export_path'):
   library.menu(QPoint())
  self.check_file(path)
  editor=library.open_material_bom(self.p);path=Path(self.tmp.name)/'toolbar.xlsx'
  with patch.object(QFileDialog,'getSaveFileName',return_value=(str(path),'')),patch.object(QMessageBox,'information'),patch('ui.main_window.remember_export_path'):editor.export_bom()
  self.check_file(path)
 def test_leaf_dirty_pending_refresh_and_cancel(self):
  node=self.window.root_item.child(0);leaf=node.child(0)
  self.assertIsNone(self.window.subtree_export_action(QMenu(self.window),leaf))
  self.window.mark_dirty();self.assertFalse(self.window.subtree_export_action(QMenu(self.window),node).isEnabled())
  with patch.object(QFileDialog,'getSaveFileName') as dialog:self.window.export_subtree(node);dialog.assert_not_called()
  self.window.dirty=False;self.window.load();node=self.window.root_item.child(0)
  self.window.on_external_change(dict(type='bom_changed',root=self.p,scope_id=self.sid))
  self.assertFalse(self.window.subtree_export_action(QMenu(self.window),node).isEnabled())
  self.window.load()
  with patch.object(QFileDialog,'getSaveFileName',return_value=('','')),patch.object(self.window.svc,'export_bom_structure') as export:self.window.export_subtree(self.window.root_item.child(0));export.assert_not_called()
