import os
os.environ['QT_QPA_PLATFORM']='offscreen'
import sys,tempfile,unittest
from pathlib import Path
from unittest.mock import patch
sys.path.insert(0,str(Path(__file__).parents[1]/'src'))
from PyQt5.QtWidgets import QApplication,QComboBox,QDialog,QMessageBox,QFileDialog
from openpyxl import load_workbook
from database.database import Database
from services.services import Services
from ui.main_window import MaterialSelectionDialog,BomWindow,MaterialWindow

class ExportSelectionTests(unittest.TestCase):
 @classmethod
 def setUpClass(cls):cls.app=QApplication.instance() or QApplication([])
 def test_filter_requires_exact_choice(self):
  def run(dialog):
   combo=dialog.findChild(QComboBox);combo.setEditText('alpha');combo.completer().setCompletionPrefix('alpha')
   self.assertEqual(combo.completer().completionCount(),1)
   with patch.object(QMessageBox,'warning') as warning:dialog.accept();warning.assert_called_once()
   combo.setEditText('BC-100 Alpha');dialog.accept();return dialog.result()
  with patch.object(MaterialSelectionDialog,'exec_',run):
   self.assertEqual(MaterialSelectionDialog.getItem(None,'选择','物料',['BC-100 Alpha','BC-200 Beta']),('BC-100 Alpha',True))
 def test_export_alternates_types_and_default_names(self):
  with tempfile.TemporaryDirectory() as tmp:
   db=Database(str(Path(tmp)/'db'));s=Services(db)
   try:
    ids=[s.create_material(dict(base_code=code,version='100',name=name,unit='pcs',manufacturer='供应商',mpn='型号')) for code,name in [('BC-','SKU名称'),('C','子项'),('A','替代一'),('B','替代二')]]
    root,child,a,b=ids;sid=s.create_sku(root)
    s.save_bom(root,[dict(material_id=child,quantity=1,unit='pcs',remark='备注内容',alternate_material_ids=[a,b]),dict(material_id=a,quantity=2,unit='pcs')])
    s.update_project({'project_name':'项目甲'})
    path=Path(tmp)/'out.xlsx';s.export_sku(sid,path);wb=load_workbook(path)
    try:
     sheet=wb['SKU对应的BOM'];headers=[c.value for c in sheet[1]]
     self.assertEqual(headers[15:20],['替代料1物料编码','替代料1版本号','替代料1名称','替代料1供应商','替代料1MPN'])
     self.assertEqual([sheet.cell(2,c).value for c in range(16,26)],['A','100','替代一','供应商','型号','B','100','替代二','供应商','型号'])
     self.assertEqual(sheet.cell(2,26).value,'备注内容');self.assertIsNone(sheet.cell(3,16).value)
     self.assertTrue(sheet.cell(2,25).alignment.wrap_text);self.assertEqual(sheet.cell(2,25).border.left.style,'thin')
     logs=list(wb['变更记录'].values);self.assertEqual(logs[0][-2:],('变更类型','变更详细内容'))
     created=next(row for row in logs[1:] if row[5]=='新建物料');self.assertIsNone(created[6])
    finally:wb.close()
    bom=BomWindow(s,root);materials=MaterialWindow(s)
    try:
     with patch.object(QFileDialog,'getSaveFileName',return_value=('','')) as chooser:
      bom.export_bom();self.assertEqual(chooser.call_args.args[2],'/home/biong/Documents/code/BOM管理系统/BC-100 SKU名称.xlsx')
      materials.export_materials();self.assertEqual(chooser.call_args.args[2],'/home/biong/Documents/code/BOM管理系统/项目甲 物料库.xlsx')
    finally:bom.close();materials.close()
   finally:db.close()
