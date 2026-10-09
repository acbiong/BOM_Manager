import tempfile,unittest,sys
from pathlib import Path
sys.path.insert(0,str(Path(__file__).parents[1]/'src'))
from database.database import Database
from services.services import Services
from openpyxl import load_workbook
class TestExports(unittest.TestCase):
 def test_database_export_sheets(self):
  f=tempfile.NamedTemporaryFile(suffix='.db'); s=Services(Database(f.name)); m=s.create_material({'base_code':'A','version':'1','name':'A','unit':'pcs'}); s.create_sku(m); out=f.name+'.xlsx'; s.export_database(out); self.assertEqual(load_workbook(out).sheetnames,['SKU清单','物料库','变更记录','A1'])
if __name__=='__main__': unittest.main()

class BomExportLayoutTests(unittest.TestCase):
 def test_hierarchy_shared_paths_fields_and_excel_format(self):
  with tempfile.TemporaryDirectory() as tmp:
   db=Database(str(Path(tmp)/'db'));s=Services(db)
   try:
    ids={code:s.create_material(dict(base_code=code,version='001',name='中文物料名称'*12,unit='pcs',grade='工业级',description='长描述'*30)) for code in ['R','P','C','D','X']}
    def relation(code,parent,qty=1):return dict(material_id=ids[code],parent_material_id=ids[parent],quantity=qty,unit='pcs',position='安装位置A',remark='业务备注\n第二行')
    s.save_bom(ids['R'],[relation('P','R'),relation('C','P',2.3),relation('D','C'),relation('X','P'),relation('C','R',4.5)])
    sid=s.create_sku(ids['R'])
    for database_export in (False,True):
     out=Path(tmp)/('all.xlsx' if database_export else 'sku.xlsx')
     if database_export:s.export_database(out)
     else:s.export_sku(sid,out)
     wb=load_workbook(out);sheet=wb['R001' if database_export else 'SKU对应的BOM']
     try:
      rows=list(sheet.values);headers=rows[0]
      self.assertEqual([r[0] for r in rows[1:]],['1','1.1','1.1.1','1.2','2','2.1'])
      self.assertEqual([r[2] for r in rows[1:]],['R','P','C','P','R','C'])
      self.assertEqual(headers[8:13],('物料位置','用量','单位','物料等级','M/P'));self.assertEqual(headers[-1],'备注')
      self.assertEqual(rows[2][8:12],('安装位置A',2.3,'pcs','工业级'));self.assertEqual(rows[2][-1],'业务备注\n第二行')
      self.assertEqual(sheet.cell(2,1).data_type,'s');self.assertEqual(sheet.cell(2,6).value,'001')
      self.assertTrue(all(5<=dimension.width<=25 for dimension in sheet.column_dimensions.values()))
      self.assertGreater(sheet.row_dimensions[2].height,22)
      centered={'层级','父项版本','物料版本号','用量','单位','物料等级','M/P'}
      for cells in sheet:
       for cell in cells:
        self.assertTrue(cell.alignment.wrap_text)
        for side in ('left','right','top','bottom'):
         edge=getattr(cell.border,side);self.assertEqual(edge.style,'thin');self.assertEqual(edge.color.rgb,'FF000000')
        if headers[cell.column-1] in centered:self.assertEqual(cell.alignment.horizontal,'center')
        if cell.row==1:self.assertTrue(cell.font.bold)
       if cells[0].row>1:self.assertEqual(cells[9].number_format,'0.0')
     finally:wb.close()
    empty=s.create_sku(s.create_material(dict(base_code='EMPTY',version='1',name='空',unit='pcs')))
    s.export_sku(empty,Path(tmp)/'empty.xlsx');wb=load_workbook(Path(tmp)/'empty.xlsx')
    self.assertEqual(wb['SKU对应的BOM'].max_row,1);self.assertTrue(wb['SKU对应的BOM']['A1'].font.bold);wb.close()
   finally:db.close()
