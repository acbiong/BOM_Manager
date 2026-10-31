import os,sys,tempfile,unittest
os.environ['QT_QPA_PLATFORM']='offscreen'
from pathlib import Path
from unittest.mock import patch
sys.path.insert(0,str(Path(__file__).parents[1]/'src'))
from services.external_bom import ExcelSource,build_external_bom,COLUMNS,REQUIRED,ALT_SUFFIXES
from services.bom_compare import compare_trees,snapshot
from services.services import Services
from database.database import Database
from ui.bom_compare import BomCompareDialog
from ui.main_window import MainWindow,MaterialHeaderMappingDialog
from config.config_service import ConfigService
from PyQt5.QtWidgets import QApplication,QDialog
from PyQt5.QtCore import Qt
from openpyxl import Workbook

class ExternalBomTests(unittest.TestCase):
 @classmethod
 def setUpClass(cls):cls.app=QApplication.instance() or QApplication([])
 def fixture(self):
  source=ExcelSource(Path(__file__).parent/'fixtures/external_bom.xls')
  try:self.assertEqual(source.names,['说明','BOM']);return source.rows('BOM')
  finally:source.close()
 def test_real_xls_unknown_materials_and_precision(self):
  rows=self.fixture();tree=build_external_bom(rows,{h:h for h in rows[0]})
  self.assertEqual(tree['name'],'外部根');self.assertEqual(tree['children'][0]['quantity'],.12345)
  self.assertEqual(tree['children'][0]['alternate_codes'],{('ALT','200')})
  self.assertTrue(all(d.status=='same' for d in compare_trees(tree,tree)))
 def test_mapping_errors_and_structure_validation_accumulate(self):
  rows=self.fixture();mapping={h:h for h in rows[0]}
  with self.assertRaisesRegex(ValueError,'必选项未映射'):build_external_bom(rows,{**mapping,'用量':''})
  broken=[rows[0],rows[1],['EXT','100','C','100','',0,'pcs','',''],['MISSING','1','D','100','孤立',1,'pcs','','']]
  with self.assertRaises(ValueError) as error:build_external_bom(broken,mapping)
  self.assertIn('第3行',str(error.exception));self.assertIn('父项物料没有对应记录',str(error.exception))
  cyclic=[rows[0],rows[1],rows[2],['NEW','100','EXT','100','外部根',1,'pcs','','']]
  with self.assertRaisesRegex(ValueError,'循环引用'):build_external_bom(cyclic,mapping)
  shared=[rows[0],rows[1],rows[2],rows[2]]
  self.assertEqual(len(build_external_bom(shared,mapping)['children']),1)
 def test_all_mapping_fields_optional_preview_and_alternates(self):
  rows=self.fixture();fields=[f for f,_ in COLUMNS]
  with tempfile.TemporaryDirectory() as tmp:
   config=ConfigService(tmp)
   dialog=MaterialHeaderMappingDialog('unused.xls',rows[0],[],required=fields,optional_fields=set(fields)-REQUIRED,preview_rows=rows[1:],alternate_suffixes=ALT_SUFFIXES,config_service=config,mapping_kind='compare_bom')
   try:
    self.assertEqual(set(dialog.combos),set(fields));self.assertEqual(set(dialog.required)-dialog.optional,REQUIRED)
    self.assertEqual(dialog.preview.item(1,4).text(),'不在库中的物料')
    dialog.add_alternate_mapping();self.assertIn('替代料1制造商',dialog.combos)
    dialog.remove_alternate_mapping('替代料1名称');self.assertNotIn('替代料1编码',dialog.combos)
    for f in REQUIRED:dialog.combos[f].setCurrentText(f)
    dialog.validate_mapping();self.assertEqual(dialog.result(),QDialog.Accepted)
    self.assertEqual(config.import_mapping('compare_bom')['物料名称'],'物料名称')
   finally:dialog.close()
 def test_sort_live_external_loading_and_no_database_writes(self):
  with tempfile.TemporaryDirectory() as tmp:
   db=Database(str(Path(tmp)/'db'));s=Services(db,ConfigService(tmp))
   r=s.create_material(dict(base_code='EXT',version='100',name='外部根',unit='pcs'));sid=s.create_sku(r)
   c=s.create_material(dict(base_code='NEW',version='100',name='不在库中的物料',unit='米'));s.save_bom(r,[dict(material_id=c,quantity=.12345,unit='米')])
   z=s.create_material(dict(base_code='AAA',version='100',name='排序',unit='pcs'));other=s.create_sku(z)
   main=MainWindow(s);main.compare_boms();dialog=main._compare_dialog
   try:
    dialog.left.setCurrentIndex(dialog.left.findData(sid));dialog.right.setCurrentIndex(dialog.right.findData(other));dialog.compare()
    main.sort_skus(1)
    self.assertEqual([dialog.left.itemData(i) for i in range(1,3)],[other,sid]);self.assertEqual(dialog.left.currentData(),sid)
    self.assertEqual(dialog.left.itemText(dialog.left.findData('load_external')),'载入其他BOM')
    before=db.conn.total_changes
    def mapping_exec(d):
     for field,combo in d.combos.items():
      if field in self.fixture()[0]:combo.setCurrentText(field)
     d.validate_mapping();return d.result()
    path=str(Path(__file__).parent/'fixtures/external_bom.xls')
    with patch('ui.bom_compare.QFileDialog.getOpenFileName',return_value=(path,'')),patch('ui.bom_compare.QInputDialog.getItem',return_value=('BOM',True)),patch.object(MaterialHeaderMappingDialog,'exec_',mapping_exec):
     dialog.right.setCurrentIndex(dialog.right.findData('load_external'));dialog.choose_source(1)
    self.assertTrue(str(dialog.right.currentData()).startswith('external:'));self.assertTrue(dialog._compared)
    self.assertEqual(db.conn.total_changes,before);self.assertEqual(len(s.materials()),3)
    self.assertEqual(dialog._rows[1].status,'modified') # only imported alternate differs
    self.assertNotIn('base_code',dialog._rows[1].changed)
    with patch('ui.main_window.MaterialDialog.exec_') as edit:
     dialog.edit_material(1,1);edit.assert_not_called()
    main.sort_skus(1);self.assertTrue(str(dialog.right.currentData()).startswith('external:'))
    self.assertTrue(dialog._compared)
   finally:dialog.close();main.close();db.close()
 def test_external_mapping_auto_compares_once_both_sides_are_ready(self):
  with tempfile.TemporaryDirectory() as tmp:
   db=Database(str(Path(tmp)/'db'));dialog=BomCompareDialog(Services(db,ConfigService(tmp)))
   path=str(Path(__file__).parent/'fixtures/external_bom.xls')
   def accept_mapping(d):
    for field,combo in d.combos.items():
     if field in self.fixture()[0]:combo.setCurrentText(field)
    d.validate_mapping();return d.result()
   try:
    with patch('ui.bom_compare.QFileDialog.getOpenFileName',return_value=(path,'')),patch('ui.bom_compare.QInputDialog.getItem',return_value=('BOM',True)),patch.object(MaterialHeaderMappingDialog,'exec_',accept_mapping):
     dialog.left.setCurrentIndex(dialog.left.findData('load_external'));dialog.choose_source(0)
     self.assertFalse(dialog._compared);self.assertEqual(dialog._rows,[])
     dialog.right.setCurrentIndex(dialog.right.findData('load_external'));dialog.choose_source(1)
     self.assertTrue(dialog._compared);self.assertTrue(dialog._rows)
     self.assertTrue(all(row.status=='same' for row in dialog._rows))
   finally:dialog.close();db.close()
 def test_xlsx_sheet_reader_and_identity_match_with_database(self):
  rows=self.fixture()
  with tempfile.TemporaryDirectory() as tmp:
   path=Path(tmp)/'bom.xlsx';wb=Workbook();wb.active.title='BOM'
   for row in rows:wb.active.append(row)
   wb.save(path);wb.close();source=ExcelSource(path)
   try:self.assertEqual(build_external_bom(source.rows('BOM'),{h:h for h in rows[0]})['children'][0]['base_code'],'NEW')
   finally:source.close()
 def test_two_external_boms_with_empty_database_and_cancel_restores_selection(self):
  from copy import deepcopy
  with tempfile.TemporaryDirectory() as tmp:
   db=Database(str(Path(tmp)/'db'));s=Services(db);dialog=BomCompareDialog(s)
   try:
    rows=self.fixture();tree=build_external_bom(rows,{h:h for h in rows[0]});changed=deepcopy(tree);changed['children'][0]['quantity']=2
    dialog.external_sources={}
    for key,label,qty in [('external:a','A',.12345),('external:b','B',2)]:
     path=Path(tmp)/(label+'.xlsx');book=Workbook();book.active.title='BOM'
     for row in rows:book.active.append(row)
     book.active.cell(3,6,qty);book.save(path);book.close()
     dialog.external_sources[key]={'label':label,'path':str(path),'sheet':'BOM','mapping':{h:h for h in rows[0]}}
    dialog.reload_choices();dialog.left.setCurrentIndex(dialog.left.findData('external:a'));dialog.right.setCurrentIndex(dialog.right.findData('external:b'))
    before=db.conn.total_changes;dialog.compare();self.assertEqual(dialog._rows[1].status,'modified')
    dialog.left.setCurrentIndex(dialog.left.findData('load_external'))
    with patch('ui.bom_compare.QFileDialog.getOpenFileName',return_value=('','')):dialog.choose_source(0)
    self.assertEqual(dialog.left.currentData(),'external:a');self.assertTrue(dialog._compared)
    self.assertEqual(db.conn.total_changes,before);self.assertEqual(s.materials(),[])
   finally:dialog.close();db.close()
 def test_compare_button_rereads_file_and_never_uses_stale_results(self):
  from openpyxl import load_workbook
  with tempfile.TemporaryDirectory() as tmp:
   db=Database(str(Path(tmp)/'db'));dialog=BomCompareDialog(Services(db));rows=self.fixture()
   try:
    for key in ('a','b'):
     path=Path(tmp)/(key+'.xlsx');book=Workbook();book.active.title='BOM'
     for row in rows:book.active.append(row)
     book.save(path);book.close()
     dialog.external_sources[key]={'label':key,'path':str(path),'sheet':'BOM','mapping':{h:h for h in rows[0]}}
    dialog.reload_choices();dialog.left.setCurrentIndex(dialog.left.findData('a'));dialog.right.setCurrentIndex(dialog.right.findData('b'))
    before=db.conn.total_changes;dialog.button.click();self.assertTrue(all(r.status=='same' for r in dialog._rows))
    path=Path(tmp)/'b.xlsx';book=load_workbook(path);book['BOM'].cell(3,6,9.5);book.save(path);book.close()
    dialog.button.click();self.assertEqual(dialog._rows[1].right['quantity'],9.5);self.assertEqual(dialog._rows[1].status,'modified')
    book=load_workbook(path);book['BOM'].cell(1,6,'Changed header');book.save(path);book.close()
    with patch('ui.bom_compare.QMessageBox.warning') as warning:
     dialog.button.click();warning.assert_called_once();self.assertFalse(dialog._compared);self.assertEqual(dialog._rows,[])
    path.unlink()
    with patch('ui.bom_compare.QMessageBox.warning') as warning:
     dialog.button.click();warning.assert_called_once();self.assertEqual(dialog._rows,[])
    self.assertEqual(db.conn.total_changes,before)
   finally:dialog.close();db.close()
 def test_full_export_roundtrip_and_exchange_by_codes(self):
  with tempfile.TemporaryDirectory() as tmp:
   db=Database(str(Path(tmp)/'db'));s=Services(db)
   root,a,b=[s.create_material(dict(base_code=c,version='1',name=c,unit='pcs')) for c in ('ROOTMAT','A','B')]
   sid=s.create_sku(root);s.save_bom(root,[dict(material_id=a,quantity=1,unit='pcs')]);s.add_alternate(s.bom(root)[0]['id'],b)
   path=Path(tmp)/'bom.xlsx';s.export_sku(sid,path);source=ExcelSource(path)
   try:rows=source.rows('ROOTMAT')
   finally:source.close()
   mapping={h:h for h in rows[0]};mapping.update({'父项物料编码':'父项编码','父项版本号':'父项版本','规格':'物料规格描述','替代料1编码':'替代料1物料编码','替代料1制造商':'替代料1供应商'})
   for alias in ('父项编码','父项版本','物料规格描述','替代料1物料编码','替代料1供应商'):mapping.pop(alias,None)
   tree=build_external_bom(rows,mapping)
   self.assertTrue(all(d.status=='same' for d in compare_trees(snapshot(s,dict(s.skus()[0])),tree)))
   s.save_bom(root,[dict(material_id=b,quantity=1,unit='pcs',alternate_material_ids=[a])])
   self.assertEqual(compare_trees(tree,snapshot(s,dict(s.skus()[0])))[1].status,'exchanged');db.close()
 def test_external_file_remembered_across_windows_and_missing_falls_back(self):
  from PyQt5.QtCore import QStandardPaths
  from shutil import copyfile
  with tempfile.TemporaryDirectory() as tmp:
   config=ConfigService(tmp);db=Database(str(Path(tmp)/'db'));s=Services(db,config)
   folder=Path(tmp)/'外部BOM';folder.mkdir();path=folder/'对比.xls';copyfile(Path(__file__).parent/'fixtures/external_bom.xls',path)
   dialog=BomCompareDialog(s)
   try:
    with patch('ui.bom_compare.QFileDialog.getOpenFileName',return_value=(str(path),'')),patch('ui.bom_compare.QInputDialog.getItem',return_value=('',False)):
     dialog.left.setCurrentIndex(dialog.left.findData('load_external'));dialog.choose_source(0)
    self.assertEqual(config.load()['database']['last_compare_bom_file'],'外部BOM/对比.xls')
   finally:dialog.close()
   s.config_service=ConfigService(tmp);dialog=BomCompareDialog(s)
   try:
    for exists in (True,False):
     if not exists:path.unlink()
     with patch('ui.bom_compare.QFileDialog.getOpenFileName',return_value=('','')) as picker:
      dialog.right.setCurrentIndex(dialog.right.findData('load_external'));dialog.choose_source(1)
     expected=str(path) if exists else QStandardPaths.writableLocation(QStandardPaths.DesktopLocation)
     self.assertEqual(picker.call_args.args[2],expected)
     self.assertEqual(config.load()['database']['last_compare_bom_file'],'外部BOM/对比.xls')
   finally:dialog.close();db.close()
 def test_sku_click_selects_whole_visible_row(self):
  from PyQt5.QtTest import QTest
  with tempfile.TemporaryDirectory() as tmp:
   db=Database(str(Path(tmp)/'db'));s=Services(db,ConfigService(tmp))
   for code in ('A','B'):s.create_sku(s.create_material(dict(base_code=code,version='1',name=code,unit='pcs')))
   main=MainWindow(s);main.show();self.app.processEvents()
   try:
    table=main.table;QTest.mouseClick(table.viewport(),Qt.LeftButton,pos=table.visualItemRect(table.item(1,2)).center())
    self.assertEqual({index.row() for index in table.selectedIndexes()},{1})
    self.assertEqual({index.column() for index in table.selectedIndexes()},{column for column in range(table.columnCount()) if not table.isColumnHidden(column)})
   finally:main.close();db.close()
