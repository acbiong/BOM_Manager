import os
os.environ['QT_QPA_PLATFORM']='offscreen'
import sys,tempfile,unittest
from pathlib import Path
from unittest.mock import patch
sys.path.insert(0,str(Path(__file__).parents[1]/'src'))
from openpyxl import Workbook
from PyQt5.QtWidgets import QApplication,QInputDialog,QFileDialog,QDialog,QMessageBox
from database.database import Database
from services.services import Services,MaterialHeaderError
from ui.main_window import MainWindow,MaterialHeaderMappingDialog,MaterialImportResultDialog,choose_material_sheet

class MaterialSheetsTests(unittest.TestCase):
 @classmethod
 def setUpClass(cls):cls.app=QApplication.instance() or QApplication([])
 def setUp(self):
  self.tmp=tempfile.TemporaryDirectory();self.path=Path(self.tmp.name)/'multi.xlsx';self.db=Database(str(Path(self.tmp.name)/'test.db'));self.s=Services(self.db)
  self.required=['物料基础编码','物料版本号','物料名称','物料规格描述','单位','物料制造商','物料MPN'];self.headers=['编码','版本','名称','描述','单位','厂家','料号']
  wb=Workbook();wb.active.title='目录';wb.active.append(['物料数据','其他工作表']);ws=wb.create_sheet('物料数据');ws.append(self.headers);ws.append(['RIGHT','1','正确物料','','pcs','厂商','MPN']);wb.save(self.path);wb.close()
 def tearDown(self):self.db.close();self.tmp.cleanup()
 def test_selected_sheet_headers_preview_and_import(self):
  with self.assertRaises(MaterialHeaderError) as ctx:self.s.import_materials(self.path,sheet_name='物料数据')
  self.assertEqual(ctx.exception.headers,self.headers)
  dialog=MaterialHeaderMappingDialog(self.path,ctx.exception.headers,ctx.exception.missing,sheet_name='物料数据')
  self.assertEqual(dialog.preview.item(0,0).text(),'RIGHT')
  for need,source in zip(self.required,self.headers):dialog.combos[need].setCurrentText(source)
  self.assertNotIn('目录',[dialog.combos[self.required[0]].itemText(i) for i in range(dialog.combos[self.required[0]].count())])
  ok,fail=self.s.import_materials(self.path,header_mapping=dialog.mapping(),sheet_name='物料数据');self.assertEqual(ok,['RIGHT']);self.assertEqual(fail,[]);dialog.close()
 def test_multisheet_requires_choice_and_invalid_sheet_never_writes(self):
  for kwargs in ({},{'sheet_name':'不存在'}):
   with self.assertRaises(ValueError):self.s.import_materials(self.path,**kwargs)
  self.assertEqual(self.s.materials(),[])
 def test_single_sheet_skips_choice_and_cancel_aborts(self):
  with patch.object(QInputDialog,'getItem',return_value=('',False)):
   with self.assertRaisesRegex(ValueError,'取消'):choose_material_sheet(self.path)
  single=Path(self.tmp.name)/'single.xlsx';wb=Workbook();wb.active.title='唯一工作表';wb.save(single);wb.close()
  with patch.object(QInputDialog,'getItem') as choose:self.assertEqual(choose_material_sheet(single),'唯一工作表');choose.assert_not_called()
 def test_main_import_passes_selected_sheet_through_mapping_retry(self):
  main=MainWindow(self.s)
  def accept(dialog):
   self.assertEqual(dialog.preview.item(0,0).text(),'RIGHT')
   for need,source in zip(self.required,self.headers):dialog.combos[need].setCurrentText(source)
   return QDialog.Accepted
  try:
   with patch.object(QFileDialog,'getOpenFileName',return_value=(str(self.path),'')),patch.object(QInputDialog,'getItem',return_value=('物料数据',True)) as choose,patch.object(MaterialHeaderMappingDialog,'exec_',accept),patch.object(MaterialImportResultDialog,'exec_',return_value=QDialog.Accepted),patch.object(QMessageBox,'warning') as warning:
    main.import_materials();warning.assert_not_called();self.assertEqual(choose.call_args.args[3],['目录','物料数据'])
   self.assertEqual([r['base_code'] for r in self.s.materials()],['RIGHT'])
  finally:main.close()
 def test_missing_required_fields_reported_in_result_for_standard_and_mapped_headers(self):
  for mapped in (False,True):
   with self.subTest(mapped=mapped):
    wb=Workbook();ws=wb.active;ws.append(self.headers if mapped else self.required)
    ws.append(['A','1','物料A','','pcs',None,''])
    ws.append(['B',None,'   ','',None,'   ',''])
    ws.append([None,'1','物料C','','pcs','厂商',''])
    wb.save(self.path);wb.close()
    mapping=dict(zip(self.required,self.headers)) if mapped else None
    ok,fail=self.s.import_materials(self.path,header_mapping=mapping)
    self.assertEqual(ok,[]);self.assertEqual([row[3] for row in fail],['必填项缺失：物料制造商','必填项缺失：物料版本号、物料名称、单位、物料制造商','必填项缺失：物料基础编码'])
    dialog=MaterialImportResultDialog(ok,fail)
    self.assertIn('第2行：A 物料A\n原因：必填项缺失：物料制造商',dialog.details.toPlainText())
    self.assertIn('物料版本号、物料名称、单位、物料制造商',dialog.details.toPlainText());dialog.close()
    self.assertEqual(self.s.materials(),[])
 def test_blank_code_and_name_ignored_before_validation(self):
  for mapped in (False,True):
   with self.subTest(mapped=mapped):
    wb=Workbook();ws=wb.active;ws.append(self.headers if mapped else self.required)
    ws.append([None,'99',None,'备注','非法单位',None,None])
    ws.append(['   ',None,'  ',None,None,None,None])
    ws.append(['ONLYCODE','1',None,'','pcs','厂商',''])
    ws.append([None,'1','只有名称','','pcs','厂商',''])
    ws.append([0,'1','数字编码','','pcs','厂商',''])
    wb.save(self.path);wb.close();progress=[]
    ok,fail=self.s.import_materials(self.path,header_mapping=dict(zip(self.required,self.headers)) if mapped else None,progress_callback=lambda n,total:progress.append((n,total)))
    self.assertEqual([r[0] for r in fail],[4,5]);self.assertEqual(progress[-1],(5,5));self.assertEqual([r['base_code'] for r in self.s.materials()],['0'])
    self.assertEqual(ok,['0'])
    self.db.conn.execute('DELETE FROM materials');self.db.conn.commit()
