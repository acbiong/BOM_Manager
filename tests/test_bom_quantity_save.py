import os,sys,tempfile,unittest,math
os.environ['QT_QPA_PLATFORM']='offscreen'
from pathlib import Path
from unittest.mock import patch
sys.path.insert(0,str(Path(__file__).parents[1]/'src'))
from database.database import Database
from services.services import Services
from ui.main_window import BomWindow,QuantityEdit
from domain.quantity import format_quantity
from PyQt5.QtWidgets import QApplication,QMessageBox,QProgressDialog
from PyQt5.QtCore import Qt

class BomQuantitySaveTests(unittest.TestCase):
 @classmethod
 def setUpClass(cls):cls.app=QApplication.instance() or QApplication([])
 def test_quantity_display_and_precision(self):
  for raw,expected in [(1,'1'),('1.23000','1.23'),(.00012345,'0.00012345'),(2.3456789,'2.3456789')]:
   self.assertEqual(format_quantity(raw),expected)
  editor=QuantityEdit();editor.setText('2.3456789');self.assertEqual(editor.text(),'2.3456789');self.assertEqual(editor.lineEdit().text(),'2.3456789')
  editor.close()
 def test_save_precision_progress_deletion_focus_and_query_budget(self):
  with tempfile.TemporaryDirectory() as tmp:
   db=Database(str(Path(tmp)/'db'));s=Services(db)
   root=s.create_material(dict(base_code='R',version='1',name='root',unit='pcs'));s.create_sku(root)
   children=[s.create_material(dict(base_code=f'C{i}',version='1',name='child',unit='pcs')) for i in range(80)]
   values=[dict(material_id=mid,quantity=.01234567,unit='pcs') for mid in children]
   events=[];s.save_bom(root,values,progress=lambda n,t:events.append(n))
   self.assertEqual(events,sorted(events));self.assertEqual(events[-1],100)
   queries=[];db.conn.set_trace_callback(queries.append);s.save_bom(root,values);db.conn.set_trace_callback(None)
   self.assertLess(sum('FROM alternates' in q for q in queries),8)
   self.assertFalse(any(q.startswith('INSERT INTO material_bom_items') for q in queries))
   w=BomWindow(s,root)
   try:
    w.show();self.app.processEvents();item=w.root_item.child(50);w.tree.setCurrentItem(item);w.tree.scrollToItem(item);self.app.processEvents()
    before=w.tree.visualItemRect(item).top();w.remove_relation(item)
    self.assertIs(w.tree.currentItem(),item);self.assertEqual(w.tree.visualItemRect(item).top(),before)
    self.assertTrue(item.data(0,Qt.UserRole)['_deleted']);self.assertEqual(w.root_item.child(0).text(1),'0.01234567')
    progress=[];original=QProgressDialog.setValue
    def set_value(dialog,n):progress.append(n);return original(dialog,n)
    with patch.object(QProgressDialog,'setValue',set_value),patch.object(QMessageBox,'information'):
     self.assertTrue(w.save())
    self.assertEqual(progress[-1],100);self.assertIn(90,progress)
    self.assertEqual(len(s.bom(root)),79);self.assertEqual(s.bom(root)[0]['quantity'],.01234567)
    path=Path(tmp)/'bom.xlsx';s.export_sku(s.skus()[0]['id'],path)
    from openpyxl import load_workbook
    wb=load_workbook(path);sheet=wb['root'];self.assertEqual(sheet.cell(3,10).value,.01234567);self.assertNotEqual(sheet.cell(3,10).number_format,'0.0');wb.close()
    for invalid in (0,-1,float('inf'),float('nan')):
     with self.assertRaises(ValueError):s.save_bom(root,[dict(material_id=children[0],quantity=invalid,unit='pcs')])
   finally:w.dirty=False;w.close();db.close()
