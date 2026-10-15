import tempfile,unittest,os,sys
from pathlib import Path
sys.path.insert(0,str(Path(__file__).parents[1]/'src'))
from database.database import Database
from services.services import Services,MaterialHeaderError
class TestServices(unittest.TestCase):
 def setUp(self):
  from types import SimpleNamespace
  self.tmp=tempfile.TemporaryDirectory(); self.f=SimpleNamespace(name=str(Path(self.tmp.name)/'bom.db')); self.s=Services(Database(self.f.name)); self.addCleanup(self.tmp.cleanup); self.addCleanup(self.s.db.close)
 def test_material_unique_and_sku(self):
  i=self.s.create_material({'base_code':'A','version':'1','name':'A','unit':'pcs'})
  with self.assertRaises(Exception): self.s.create_material({'base_code':'A','version':'1','name':'dup','unit':'pcs'})
  self.s.create_sku(i); self.assertEqual(len(self.s.skus()),1)
 def test_bom_and_idle(self):
  a=self.s.create_material({'base_code':'A','version':'1','name':'A','unit':'pcs'}); b=self.s.create_material({'base_code':'B','version':'1','name':'B','unit':'pcs'}); c=self.s.create_material({'base_code':'C','version':'1','name':'C','unit':'pcs'}); self.s.save_bom(a,[{'material_id':b,'quantity':2,'unit':'pcs','level_no':1},{'material_id':c,'quantity':1,'unit':'pcs','level_no':2,'parent_item_id':1}]); rows=self.s.bom(a); self.assertEqual([r['material_id'] for r in rows],[b,c]); self.assertEqual([r['level_no'] for r in rows],[1,2]); flags={r['id']:r['has_bom'] for r in self.s.materials()}; self.assertEqual(flags[a],'Y'); self.assertEqual(flags[b],'Y'); self.assertEqual(flags[c],'N')
  self.s.save_bom(b,[]); self.assertEqual({r['id']:r['has_bom'] for r in self.s.materials()}[b],'N')
 def test_deleting_last_child_notifies_parent_material(self):
  from infrastructure.events import WindowRefreshBus
  root=self.s.create_material({'base_code':'ROOT','version':'1','name':'根','unit':'pcs'}); parent=self.s.create_material({'base_code':'PARENT','version':'1','name':'父物料','unit':'pcs'}); child=self.s.create_material({'base_code':'CHILD','version':'1','name':'子物料','unit':'pcs'}); self.s.save_bom(root,[{'material_id':parent,'quantity':1,'unit':'pcs'},{'material_id':child,'quantity':1,'unit':'pcs','parent_item_id':1}]); events=[]; WindowRefreshBus.subscribe(events.append); self.s.save_bom(root,[{'material_id':parent,'quantity':1,'unit':'pcs'}]); WindowRefreshBus.unsubscribe(events.append); self.assertIn(parent,[e.get('root') for e in events]); self.assertEqual({r['id']:r['has_bom'] for r in self.s.materials()}[parent],'N')
 def test_delete_bom_item_publishes_parent_update(self):
  from infrastructure.events import WindowRefreshBus
  root=self.s.create_material({'base_code':'ROOT2','version':'1','name':'根2','unit':'pcs'}); child=self.s.create_material({'base_code':'CHILD2','version':'1','name':'子2','unit':'pcs'}); self.s.save_bom(root,[{'material_id':child,'quantity':1,'unit':'pcs'}]); events=[]; WindowRefreshBus.subscribe(events.append); item=self.s.bom(root)[0]['id']; self.s.delete_bom_item(item); WindowRefreshBus.unsubscribe(events.append); self.assertIn({'type':'bom_changed','root':root,'shared':True},events); self.assertEqual({r['id']:r['has_bom'] for r in self.s.materials()}[root],'N')
 def test_import_export(self):
  from openpyxl import Workbook,load_workbook
  p=self.f.name+'.xlsx'; w=Workbook(); w.active.append(['物料基础编码','物料版本号','物料名称','物料规格描述','单位','物料制造商','物料MPN']); w.active.append(['X','1','测试','','pcs','厂商','M']); w.save(p); ok,fail=self.s.import_materials(p); self.assertEqual(ok,['X']); self.assertEqual(fail,[]); out=self.f.name+'-out.xlsx'; self.s.export_materials(out); self.assertEqual(load_workbook(out).sheetnames,['变更记录','物料库']); os.unlink(p); os.unlink(out)
 def test_delete_reference_blocked(self):
  a=self.s.create_material({'base_code':'A','version':'1','name':'A','unit':'pcs'}); b=self.s.create_material({'base_code':'B','version':'1','name':'B','unit':'pcs'}); self.s.save_bom(a,[{'material_id':b,'quantity':1,'unit':'pcs'}])
  with self.assertRaises(ValueError): self.s.delete_material(b)
 def test_bom_import_and_project(self):
     from openpyxl import Workbook
     p=self.f.name+'-bom.xlsx'; a=self.s.create_material({'base_code':'A','version':'1','name':'A','unit':'pcs'}); b=self.s.create_material({'base_code':'B','version':'1','name':'B','unit':'pcs'}); w=Workbook(); w.active.append(['父项物料编码','父项版本号','物料基础编码','物料版本号','物料数量','位置']); w.active.append(['A','1','B','1','3','P1']); w.save(p); self.assertEqual(self.s.import_bom(p,a),1); self.assertEqual(self.s.bom(a)[0]['quantity'],3); self.s.update_project({'project_name':'项目'}); self.assertEqual(self.s.project()['project_name'],'项目'); os.unlink(p)
 def test_bom_consistency_and_sku_version(self):
  from openpyxl import Workbook
  a=self.s.create_material({'base_code':'A','version':'1','name':'A','unit':'pcs'}); b=self.s.create_material({'base_code':'B','version':'1','name':'B','unit':'pcs'}); self.s.save_bom(a,[{'material_id':b,'quantity':2,'unit':'pcs','position':'P'}]); same,detail=self.s.validate_bom_rows(a,[{'material_id':b,'quantity':2,'unit':'pcs','position':'P'}]); self.assertTrue(same); sid=self.s.create_sku(a); self.s.change_sku_version(sid,'2'); self.assertEqual(self.s.skus()[0]['version'],'2')
 def test_sqlite_schema_validation(self):
  ok,missing=self.s.db.validate_schema(); self.assertTrue(ok); self.assertEqual(missing,[])
  self.assertIsNotNone(self.s.db.conn.execute("SELECT 1 FROM sqlite_master WHERE type='table' AND name='material_bom_items'").fetchone())
 def test_existing_sqlite_database_is_reused_without_overwrite(self):
  """打开已有数据库应保留数据，不应以空库覆盖。"""
  existing=self.f.name+'-existing.db'
  first=Database(existing); Services(first).create_material({'base_code':'KEEP','version':'1','name':'已有物料','unit':'pcs'}); first.close()
  reopened=Database(existing)
  self.assertEqual(reopened.conn.execute('SELECT name FROM materials WHERE base_code=?',('KEEP',)).fetchone()['name'],'已有物料')
  reopened.close()
  new_path=self.f.name+'-new.db'; created=Database(new_path)
  self.assertTrue(created.validate_schema()[0]); created.close()
 def test_alternate_crud(self):
  a=self.s.create_material({'base_code':'A','version':'1','name':'A','unit':'pcs'}); b=self.s.create_material({'base_code':'B','version':'1','name':'B','unit':'pcs'}); c=self.s.create_material({'base_code':'C','version':'1','name':'C','unit':'pcs'}); self.s.save_bom(a,[{'material_id':b,'quantity':1,'unit':'pcs'}]); bid=self.s.bom(a)[0]['id']; self.s.add_alternate(bid,c); self.assertEqual(len(self.s.alternates(bid)),1); self.s.remove_alternate(self.s.alternates(bid)[0]['id']); self.assertEqual(len(self.s.alternates(bid)),0)
 def test_recursive_delete_and_event_bus(self):
  from infrastructure.events import WindowRefreshBus
  a=self.s.create_material({'base_code':'A','version':'1','name':'A','unit':'pcs'}); b=self.s.create_material({'base_code':'B','version':'1','name':'B','unit':'pcs'}); c=self.s.create_material({'base_code':'C','version':'1','name':'C','unit':'pcs'}); events=[]; WindowRefreshBus.subscribe(events.append); self.s.save_bom(a,[{'material_id':b,'quantity':1,'unit':'pcs','level_no':1},{'material_id':c,'quantity':1,'unit':'pcs','level_no':2,'parent_item_id':1}]); self.assertTrue(events); first=self.s.bom(a)[0]['id']; self.s.delete_bom_item(first); self.assertEqual(len(self.s.bom(a)),0); WindowRefreshBus.unsubscribe(events.append)
 def test_material_change_event(self):
  from infrastructure.events import WindowRefreshBus
  events=[]; WindowRefreshBus.subscribe(events.append); self.s.create_material({'base_code':'E','version':'1','name':'事件','unit':'pcs'}); self.assertTrue(any(e.get('type')=='material_changed' for e in events)); WindowRefreshBus.unsubscribe(events.append)
 def test_duplicate_import_overwrite_or_skip(self):
  from openpyxl import Workbook
  p=self.f.name+'-dup.xlsx'; self.s.create_material({'base_code':'D','version':'1','name':'旧','unit':'pcs'}); w=Workbook(); w.active.append(['物料基础编码','物料版本号','物料名称','物料规格描述','单位','物料制造商','物料MPN']); w.active.append(['D','1','新','','pcs','厂商','M']); w.save(p); ok,fail=self.s.import_materials(p,lambda d,o:'overwrite'); self.assertEqual(ok,['D']); self.assertEqual(self.s.materials()[0]['name'],'新')
 def test_identical_duplicate_auto_skip(self):
  from openpyxl import Workbook
  p=self.f.name+'-same.xlsx'; self.s.create_material({'base_code':'S','version':'1','name':'相同','description':'','unit':'pcs','manufacturer':'厂','mpn':'M'}); w=Workbook(); w.active.append(['物料基础编码','物料版本号','物料名称','物料规格描述','单位','物料制造商','物料MPN']); w.active.append(['S','1','相同','','pcs','厂','M']); w.save(p); ok,fail=self.s.import_materials(p,lambda d,o:'cancel'); self.assertEqual(ok,[]); self.assertEqual(len(fail),1); self.assertIn('自动跳过',fail[0][3])
 def test_material_import_reports_headers_for_mapping(self):
  from openpyxl import Workbook
  p=self.f.name+'-mapping.xlsx'; w=Workbook(); w.active.append(['编码','版本','名称','描述','单位','厂家','料号']); w.active.append(['MAP','1','映射物料','','pcs','厂商','M']); w.save(p)
  with self.assertRaises(MaterialHeaderError) as ctx:self.s.import_materials(p)
  self.assertEqual(ctx.exception.headers,['编码','版本','名称','描述','单位','厂家','料号'])
  required=['物料基础编码','物料版本号','物料名称','物料规格描述','单位','物料制造商','物料MPN']; ok,fail=self.s.import_materials(p,header_mapping=dict(zip(required,ctx.exception.headers)))
  self.assertEqual(ok,['MAP']); self.assertEqual(fail,[])
if __name__=='__main__': unittest.main()
