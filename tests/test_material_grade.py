import os
os.environ['QT_QPA_PLATFORM']='offscreen'
import sys,tempfile,unittest
from pathlib import Path
sys.path.insert(0,str(Path(__file__).parents[1]/'src'))
from openpyxl import Workbook,load_workbook
from PyQt5.QtCore import Qt
from PyQt5.QtWidgets import QApplication,QLabel,QDialog
from database.database import Database
from services.services import Services
from ui.main_window import MaterialWindow,MaterialDialog,BomWindow,MaterialHeaderMappingDialog,MaterialConflictDialog

class MaterialGradeTests(unittest.TestCase):
 @classmethod
 def setUpClass(cls):cls.app=QApplication.instance() or QApplication([])
 def setUp(self):
  self.tmp=tempfile.TemporaryDirectory();self.path=Path(self.tmp.name);self.db=Database(str(self.path/'db'));self.s=Services(self.db)
 def tearDown(self):self.db.close();self.tmp.cleanup()
 def create(self,code,grade=''):
  return self.s.create_material(dict(base_code=code,version='1',name=code,unit='pcs',manufacturer='厂',grade=grade))
 def test_legacy_migration_default_and_version_copy(self):
  mid=self.create('A');self.db.conn.execute('ALTER TABLE materials DROP COLUMN grade');self.db.conn.commit();self.db.close()
  self.db=Database(str(self.path/'db'));self.s=Services(self.db)
  self.assertEqual(self.s.material_snapshot(mid)['grade'],'')
  self.s.update_material(mid,dict(self.s.material_snapshot(mid),grade='工业级'))
  parent=self.create('P');self.s.save_bom(parent,[dict(material_id=mid,quantity=1,unit='pcs')])
  new=self.s.change_material_version(mid,'2',set())
  versions=[r for r in self.s.materials() if r['base_code']=='A']
  self.assertEqual(len(versions),2);self.assertEqual({r['grade'] for r in versions},{'工业级'})
 def test_excel_optional_mapped_conflict_and_export(self):
  headers=['物料基础编码','物料版本号','物料名称','物料规格描述','单位','物料制造商','物料MPN']
  path=self.path/'in.xlsx'
  def write(grade_header=None,grade=None):
   wb=Workbook();wb.active.append(headers+([grade_header] if grade_header else []));wb.active.append(['A','1','A','','pcs','厂','']+([grade] if grade_header else []));wb.save(path);wb.close()
  write();self.assertEqual(self.s.import_materials(path)[0],['A'])
  write('等级分类','医疗级');mapping={h:h for h in headers};mapping['物料等级']='等级分类'
  dialog=MaterialHeaderMappingDialog(path,headers+['等级分类'],[])
  dialog.combos['物料等级'].setCurrentText('等级分类');dialog.validate_mapping();self.assertEqual(dialog.result(),QDialog.Accepted);dialog.close()
  def overwrite(incoming,old):
   conflict=MaterialConflictDialog(incoming,old);self.assertEqual(conflict.table.item(1,5).text(),'医疗级');conflict.close();return 'overwrite'
  self.assertEqual(self.s.import_materials(path,overwrite,header_mapping=mapping)[0],['A'])
  self.assertIn('自动跳过',self.s.import_materials(path,header_mapping=mapping)[1][0][3])
  write();self.s.import_materials(path);self.assertEqual(self.s.materials()[0]['grade'],'医疗级')
  out=self.path/'out.xlsx';self.s.export_materials(out);wb=load_workbook(out);rows=list(wb['物料库'].values);wb.close()
  self.assertEqual(rows[0][4:7],('单位','物料等级','物料制造商'));self.assertEqual(rows[1][5],'医疗级')
  write('物料等级','');self.s.import_materials(path,lambda *_:'overwrite');self.assertEqual(self.s.materials()[0]['grade'],'')
 def test_edit_bom_refresh_and_alignment_after_table_refresh(self):
  parent=self.create('P');child=self.create('C','A级');self.s.save_bom(parent,[dict(material_id=child,quantity=1,unit='pcs')])
  material=MaterialWindow(self.s);bom=BomWindow(self.s,parent)
  try:
   for _ in range(2):
    material.refresh()
    for col in [1,4,8,9,10]:self.assertEqual(material.t.item(0,col).textAlignment(),Qt.AlignCenter)
   self.assertEqual(material._column_labels[5],'物料等级')
   item=bom.root_item.child(0);self.assertEqual(item.textAlignment(2),Qt.AlignCenter)
   bom.show_item_info(item);self.assertIsInstance(bom.fields['grade'],QLabel);self.assertEqual(bom.fields['grade'].text(),'A级')
   from unittest.mock import patch
   def edit(dialog):dialog.e['grade'].setText('B级');return QDialog.Accepted
   with patch.object(MaterialDialog,'exec_',edit):material.edit_current(0,5)
   self.assertEqual(self.s.material_snapshot(child)['grade'],'B级');self.assertTrue(bom.refresh_action.isEnabled())
   bom.refresh_external();bom.show_item_info(bom.root_item.child(0));self.assertEqual(bom.fields['grade'].text(),'B级')
   self.assertTrue(any('物料等级：A级 → B级' in r[0] for r in self.db.conn.execute('SELECT detail FROM change_logs')))
  finally:material.close();bom.dirty=False;bom.close()
