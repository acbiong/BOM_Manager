"""采购BOM实际Excel读回及导出状态联动回归。"""
import os
os.environ['QT_QPA_PLATFORM']='offscreen'
import sys,tempfile,unittest
from pathlib import Path
from unittest.mock import patch
sys.path.insert(0,str(Path(__file__).parents[1]/'src'))
from openpyxl import load_workbook
from PyQt5.QtWidgets import QApplication,QFileDialog,QMessageBox,QMenu
from PyQt5.QtCore import QPoint
from database.database import Database
from services.services import Services
from ui.main_window import MainWindow

class PurchaseExportTests(unittest.TestCase):
 @classmethod
 def setUpClass(cls):cls.app=QApplication.instance() or QApplication([])
 def setUp(self):
  self.tmp=tempfile.TemporaryDirectory();self.path=Path(self.tmp.name);self.db=Database(str(self.path/'bom.db'));self.s=Services(self.db)
  self.ids={code:self.s.create_material(dict(base_code=code,version='1',name=code,description='规格'*70,unit='pcs',grade='A',manufacturer='厂家',mpn='MPN-'+code,mp_flag='P')) for code in ('ROOT','A','B','C','D','X','Y')}
  self.sid=self.s.create_sku(self.ids['ROOT']);self.main=None
 def tearDown(self):
  if self.main:
   for w in list(self.main._material_bom_windows.values()):w.dirty=False;w.close()
   self.main.close()
  self.db.close();self.tmp.cleanup()
 def item(self,code,quantity,remark=''):
  return dict(material_id=self.ids[code],quantity=quantity,unit='pcs',remark=remark)
 def export(self,root=None):
  path=self.path/'采购.xlsx';self.s.export_purchase_bom(root or self.ids['ROOT'],path)
  return load_workbook(path)
 def test_all_levels_shared_paths_versions_alternates_and_styles(self):
  self.s.save_bom(self.ids['C'],[self.item('D',.5)])
  self.s.save_bom(self.ids['A'],[self.item('C',1.5,'A分支备注')])
  self.s.save_bom(self.ids['B'],[self.item('C',2,'B分支备注')])
  version2=self.s.create_material(dict(base_code='C',version='2',name='C新版本',unit='pcs'))
  self.s.save_bom(self.ids['ROOT'],[self.item('A',2),self.item('B',3),self.item('C',.5),dict(material_id=version2,quantity=1,unit='pcs')])
  for parent,alt in [('A','X'),('B','Y')]:
   edge=self.db.conn.execute('SELECT id FROM material_bom_items WHERE parent_material_id=? AND child_material_id=?',(self.ids[parent],self.ids['C'])).fetchone()[0]
   self.s.add_alternate(edge,self.ids[alt])
  wb=self.export()
  try:
   sh=wb.active;headers=[c.value for c in sh[1]]
   expected=['序号','物料编码','物料名称','规格','用量','单位','物料等级','M/P','制造商','MPN']
   expected += [f'替代料{n}{f}' for n in (1,2) for f in ('物料编码','物料名称','规格','制造商','MPN')]+['备注']
   self.assertEqual(headers,expected)
   rows={row[1]:dict(zip(headers,row)) for row in sh.iter_rows(min_row=2,values_only=True)}
   self.assertEqual(set(rows),{'A1','B1','C1','D1','C2'})
   self.assertEqual({code:row['用量'] for code,row in rows.items()},{'A1':2,'B1':3,'C1':9.5,'D1':4.8,'C2':1})
   self.assertEqual(rows['C1']['M/P'],'P');self.assertEqual(rows['C1']['物料等级'],'A')
   self.assertEqual(rows['C1']['替代料1物料编码'],'X1');self.assertEqual(rows['C1']['替代料2物料编码'],'Y1')
   self.assertEqual(rows['C1']['替代料1规格'],'规格'*70)
   self.assertIn('A分支备注\nB分支备注',rows['C1']['备注'])
   self.assertIn('ROOT1/A1/C1',rows['C1']['备注']);self.assertIn('无替代料',rows['C1']['备注'])
   self.assertEqual([sh.cell(r,1).value for r in range(2,sh.max_row+1)],list(range(1,6)))
   self.assertTrue(all(5<=d.width<=60 for d in sh.column_dimensions.values()))
   for cells in sh:
    for cell in cells:
     self.assertTrue(cell.alignment.wrap_text)
     for side in ('left','right','top','bottom'):
      self.assertEqual(getattr(cell.border,side).style,'thin');self.assertEqual(getattr(cell.border,side).color.rgb,'FF000000')
     if headers[cell.column-1] in ('序号','用量','单位','物料等级','M/P'):self.assertEqual(cell.alignment.horizontal,'center')
     if cell.row==1:self.assertTrue(cell.font.bold)
    if cells[0].row>1:self.assertEqual(cells[4].number_format,'0.0')
  finally:wb.close()
 def test_empty_bom_and_non_sku_root(self):
  wb=self.export()
  try:self.assertEqual(wb.active.max_row,1);self.assertEqual(wb.active.max_column,11)
  finally:wb.close()
  self.s.save_bom(self.ids['A'],[self.item('C',1.5)])
  wb=self.export(self.ids['A'])
  try:self.assertEqual(wb.active.cell(2,2).value,'C1');self.assertEqual(wb.active.cell(2,5).value,1.5)
  finally:wb.close()
 def test_decimal_products_rounded_only_after_aggregation(self):
  self.s.save_bom(self.ids['A'],[self.item('C',.5)])
  self.s.save_bom(self.ids['B'],[self.item('C',.5)])
  self.s.save_bom(self.ids['ROOT'],[self.item('A',.5),self.item('B',.5)])
  wb=self.export()
  try:
   rows={r[1]:r[4] for r in wb.active.iter_rows(min_row=2,values_only=True)}
   self.assertEqual(rows['C1'],.5) # .25 + .25, not .3 + .3
  finally:wb.close()
 def test_editor_states_match_regular_export(self):
  self.s.save_bom(self.ids['ROOT'],[self.item('A',1)])
  self.main=MainWindow(self.s);w=self.main.open_bom_id(self.sid)
  self.assertTrue(w.purchase_export_action.isEnabled())
  w.show_item_info(w.root_item.child(0));w.fields['quantity'].setValue(2)
  self.assertFalse(w.purchase_export_action.isEnabled())
  with patch.object(QFileDialog,'getSaveFileName') as choose:w.export_purchase_bom();choose.assert_not_called()
  with patch.object(QMessageBox,'information'):self.assertTrue(w.save())
  self.assertTrue(w.purchase_export_action.isEnabled())
  self.s.save_bom(self.ids['ROOT'],[self.item('A',3)])
  self.assertFalse(w.purchase_export_action.isEnabled());w.refresh_external();self.assertFalse(w.purchase_export_action.isEnabled())
  with patch.object(self.s,'save_bom',side_effect=ValueError('测试保存失败')),patch.object(QMessageBox,'warning'):self.assertFalse(w.save())
  self.assertFalse(w.purchase_export_action.isEnabled())
  with patch.object(QMessageBox,'information'):self.assertTrue(w.save())
  output=self.path/'窗口采购.xlsx'
  with patch.object(QFileDialog,'getSaveFileName',return_value=(str(output),'')) as choose,patch.object(QMessageBox,'information'):
   w.export_purchase_bom();self.assertTrue(output.exists());self.assertIn('ROOT1 ROOT 采购BOM.xlsx',choose.call_args.args[2])
 def test_sku_context_menu_and_cancel_export(self):
  self.main=MainWindow(self.s);output=self.path/'SKU采购.xlsx'
  def select_purchase(menu,*args):return next(a for a in menu.actions() if a.text()=='采购BOM导出')
  with patch.object(QMenu,'exec_',select_purchase),patch.object(QFileDialog,'getSaveFileName',return_value=(str(output),'')),patch.object(QMessageBox,'information'):
   self.main.sku_context_menu(QPoint(5,5))
  self.assertTrue(output.exists())
  with patch.object(QFileDialog,'getSaveFileName',return_value=('','')),patch.object(self.s,'export_purchase_bom') as exporter:
   self.main.export_sku_id(self.sid,purchase=True);exporter.assert_not_called()
