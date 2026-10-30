"""Additional annotated workflows, rendered only against the build-time DB copy."""
from pathlib import Path
from unittest.mock import patch
from PyQt5.QtCore import Qt,QPoint,QRect
from PyQt5.QtGui import QColor
from PyQt5.QtWidgets import QApplication,QMenu,QComboBox,QLineEdit,QDialogButtonBox,QToolBar,QPushButton

def generate(e):
    out,main,materials,bom,svc=[e[k] for k in ('out','main','materials','bom','svc')]
    snap,shown,dialog,ok,bar,action=[e[k] for k in ('snap','shown','dialog','ok','bar','action')]
    from ui.main_window import (CreateMaterialDialog,MaterialVersionDialog,MaterialUsageDialog,
        MaterialSelectionDialog,MaterialHeaderMappingDialog)
    from ui.interactive_windows import QFileDialog,QMessageBox,QDialog,QColorDialog,QInputDialog
    from ui.replacement_selection import ReplacementSelectionDialog
    from ui.material_clipboard import copy_material
    from ui.bom_colors import ColorMenu
    def accept_button(w):
        buttons=w.findChild(QDialogButtonBox)
        return next(b for b in buttons.buttons() if buttons.buttonRole(b)==QDialogButtonBox.AcceptRole)
    def table_rect(w,table,row,column=0):
        table.scrollToItem(table.item(row,column));QApplication.processEvents()
        rect=table.visualItemRect(table.item(row,column));rect.translate(table.viewport().mapTo(w,QPoint()));return rect
    def menu_shot(invoke,entries):
        found=set()
        def display(menu,*args):
            menu.setStyleSheet('QMenu {font-size:12px;} QMenu::item {padding:2px 12px;}')
            shown(menu)
            for label,key in entries:
                a=action(menu,label)
                if not a.isEnabled():raise RuntimeError('教程目标菜单禁用：'+key)
                snap(key,menu,rect=menu.actionGeometry(a));found.add(key)
            menu.hide();return None
        with patch.object(QMenu,'exec_',display):invoke()
        if found!={key for _,key in entries}:raise RuntimeError('未打开目标菜单')
    def question_shot(invoke,key):
        def question(parent,title,text,buttons=QMessageBox.Yes|QMessageBox.No,default=QMessageBox.No):
            box=dialog(QMessageBox(parent));box.setWindowTitle(title);box.setText(text);box.setStandardButtons(buttons);shown(box)
            target=box.button(QMessageBox.Yes) or box.buttons()[0];snap(key,box,target);box.hide();return QMessageBox.No
        with patch.object(QMessageBox,'question',question):invoke()
    def file_shots(title,path,keys):
        f=QFileDialog(main,title,str(Path(path).parent),'所有文件 (*)');f.setOption(QFileDialog.DontUseNativeDialog,True);f.setAcceptMode(QFileDialog.AcceptOpen);f.selectFile(str(path));dialog(f)
        snap(keys[0],f,f.findChild(QLineEdit,'fileNameEdit'));snap(keys[1],f,accept_button(f));f.hide()
    def select_node(node):
        bom.tree.clearSelection();bom.tree.setCurrentItem(node);bom.tree.scrollToItem(node);bom.show_item_info(node);shown(bom)
    def tree_menu(node,entries):
        select_node(node);menu_shot(lambda:bom.context_menu(bom.tree.visualItemRect(node).center()),entries)
    # Unlock only isolated snapshot rows so every editing tutorial is actionable.
    for sku in list(svc.skus()):
        if sku['locked']:svc.set_sku_locked(sku['id'],False)
    bom.dirty=False;bom.load();bom.apply_lock_state();bom.tree.expandToDepth(1)
    bar('db_import_menu',main,'文件');menu=action(main.menuBar(),'文件').menu();shown(menu)
    snap('db_import_action',menu,rect=menu.actionGeometry(action(menu,'导入项目数据库')));menu.hide()
    file_shots('导入项目数据库',Path(svc.db.path),('db_import_file','db_import_ok'))
    main.refresh_project_menu();shown(main.history_menu);snap('db_import_history',main.history_menu,rect=main.history_menu.actionGeometry(main.history_menu.actions()[0]));main.history_menu.hide()
    used=next(dict(r) for r in svc.materials() if r['idle']=='N' and not any(s['material_id']==r['id'] for s in svc.skus()))
    def library_menu(mid,entries):
        materials.locate_material(mid);shown(materials)
        row=next(i for i in range(materials.t.rowCount()) if materials.t.item(i,0).data(Qt.UserRole)==mid)
        menu_shot(lambda:materials.menu(materials.t.visualItemRect(materials.t.item(row,0)).center()),entries)
    library_menu(used['id'],[('类似新增物料','similar_menu'),('物料版本变更','revision_menu'),('物料反查','reverse_menu')])
    similar=dialog(CreateMaterialDialog(materials,data={**used,'base_code':'','version':''}))
    snap('similar_code',similar,similar.e['base_code']);snap('similar_version',similar,similar.e['version']);snap('similar_name',similar,similar.e['name']);ok('similar_ok',similar)
    refs=svc.material_version_references(used['id']);revision=dialog(MaterialVersionDialog(used,refs,materials))
    revision.version.setText(str(int(used['version'])+1) if used['version'].isdigit() else used['version']+'-NEW')
    snap('revision_value',revision,revision.version);snap('revision_refs',revision,revision.references);ok('revision_ok',revision)
    usage=dialog(MaterialUsageDialog(used,svc.material_usage(used['id']),materials));snap('reverse_rows',usage,usage.table)
    snap('reverse_open',usage,rect=table_rect(usage,usage.table,0));usage.hide()
    pos=main.table.visualItemRect(main.table.item(0,1)).center();sid=int(main.table.item(0,6).text())
    menu_shot(lambda:main.sku_context_menu(pos),[('删除SKU','delete_sku_menu')]);question_shot(lambda:main.delete_sku(sid),'delete_sku_confirm')
    idle=next(dict(r) for r in svc.materials() if r['idle']=='Y');library_menu(idle['id'],[('删除物料','material_delete_menu')])
    question_shot(lambda:QMessageBox.question(materials,'确认','删除物料？'),'material_delete_confirm')
    node=bom.root_item.child(0);select_node(node)
    tree_menu(node,[('删除物料','delete_node_menu'),('替换物料','replace_menu'),('复制物料','clipboard_copy'),('剪切物料','clipboard_cut'),('新增替代料','alternate_menu')])
    question_shot(lambda:bom.confirm_delete(node),'delete_node_confirm')
    bom.remove_relation(node);tree_menu(node,[('取消删除','delete_node_restore')]);bom.restore_relation(node)
    def choose_shots(invoke,first,confirm):
        def capture(w):
            dialog(w);combo=w.findChild(QComboBox)
            if combo.count()==0:raise RuntimeError('物料候选为空：'+first)
            combo.setCurrentIndex(0);snap(first,w,combo);snap(confirm,w,accept_button(w));w.hide();return QDialog.Rejected
        with patch.object(MaterialSelectionDialog,'exec_',capture):invoke()
    choose_shots(lambda:bom.replace_material(node),'replace_input','replace_ok')
    group=dict(nodes=[node],paths=[bom.root_item.text(0)+' / '+node.text(0)],current=True,reason='')
    # Conditional multi-occurrence dialog, populated with a real relationship.
    replacement=dialog(ReplacementSelectionDialog([group],bom));snap('replace_matches',replacement,replacement.table);ok('replace_matches_ok',replacement)
    question_shot(lambda:QMessageBox.question(bom,'复制原子BOM','目标物料没有BOM结构，是否将原物料的子BOM复制到目标物料？\n选择“否”仅替换物料，不复制子BOM；点击保存后才写入数据库。',QMessageBox.Yes|QMessageBox.No|QMessageBox.Cancel,QMessageBox.Cancel),'replace_subtree')
    # Prefer an existing same-base version; otherwise add one in the demo copy.
    version_node=next((n for _,n in bom.expansion_nodes() if n is not bom.root_item and bom.material_versions(n)),node)
    if not bom.material_versions(version_node):
        data=svc.material_snapshot(version_node.data(0,Qt.UserRole)['material_id']);svc.create_material({**data,'version':data['version']+'-HELP'})
    tree_menu(version_node,[('物料版本更新','bom_revision_menu')]);choose_shots(lambda:bom.update_material_version(version_node),'bom_revision_choice','bom_revision_ok')
    select_node(node);choose_shots(bom.add_alternate,'alternate_choice','alternate_ok')
    alt=next(r for r in svc.materials() if r['id'] not in bom.replacement_blocked_materials(node) and r['id'] not in bom.alternate_ids(node))
    bom._alternate_edits[bom.alternate_key(node)]=bom.alternate_ids(node)+[alt['id']];bom.mark_dirty();bom.show_item_info(node)
    table=bom.alt_table;row=table.rowCount()-1;shown(bom);bom.info_scroll.ensureWidgetVisible(table);QApplication.processEvents();snap('alternate_select',bom,rect=table_rect(bom,table,row))
    menu_shot(lambda:bom.alternate_menu(table.visualItemRect(table.item(row,0)).center()),[('切换成主物料','alternate_promote')])
    # A library item not already under the destination provides an allowed paste.
    copy_material(svc,alt['id']);tree_menu(node,[('粘贴为同级物料','clipboard_peer'),('粘贴为子级物料','clipboard_child')])
    siblings=[bom.root_item.child(i) for i in range(bom.root_item.childCount())]
    up=next(n for n in siblings if bom.can_move_item(n,-1));down=next(n for n in siblings if bom.can_move_item(n,1))
    tree_menu(up,[('物料上移','move_up')]);tree_menu(down,[('物料下移','move_down')])
    tree_menu(node,[('设置颜色','color_menu')]);colors=ColorMenu(bom,[node],bom);shown(colors)
    for label,key in [('浅蓝','color_default'),('设置更多颜色','color_more')]:snap(key,colors,rect=colors.actionGeometry(action(colors,label)))
    colors.hide();color=QColorDialog(QColor('#4169a1'),bom);color.setOption(QColorDialog.DontUseNativeDialog,True);dialog(color)
    # Mark the color selection region, excluding the bottom confirmation buttons.
    snap('color_pick',color,rect=QRect(10,10,color.width()-20,color.height()-65));snap('color_ok',color,accept_button(color));color.hide()
    question_shot(lambda:QMessageBox.question(bom,'保存颜色','是否将此颜色保存到“设置颜色”菜单？',QMessageBox.Yes|QMessageBox.No,QMessageBox.Yes),'color_remember')
    from ui.bom_compare import BomCompareDialog
    bar('compare_menu',main,'BOM对比');compare=dialog(BomCompareDialog(svc,main));compare.left.setCurrentIndex(1);snap('compare_left',compare,compare.left)
    compare.right.setCurrentIndex(min(2,compare.right.count()-2));snap('compare_right',compare,compare.right);snap('compare_run',compare,compare.button);compare.compare();shown(compare)
    snap('compare_diff',compare,compare.differences_only)
    table=compare.tables[0];index=next(table.model().index(i,1) for i in range(table.rowCount()) if table.item(i,1) is not None)
    menu_shot(lambda:compare.context_menu(table,index,table.viewport().mapToGlobal(QPoint(30,40))),[('打开BOM进行编辑','compare_edit')]);compare.close()
    # Import a leaf's children; root in the file must identify this selected leaf.
    leaf=next(n for _,n in bom.expansion_nodes() if n is not bom.root_item and n.childCount()==0 and not bom.svc.bom(n.data(0,Qt.UserRole)['material_id']))
    tree_menu(leaf,[('导入BOM','bom_import_menu')]);target=svc.material_snapshot(leaf.data(0,Qt.UserRole)['material_id'])
    from openpyxl import Workbook
    fields=svc.BOM_HEADERS;rows=[['ROOT','ROOT',target['base_code'],target['version'],1,'']]
    book=Workbook();book.active.title='BOM';book.active.append(fields);book.active.append(rows[0]);book.create_sheet('填写说明');path=out/'BOM导入示例.xlsx';book.save(path);book.close()
    file_shots('导入BOM',path,('bom_import_file','bom_import_open'))
    sheet=QInputDialog(bom);sheet.setWindowTitle('选择BOM工作表');sheet.setLabelText('请选择BOM所在的工作表（Sheet）：');sheet.setComboBoxItems(['BOM','填写说明']);dialog(sheet);snap('bom_import_sheet',sheet,sheet.findChild(QComboBox));ok('bom_import_sheet_ok',sheet)
    mapping=dialog(MaterialHeaderMappingDialog(str(path),fields,[],bom,sheet_name='BOM',required=fields,preview_rows=rows))
    for i,field in enumerate(fields):
        mapping.mapping_table.scrollToItem(mapping.mapping_table.item(0,i));QApplication.processEvents();snap('bom_mapping_'+str(i),mapping,mapping.combos[field])
    add=next(b for b in mapping.findChildren(QPushButton) if b.text()=='新增替代料');snap('bom_mapping_alt',mapping,add);ok('bom_mapping_ok',mapping);path.unlink()
    from ui.main_window import BomImportIssuesDialog
    from services.services import BomImportError
    parent=bom.root_item.text(0);child=target['base_code']+target['version']+' '+target['name']
    comparison=dict(parent=parent,material=child,field='位置',left_path=parent+' / 示例分支一',right_path=parent+' / 示例分支二',left_line=12,right_line=34,left='（空）',right='/')
    mismatch=dialog(BomImportIssuesDialog(BomImportError(['示例：两个分支的子BOM位置不一致，禁止导入'],[comparison]),bom))
    snap('bom_import_differences',mismatch,mismatch.table);mismatch.hide()
    bom.dirty=False;bom.load()
