#!/usr/bin/env python3
"""Read the supplied BrainAI Excel; validate imports only in a temporary database."""
import argparse,hashlib,json,sys,tempfile
from pathlib import Path
ROOT=Path(__file__).resolve().parents[1];sys.path.insert(0,str(ROOT/'src'))
from openpyxl import load_workbook
from database.database import Database
from services.services import Services,BomImportError


def main():
 parser=argparse.ArgumentParser();parser.add_argument('--source',required=True);parser.add_argument('--report',required=True);args=parser.parse_args()
 source=Path(args.source);before=hashlib.sha256(source.read_bytes()).hexdigest()
 book=load_workbook(source,data_only=True);sheet=book.worksheets[0]
 mapping=dict(zip(Services.BOM_HEADERS,['父项编码','父项版本','物料编码','物料版本','BOM节点数量','位号']))
 report={}
 try:
  with tempfile.TemporaryDirectory(dir=ROOT/'.tmp') as tmp:
   db=Database(str(Path(tmp)/'fixture.db'));svc=Services(db);ids={}
   try:
    for row in sheet.iter_rows(min_row=2,values_only=True):
     if not row[3]:continue
     key=(str(row[3]),str(row[4]))
     if key not in ids:ids[key]=svc.create_material(dict(base_code=key[0],version=key[1],name=str(row[5] or ''),unit='pcs'))
    root=ids[str(sheet.cell(2,4).value),str(sheet.cell(2,5).value)]
    snapshot=list(db.conn.iterdump())
    try:svc.import_bom(source,root,header_mapping=mapping,persist=False)
    except BomImportError as error:report['original_issues']=error.issues
    else:raise AssertionError('Original position differences should block import')
    assert len(report['original_issues'])==3,report['original_issues']
    assert all('子BOM不一致' in issue and '位置' in issue for issue in report['original_issues'])
    assert snapshot==list(db.conn.iterdump())
    # Controlled copy only: make the three positions match for the success case.
    for left,right in [(224,246),(225,247),(226,248)]:sheet.cell(right,8).value=sheet.cell(left,8).value
    corrected=Path(tmp)/'same-subtrees.xlsx';book.save(corrected)
    nodes=svc.import_bom(corrected,root,header_mapping=mapping,persist=False)
    subtrees=[]
    def walk(items):
     for node in items:
      if node['data']['base_code']=='BC1-9900149':subtrees.append(node['children'])
      walk(node['children'])
    walk(nodes)
    assert len(subtrees)==2 and len(subtrees[0])==len(subtrees[1])==15
    assert snapshot==list(db.conn.iterdump())
    report.update(corrected_copy_preview=True,shared_occurrences=2,children_per_occurrence=15,database_unchanged=True)
   finally:db.close()
 finally:book.close()
 report['source_unchanged']=before==hashlib.sha256(source.read_bytes()).hexdigest();assert report['source_unchanged']
 report['source_sha256']=before
 Path(args.report).write_text(json.dumps(report,ensure_ascii=False,indent=2));print(json.dumps(report,ensure_ascii=False,indent=2))
if __name__=='__main__':main()
