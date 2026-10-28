#!/usr/bin/env python3
"""Measure GUI operations on a read-only backup, never on the user's live DB."""
import os,sys,argparse,tempfile,sqlite3,hashlib,json
from pathlib import Path
os.environ['QT_QPA_PLATFORM']='offscreen'
ROOT=Path(__file__).resolve().parents[1];sys.path.insert(0,str(ROOT/'src'))
from PyQt5.QtWidgets import QApplication
from database.database import Database
from services.services import Services
from ui.main_window import MainWindow,MaterialWindow,BomWindow
from ui.operation_progress import operation,RECENT_OPERATIONS
from ui.bom_compare import BomCompareDialog

def main():
 parser=argparse.ArgumentParser();parser.add_argument('--source',required=True);parser.add_argument('--report',required=True);parser.add_argument('--screenshot');args=parser.parse_args()
 from ui.operation_progress import OperationProgress
 original_update=OperationProgress.update
 captured=[]
 def observe(controller,*args,**kwargs):
  original_update(controller,*args,**kwargs)
  if parsed.screenshot and not captured and controller.dialog is not None and controller.dialog.isVisible():
   controller.dialog.grab().save(parsed.screenshot);captured.append(True)
 parsed=args
 OperationProgress.update=observe
 source=Path(args.source).resolve();before=hashlib.sha256(source.read_bytes()).hexdigest();app=QApplication([])
 with tempfile.TemporaryDirectory(dir=ROOT/'.tmp') as tmp:
  path=Path(tmp)/'db';a=sqlite3.connect(source.as_uri()+'?mode=ro',uri=True);b=sqlite3.connect(path)
  try:a.backup(b)
  finally:a.close();b.close()
  db=Database(str(path));s=Services(db);skus=list(s.skus());largest=max(skus,key=lambda r:len(s.for_sku(r['id']).bom(r['material_id'])));sid=largest['id'];scoped=s.for_sku(sid);mid=largest['material_id']
  RECENT_OPERATIONS.clear();main=MainWindow(s);main.show();main.show_materials();main.open_bom_id(sid)
  main.refresh();main._materials.refresh()
  outputs={}
  for name,func in [('导出SKU清单',s.export_sku_list),('导出物料库',s.export_materials),('导出资料库',s.export_database),('导出BOM',lambda p:scoped.export_bom_structure(mid,p)),('导出采购BOM',lambda p:scoped.export_purchase_bom(mid,p))]:
   output=str(Path(tmp)/(name+'.xlsx'));operation(name)(lambda:func(output))();outputs[name]=output
  main.compare_sku_bom(sid)
  compare=next(w for w in app.topLevelWidgets() if isinstance(w,BomCompareDialog));compare.right.setCurrentIndex(compare.right.findData(sid));compare.compare()
  # Exercise sheet scan/validation/import with existing rows, safely skipping identical rows.
  operation('物料导入扫描')(lambda:s.import_materials(outputs['导出物料库'],lambda *args:'skip',sheet_name='物料库'))()
  mapping={'父项物料编码':'父项编码','父项版本号':'父项版本','物料基础编码':'物料基础编码','物料版本号':'物料版本号','物料数量':'用量','位置':'物料位置'}
  from openpyxl import load_workbook
  book=load_workbook(outputs['导出BOM'],read_only=True);sheet_name=next(ws.title for ws in book if '父项编码' in next(ws.iter_rows(values_only=True),()));book.close()
  operation('BOM导入扫描')(lambda:scoped.import_bom(outputs['导出BOM'],mid,sheet_name=sheet_name,header_mapping=mapping,persist=False,conflict_resolver=lambda _: 'system'))()
  records=list(RECENT_OPERATIONS)
  assert db.conn.execute('PRAGMA integrity_check').fetchone()[0]=='ok'
  for window in app.topLevelWidgets():
   if isinstance(window,BomWindow):window.dirty=False
   window.close()
  db.close()
 assert hashlib.sha256(source.read_bytes()).hexdigest()==before
 report={'source_sha256':before,'operations':records,'slow_without_progress':[r for r in records if r['ms']>100 and not r['progress_shown']]}
 Path(args.report).write_text(json.dumps(report,ensure_ascii=False,indent=2));print(json.dumps(report,ensure_ascii=False,indent=2))
if __name__=='__main__':main()
